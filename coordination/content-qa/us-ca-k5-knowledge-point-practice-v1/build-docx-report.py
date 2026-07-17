#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PACKAGE_DIR = Path(__file__).resolve().parent
OUTPUT = PACKAGE_DIR / "qa-report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)
LIGHT_GRAY = "F2F4F7"
BORDER = "A6A6A6"


def read_json(name: str):
    return json.loads((PACKAGE_DIR / name).read_text(encoding="utf-8"))


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_paragraph(doc, text="", size=11, color=BLACK, bold=False, italic=False, after=6, before=0):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 8
    elif level == 2:
        size, color, before, after = 13, BLUE, 12, 6
    else:
        size, color, before, after = 12, DARK_BLUE, 8, 4
    set_paragraph_spacing(p, before=before, after=after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=True)
    return p


def add_rule(paragraph, color="2E74B5", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_text(cell, text, bold=False, color=BLACK, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)
    p = cell.paragraphs[0]
    p.alignment = align
    set_paragraph_spacing(p, before=0, after=0, line=1.10)
    run = p.add_run(str(text))
    set_run_font(run, size=10.5, color=color, bold=bold)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    total = sum(widths)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Pt(width / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, value in [("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")]:
                node = tc_mar.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    tc_mar.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, fill=LIGHT_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, align=WD_ALIGN_PARAGRAPH.CENTER if compact or i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def grade_table(pack):
    order = ["K", "P1", "P2", "P3", "P4", "P5"]
    topic_counts = {}
    for row in pack["topicCoverage"]:
        topic_counts[row["grade"]] = topic_counts.get(row["grade"], 0) + 1
    return [
        [grade, topic_counts.get(grade, 0), pack["counts"]["gradeCounts"].get(grade, 0)]
        for grade in order
    ]


def main():
    pack = read_json("question-pack.json")
    round1 = read_json("qa-round-1.json")
    round2 = read_json("qa-round-2.json")
    solvability = read_json("solvability-audit.json")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(header, after=0)
    run = header.add_run("S18 QA Report | Candidate Only")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(footer, after=0)
    run = footer.add_run("MAIS-MVP California K-G5 Knowledge-Point Practice Candidate")
    set_run_font(run, size=9, color=MUTED)

    add_paragraph(doc, "S18 QA REPORT", size=23, color=BLACK, bold=True, after=4)
    add_paragraph(doc, "California K-G5 New RAG Knowledge-Point Practice Candidate", size=14, color=MUTED, after=14)
    metadata = [
        ("Package", pack["packageId"]),
        ("Date", "June 22, 2026"),
        ("Owners", "S21 content pipeline generation; S18 curriculum QA"),
        ("Decision", "candidate-only-two-round-qa-pass"),
        ("Live status", "Not integrated into live student practice in this dialogue"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=2)
        r = p.add_run(f"{label}: ")
        set_run_font(r, bold=True)
        r = p.add_run(value)
        set_run_font(r)
    rule = doc.add_paragraph()
    set_paragraph_spacing(rule, after=12)
    add_rule(rule)

    add_heading(doc, "Executive Verdict", 1)
    add_paragraph(
        doc,
        "The candidate package passes two internal QA rounds for topic coverage, deterministic answer validity, source-distance boundaries, bilingual field completeness, and no-live-integration status. It should remain candidate-only until a future owner assignment routes it through S23, S04/S05, S11, and S22.",
    )

    add_heading(doc, "Package Scope", 1)
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Knowledge-point topics", len(pack["topicCoverage"])],
            ["Practice questions", len(pack["questions"])],
            ["Questions per topic", pack["counts"]["questionsPerKnowledgePoint"]],
            ["Textbook lesson topics", "29"],
            ["Grade 1 micro-lesson topics", "12"],
            ["Integration status", pack["integrationStatus"]],
        ],
        [3100, 6260],
    )

    add_heading(doc, "Grade Coverage", 1)
    add_table(
        doc,
        ["Grade", "Topics", "Questions"],
        grade_table(pack),
        [1500, 2000, 2000],
        compact=True,
    )

    doc.add_page_break()
    add_heading(doc, "Two-Round QA Evidence", 1)
    add_table(
        doc,
        ["QA Round", "Method", "Accepted", "Errors", "Warnings", "Status"],
        [
            ["Round 1", "Inventory, answer-key, MC uniqueness, source-visible scan", f"{round1['counts']['acceptedRows']}/{round1['counts']['totalQuestions']}", round1["counts"]["errorCount"], round1["counts"]["warningCount"], round1["status"]],
            ["Round 2", "Coverage, grade fit, bilingual fields, no-live boundary", f"{round2['counts']['acceptedRows']}/{round2['counts']['totalQuestions']}", round2["counts"]["errorCount"], round2["counts"]["warningCount"], round2["status"]],
            ["Solvability", "Stored, independent, and computed answer comparison", f"{solvability['counts']['deterministicPass']}/{solvability['counts']['totalQuestions']}", 0, 0, solvability["status"]],
        ],
        [1300, 3550, 1350, 850, 900, 900],
        compact=True,
    )

    add_heading(doc, "Source-Safety Boundary", 1)
    add_paragraph(
        doc,
        "The package uses public standards identifiers and existing MAIS-authored knowledge-point metadata only. It does not commit raw textbook text, private corpus chunks, IXL exercise text, IXL preview wording, official standards prose, released assessment items, screenshots, answer keys, or copied diagrams.",
    )

    add_heading(doc, "Representative Candidate Rows", 1)
    sample_rows = [
        [q["grade"], q["topicId"], q["type"], q["difficulty"], q["answer"]]
        for q in pack["questions"][:4]
    ]
    add_table(
        doc,
        ["Grade", "Topic ID", "Type", "Difficulty", "Answer"],
        sample_rows,
        [900, 4200, 1450, 1300, 1510],
    )

    doc.add_page_break()
    add_heading(doc, "Release Boundary And Next Gates", 1)
    next_gates = [
        ("S23 integration and promotion", "Create a candidate-to-live mapping and scope plan if the owner later asks to integrate."),
        ("S04 practice lead", "Receive explicit live practice data assignment before any app question-bank import."),
        ("S05 lesson lead", "Receive explicit assignment before attaching practiceQuestionIds to live lessons."),
        ("S11 QA", "Run route, practice filter, and regression evidence after any live integration."),
        ("S22 release engineering", "Run build and release readiness gates before production deployment.")
    ]
    add_table(doc, ["Owner", "Required Gate"], next_gates, [2500, 6860])

    add_heading(doc, "Risks", 1)
    risks = [
        "This is a topic-adapted seed bank, not a complete production bank.",
        "Chinese fields are deterministic functional translations and should receive S09 language polish before public use.",
        "No answer-critical visuals are included; future visual rows need deterministic SVG/exact-layer review.",
        "Public copy must not claim a complete or official California curriculum from this candidate-only package.",
    ]
    for risk in risks:
        add_paragraph(doc, risk)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Evidence File Index", 1)
    artifact_rows = [
        ["question-pack.json", f"{len(pack['questions'])}-row candidate package"],
        ["topic-coverage.json", "41 target knowledge-point mapping records"],
        ["qa-round-1.json", "Round 1 deterministic QA evidence"],
        ["qa-round-2.json", "Round 2 independent QA evidence"],
        ["solvability-audit.json/csv/md", "Answer consistency evidence"],
        ["manual-review-results.csv", "Row-level candidate-only acceptance ledger"],
        ["s18-two-round-qa-decision.md", "Candidate-only S18 decision note"],
    ]
    add_table(doc, ["Artifact", "Purpose"], artifact_rows, [3200, 6160])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
