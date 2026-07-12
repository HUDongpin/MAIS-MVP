#!/usr/bin/env python3
"""Generate a polished DOCX guide from the committed U.S. math safe-RAG data."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUT_PATH = PROJECT_ROOT / "coordination" / "content-qa" / "2026-06-01-S18-us-top11-math-rag-standards-guide.docx"
BUILD_DIR = PROJECT_ROOT / ".tmp" / "us-rag-docx-build"
REPORT_DATE = "2026-06-01"
GUIDE_TITLE = "MAIS U.S. Top-11 Math RAG"
GUIDE_HEADER = "MAIS U.S. Top-11 Math RAG Standards Guide"
GUIDE_STATE_LIST = "California, Texas, Florida, New York, Pennsylvania, Illinois, Ohio, Georgia, North Carolina, Michigan, and Arkansas"

BUNDLED_NODE = Path("/Users/dongpinhu/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node")

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(94, 108, 132)
GOLD = RGBColor(122, 90, 0)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)

FILL_BLUE = "E8EEF5"
FILL_LIGHT = "F4F6F9"
FILL_NAVY = "0B2545"
FILL_GOLD = "FFF4CC"
BORDER = "CBD5E1"

GRADE_ORDER = ["P1", "P2", "P3", "P4", "P5", "P6", "S1", "S2", "S3", "S4", "S5", "S6"]


def run_command(args: list[str], cwd: Path) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(args, cwd=cwd, text=True, capture_output=True, check=False)
    if completed.returncode != 0:
        raise RuntimeError(
            "Command failed:\n"
            + " ".join(args)
            + "\n\nSTDOUT:\n"
            + completed.stdout
            + "\n\nSTDERR:\n"
            + completed.stderr
        )
    return completed


def node_path() -> str:
    configured = os.environ.get("CODEX_NODE")
    if configured:
        return configured
    if BUNDLED_NODE.exists():
        return str(BUNDLED_NODE)
    return "node"


def compile_and_extract_us_rag() -> dict[str, Any]:
    if BUILD_DIR.exists():
        shutil.rmtree(BUILD_DIR)
    BUILD_DIR.mkdir(parents=True, exist_ok=True)

    tsc = PROJECT_ROOT / "node_modules" / "typescript" / "bin" / "tsc"
    run_command(
        [
            node_path(),
            str(tsc),
            "-p",
            "tsconfig.rag.json",
            "--outDir",
            str(BUILD_DIR.relative_to(PROJECT_ROOT)),
            "--noEmit",
            "false",
            "--incremental",
            "false",
        ],
        PROJECT_ROOT,
    )

    compiled_data = BUILD_DIR / "data" / "rag" / "usMath.js"
    js = r"""
const data = require(process.argv[1]);
const gradeOrder = ["K","P1","P2","P3","P4","P5","P6","S1","S2","S3","S4","S5","S6"];
const profiles = data.unitedStatesMathStateProfiles
  .slice()
  .sort((a, b) => a.populationRank - b.populationRank);

function countsBy(items, key) {
  return items.reduce((acc, item) => {
    const value = typeof key === "function" ? key(item) : item[key];
    acc[value] = (acc[value] || 0) + 1;
    return acc;
  }, {});
}

function unique(values) {
  return Array.from(new Set(values.filter(Boolean)));
}

const sourceById = Object.fromEntries(data.unitedStatesMathSourceRegistry.map((source) => [source.id, source]));

const states = profiles.map((profile) => {
  const allCards = data.unitedStatesMathSafeCards.filter((card) => card.state === profile.state);
  const overviewCards = data.unitedStatesMathGradeOverviewCards.filter((card) => card.state === profile.state);
  const standardsCards = data.unitedStatesMathStandardsLibraryCards.filter((card) => card.state === profile.state);
  const textbookCards = data.unitedStatesMathTextbookCompatibilityCards.filter((card) => card.state === profile.state);
  const examCards = data.unitedStatesMathExamPatternCards.filter((card) => card.state === profile.state);
  const sourceIds = unique(allCards.flatMap((card) => card.sourceIds));

  const grades = gradeOrder.map((grade) => {
    const standards = standardsCards
      .filter((card) => card.grade === grade)
      .map((card) => ({
        id: card.standardIds[0],
        domain: card.domainTags[0],
        cluster: card.clusterTags[0],
        concepts: card.conceptIds,
        competencies: card.competencyTags,
        topics: card.topicIds,
      }));
    const overview = overviewCards.find((card) => card.grade === grade);
    const textbook = textbookCards.find((card) => card.grade === grade);
    const exam = examCards.find((card) => card.grade === grade);
    return {
      grade,
      usGradeLabel: overview ? overview.usGradeLabel : grade,
      standards,
      topics: unique(standards.flatMap((entry) => entry.topics)),
      concepts: unique(standards.flatMap((entry) => entry.concepts)),
      competencies: unique(standards.flatMap((entry) => entry.competencies)),
      overviewSummary: overview ? overview.safeSummary : "",
      textbookNote: textbook ? textbook.textbookCompatibilityNotes.join(" ") : "",
      examItemTags: exam ? exam.itemTypeTags : [],
      examNote: exam ? exam.examPatternNotes.join(" ") : "",
    };
  });

  return {
    profile,
    counts: countsBy(allCards, "cardKind"),
    totalCards: allCards.length,
    sourceIds,
    sources: sourceIds.map((id) => sourceById[id]).filter(Boolean),
    grades,
  };
});

const result = {
  generatedAt: new Date().toISOString().slice(0, 10),
  profiles,
  states,
  totals: {
    stateCount: profiles.length,
    sourceCount: data.unitedStatesMathSourceRegistry.length,
    rawCorpusAllowedCount: data.unitedStatesMathSourceRegistry.filter((source) => source.rawCorpusAllowed).length,
    cards: data.unitedStatesMathSafeCards.length,
    byKind: countsBy(data.unitedStatesMathSafeCards, "cardKind"),
    byLibraryLane: countsBy(data.unitedStatesMathSafeCards, "libraryLane"),
  },
  sourceRegistry: data.unitedStatesMathSourceRegistry,
};

console.log(JSON.stringify(result));
"""
    extracted = run_command([node_path(), "-e", js, str(compiled_data)], PROJECT_ROOT)
    data = json.loads(extracted.stdout)
    shutil.rmtree(BUILD_DIR, ignore_errors=True)
    return data


def set_run_font(run, *, name: str = "Calibri", size: float | None = None, color: RGBColor | None = None, bold: bool | None = None, italic: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(paragraph, *, before: float = 0, after: float = 6, line_spacing: float = 1.25, keep_next: bool = False, alignment: WD_ALIGN_PARAGRAPH | None = None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    if alignment is not None:
        paragraph.alignment = alignment
    if keep_next:
        p_pr = paragraph._p.get_or_add_pPr()
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def add_para(doc: Document, text: str = "", *, size: float = 11, color: RGBColor = BLACK, bold: bool = False, italic: bool = False, before: float = 0, after: float = 6, line_spacing: float = 1.25, align: WD_ALIGN_PARAGRAPH | None = None, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    set_paragraph_format(paragraph, before=before, after=after, line_spacing=line_spacing, alignment=align)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    style = f"Heading {level}"
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    set_paragraph_format(
        paragraph,
        before={1: 18, 2: 14, 3: 10}.get(level, 8),
        after={1: 10, 2: 7, 3: 5}.get(level, 4),
        line_spacing=1.15,
        keep_next=True,
    )
    return paragraph


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_sep, fld_text, fld_char_end])


def paragraph_border_bottom(paragraph, color: str = "2E74B5", size: str = "12", space: str = "8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run(GUIDE_HEADER)
    set_run_font(run, size=9, color=MUTED, bold=True)
    paragraph_border_bottom(header, color="D9E2EC", size="6", space="4")

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"Prepared {REPORT_DATE} | Page ")
    set_run_font(run, size=9, color=MUTED)
    add_page_number(footer)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, size: float = 9, color: RGBColor = BLACK, bold: bool = False, fill: str | None = None, align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT, line_spacing: float = 1.08):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    set_paragraph_format(paragraph, before=0, after=0, line_spacing=line_spacing, alignment=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def ensure_tbl_pr(table):
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    return tbl_pr


def replace_child(parent, tag: str, child):
    existing = parent.find(qn(tag))
    if existing is not None:
        parent.remove(existing)
    parent.append(child)


def set_table_borders(table, color: str = BORDER):
    tbl_pr = ensure_tbl_pr(table)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    replace_child(tbl_pr, "w:tblBorders", borders)


def set_table_margins(table, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120):
    tbl_pr = ensure_tbl_pr(table)
    margins = OxmlElement("w:tblCellMar")
    for name, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        element = OxmlElement(f"w:{name}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    replace_child(tbl_pr, "w:tblCellMar", margins)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = ensure_tbl_pr(table)

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    replace_child(tbl_pr, "w:tblW", tbl_w)

    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    replace_child(tbl_pr, "w:tblInd", tbl_ind)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    replace_child(tbl_pr, "w:tblLayout", layout)

    grid = table._tbl.tblGrid
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    table._tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")

    set_table_margins(table)
    set_table_borders(table)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int], *, header_fill: str = FILL_BLUE, font_size: float = 8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, size=8.5, color=NAVY, bold=True, fill=header_fill, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], value, size=font_size, color=BLACK, align=align)
    add_para(doc, "", after=4)
    return table


def add_callout(doc: Document, label: str, text: str, *, fill: str = FILL_LIGHT, accent: RGBColor = DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA], indent_dxa=0)
    set_table_borders(table, color="D9E2EC")
    set_table_margins(table, top=120, bottom=120, start=180, end=180)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    set_paragraph_format(paragraph, before=0, after=0, line_spacing=1.15)
    run = paragraph.add_run(label + " ")
    set_run_font(run, size=10, color=accent, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=10, color=BLACK)
    add_para(doc, "", after=6)


def comma_list(values: list[str], limit: int | None = None) -> str:
    cleaned = [str(value) for value in values if value]
    if limit is not None and len(cleaned) > limit:
        return ", ".join(cleaned[:limit]) + f", +{len(cleaned) - limit} more"
    return ", ".join(cleaned)


def standards_cell(standards: list[dict[str, Any]]) -> str:
    return "\n".join(f"{entry['id']} - {entry['domain']} ({entry['cluster']})" for entry in standards)


def add_cover(doc: Document, data: dict[str, Any]):
    add_para(doc, "CURRICULUM REFERENCE GUIDE", size=10, color=GOLD, bold=True, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, GUIDE_TITLE, size=28, color=NAVY, bold=True, after=6, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
    add_para(doc, "State Standards Safe-Card Display", size=17, color=DARK_BLUE, bold=True, after=18, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.1)
    add_para(
        doc,
        f"A polished DOCX output of the committed U.S. math RAG files for {GUIDE_STATE_LIST}.",
        size=11,
        color=MUTED,
        after=24,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.25,
    )

    totals = data["totals"]
    rows = [
        ["States", str(totals["stateCount"]), "Eleven-state owner-priority safe-RAG set"],
        ["Safe cards", str(totals["cards"]), "Committed standards-safe abstractions"],
        ["Standards-family cards", str(totals["byKind"]["standards"]), "Identifier/domain/topic coverage"],
        ["Raw corpus allowed", str(totals["rawCorpusAllowedCount"]), "No raw official, assessment, textbook, or OER body text"],
    ]
    add_table(doc, ["Metric", "Value", "Meaning"], rows, [2200, 1400, 5760], header_fill=FILL_NAVY, font_size=9)
    for cell in doc.tables[-1].rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = WHITE

    add_callout(
        doc,
        "Source boundary:",
        "This guide displays state profiles, standard identifiers, domain/topic tags, card counts, source metadata, and MAIS-authored summaries only. It does not reproduce official standards wording, released assessment items, textbook content, figures, tables, rubrics, or source passages.",
        fill=FILL_GOLD,
        accent=GOLD,
    )
    add_para(doc, f"Prepared for Dr. Peter Hu | S18 Curriculum QA | {REPORT_DATE}", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()


def add_overview(doc: Document, data: dict[str, Any]):
    add_heading(doc, "1. Eleven-State Coverage Dashboard", 1)
    add_para(
        doc,
        "The committed U.S. RAG layer is organized as safe cards. Each state has grade-overview, standards-family, textbook-compatibility, and exam-pattern cards, while all raw corpus flags remain off. Arkansas is included as the eleventh owner-priority state expansion.",
        size=10.5,
        color=BLACK,
        after=8,
    )
    rows = []
    for state in data["states"]:
        profile = state["profile"]
        rows.append(
            [
                f"{profile['state']}\n{profile['displayName']}",
                f"Phase {profile['statePriorityPhase']}\nRank {profile['populationRank']}",
                profile["standardsName"],
                profile["commonCoreStatus"].replace("-", " "),
                profile["assessmentProgram"],
                f"{state['totalCards']} cards\n{state['counts'].get('standards', 0)} standards",
            ]
        )
    add_table(
        doc,
        ["State", "Priority", "Standards System", "CCSS Status", "Assessment Metadata", "RAG Coverage"],
        rows,
        [900, 950, 2450, 1250, 2450, 1360],
        header_fill=FILL_BLUE,
        font_size=8,
    )

    add_heading(doc, "2. Safe-RAG Use Guardrails", 1)
    rows = [
        ["Allowed display", "State/profile metadata, standard IDs, grade/domain structure, MAIS topic tags, competency tags, abstract assessment-design signals, source registry metadata, and MAIS-authored summaries."],
        ["Not allowed", "Official standards wording, lightly rewritten standards text, released/sample item stems, answer keys, scoring rubrics, diagrams, screenshots, textbook body text, examples, exercises, teacher notes, or assessment-bank content."],
        ["Generation rule", "Use the cards as alignment targets and design constraints. Generate fresh MAIS-authored contexts, values, prompts, hints, explanations, diagrams, and distractors."],
        ["Commercial readiness", "The RAG layer is safe-card metadata. Commercial pilots still require counsel review for source use and written authorization before any raw OER or publisher content is ingested."],
    ]
    add_table(doc, ["Policy Area", "Working Rule"], rows, [1900, 7460], header_fill=FILL_BLUE, font_size=9)
    doc.add_page_break()


def add_state_section(doc: Document, state: dict[str, Any], display_index: int):
    profile = state["profile"]
    add_heading(doc, f"{display_index}. {profile['displayName']} ({profile['state']})", 1)
    add_para(
        doc,
        f"{profile['standardsName']} | {profile['standardsVersion']}",
        size=10.5,
        color=DARK_BLUE,
        bold=True,
        after=4,
    )
    add_callout(doc, "No-endorsement notice:", profile["noEndorsementNotice"], fill=FILL_LIGHT, accent=DARK_BLUE)

    profile_rows = [
        ["Curriculum track", profile["curriculumTrack"]],
        ["Priority phase", f"Phase {profile['statePriorityPhase']} | Population rank {profile['populationRank']}"],
        ["Standards prefix", profile["standardPrefix"]],
        ["Common Core relation", f"{profile['commonCoreStatus'].replace('-', ' ')} | Crosswalk: {profile['crosswalkRelationToCcss']}"],
        ["Adoption policy", profile["adoptionPolicy"]],
        ["Materials policy", profile["materialsPolicy"]],
        ["Assessment metadata", profile["assessmentProgram"]],
    ]
    add_table(doc, ["Profile Field", "RAG Value"], profile_rows, [2100, 7260], header_fill=FILL_BLUE, font_size=8.7)

    add_heading(doc, "Standards-Family Matrix", 2)
    add_para(
        doc,
        "Each row lists the committed identifier/domain coverage for one MAIS grade mapping. These are safe-card identifiers and MAIS-authored tags, not official standards text.",
        size=9.5,
        color=MUTED,
        after=5,
    )
    rows = []
    for grade in state["grades"]:
        exam_text = comma_list(grade["examItemTags"], 4) if grade["examItemTags"] else "No exam-pattern card in v1 for this grade"
        rows.append(
            [
                f"{grade['grade']}\n{grade['usGradeLabel']}",
                standards_cell(grade["standards"]),
                comma_list(grade["topics"], 5),
                exam_text,
            ]
        )
    add_table(
        doc,
        ["Grade", "Standard Families / Domains", "MAIS Topic Links", "Assessment Design Signals"],
        rows,
        [900, 4320, 2260, 1880],
        header_fill=FILL_BLUE,
        font_size=7.6,
    )

    add_heading(doc, "State Sources in the Safe Registry", 2)
    source_rows = []
    for source in state["sources"]:
        if source["state"] == "US":
            continue
        source_rows.append(
            [
                source["id"],
                f"{source['owner']}\n{source['sourceKind'].replace('-', ' ')}",
                source["repositoryRetention"].replace("-", " "),
                "Yes" if source["safeCardAllowed"] else "No",
                "Yes" if source["rawCorpusAllowed"] else "No",
            ]
        )
    add_table(
        doc,
        ["Source ID", "Owner / Kind", "Retention", "Safe Card", "Raw Corpus"],
        source_rows,
        [2200, 3000, 1900, 1100, 1160],
        header_fill=FILL_BLUE,
        font_size=7.6,
    )
    doc.add_page_break()


def add_appendix(doc: Document, data: dict[str, Any]):
    add_heading(doc, "Appendix A. Card Count and Source Registry Summary", 1)
    totals = data["totals"]
    rows = [
        ["Grade overview", str(totals["byKind"].get("grade-overview", 0)), "One per state per MAIS grade mapping."],
        ["Standards-family", str(totals["byKind"].get("standards", 0)), "Standard identifiers, domain tags, concept tags, and topic links."],
        ["Textbook compatibility", str(totals["byKind"].get("textbook-compatibility", 0)), "Publisher-neutral implementation bridge; authorization required for publisher-specific mapping."],
        ["Exam pattern", str(totals["byKind"].get("exam-pattern", 0)), "Assessment readiness metadata only; no released items or scoring language."],
    ]
    add_table(doc, ["Card Type", "Count", "Document Meaning"], rows, [2100, 1200, 6060], header_fill=FILL_BLUE, font_size=9)

    lane_counts = totals["byLibraryLane"]
    rows = [[lane.replace("-", " "), str(count)] for lane, count in sorted(lane_counts.items())]
    add_table(doc, ["Library Lane", "Cards"], rows, [4200, 5160], header_fill=FILL_BLUE, font_size=9)

    add_heading(doc, "Appendix B. Global Policy Sources", 1)
    global_rows = []
    for source in data["sourceRegistry"]:
        if source["state"] == "US":
            global_rows.append(
                [
                    source["id"],
                    source["title"],
                    source["allowedUse"],
                    source["verbatimLimit"],
                ]
            )
    add_table(
        doc,
        ["Source ID", "Title", "Allowed Use", "Verbatim Limit"],
        global_rows,
        [1850, 2650, 3250, 1610],
        header_fill=FILL_BLUE,
        font_size=7.8,
    )

    add_heading(doc, "Appendix C. Provenance Note", 1)
    add_para(
        doc,
        "Generated from `data/rag/usMath.ts`, related U.S. RAG retrieval exports, the project decision record dated 2026-05-23, and the 2026-06-01 Arkansas Top-11 expansion. The document intentionally summarizes standards as identifiers and metadata because the committed RAG layer is designed to avoid protected source-body retention.",
        size=10,
        color=BLACK,
        after=6,
    )


def remove_trailing_page_break(doc: Document):
    body = doc._body._element
    paragraphs = body.findall(qn("w:p"))
    if not paragraphs:
        return
    last = paragraphs[-1]
    for br in last.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            body.remove(last)
            return


def build_docx(data: dict[str, Any]):
    doc = Document()
    configure_document(doc)
    add_cover(doc, data)
    add_overview(doc, data)
    for index, state in enumerate(data["states"], 1):
        add_state_section(doc, state, index)
    add_appendix(doc, data)
    remove_trailing_page_break(doc)
    doc.save(OUT_PATH)


def main():
    data = compile_and_extract_us_rag()
    build_docx(data)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
