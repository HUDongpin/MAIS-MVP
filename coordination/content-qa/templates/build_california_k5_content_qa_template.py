#!/usr/bin/env python3
"""Build the MAIS California K-Grade 5 lesson content-QA DOCX template."""

from __future__ import annotations

import json
import os
import re
import subprocess
import tempfile
import zipfile
from argparse import ArgumentParser
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor, Twips
from lxml import etree


EXPECTED_GRADE_COUNTS = {"K": 6, "P1": 16, "P2": 4, "P3": 5, "P4": 5, "P5": 5}
EXPECTED_STANDARD_COUNTS = {"K": 22, "P1": 21, "P2": 26, "P3": 25, "P4": 28, "P5": 26}
TARGET_TOPIC_ID = "us-ca-math-p1-1-oa-add-subtract"
LESSON_BASE_URL = "https://mais.hk/student/lessons"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
ACCENT_BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK_BLUE = "1F4E79"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
MUTED = "5B6573"
RISK_RED = "9B1C1C"
CAUTION_GOLD = "7A5A00"
POSITIVE_BLUE = "1F3A5F"


@dataclass(frozen=True)
class LessonRecord:
    topic_id: str
    url: str
    source_lesson_id: str
    grade_id: str
    grade_label: str
    lesson_type: str
    knowledge_point_code: str
    title: str
    source_title: str
    runtime_title: str
    expected_display_title: str
    domain_id: str
    domain_title: str
    cluster_id: str
    cluster_title: str
    standard_ids: tuple[str, ...]
    estimated_minutes: int
    source_practice_count: int
    runtime_block_types: tuple[str, ...]
    production_ready: bool
    review_status: str
    integration_status: str
    approval_status: str
    difficulty: str
    language_variant: str
    source_safety_status: str
    package_id: str


@dataclass(frozen=True)
class StandardCatalogRow:
    grade_id: str
    grade_label: str
    domain_id: str
    domain_title: str
    cluster_id: str
    cluster_title: str
    standard_ids: tuple[str, ...]


@dataclass(frozen=True)
class BuildSummary:
    output_path: Path
    lesson_count: int
    standard_count: int
    content_control_count: int
    table_count: int


def _load_core_package(project_root: Path) -> dict[str, Any]:
    package_path = (
        project_root
        / "data"
        / "generated-content"
        / "us-ca-math-k-g5-textbooks-v1"
        / "lessons.json"
    )
    return json.loads(package_path.read_text(encoding="utf-8"))


def _load_runtime_snapshot(project_root: Path) -> dict[str, Any]:
    tsx = project_root / "node_modules" / ".bin" / "tsx"
    if not tsx.exists():
        raise RuntimeError(
            "The lesson inventory requires node_modules/.bin/tsx. Run npm ci in the worktree first."
        )

    script = r'''
import {
  californiaK5TextbookLessonSeeds,
  californiaElementaryMicroLessonSeeds,
  californiaK5TextbookLessonCoverageRecords,
  californiaElementaryMicroLessonCoverageRecords
} from "./data/usCaliforniaLessons";
import { californiaElementaryMicroLessonSpecs } from "./data/usCaliforniaMicroLessons";

const coverage = new Map(
  [...californiaK5TextbookLessonCoverageRecords, ...californiaElementaryMicroLessonCoverageRecords]
    .map((record) => [record.topicId, record])
);
const microSpecs = new Map(
  californiaElementaryMicroLessonSpecs.map((record) => [record.topicId, record])
);
const projectSeed = (seed) => ({
  topicId: seed.topicId,
  title: seed.title?.en ?? "",
  estimatedMinutes: seed.estimatedMinutes,
  practiceQuestionIds: seed.practiceQuestionIds ?? [],
  productionReady: Boolean(seed.productionReady),
  blockTypes: (seed.blocks ?? []).map((block) => block.type),
  coverage: coverage.get(seed.topicId) ?? null,
  microSpec: microSpecs.get(seed.topicId) ?? null
});
console.log(JSON.stringify({
  core: californiaK5TextbookLessonSeeds.map(projectSeed),
  micro: californiaElementaryMicroLessonSeeds.map(projectSeed)
}));
'''
    completed = subprocess.run(
        [str(tsx), "-e", script],
        cwd=project_root,
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(completed.stdout)


def _knowledge_point_code(title: str) -> str:
    return title.split(maxsplit=1)[0] if title.strip() else ""


def _expected_display_title(title: str) -> str:
    """Mirror the current California display-prefix cleanup for inventory evidence."""
    without_code = re.sub(r"^[K1-5]-[A-Z]+\.\d+\s+", "", title, flags=re.IGNORECASE)
    return re.sub(
        r"^(?:Kindergarten|Grade\s+[1-5])\s+", "", without_code, flags=re.IGNORECASE
    ).strip()


def load_lesson_inventory(project_root: Path) -> list[LessonRecord]:
    project_root = project_root.resolve()
    package = _load_core_package(project_root)
    runtime = _load_runtime_snapshot(project_root)

    runtime_core = {row["topicId"]: row for row in runtime["core"]}
    core_records: list[LessonRecord] = []
    for lesson in package["lessons"]:
        metadata = lesson["metadata"]
        topic_id = metadata["topicId"]
        seed = runtime_core.get(topic_id)
        if seed is None:
            raise RuntimeError(f"Core lesson is missing from the runtime inventory: {topic_id}")
        title = seed["title"] or lesson["studentLesson"]["en"]["title"]
        source_title = lesson["studentLesson"]["en"]["title"]
        core_records.append(
            LessonRecord(
                topic_id=topic_id,
                url=f"{LESSON_BASE_URL}/{topic_id}",
                source_lesson_id=lesson["id"],
                grade_id=metadata["grade"],
                grade_label=metadata["usGradeLabel"],
                lesson_type="Core",
                knowledge_point_code=_knowledge_point_code(title),
                title=title,
                source_title=source_title,
                runtime_title=title,
                expected_display_title=_expected_display_title(title),
                domain_id=metadata["domainId"],
                domain_title=metadata["domainTitle"],
                cluster_id=metadata["clusterId"],
                cluster_title=metadata["clusterTitle"],
                standard_ids=tuple(metadata["standardIds"]),
                estimated_minutes=int(metadata["estimatedMinutes"]),
                source_practice_count=len(seed["practiceQuestionIds"]),
                runtime_block_types=tuple(seed["blockTypes"]),
                production_ready=bool(seed["productionReady"]),
                review_status=lesson.get("reviewStatus", "unknown"),
                integration_status=lesson.get("integrationStatus", "unknown"),
                approval_status=lesson.get("approval", {}).get("status", "not-recorded"),
                difficulty=metadata.get("difficulty", "not-recorded"),
                language_variant=metadata.get("languageVariant", "not-recorded"),
                source_safety_status=metadata.get("sourceSafetyStatus", "not-recorded"),
                package_id=lesson.get("packageId", package.get("packageId", "unknown")),
            )
        )

    micro_records: list[LessonRecord] = []
    for seed in runtime["micro"]:
        spec = seed.get("microSpec")
        if not spec:
            raise RuntimeError(f"Micro lesson is missing its source specification: {seed['topicId']}")
        topic_id = spec["topicId"]
        micro_records.append(
            LessonRecord(
                topic_id=topic_id,
                url=f"{LESSON_BASE_URL}/{topic_id}",
                source_lesson_id=topic_id,
                grade_id=spec["grade"],
                grade_label=spec["usGradeLabel"],
                lesson_type="Micro",
                knowledge_point_code=spec["knowledgePointCode"],
                title=seed["title"] or spec["maisTitle"],
                source_title=spec["maisTitle"],
                runtime_title=seed["title"] or spec["maisTitle"],
                expected_display_title=_expected_display_title(seed["title"] or spec["maisTitle"]),
                domain_id=spec["domainId"],
                domain_title=spec["domainTitle"],
                cluster_id=spec["strandId"],
                cluster_title=spec["strandTitle"],
                standard_ids=tuple(spec["standardIds"]),
                estimated_minutes=int(spec["estimatedMinutes"]),
                source_practice_count=len(seed["practiceQuestionIds"]),
                runtime_block_types=tuple(seed["blockTypes"]),
                production_ready=bool(seed["productionReady"]),
                review_status="runtime-source-review-required",
                integration_status="runtime-live",
                approval_status="not-recorded-in-micro-source",
                difficulty=spec.get("difficulty", "not-recorded"),
                language_variant="en-primary",
                source_safety_status="MAIS-authored-original; verify source-policy evidence",
                package_id=spec["packageId"],
            )
        )

    inventory = core_records + micro_records
    grade_counts = {grade: sum(row.grade_id == grade for row in inventory) for grade in EXPECTED_GRADE_COUNTS}
    if grade_counts != EXPECTED_GRADE_COUNTS:
        raise RuntimeError(f"California K-5 route count changed: {grade_counts}")
    topic_ids = [row.topic_id for row in inventory]
    if len(topic_ids) != len(set(topic_ids)):
        raise RuntimeError("California K-5 route inventory contains duplicate topic IDs")
    return inventory


def load_practice_standard_mappings(project_root: Path) -> dict[str, str]:
    """Return and validate the current practice-pack canonical-to-MAIS map."""
    pack_path = (
        project_root.resolve()
        / "data"
        / "generated-content"
        / "us-ca-k5-knowledge-point-practice-v1"
        / "question-pack.json"
    )
    pack = json.loads(pack_path.read_text(encoding="utf-8"))
    mappings: dict[str, str] = {}
    for question in pack["questions"]:
        canonical_ids = question.get("canonicalStandardIds", [])
        mais_ids = question.get("maisStandardIds", [])
        if len(canonical_ids) != len(mais_ids):
            raise RuntimeError(
                f"Practice item {question.get('id', 'unknown')} has misaligned canonical and MAIS standard arrays"
            )
        for canonical_id, mais_id in zip(canonical_ids, mais_ids, strict=True):
            expected = f"CA.CCSS.Math.{canonical_id}"
            if mais_id != expected:
                raise RuntimeError(
                    f"Practice standard mapping changed for {canonical_id}: expected {expected}, found {mais_id}"
                )
            prior = mappings.setdefault(canonical_id, mais_id)
            if prior != mais_id:
                raise RuntimeError(f"Conflicting practice mapping for {canonical_id}")
    if len(mappings) != 148:
        raise RuntimeError(f"Expected 148 practice-pack standard mappings; found {len(mappings)}")
    return mappings


def load_standard_catalog(project_root: Path) -> list[StandardCatalogRow]:
    package = _load_core_package(project_root.resolve())
    catalog = [
        StandardCatalogRow(
            grade_id=lesson["metadata"]["grade"],
            grade_label=lesson["metadata"]["usGradeLabel"],
            domain_id=lesson["metadata"]["domainId"],
            domain_title=lesson["metadata"]["domainTitle"],
            cluster_id=lesson["metadata"]["clusterId"],
            cluster_title=lesson["metadata"]["clusterTitle"],
            standard_ids=tuple(lesson["metadata"]["standardIds"]),
        )
        for lesson in package["lessons"]
    ]
    standards = [standard_id for row in catalog for standard_id in row.standard_ids]
    if len(standards) != 148 or len(set(standards)) != 148:
        raise RuntimeError(
            f"Expected 148 unique K-5 standards in the core package; found {len(standards)} total and {len(set(standards))} unique"
        )
    grade_counts = {
        grade: sum(row.grade_id == grade for row in catalog for _ in row.standard_ids)
        for grade in EXPECTED_STANDARD_COUNTS
    }
    if grade_counts != EXPECTED_STANDARD_COUNTS:
        raise RuntimeError(f"California K-5 standard count changed: {grade_counts}")
    practice_mappings = load_practice_standard_mappings(project_root)
    if set(standards) != set(practice_mappings):
        missing = sorted(set(standards) - set(practice_mappings))
        unexpected = sorted(set(practice_mappings) - set(standards))
        raise RuntimeError(
            f"Lesson and practice standard indexes diverged; missing={missing}, unexpected={unexpected}"
        )
    return catalog


def _set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang TC")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def _set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA} DXA: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _format_cell_text(cell, *, size: float = 8.5, bold: bool = False, color: str = "000000") -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            _set_run_font(run, size=size, color=color, bold=bold)


def _add_table(document: DocumentObject, headers: list[str], widths_dxa: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
        _set_cell_shading(cell, HEADER_FILL)
        _format_cell_text(cell, size=8.5, bold=True, color=INK_BLUE)
    _set_repeat_table_header(table.rows[0])
    _set_table_geometry(table, widths_dxa)
    return table


def _add_text_row(table, values: list[str], *, size: float = 8.25, fill: str | None = None):
    row = table.add_row()
    for cell, value in zip(row.cells, values, strict=True):
        cell.text = value
        if fill:
            _set_cell_shading(cell, fill)
        _format_cell_text(cell, size=size)
    _set_row_cant_split(row)
    return row


def _add_sdt(paragraph, tag: str, initial_text: str, *, size: float = 8.5, color: str = MUTED, bold: bool = False) -> None:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", tag):
        raise ValueError(f"Invalid content-control tag: {tag}")
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    tag_node = OxmlElement("w:tag")
    tag_node.set(qn("w:val"), tag)
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), tag)
    text_prop = OxmlElement("w:text")
    sdt_pr.extend([tag_node, alias, text_prop])
    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "PingFang TC")
    run_pr.append(fonts)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    run_pr.append(size_node)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    run_pr.append(color_node)
    if bold:
        run_pr.append(OxmlElement("w:b"))
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = initial_text
    run.append(text_node)
    sdt_content.append(run)
    sdt.append(sdt_pr)
    sdt.append(sdt_content)
    paragraph._p.append(sdt)


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _set_control_cell(cell, tag: str, initial_text: str, *, size: float = 8.25, color: str = MUTED) -> None:
    paragraph = cell.paragraphs[0]
    _clear_paragraph(paragraph)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    _add_sdt(paragraph, tag, initial_text, size=size, color=color)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.extend([color, underline])
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_callout(document: DocumentObject, label: str, text: str, *, fill: str = CALLOUT_FILL) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)
    label_run = paragraph.add_run(f"{label}: ")
    _set_run_font(label_run, size=10, color=INK_BLUE, bold=True)
    text_run = paragraph.add_run(text)
    _set_run_font(text_run, size=10, color="20252B")


def _add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, result, end):
        run = OxmlElement("w:r")
        run.append(node)
        paragraph._p.append(run)


def _configure_styles(document: DocumentObject) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    style_tokens = {
        "Title": (24, INK_BLUE, 0, 4),
        "Subtitle": (13, MUTED, 0, 12),
        "Heading 1": (16, ACCENT_BLUE, 18, 10),
        "Heading 2": (13, ACCENT_BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in style_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def _configure_sections(document: DocumentObject) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        header_paragraph = header.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_paragraph.paragraph_format.space_after = Pt(0)
        run = header_paragraph.add_run("MAIS | A18 California K-5 Content QA")
        _set_run_font(run, size=8.5, color=MUTED, bold=True)

        footer = section.footer
        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_paragraph.paragraph_format.space_before = Pt(0)
        label = footer_paragraph.add_run("Internal QA template | Page ")
        _set_run_font(label, size=8, color=MUTED)
        _add_field(footer_paragraph, " PAGE ")


def _add_title_block(document: DocumentObject) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("A18 CONTENT QUALITY ASSURANCE")
    _set_run_font(run, size=9, color=ACCENT_BLUE, bold=True)

    title = document.add_paragraph(style="Title")
    title.add_run("California K-Grade 5 Lesson Content QA Template")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Reusable workbook for standards alignment, mathematical accuracy, rendered-content parity, language quality, evidence, issue tracking, and retest control"
    )

    metadata = _add_table(document, ["Document control", "Value"], [1800, 7560])
    fields = [
        ("Document version", "DOC_VERSION", "1.0"),
        ("Campaign ID", "CAMPAIGN_ID", f"CA-K5-{date.today().strftime('%Y%m%d')}"),
        ("Review lead", "REVIEW_LEAD", "[enter A18 reviewer]"),
        ("Independent reviewer", "INDEPENDENT_REVIEWER", "[enter curriculum/SME reviewer]"),
        ("Campaign state", "CAMPAIGN_STATE", "NOT TESTED"),
    ]
    for label, tag, initial in fields:
        row = metadata.add_row()
        row.cells[0].text = label
        _format_cell_text(row.cells[0], size=9, bold=True, color=INK_BLUE)
        _set_control_cell(row.cells[1], tag, initial, size=9)
        _set_row_cant_split(row)
    _set_table_geometry(metadata, [1800, 7560])

    _add_callout(
        document,
        "Boundary",
        "This internal MAIS instrument is mapped to California Department of Education criteria. It does not represent CDE adoption, official endorsement, or a statewide completeness determination.",
    )


def _add_numbered_steps(document: DocumentObject, steps: list[str]) -> None:
    for step in steps:
        paragraph = document.add_paragraph(step, style="List Number")
        paragraph.paragraph_format.keep_together = True


def _add_bullets(document: DocumentObject, bullets: list[str]) -> None:
    for bullet in bullets:
        document.add_paragraph(bullet, style="List Bullet")


def _add_page_heading(document: DocumentObject, text: str) -> None:
    heading = document.add_heading(text, level=1)
    heading.paragraph_format.page_break_before = True


def _add_front_matter(document: DocumentObject) -> None:
    document.add_heading("1. Purpose and Operating Instructions", level=1)
    document.add_paragraph(
        "Use this workbook to run a consistent content-QA campaign across the 41 currently exposed California Kindergarten through Grade 5 lesson pages. The master register tracks coverage; copy the detailed form for each lesson reviewed."
    )
    _add_numbered_steps(
        document,
        [
            "Create the review record ID in the form CA-K5-<lesson-slug>-<YYYYMMDD> and record the exact production URL, build, viewport, language, and Asia/Hong_Kong timestamp.",
            "Capture account pre-state and post-state. Opening the lesson may record a lesson-start event even when the reviewer does not answer a question. Never submit an answer or invoke lesson completion on Student Jon. Use a designated resettable QA account for answer, completion, reset, or persistence tests, then restore its state.",
            "Independently map each claimed standard to exact lesson-block evidence and derive every mathematical answer before comparing it with the page.",
            "Review authored source versus rendered output, then test student-visible media, practice, feedback, accessibility, and each supported locale in scope.",
            "Log evidence as <record-id>-E### and issues as <record-id>-I###. Complete retest and all required downstream gates before changing the overall disposition.",
        ],
    )

    document.add_heading("1.1 Status and Evidence Rules", level=2)
    status_table = _add_table(document, ["Status", "Use", "Completion rule"], [1200, 3600, 4560])
    status_rows = [
        ("PASS", "Criterion is satisfied.", "Evidence is mandatory."),
        ("FAIL", "Criterion is not satisfied.", "Link an issue ID and corrective action."),
        ("BLOCKED", "Review cannot proceed because prerequisite evidence or access is unavailable.", "Name the blocker owner and next action."),
        ("N/A", "Criterion is outside the approved lesson or locale scope.", "A rationale is mandatory."),
        ("NOT TESTED", "Criterion has not yet been evaluated.", "Mandatory NOT TESTED checks prevent final approval."),
    ]
    for row in status_rows:
        _add_text_row(status_table, list(row), size=8.75)
    _set_table_geometry(status_table, [1200, 3600, 4560])

    document.add_heading("1.2 Severity and Release Gates", level=2)
    severity = _add_table(document, ["Level", "Meaning", "Required action"], [1200, 3750, 4410])
    for row in (
        ("P0", "Stop-the-line mathematical, answer, safety, source-rights, or broad production harm.", "Stop release and route immediately to the accountable owner."),
        ("P1", "Serious lesson, grade, standards, or student-visible feature failure.", "Hold the affected lesson or grade until verified."),
        ("P2", "Limited clarity, accessibility, terminology, or polish issue with recoverable impact.", "Conditional acceptance only with owner and due date."),
        ("Observation", "Non-defect note or improvement opportunity.", "Track separately; it does not replace a criterion status."),
    ):
        _add_text_row(severity, list(row), size=8.75)
    _set_table_geometry(severity, [1200, 3750, 4410])
    _add_callout(
        document,
        "Gate rule",
        "Any P0, unresolved P1, failed mandatory content/alignment criterion, or mandatory BLOCKED/NOT TESTED check prevents CONTENT PASS. Content approval does not equal production release approval.",
        fill="FFF4F2",
    )

    document.add_heading("1.3 Official Standards Basis", level=2)
    document.add_paragraph(
        "Use official CDE sources for identifiers, structure, California additions, framework connections, and evaluation categories. Record source roles and access dates; do not copy official standard prose, framework examples, proprietary item text, screenshots, or exercise layouts into MAIS content."
    )
    sources = [
        ("CDE 2025 evaluation-map guidance", "https://www.cde.ca.gov/ci/ma/im/evalmapinstructmath2025.asp"),
        ("CDE standards-map guidance", "https://www.cde.ca.gov/ci/ma/im/cccsssminstructmath2025.asp"),
        ("California Mathematics Framework Chapter 13", "https://www.cde.ca.gov/ci/ma/cf/documents/mathframeworkch13.pdf"),
        ("California Mathematics Framework Chapter 6: K-5", "https://www.cde.ca.gov/ci/ma/cf/documents/mathframeworkch6.pdf"),
        ("CDE California Content Standards search", "https://www2.cde.ca.gov/cacs/math"),
    ]
    for label, url in sources:
        paragraph = document.add_paragraph(style="List Bullet")
        _add_hyperlink(paragraph, label, url)


def _add_tracker(document: DocumentObject, inventory: list[LessonRecord]) -> None:
    _add_page_heading(document, "2. 41-lesson master coverage register")
    document.add_paragraph(
        "This register is generated from the current 29 core lesson records and 12 Grade 1 micro-lesson records. Initial status is NOT TESTED; metadata presence is not QA approval."
    )
    table = _add_table(
        document,
        ["#", "Grade", "Type", "Lesson / production slug", "Claimed standards", "Status", "Reviewer / date"],
        [550, 800, 650, 3330, 1650, 950, 1430],
    )
    for index, lesson in enumerate(inventory, start=1):
        row = table.add_row()
        row.cells[0].text = str(index)
        row.cells[1].text = "K" if lesson.grade_id == "K" else lesson.grade_id.removeprefix("P")
        row.cells[2].text = lesson.lesson_type
        row.cells[3].text = (
            f"{lesson.knowledge_point_code} {lesson.expected_display_title}\n{lesson.topic_id}"
        )
        row.cells[4].text = ", ".join(lesson.standard_ids)
        for cell in row.cells[:5]:
            _format_cell_text(cell, size=7.6)
        _set_control_cell(row.cells[5], f"TRACK_STATUS_{index:02d}", "NOT TESTED", size=7.6, color=CAUTION_GOLD)
        _set_control_cell(
            row.cells[6], f"TRACK_REVIEWER_DATE_{index:02d}", "[reviewer / HKT date]", size=7.4
        )
        if index % 2 == 0:
            for cell in row.cells:
                _set_cell_shading(cell, "FAFBFC")
        _set_row_cant_split(row)
    _set_table_geometry(table, [550, 800, 650, 3330, 1650, 950, 1430])


CALIBRATION_SURFACES = [
    ("Lifecycle", "Source approval, candidate/live state, and downstream gates agree", "Source remains candidate-only and approved-for-review while runtime is productionReady; reconcile before disposition"),
    ("Title layers", "Source, runtime, and expected student display title are recorded", "Re-verify prefix cleanup and the actual visible title at the tested build"),
    ("Authored/rendered order", "Objectives, launch, concept, worked, guided, independent, repair, exit, teacher support", "Inventory actual order separately; authored extension and exit evidence are not assumed visible"),
    ("Concept explanation", "Source concept and learning goal", "Concept section and 8 + 3 = 11 representation; re-check wording and AI Tutor entry"),
    ("Concept audio", "Narration matches current visible concept text and pronunciation", "36.192-second worktree asset; capture deployed asset, duration, transcript, and build independently"),
    ("Worked example", "Prompt, answer, reasoning, representation", "7 + 4 = 11 with answer reveal/check and illustration/caption; re-derive and re-observe"),
    ("Checklist/self-check", "Authored guided-practice intent remains visible and measurable", "Student UI may show three generic quick checks instead of authored guided prompts; compare exact text"),
    ("Visualization", "Topic-fit model, exact mathematics, instructions, and visible safeguards", "Number-line lab seed plus safeguard/read-me claims; verify actual note visibility and accessibility"),
    ("Illustrations", "Asset math, provenance, caption, alt text, and locale agree", "Runtime assets may differ from text-only source plan; compare exact PNG/SVG and localization"),
    ("Practice", "Every rendered item, primary standard, grader, feedback, and reset", "Source-linked practice IDs: 8; rendered pager observed: 5; target prompts include 15 - 5 = 10, 19 - 3 = 16, 5 + 1 = 6, 17 - 4 = 13, 2 + 3 = 5"),
    ("Localization", "Genuine English, Traditional Chinese, Simplified Chinese", "Review each locale independently; duplicated English is not localization"),
    ("Teacher support", "Teacher-only evidence is accurate and role-gated", "Confirm the teacher guide is absent from the student view without treating correct role gating as a defect"),
    ("Navigation and help", "Lesson navigation, progress, AI support", "Re-test unit directory, return navigation, progress cues, and AI Tutor/Nova entry"),
    ("Start-event side effect", "Calibration has pre/post-state evidence", "Opening a logged-in lesson may post action=start; treat read-only as no answer/completion, not necessarily state-free"),
    ("Completion gate", "Completion reflects intended learning evidence", "Do not exercise on Student Jon; use a resettable QA account to verify whether completion depends on checklist/practice evidence"),
]


def _add_calibration_record(document: DocumentObject, target: LessonRecord) -> None:
    _add_page_heading(document, "3. Calibration Record - Grade 1 Add/Subtract")
    _add_callout(
        document,
        "Calibration only",
        "The fields below preserve current source and read-only rendered-page observations as reviewer prompts. Every criterion remains NOT TESTED; this section is not a pass verdict or defect report.",
    )

    metadata = _add_table(document, ["Field", "Current calibration metadata"], [2100, 7260])
    rows = [
        ("Production URL", target.url),
        ("Slug", target.topic_id),
        ("Source lesson ID", target.source_lesson_id),
        ("Grade / type", f"{target.grade_label} / {target.lesson_type}"),
        ("Knowledge point", target.knowledge_point_code),
        ("Domain / cluster", f"{target.domain_id} {target.domain_title} / {target.cluster_id} {target.cluster_title}"),
        ("Claimed standards", ", ".join(target.standard_ids)),
        ("MAIS standard IDs", ", ".join(f"CA.CCSS.Math.{standard_id}" for standard_id in target.standard_ids)),
        ("Source title", target.source_title),
        ("Runtime title", target.runtime_title),
        ("Expected display title", target.expected_display_title),
        ("Difficulty / duration / language", f"{target.difficulty} / {target.estimated_minutes} minutes / {target.language_variant}"),
        ("Lifecycle", f"source review={target.review_status}; source integration={target.integration_status}; source approval={target.approval_status}; runtime productionReady={str(target.production_ready).lower()}"),
        ("Source safety", target.source_safety_status),
        ("Runtime blocks", ", ".join(target.runtime_block_types)),
        ("Practice-count reconciliation", "Source-linked practice IDs: 8; rendered pager observed during calibration: 5"),
        ("Account boundary", "Student Jon - account name only. Opening may record a start event. Do not submit answers or complete; capture pre/post-state and use a resettable QA account for stateful tests."),
    ]
    for label, value in rows:
        row = _add_text_row(metadata, [label, value], size=8.75)
        _format_cell_text(row.cells[0], size=8.75, bold=True, color=INK_BLUE)
    _set_table_geometry(metadata, [2100, 7260])

    document.add_heading("3.1 Calibration Surface Inventory", level=2)
    surface_table = _add_table(
        document,
        ["Surface", "Source expectation", "Current observation to re-verify", "Status", "Evidence / issue"],
        [1400, 2450, 2900, 1050, 1560],
    )
    for index, (surface, expectation, observation) in enumerate(CALIBRATION_SURFACES, start=1):
        row = surface_table.add_row()
        row.cells[0].text = surface
        row.cells[1].text = expectation
        row.cells[2].text = observation
        for cell in row.cells[:3]:
            _format_cell_text(cell, size=8.1)
        _set_control_cell(row.cells[3], f"CAL_STATUS_{index:02d}", "NOT TESTED", size=8, color=CAUTION_GOLD)
        _set_control_cell(row.cells[4], f"CAL_EVIDENCE_{index:02d}", "[evidence / issue]", size=7.8)
        _set_row_cant_split(row)
    _set_table_geometry(surface_table, [1400, 2450, 2900, 1050, 1560])


FORM_METADATA_FIELDS = [
    ("Review record ID", "FORM_REVIEW_RECORD_ID", "CA-K5-<lesson-slug>-<YYYYMMDD>"),
    ("Lesson URL", "FORM_LESSON_URL", "https://mais.hk/student/lessons/<lesson-slug>"),
    ("Lesson slug", "FORM_LESSON_SLUG", "[enter exact slug]"),
    ("Source lesson ID", "FORM_SOURCE_LESSON_ID", "[enter source record ID]"),
    ("Grade / core-micro", "FORM_GRADE_TYPE", "[enter grade and type]"),
    ("Knowledge point", "FORM_KNOWLEDGE_POINT", "[enter MAIS knowledge-point code]"),
    ("Domain / cluster", "FORM_DOMAIN_CLUSTER", "[enter canonical IDs]"),
    ("Canonical standard IDs", "FORM_STANDARD_IDS", "[enter claimed IDs]"),
    ("MAIS standard IDs", "FORM_MAIS_STANDARD_IDS", "[enter CA.CCSS.Math.* IDs]"),
    ("Package / version", "FORM_PACKAGE_VERSION", "[enter package and version]"),
    ("Difficulty / minutes / language", "FORM_DIFFICULTY_MINUTES_LANGUAGE", "[enter current source metadata]"),
    ("Source approval / safety", "FORM_SOURCE_APPROVAL_SAFETY", "[record approval, rights, originality, attribution, last checked]"),
    ("Source/runtime lifecycle", "FORM_LIFECYCLE", "[record candidate, review, integration, and runtime states]"),
    ("Host / deployment / build", "FORM_BUILD", "[enter host and build evidence]"),
    ("Timestamp", "FORM_TIMESTAMP", "[YYYY-MM-DD HH:MM Asia/Hong_Kong]"),
    ("Browser / OS / viewport", "FORM_BROWSER", "[enter browser, OS, and exact viewport]"),
    ("Account", "FORM_ACCOUNT", "[account name only - never password]"),
    ("Pre-state / post-state", "FORM_STATE", "[record and restore progress state]"),
    ("Primary reviewer", "FORM_PRIMARY_REVIEWER", "[enter A18 reviewer]"),
    ("Independent reviewer", "FORM_INDEPENDENT_REVIEWER", "[enter curriculum/SME reviewer]"),
]


CRITERIA_BY_CATEGORY = {
    "Category 1 - Mathematics Content and Alignment (all mandatory)": [
        "Claimed canonical and MAIS standard IDs are valid for the stated grade/domain.",
        "Each claimed standard has exact lesson-block evidence; metadata alone is not accepted.",
        "Primary versus supporting standards and introduced/practiced/assessed roles are recorded.",
        "Learning goals, examples, practice, and exit evidence align to the same intended mathematics.",
        "All prompts, worked steps, representations, answers, units, labels, and feedback are mathematically correct.",
        "Every answer is independently derived and ambiguous or multi-answer prompts define accepted responses.",
        "Representations are internally consistent and support grade-appropriate reasoning rather than decoration.",
        "Relevant Mathematical Practice behaviors are authentic and evidenced in student actions.",
        "Prerequisites, progression, and grade-level cognitive demand are appropriate.",
        "The lesson coheres around a meaningful K-5 big idea rather than isolated standard labels.",
    ],
    "Category 2 - Program Organization": [
        "Authored objectives, concept, worked example, guided practice, independent practice, remediation, exit ticket, and teacher support are inventoried.",
        "Rendered blocks match the authored inventory or every omission/addition/order change is explained.",
        "Lesson sequence moves from explanation and model to guided reasoning, independent application, and repair.",
        "Estimated duration is plausible for the rendered content and intended learner age.",
        "Navigation, headings, labels, and progress cues are coherent and use student-facing language.",
    ],
    "Category 3 - Assessment": [
        "Self-check or guided-practice tasks measure the stated goal rather than generic completion behavior.",
        "Every student-visible practice item is solvable, grade-appropriate, and linked to an independently verified answer.",
        "Accepted-answer behavior, correct feedback, wrong feedback, retry, reset, and answer reveal are accurate.",
        "Practice difficulty and item variety provide useful evidence of learning without repeating one surface form.",
        "Exit evidence and completion behavior do not award mastery without the intended learning checks.",
    ],
    "Category 4 - Access and Equity": [
        "Vocabulary, sentence length, contexts, and visual load are appropriate for the grade and avoid unnecessary barriers or bias.",
        "English, Traditional Chinese, and Simplified Chinese are reviewed separately; duplicated English is not accepted as translation.",
        "Mathematical terminology and symbols preserve meaning across each locale in release scope.",
        "Illustrations, captions, and alt text describe the same mathematics and contain no misleading count, label, color, or unit.",
        "Concept audio matches the currently rendered text and has accurate pronunciation, pacing, and control labels.",
        "Keyboard operation, focus order, headings, control names, contrast, zoom, and responsive layout support access.",
    ],
    "Category 5 - Instructional Planning and Support": [
        "Misconceptions are mathematically plausible and each repair move addresses the actual error mechanism.",
        "Teacher notes identify goals, evidence, likely errors, and actionable support without exposing internal generator jargon.",
        "Extension, remediation, and AI Tutor entry preserve the lesson goal and use appropriate safeguards.",
        "Source originality, rights, attribution, source-distance, and last-checked evidence are documented.",
        "Candidate, review, integration, productionReady, and live-page states are reconciled before any promotion recommendation.",
    ],
}


def _add_form_metadata(document: DocumentObject) -> None:
    document.add_heading("4.1 Review Run Metadata", level=2)
    table = _add_table(document, ["Field", "Reviewer entry"], [2300, 7060])
    for label, tag, initial in FORM_METADATA_FIELDS:
        row = table.add_row()
        row.cells[0].text = label
        _format_cell_text(row.cells[0], size=8.7, bold=True, color=INK_BLUE)
        _set_control_cell(row.cells[1], tag, initial, size=8.6)
        _set_row_cant_split(row)
    _set_table_geometry(table, [2300, 7060])


def _add_standard_evidence_form(document: DocumentObject) -> None:
    document.add_heading("4.2 Standard-by-Standard Evidence", level=2)
    document.add_paragraph(
        "Add rows as needed. Each claimed standard must identify whether it is primary or supporting, where it is introduced/practiced/assessed, and the exact source and rendered evidence."
    )
    table = _add_table(
        document,
        ["Canonical / MAIS ID", "Role", "I/P/A", "Exact block evidence", "Alignment status", "Issue / action"],
        [1750, 900, 750, 2900, 1200, 1860],
    )
    for index in range(1, 9):
        row = table.add_row()
        _set_control_cell(row.cells[0], f"STD_{index:02d}_IDS", "[canonical / MAIS IDs]", size=7.8)
        _set_control_cell(row.cells[1], f"STD_{index:02d}_ROLE", "[primary/supporting]", size=7.6)
        _set_control_cell(row.cells[2], f"STD_{index:02d}_IPA", "[I/P/A]", size=7.8)
        _set_control_cell(row.cells[3], f"STD_{index:02d}_EVIDENCE", "[source + rendered evidence]", size=7.8)
        _set_control_cell(row.cells[4], f"STD_{index:02d}_STATUS", "NOT TESTED", size=7.8, color=CAUTION_GOLD)
        _set_control_cell(row.cells[5], f"STD_{index:02d}_ISSUE", "[issue / action]", size=7.8)
        _set_row_cant_split(row)
    _set_table_geometry(table, [1750, 900, 750, 2900, 1200, 1860])


def _add_criteria_forms(document: DocumentObject) -> None:
    document.add_heading("4.3 CDE-Mapped Page-Level Review Criteria", level=2)
    criterion_index = 0
    for category, criteria in CRITERIA_BY_CATEGORY.items():
        document.add_heading(category, level=3)
        table = _add_table(
            document,
            ["ID", "Mandatory", "Criterion", "Status", "Evidence / rationale", "Issue ID"],
            [550, 1400, 3200, 1100, 2150, 960],
        )
        for criterion in criteria:
            criterion_index += 1
            row = table.add_row()
            row.cells[0].text = f"C{criterion_index:02d}"
            row.cells[1].text = "YES" if category.startswith("Category 1") else "SCOPE"
            row.cells[2].text = criterion
            for cell in row.cells[:3]:
                _format_cell_text(cell, size=7.75)
            _set_control_cell(row.cells[3], f"CRITERION_{criterion_index:02d}_STATUS", "NOT TESTED", size=7.7, color=CAUTION_GOLD)
            _set_control_cell(row.cells[4], f"CRITERION_{criterion_index:02d}_EVIDENCE", "[evidence or N/A rationale]", size=7.5)
            _set_control_cell(row.cells[5], f"CRITERION_{criterion_index:02d}_ISSUE", "[issue ID]", size=7.5)
            _set_row_cant_split(row)
        _set_table_geometry(table, [550, 1400, 3200, 1100, 2150, 960])


def _add_locale_form(document: DocumentObject) -> None:
    document.add_heading("4.4 Locale Review", level=2)
    table = _add_table(document, ["Locale", "Qualified reviewer", "Status", "Evidence and terminology notes"], [1500, 1900, 1100, 4860])
    locale_rows = [
        ("English", "EN"),
        ("Traditional Chinese", "ZH_HANT"),
        ("Simplified Chinese", "ZH_HANS"),
    ]
    for locale, tag in locale_rows:
        row = table.add_row()
        row.cells[0].text = locale
        _format_cell_text(row.cells[0], size=8.5, bold=True)
        _set_control_cell(row.cells[1], f"LOCALE_{tag}_REVIEWER", "[qualified reviewer]", size=8.2)
        _set_control_cell(row.cells[2], f"LOCALE_{tag}_STATUS", "NOT TESTED", size=8.2, color=CAUTION_GOLD)
        _set_control_cell(row.cells[3], f"LOCALE_{tag}_EVIDENCE", "[translation, terminology, display evidence]", size=8.1)
        _set_row_cant_split(row)
    _set_table_geometry(table, [1500, 1900, 1100, 4860])


def _add_practice_worksheet(document: DocumentObject) -> None:
    document.add_heading("4.5 Student-Visible Practice Item Worksheet", level=2)
    document.add_paragraph(
        "Record every rendered practice item. Add rows when more than ten items are visible. Derive the answer before comparing with the grader; test accepted answers, correct/wrong feedback, retry, and reset."
    )
    table = _add_table(
        document,
        ["Item", "Rendered prompt / location", "Primary standard(s) for this item", "Independent answer", "Status", "Evidence", "Issue ID"],
        [650, 2000, 1200, 1250, 1000, 2100, 1160],
    )
    for index in range(1, 11):
        row = table.add_row()
        row.cells[0].text = str(index)
        _format_cell_text(row.cells[0], size=8.2, bold=True)
        _set_control_cell(row.cells[1], f"PRACTICE_{index:02d}_PROMPT", "[prompt / pager position]", size=7.8)
        _set_control_cell(row.cells[2], f"PRACTICE_{index:02d}_PRIMARY_STANDARDS", "[item-specific IDs]", size=7.6)
        _set_control_cell(row.cells[3], f"PRACTICE_{index:02d}_ANSWER", "[derived answer]", size=7.8)
        _set_control_cell(row.cells[4], f"PRACTICE_{index:02d}_STATUS", "NOT TESTED", size=7.7, color=CAUTION_GOLD)
        _set_control_cell(row.cells[5], f"PRACTICE_{index:02d}_EVIDENCE", "[feedback + evidence ID]", size=7.7)
        _set_control_cell(row.cells[6], f"PRACTICE_{index:02d}_ISSUE", "[issue ID]", size=7.7)
        _set_row_cant_split(row)
    _set_table_geometry(table, [650, 2000, 1200, 1250, 1000, 2100, 1160])


def _add_detailed_form(document: DocumentObject) -> None:
    _add_page_heading(document, "4. Reusable Detailed Lesson Review Form")
    _add_callout(
        document,
        "Completion rule",
        "Copy this form for each lesson. A completed form has no blank statuses; every PASS has evidence, every N/A has rationale, and every FAIL links to an issue and retest path.",
    )
    _add_form_metadata(document)
    _add_standard_evidence_form(document)
    _add_criteria_forms(document)
    _add_locale_form(document)
    _add_practice_worksheet(document)


def _add_issue_retest_and_signoff(document: DocumentObject) -> None:
    _add_page_heading(document, "5. Consolidated Issue Register")
    document.add_paragraph(
        "Lifecycle: New -> Triaged -> In fix -> Ready for retest -> Verified/Closed. Reopened and Accepted risk require explicit evidence and authorization."
    )
    issue_table = _add_table(
        document,
        ["Issue ID", "Severity", "State", "Expected / actual / reproduction", "Owner / due", "Evidence / retest"],
        [1150, 850, 1050, 3150, 1400, 1760],
    )
    for index in range(1, 9):
        row = issue_table.add_row()
        _set_control_cell(row.cells[0], f"ISSUE_{index:02d}_ID", "<record-id>-I###", size=7.7)
        _set_control_cell(row.cells[1], f"ISSUE_{index:02d}_SEVERITY", "[P0/P1/P2/Obs]", size=7.5)
        _set_control_cell(row.cells[2], f"ISSUE_{index:02d}_STATE", "New", size=7.7)
        _set_control_cell(row.cells[3], f"ISSUE_{index:02d}_DETAIL", "[user harm, exact expected/actual, reproducible steps, blast radius]", size=7.5)
        _set_control_cell(row.cells[4], f"ISSUE_{index:02d}_OWNER", "[owner / due date]", size=7.5)
        _set_control_cell(row.cells[5], f"ISSUE_{index:02d}_EVIDENCE", "[evidence IDs / retest link]", size=7.5)
        _set_row_cant_split(row)
    _set_table_geometry(issue_table, [1150, 850, 1050, 3150, 1400, 1760])

    document.add_heading("5.1 Retest History", level=2)
    retest = _add_table(
        document,
        ["Issue ID", "Date / build", "Verifier", "Original reproduction", "Adjacent regression", "Result / evidence"],
        [1100, 1300, 1050, 2200, 2200, 1510],
    )
    for index in range(1, 7):
        row = retest.add_row()
        _set_control_cell(row.cells[0], f"RETEST_{index:02d}_ISSUE", "[issue ID]", size=7.7)
        _set_control_cell(row.cells[1], f"RETEST_{index:02d}_BUILD", "[HKT date / build]", size=7.5)
        _set_control_cell(row.cells[2], f"RETEST_{index:02d}_VERIFIER", "[verifier]", size=7.5)
        _set_control_cell(row.cells[3], f"RETEST_{index:02d}_ORIGINAL", "[original repro result]", size=7.5)
        _set_control_cell(row.cells[4], f"RETEST_{index:02d}_ADJACENT", "[adjacent regression result]", size=7.5)
        _set_control_cell(row.cells[5], f"RETEST_{index:02d}_RESULT", "[Verified/Reopened + evidence]", size=7.5)
        _set_row_cant_split(row)
    _set_table_geometry(retest, [1100, 1300, 1050, 2200, 2200, 1510])

    document.add_heading("5.2 Overall Disposition and Gate Sign-Off", level=2)
    document.add_paragraph(
        "Controlled overall disposition values: BLOCKED / NEEDS REPAIR / CONDITIONAL PASS / CONTENT PASS."
    )
    disposition = _add_table(document, ["Decision field", "Controlled entry"], [2500, 6860])
    for label, tag, initial in (
        ("Overall content disposition", "FORM_OVERALL_DISPOSITION", "BLOCKED"),
        ("Open P0 / P1 / P2 counts", "FORM_OPEN_ISSUE_COUNTS", "[enter counts]"),
        ("Conditional owner / due date", "FORM_CONDITIONAL_OWNER", "[required for conditional pass]"),
        ("Final rationale", "FORM_FINAL_RATIONALE", "[summarize evidence, limitations, and unresolved risk]"),
    ):
        row = disposition.add_row()
        row.cells[0].text = label
        _format_cell_text(row.cells[0], size=8.7, bold=True, color=INK_BLUE)
        _set_control_cell(row.cells[1], tag, initial, size=8.5)
        _set_row_cant_split(row)
    _set_table_geometry(disposition, [2500, 6860])

    signoff = _add_table(document, ["Gate owner", "Name / date", "Decision and evidence"], [2200, 2200, 4960])
    signoff_rows = [
        ("A18 content acceptance", "SIGNOFF_A18_NAME", "SIGNOFF_A18_DECISION"),
        ("A09 language/accessibility", "SIGNOFF_A09_NAME", "SIGNOFF_A09_DECISION"),
        ("A23 promotion decision", "SIGNOFF_A23_NAME", "SIGNOFF_A23_DECISION"),
        ("A05 lesson integration", "SIGNOFF_A05_NAME", "SIGNOFF_A05_DECISION"),
        ("A11 regression", "SIGNOFF_A11_NAME", "SIGNOFF_A11_DECISION"),
        ("A22 release readiness", "SIGNOFF_A22_NAME", "SIGNOFF_A22_DECISION"),
    ]
    for owner, name_tag, decision_tag in signoff_rows:
        row = signoff.add_row()
        row.cells[0].text = owner
        _format_cell_text(row.cells[0], size=8.3, bold=True)
        _set_control_cell(row.cells[1], name_tag, "[name / HKT date]", size=8)
        _set_control_cell(row.cells[2], decision_tag, "[decision / evidence ID]", size=8)
        _set_row_cant_split(row)
    _set_table_geometry(signoff, [2200, 2200, 4960])


def _add_owner_routing(document: DocumentObject) -> None:
    document.add_heading("5.3 Issue Owner Routing", level=2)
    routing = _add_table(document, ["Finding type", "Primary owner", "Required collaborator / gate"], [3500, 1800, 4060])
    rows = [
        ("Curriculum, mathematical answer, standards, source safety", "A18", "Independent SME; A23 before promotion"),
        ("Lesson rendering, authored-versus-visible parity", "A05", "A18 content decision; A11 regression"),
        ("Practice item or grader behavior", "A04", "A18 answer QA; A11 regression"),
        ("Visualization mathematics, interaction, safeguard", "A06", "A18 topic-fit review; A11 regression"),
        ("AI Tutor behavior", "A07", "A18 pedagogy; A11 regression"),
        ("Translation, terminology, accessibility copy", "A09", "Qualified language reviewer; A18 sign-off"),
        ("Promotion/integration sequencing", "A23", "A18, A05/A04, A11, A22"),
        ("Build, deployment, production parity", "A22", "A11 regression; relevant feature owner"),
    ]
    for row in rows:
        _add_text_row(routing, list(row), size=8.25)
    _set_table_geometry(routing, [3500, 1800, 4060])


def _add_standards_appendix(document: DocumentObject, catalog: list[StandardCatalogRow]) -> None:
    _add_page_heading(document, "Appendix A. Standards Metadata Appendix - 148 Identifiers")
    document.add_paragraph(
        "The 29 core lesson records collectively list 148 unique fine-grained K-Grade 5 identifiers. This appendix is generated from package metadata and is a review index only; it does not establish instructional sufficiency, adoption status, or QA acceptance."
    )
    counts = _add_table(document, ["Grade", "Unique identifiers", "Review note"], [1500, 1800, 6060])
    grade_labels = {"K": "Kindergarten", "P1": "Grade 1", "P2": "Grade 2", "P3": "Grade 3", "P4": "Grade 4", "P5": "Grade 5"}
    for grade, count in EXPECTED_STANDARD_COUNTS.items():
        _add_text_row(counts, [grade_labels[grade], str(count), "Metadata coverage; require per-standard lesson evidence"], size=8.5)
    _set_table_geometry(counts, [1500, 1800, 6060])

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)

    table = _add_table(document, ["Grade", "Domain", "Cluster", "Canonical IDs / MAIS prefix rule"], [1200, 1850, 2100, 4210])
    for row_data in catalog:
        canonical = ", ".join(row_data.standard_ids)
        mais = ", ".join(f"CA.CCSS.Math.{standard_id}" for standard_id in row_data.standard_ids)
        row = _add_text_row(
            table,
            [
                row_data.grade_label,
                f"{row_data.domain_id}\n{row_data.domain_title}",
                f"{row_data.cluster_id}\n{row_data.cluster_title}",
                f"Canonical: {canonical}\nMAIS: {mais}",
            ],
            size=7.55,
        )
        _set_row_cant_split(row)
    _set_table_geometry(table, [1200, 1850, 2100, 4210])

    progression_heading = document.add_heading("Appendix B. K-Grade 5 Progression Prompts", level=1)
    progression_heading.paragraph_format.page_break_before = True
    progression = _add_table(document, ["Grade", "MAIS-authored progression prompt for QA"], [1200, 8160])
    prompts = [
        ("Kindergarten", "Move from concrete/picture counting toward number words, quantities, early composition/decomposition, measurement language, and basic shape reasoning."),
        ("Grade 1", "Deepen addition/subtraction situations, place value to 120, measurement/data, and shape partitioning through equations, drawings, counters, number lines, and spoken reasoning."),
        ("Grade 2", "Build fluency within 20, base-ten reasoning to 1000, standard-unit measurement, time/money/data, arrays, and equal partitions while addressing regrouping and unit confusion."),
        ("Grade 3", "Develop multiplication/division meaning, area, fractions on number lines, perimeter, elapsed time, scaled graphs, and quadrilateral categories from models before symbols."),
        ("Grade 4", "Generalize multi-digit operations, factors/multiples, fractions/decimals, conversion, angle measure, and geometric classification with estimation and reasonableness checks."),
        ("Grade 5", "Prepare for middle school through decimal operations, fraction multiplication/division, volume, coordinate graphing, numerical expressions, and pattern relationships using models before algorithms."),
    ]
    for row in prompts:
        _add_text_row(progression, list(row), size=8.5)
    _set_table_geometry(progression, [1200, 8160])


def _count_content_controls(docx_path: Path) -> tuple[int, int]:
    with zipfile.ZipFile(docx_path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return len(root.xpath(".//w:sdt", namespaces=ns)), len(root.xpath(".//w:tbl", namespaces=ns))


def _strip_revision_identifiers(docx_path: Path) -> None:
    """Remove transient Word revision-session IDs without changing content."""
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with tempfile.NamedTemporaryFile(
        prefix=f".{docx_path.stem}.", suffix=".docx", dir=docx_path.parent, delete=False
    ) as temp_file:
        temp_path = Path(temp_file.name)
    try:
        with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as destination:
            for member in source.infolist():
                payload = source.read(member.filename)
                if member.filename == "word/document.xml":
                    root = etree.fromstring(payload)
                    for element in root.iter():
                        for attribute_name in list(element.attrib):
                            qualified = etree.QName(attribute_name)
                            if (
                                qualified.namespace == word_namespace
                                and qualified.localname.startswith("rsid")
                            ):
                                del element.attrib[attribute_name]
                    payload = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                destination.writestr(member, payload)
        os.replace(temp_path, docx_path)
    finally:
        temp_path.unlink(missing_ok=True)


def build_document(project_root: Path, output_path: Path) -> BuildSummary:
    project_root = project_root.resolve()
    output_path = output_path.resolve()
    inventory = load_lesson_inventory(project_root)
    catalog = load_standard_catalog(project_root)
    target = next(row for row in inventory if row.topic_id == TARGET_TOPIC_ID)

    document = Document()
    _configure_styles(document)
    _configure_sections(document)
    document.core_properties.title = "California K-Grade 5 Lesson Content QA Template"
    document.core_properties.subject = "Internal MAIS A18 content quality assurance workbook"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""

    _add_title_block(document)
    _add_front_matter(document)
    _add_tracker(document, inventory)
    _add_calibration_record(document, target)
    _add_detailed_form(document)
    _add_issue_retest_and_signoff(document)
    _add_owner_routing(document)
    _add_standards_appendix(document, catalog)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    _strip_revision_identifiers(output_path)
    content_controls, table_count = _count_content_controls(output_path)
    return BuildSummary(
        output_path=output_path,
        lesson_count=len(inventory),
        standard_count=sum(len(row.standard_ids) for row in catalog),
        content_control_count=content_controls,
        table_count=table_count,
    )


def main() -> None:
    parser = ArgumentParser(description=__doc__)
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(__file__).resolve().parents[3],
        help="MAIS-MVP worktree root",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).with_name("MAIS_CA-Math_K-5_Content_QA_Template.docx"),
        help="Output DOCX path",
    )
    args = parser.parse_args()
    summary = build_document(args.project_root, args.output)
    print(
        json.dumps(
            {
                "output": str(summary.output_path),
                "lessonCount": summary.lesson_count,
                "standardCount": summary.standard_count,
                "contentControlCount": summary.content_control_count,
                "tableCount": summary.table_count,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
