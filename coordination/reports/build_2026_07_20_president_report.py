from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-20-president-report.docx")
REPORT_DATE = "2026-07-20"
WINDOW = "2026-07-19 06:00 to 2026-07-20 06:00 Asia/Shanghai (UTC+8)"
UTC_WINDOW = "2026-07-18 22:00 to 2026-07-19 22:00 UTC"
FORMAT = "President Report Business Brief v1"

# STSong is available to the local LibreOffice renderer and supports Simplified Chinese.
FONT = "STSong"
BLUE = RGBColor(46, 116, 181)
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"
RISK_FILL = "FCE4D6"


def set_run_font(run, size=11, bold=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    language = rpr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        rpr.append(language)
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "zh-CN")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_widths(table, widths):
    if sum(widths) != 9360:
        raise ValueError("All report tables must use 9,360 DXA.")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag, value in (("tblW", "9360"), ("tblInd", "120")):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_pr.append(node)
        node.set(qn("w:type"), "dxa")
        node.set(qn("w:w"), value)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)


def keep_row(row, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)
    if header:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)


def add_table(doc, headers, rows, widths, body_size=7.4, row_fills=None, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_widths(table, widths)
    keep_row(table.rows[0], header=True)
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        set_run_font(p.add_run(value), size=8.7, bold=True)
        shade(cell, header_fill)
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        keep_row(row)
        fill = (row_fills or {}).get(row_index)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i < 2 and len(str(value)) <= 22:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(str(value)), size=body_size)
            if fill:
                shade(cell, fill)
    set_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def para(doc, text="", size=10, bold=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), size=15, bold=True, color=BLUE)
    return p


def bullet(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    set_run_font(p.add_run(text), size=9.5)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    for kind, text in (("begin", None), ("instrText", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
        if kind in ("begin", "separate", "end"):
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText" if kind == "instrText" else "w:t")
            if kind == "instrText":
                node.set(qn("xml:space"), "preserve")
            node.text = text
        run._r.append(node)


def setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for name in ("Normal", "Heading 1"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("MAIS-MVP | Daily coordination report | A01-A25"), size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run(f"{REPORT_DATE} | {FORMAT} | Page "), size=8.5, color=MUTED)
    add_page_field(footer)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)


def agent_rows():
    rows = {
        "A01": ("提交证据", "login/register funnel 与 demo CTA 路由提交；无正式交接。", "A01/A12 补 auth-shell 回归与 handoff。"),
        "A02": ("提交证据", "Knowledge Galaxy dashboard 变更；无 A02 log。", "A02/A15 确认 adaptive UI 回归。"),
        "A03": ("无新增日志", "未发现窗口内正式 A03 handoff。", "保持现状。"),
        "A04": ("提交证据", "Practice Island、stars API、kid mode、read-aloud 提交。", "A04/A11 补 targeted practice E2E。"),
        "A05": ("提交证据", "CCSS textbook port：270 interactive lessons。", "A05/A18/A11 复核内容与 lesson routes。"),
        "A06": ("提交证据", "CCSS lab catalog、184 signature benches 与 probes。", "A06/A11/A22 清洁工作树验证。"),
        "A07": ("无新增日志", "未发现窗口内正式 A07 handoff。", "保持 provider 改动冻结。"),
        "A08": ("无新增日志", "未发现窗口内正式 A08 handoff。", "共享类型改动由 A08 核对。"),
        "A09": ("无新增日志", "未发现窗口内正式 A09 handoff。", "auth/lesson 新文案后续可及性检查。"),
        "A10": ("报告连续性", "仅有前一日报告日志；本报告在 cutoff 后生成。", "A10 维持证据/格式一致性。"),
        "A11": ("无新增日志", "没有 A11 full-regression handoff。", "A11 拆分并执行 clean-source gate。"),
        "A12": ("提交证据", "credentials-only login / privacy funnel；无正式 A12 log。", "A12/A11 验证 auth/session contract。"),
        "A13": ("WIP 提交证据", "Teacher Console redesign 在 all-ref/WIP；非 main 验收。", "A13 补 scope、checks 与 handoff。"),
        "A14": ("无新增日志", "仅见 parent E2E 的当前 dirty path。", "A14/A11 不将 dirty 测试视为验收。"),
        "A15": ("提交证据", "Galaxy/adaptive surface 变更；无 A15 regression log。", "A15/A02 补 recommendation semantics。"),
        "A16": ("无新增日志", "未发现窗口内正式 A16 handoff。", "保持现状。"),
        "A17": ("WIP 证据", "当前 gamification paths 处于 dirty 状态。", "A17/A20 分离 reward 与 game slices。"),
        "A18": ("提交证据", "CCSS lesson port 已入 main；无 fresh content-QA signoff。", "A18 独立内容 QA 后再 promotion。"),
        "A19": ("无新增日志", "未发现窗口内环境配置交接。", "保持 secrets/env 不变。"),
        "A20": ("WIP 证据", "Adventure/Fishing game paths 尚为 dirty/WIP。", "A20/A22 先隔离可发布 slice。"),
        "A21": ("无新增日志", "未发现窗口内 fresh candidate handoff。", "不得绕过 A18/A23 gate。"),
        "A22": ("无新增日志", "无 clean-worktree build/Playwright handoff。", "A22 建立 clean release source。"),
        "A23": ("无新增日志", "未见窗口内 promotion decision。", "维持 candidate-to-live gate。"),
        "A24": ("无新增日志", "未见窗口内 exact-layer signoff。", "有需要时按链路交接。"),
        "A25": ("A10代跑清单", "A10 触发 required dirty-map；70 entries、15 unmapped、0 ambiguous。", "需 A25 正式 ownership/slice handoff。"),
    }
    return [(f"A{i:02d}", *rows[f"A{i:02d}"]) for i in range(1, 26)]


def build():
    doc = Document()
    setup(doc)
    prepared = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M CST")
    para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    add_table(doc, ["Field", "Value"], [
        ("Report date", REPORT_DATE),
        ("Scheduled report time", "06:00 Asia/Shanghai (UTC+8)"),
        ("Prepared during run", prepared),
        ("Reporting window", WINDOW),
        ("UTC filter window", UTC_WINDOW),
        ("Prepared by", "A10 reporting automation; A01-A25 formal scope"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Format", FORMAT),
        ("Evidence basis", "Session/blocker/decision review; main Git window; safe fresh checks; current dirty-map snapshot."),
    ], [2200, 7160], body_size=8.0, header_fill=LIGHT_BLUE_FILL)

    heading(doc, "Chinese Executive Summary")
    add_table(doc, ["序号", "重点"], [
        ("1", "窗口内 main 有 14 个提交、564 个变更路径（108,156 insertions / 1,607 deletions）。重点包括登录注册漏斗、Practice Island、完整 CCSS 知识星系、CCSS 可视化目录，以及 270 个互动教材课程。"),
        ("2", "`npm run type-check` 通过，`npm run test:analytics` 27/27 通过；但 production build 与 broad Playwright 未运行，不能把本窗口提交视为 release-ready。"),
        ("3", "窗口内没有新的 A01-A25 正式 feature handoff。提交能证明代码移动，不能替代对应 owner 的验收、内容 QA 或 A11/A22 发布门禁。"),
        ("4", "报告生成时 required dirty-map 为 70 entries，其中 15 个尚未映射 owner；root release frozen，当前 dirty root 不可作为生产发布源。"),
        ("5", "最需要的管理动作是：A25 将 15 个未映射项与 WIP 切成责任包；A11/A22 在 clean worktree 运行 build/E2E；A01/A04/A05/A06/A12/A15 对相应提交补正式 handoff。"),
    ], [700, 8660], body_size=8.0)

    heading(doc, "English Executive Summary")
    add_table(doc, ["No.", "Summary"], [
        ("1", "Main recorded 14 in-window commits across 564 paths (108,156 insertions and 1,607 deletions), covering login/register, Practice Island, the CCSS knowledge galaxy, the visualization catalog, and 270 interactive CCSS lessons."),
        ("2", "Fresh type-check passed and analytics passed 27/27. A production build and broad Playwright gate were not run, so the window's commits are not release-ready evidence."),
        ("3", "No new formal A01-A25 feature handoffs were recorded. Git commits prove movement, not owner acceptance, curriculum QA, or A11/A22 release clearance."),
        ("4", "The post-cutoff required dirty map contains 70 entries, including 15 without an owner. The root release remains frozen and cannot be a production source."),
        ("5", "The immediate management need is ownership/slicing by A25, a clean-worktree build/E2E gate by A11/A22, and formal handoffs from the owners of the changed surfaces."),
    ], [700, 8660], body_size=8.0)

    heading(doc, "报告窗口摘要（中文）")
    add_table(doc, ["项目", "结果"], [
        ("正式窗口", WINDOW),
        ("主线提交", "14 commits；首个 `b4cb36c7b3`（CCSS knowledge galaxy），末个 `8e2b520176`（CCSS textbook worktree merge）。"),
        ("主线文件变化", "564 paths；482 components、26 data、25 lib、11 tests、7 app、7 scripts；108,156 insertions / 1,607 deletions。"),
        ("正式会话日志", "未发现窗口内新增的 A01-A25 feature handoff；A10 仅有前一日的 report continuity log。"),
        ("Blockers / decisions", "没有窗口内新增 blocker 或 decision record；历史 blocker 不得自动视为本次完成或解除。"),
        ("All-ref / WIP", "A13 Teacher Console redesign 与 A20/A17 游戏相关工作存在 all-ref 或 dirty 证据，未计为 main 已验收交付。"),
        ("空窗口 fallback", "不适用：存在 main 提交活动；但 formal owner handoff 为空。"),
    ], [2450, 6910], body_size=7.25)

    heading(doc, "整体项目进展（中文）")
    bullet(doc, "产品面：登录/注册新增 credentials-only 与 privacy-first funnel 变更；Practice Island 加入 stars API、kid mode 与 read-aloud；知识星系扩展到完整 CCSS star map、ignition bursts 与 constellation badges。")
    bullet(doc, "课程与可视化面：CCSS-style Visualization Lab catalog、184 个 signature benches 与 runtime probes 已有提交证据；CCSS textbook port 宣称 270 个 interactive lessons，仍需 A05/A18/A11 的独立验收链。")
    bullet(doc, "治理面：本报告的 A10 只汇总证据。没有新增 formal handoff 的情况意味着当前优先级应从继续堆叠改动，转为按 owner 补 checks、签核与可发布切片。")
    bullet(doc, "发布面：A25 规则要求 dirty root 不发布。当前 70-entry dirty snapshot、15 个未映射 owner 及缺少 clean build/E2E 都使发布状态为红色。")

    heading(doc, "A01-A25 会话状态表（中文）")
    fills = {0: PENDING_FILL, 1: PENDING_FILL, 3: PENDING_FILL, 4: PENDING_FILL, 5: PENDING_FILL, 9: PENDING_FILL, 10: RISK_FILL, 11: PENDING_FILL, 12: PENDING_FILL, 13: PENDING_FILL, 14: PENDING_FILL, 16: PENDING_FILL, 17: PENDING_FILL, 19: PENDING_FILL, 20: PENDING_FILL, 21: RISK_FILL, 22: PENDING_FILL, 23: PENDING_FILL, 24: RISK_FILL}
    add_table(doc, ["ID", "状态", "窗口内活动", "下一步 / 风险"], agent_rows(), [620, 1320, 3730, 3690], body_size=6.05, row_fills=fills)

    heading(doc, "已完成工作（中文）")
    bullet(doc, "提交证据：CCSS Knowledge Galaxy（`b4cb36c7b3`）、Visualization Lab catalog / signature benches、Practice Island、credentials-only login / register funnel，以及 CCSS textbook port（`bcdf95b7ff`）均已进入 main。")
    bullet(doc, "当前基础验证：A10 fresh `npm run type-check` 通过；`npm run test:analytics` 27/27 通过；`git diff --check`、index diff check 通过，且无已暂存变更。")
    bullet(doc, "A10 按报告门禁刷新 A25 dirty-map，并输出本次 DOCX；未改动 feature code、Git history、部署或 secret 配置。")

    heading(doc, "进行中工作（中文）")
    bullet(doc, "A13 Teacher Console redesign 及 A17/A20 游戏流目前只见 WIP/all-ref/dirty evidence；需要各 owner 的明确文件范围、目标检查与 handoff。")
    bullet(doc, "A01/A04/A05/A06/A12/A15 所映射的 main 变更需补 owner-level acceptance：auth、practice、lesson/content、visualization 与 adaptive semantics 不应只凭 commit message 关闭。")
    bullet(doc, "A25 需将 15 个 unmapped paths 分派到现有 A01-A25 owner，并把 70-entry root 状态拆为 reviewable clean slices。")

    heading(doc, "阻塞项（中文）")
    add_table(doc, ["阻塞项", "证据", "Owner / safe next action"], [
        ("无法从 dirty root 发布", "A25 policy: root release frozen；post-cutoff dirty-map 70 entries。", "A25/A22：只使用 clean worktree、clean clone 或 reviewed clean slice。"),
        ("15 个未映射 owner", "required dirty-map: 15 unmapped、0 ambiguous。", "A25：建立 owner/pathspec 与 package routing。"),
        ("发布门禁不完整", "build 与 broad Playwright 未运行；无 A11/A22 final handoff。", "A11/A22：在 clean source 运行并记录结果。"),
        ("正式验收日志缺失", "窗口内无 fresh A01-A25 feature handoff。", "相应 owner：补 files/checks/risk/next-step handoff。"),
        ("历史 assertion script 缺失", "`coordination/release-intake/assert-no-staged-changes.mjs` 不存在；直接 Git 已确认无 staged changes。", "A10/A25：仅在确认属于 tooling contract 后再决定是否恢复。"),
    ], [1900, 4200, 3260], body_size=6.65, row_fills={0: RISK_FILL, 1: RISK_FILL, 2: RISK_FILL, 3: PENDING_FILL, 4: PENDING_FILL})

    heading(doc, "风险（中文）")
    bullet(doc, "规模风险：564 paths、482 个 components 的单窗口变更横跨 auth、课程、practice、visualization、shared types 与测试；任何单一通过的 TypeScript 检查都不足以覆盖功能行为。")
    bullet(doc, "归属风险：没有 fresh formal handoff 时，无法可靠判断每一流是否完成、被谁验收、是否存在未记录的 cross-file 冲突。")
    bullet(doc, "内容风险：270 CCSS lessons 是提交描述，不等同于 A18 curriculum QA 或 A23 candidate-to-live decision；不得自动视为最终上线内容。")
    bullet(doc, "发布风险：dirty root 与未映射路径可能混入不同 owner 的 WIP；A22 必须从 clean source 验证，而不是通过删除或忽略这些路径来取得假绿。")

    heading(doc, "测试 / 构建状态（中文）")
    add_table(doc, ["检查", "结果", "解释"], [
        ("npm run type-check", "PASS", "A10 在当前 checkout 运行；仅 TypeScript gate。"),
        ("npm run test:analytics", "PASS 27/27", "学习分析/自适应的既有 suite 绿色。"),
        ("git diff --check", "PASS", "工作树和 index diff checks 均无 whitespace errors。"),
        ("已暂存变更", "NONE", "用 direct Git check 确认；旧 assertion script path 缺失。"),
        ("release:dirty-map", "RED release intake", "70 entries、15 unmapped、0 ambiguous；root release frozen。"),
        ("npm run build", "Not run", "A22 clean-worktree release gate，不能由 dirty root 报告替代。"),
        ("Broad Playwright", "Not run", "A11/A22 尚无 final regression handoff。"),
    ], [2200, 2100, 5060], body_size=7.0, row_fills={0: OK_FILL, 1: OK_FILL, 2: OK_FILL, 3: OK_FILL, 4: RISK_FILL, 5: PENDING_FILL, 6: PENDING_FILL})

    heading(doc, "变更文件（中文）")
    add_table(doc, ["类别", "窗口内文件证据", "说明"], [
        ("提交总量", "564 paths；108,156 insertions / 1,607 deletions。", "main 的精确时间窗 diff。"),
        ("组件与可视化", "482 components；含 `AdaptiveKnowledgeGalaxy.tsx`、Visualization Lab catalog 与 signature benches。", "A02/A06 需补 route/browser 验证。"),
        ("课程与数据", "26 data paths；CCSS textbook/lesson 相关文件。", "A05/A18 的课程 QA 尚未由 formal log 证明。"),
        ("逻辑与共享类型", "25 lib paths、`types/index.ts`、auth / practice supporting code。", "A08/A12/A15 需检查 cross-surface semantics。"),
        ("测试与配置", "11 test paths、`playwright.config.ts`、7 scripts、`package.json`。", "A11/A22 需确认 clean gate 与 harness behavior。"),
        ("报告后当前状态", "70 dirty entries，15 unmapped；本报告及构建脚本属于 report-only artifacts。", "后截止时间的状态，只作 release-risk context。"),
    ], [2050, 4500, 2810], body_size=6.85)

    heading(doc, "明日优先事项（中文）")
    bullet(doc, "A25：先为 15 个未映射路径指定 A01-A25 owner，并把 70-entry dirty root 拆成 runtime、tests、docs、release-hygiene 与 WIP clean slices。", numbered=True)
    bullet(doc, "A11/A22：在 isolated clean worktree 运行 `npm run type-check`、production build、目标 Playwright（auth、practice、lesson、visualization），并记录最终 release gate。", numbered=True)
    bullet(doc, "A01/A04/A05/A06/A12/A15：对各自 main 变更补正式 handoff，列明实际文件、checks、遗留风险与跨 owner 依赖。", numbered=True)
    bullet(doc, "A13/A17/A20：把 Teacher Console 与游戏 WIP 从可发布 slice 分离；未经 owner handoff 与 A11/A22 checks 不进入 release planning。", numbered=True)

    heading(doc, "需要 Owner 决策（中文）")
    bullet(doc, "请确认：在 A25 完成 owner mapping 及 A11/A22 clean build/E2E 前，是否继续保持当前发布冻结，而不是从 dirty root 部署。")
    bullet(doc, "请指定优先验收顺序：登录/注册、Practice Island、CCSS lessons、Visualization Lab、Knowledge Galaxy 五个已提交 surface 中，哪一个先获得正式 owner + QA + regression handoff。")
    bullet(doc, "请决定 A13 Teacher Console 与 A17/A20 游戏 WIP 是否进入下一批独立 review packages，或继续保留为未发布的 in-flight work。")
    bullet(doc, "请授权 A25 在既有 A01-A25 模型内为 15 个 unmapped paths 提出具体 owner/pathspec 修订；不应新建 A26+ role。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
