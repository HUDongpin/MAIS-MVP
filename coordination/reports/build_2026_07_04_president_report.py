from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-04-president-report.docx")

FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    hdr = table.rows[0].cells
    repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        hdr[idx].text = ""
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, bold=True, color=BLACK)
        set_cell_shading(hdr[idx], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=body_size, color=BLACK)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DEEP_BLUE),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-04"),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Reporting window", "2026-07-03 08:00 to 2026-07-04 08:00 Asia/Hong_Kong"),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL)


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "本窗口有 fresh AI session activity：A04/A11 关闭一个 `/about` logged-out progress 展示问题，A11 完成两批 bug report 核验并验证 20 个当前真实问题的局部修复，A06 持续推进 Manim v2 source-architecture owner-gate 证据链，A25 完成 dirty-worktree 治理状态更新。"),
            ("2", "报告时 fresh checks 通过：`npm run release:dirty-map -- --assert-current --max-age-minutes 240` 显示 A25 dirty map current、4206 expanded status entries；`npm run type-check -- --pretty false` 通过；`npm run test:analytics` 通过 27/27；`git diff --check` 无输出。"),
            ("3", "A25 当前闭环处于“验证 / validate”阶段：切片和提取已完成，合并与清理仍被 71 条 pending canonical authorization rows、owner input not ready、compose deletion confirmation hold、dirty release source 与 strict lifecycle gate 阻塞。"),
            ("4", "A22 release-source clean gate 仍为 blocked；root main 有 1455 collapsed dirty status entries、4206 expanded dirty entries，不可作为常规发布源。任何 deploy 仍需 clean worktree、clean clone、reviewed clean release slice 或 owner-approved pruned staging package。"),
            ("5", "今天最需要 Dr. Peter Hu 决策的是：确认 compose worktree deletion 是否已完成；审核/批准/拒绝/暂缓 71 条 canonical authorization rows；指派 A11/A18/A22 补齐 A06 Manim v2 canonical owner reports；决定 A11/A04 bug fix 是否进入 clean release slice。"),
        ],
        [700, 8660],
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "Fresh window activity exists across A04, A06, A11, and A25. A04/A11 fixed the public `/about` logged-out progress state; A11 verified external bug reports and locally verified 20 real/current fixes; A06 advanced Manim v2 owner-gate evidence; A25 hardened dirty-worktree governance."),
            ("2", "Fresh report checks passed: the A25 dirty map is current with 4,206 expanded entries; root type-check passed; analytics tests passed 27/27; tracked diff whitespace checks passed."),
            ("3", "The project is not release-clean. A25 reports the closure loop is still in validate; merge and cleanup are blocked by pending owner authorization, owner input readiness, a compose-deletion hold, and dirty release/lifecycle gates."),
            ("4", "A11 and A04 produced valuable local quality improvements, but those fixes are dirty-root evidence, not clean release proof. A22 must still consume a clean source path before any production claim."),
            ("5", "Owner attention should focus on authorization rows, the compose deletion confirmation, A06 Manim v2 owner reports, and clean-slice packaging for the verified bug fixes."),
        ],
        [700, 8660],
    )

    add_heading(doc, "汇报窗口摘要", 1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("窗口", "2026-07-03 08:00 至 2026-07-04 08:00 Asia/Hong_Kong。"),
            ("活跃 session logs", "A04 bug-202 `/about` public practice showcase fix；A06 Manim v2 source-architecture continuation；A11 bug verification/fix verification；A25 dirty-worktree governance status。"),
            ("无 fresh 文件", "`coordination/blockers/` 与 `coordination/decisions/` 在窗口内无新增文件；blocking evidence 主要来自 A25 release-intake artifacts 与 A06/A11 handoff reports。"),
            ("提交记录", "窗口内有 1 个 Git commit：`e909992b0 ci: make backup snapshot checks push-safe`，改动 `.github/workflows/ci.yml`。"),
            ("No assigned work flag", "No assigned work in this reporting window: No。虽然没有单一 nightly assignment 文件，窗口内有明确 session activity 与验证产物。"),
        ],
        [1900, 7460],
    )

    add_heading(doc, "整体项目进展", 1)
    add_para(doc, "产品质量方面，A11 将两个外部 bug report 拆成可执行结论：20260704 report 的 6 个 claims 中仅 Bug 202 `/about` public practice showcase 被确认需要立即修；20260703 Bug 167-202 中 20 个真实/当前问题已被本地验证为已修复。")
    add_para(doc, "工程治理方面，A25 已把 dirty-root 工作推进到更清晰的闭环状态：slice 与 extract 完成，validate active，owner blocker reports 当前为 10/10 valid、0 pending；但 merge/cleanup 仍不能开始。")
    add_para(doc, "Visualization/Manim 方面，A06 完成一系列 source-architecture、owner-gate、source-ledger mismatch 与 final-closure guard 传播，使未来 Manim v2/NANIM 2.0 invocation 保持 one-topic/one-concept-cluster/one-review-slice 的边界；最终完成仍依赖 A11/A18/A22。")

    add_heading(doc, "A01-A25 会话状态表", 1)
    status_rows = [
        ("A01", "App shell", "无 fresh log", "A11/A04 已修 `/about` public showcase 误显 progress；若进入 release slice，需 A01/A09 复核 shell/copy 影响。", "无直接 blockers；受 dirty-root release gate 影响。", "未单独运行。"),
        ("A02", "Dashboard", "无 fresh log", "A25 指向 A02/A15 dashboard-adaptive closure；A15 检查发现 dashboard/adaptive package 仍依赖 shared contracts。", "A08/A12/A07/A10/A22 前置依赖。", "A25 A15 worktree `test:analytics` 编译失败。"),
        ("A03", "Roadmap", "A25 routed blocker", "A25 在 A03 worktree 记录 roadmap closure blocker。", "A03 narrow `data/topics.ts` 范围不足以解决 GradeId/difficulty/curriculum shared drift。", "A03 worktree type-check failed，184 A03-scoped errors。"),
        ("A04", "完成", "修复 `/about` logged-out public mission showcase：Ready/Preview set 替代 In progress/Round progress，并加 source regression。", "A04 broader practice package仍需 shared contract/game/content coordination。", "Source regression 18/18；local browser `/about` proof passed。"),
        ("A05", "被 A11 消费", "A11 bug-fix slice 涉及 worked-example exact SVG scenes and lesson visual fidelity。", "Clean release slice 与 A18/A21/A24 content/exact-layer gate 未完成。", "A11 focused suite 109/109；root type-check passed。"),
        ("A06", "高活跃；未最终验收", "Manim v2 source-architecture handoff、browser bridge、owner-gate/report guards、source mismatch propagation guards持续完成。", "缺 A11 browser report、A18 teaching final decisions、A22 clean release gate、final 4/4 objective audit。", "多轮 focused/adjacent/Manim tests passed；root type-check report时 passed。"),
        ("A07", "A25 routed blocker", "A25 记录 AI tutor package blocker。", "缺 A12 userStore/aiGovernance exports、A07 provider contract alignment、A08 shared type coordination。", "A07 worktree type-check failed，28 A07-scoped lines。"),
        ("A08", "A25 shared dependency", "Shared schema/state remains a critical dependency for A03/A04/A07/A12/A15/A18 packages。", "Shared type/curriculum/difficulty/storage drift blocks owner packages。", "Root type-check passed；isolated packages still fail。"),
        ("A09", "无 fresh log", "Copy/accessibility owner may need to review public wording and US grade labels after package routing。", "No direct blocker; depends on owner package scope. ", "未单独运行。"),
        ("A10", "报告完成", "生成本 president report；读取 AGENTS、automation memory、session logs、reports、release-intake evidence。", "不涉及 feature code；build 未跑。", "DOCX validation/render QA performed separately。"),
        ("A11", "高活跃", "完成 20260704 deliverable bug verification；完成 Bug 167-202 real/current verification；验证 20 项局部修复。", "Dirty-root verification不是 release proof；部分 production probes inconclusive。", "109/109 focused tests；`npm run type-check` passed；browser smoke passed。"),
        ("A12", "A25 routed blocker", "A25 修正 A12 recommended worktree 到 A08-A12 shared-contract closure。", "A12 blocked pending A07/A12 split and shared exports/userStore/API contract remediation。", "A25 routing gates passed after repair。"),
        ("A13", "A25 routed blocker", "A25 记录 A13/A14 console blocker；A11修复 teacher assignment filter行为。", "A13/A14 local fix depends on A08/A12/A10/A22 shared/package baseline。", "A13/A14 worktree type-check failed。"),
        ("A14", "A25 routed blocker", "Parent console/profile/avatar lifecycle included in A25 closure sequencing。", "A14 profile/avatar worktree lifecycle still needs owner decisions。", "未单独运行。"),
        ("A15", "A25 routed blocker", "A25 修正 A15 dashboard/adaptive routing到 A02/A15 worktree。", "Adaptive checks blocked by A08 schema、A12 storage/API、A07 provider-name alignment、A10/A22 package baseline。", "A15 worktree `test:analytics` failed during compile。"),
        ("A16", "无 fresh log", "No fresh research-session activity in window。", "无 fresh blocker。", "未运行。"),
        ("A17", "A20 package dependency", "A17/A20 game/motivation closure被 A25 检查。", "A20 blocker report row routing mismatch已记录；game package需正确 surfaced。", "A20 worktree type-check failed earlier on shared blockers。"),
        ("A18", "A25/A06 dependency", "A18/A21 content-evidence closure被 A25检查；A06等待 A18 Manim teaching decisions。", "A18不能独自解决跨 A06/A08/A12/A13/A20/content failures；Manim final teaching decisions missing。", "A18/A21 worktree `test:analytics` failed during compile。"),
        ("A19", "无 fresh log", "No fresh API env/config work in window。", "无 fresh blocker；provider/env changes仍须 A19-owned redacted process。", "未运行。"),
        ("A20", "A25 routed blocker", "A25 记录 A17/A20 game/motivation blocker；A11 report route includes game/practice items。", "A20 package row/routing and shared type gate block local remediation。", "A20 worktree type-check failed before A20-scoped errors were reached。"),
        ("A21", "A18 dependency", "A18/A21 content package dirty scope inspected by A25。", "Content/RAG promotion must wait for A18 QA, shared contracts, and release gates。", "A18/A21 worktree `test:analytics` failed during compile。"),
        ("A22", "Release gate blocked", "A22 release-source evidence consumed by A25: dirty-root deploy blocked closed; no deploy command executed。", "Release source clean=no；root status 4206 expanded entries；clean-source path required。", "No-dirty-root deploy probes passed as failed-closed。"),
        ("A23", "无 fresh log", "No candidate-to-live integration log in window。", "Content packages still require promotion planning before live integration。", "未运行。"),
        ("A24", "A05/A18 dependency", "A11 worked-example fixes use exact visual fidelity path; A24 exact-layer remains dependent on A18/A21 package decisions。", "No independent A24 signoff in window。", "未单独运行。"),
        ("A25", "高活跃", "Dirty-worktree governance status updated；owner blocker reports closed to 0 pending；A12/A15 routing repaired；authorization frontier pre-checks ready 71/71。", "71 pending canonical authorization rows；owner input not ready；compose deletion confirmation hold；0 cleanup/executable rows。", "Dirty map current; no staged changes; remediation gates current in A25 artifacts。"),
    ]
    roles = {
        "A01": "App shell",
        "A02": "Dashboard",
        "A03": "Roadmap",
        "A04": "Practice",
        "A05": "Lesson",
        "A06": "Visualization",
        "A07": "AI tutor",
        "A08": "State/analytics",
        "A09": "Copy/i18n/a11y",
        "A10": "Reporting",
        "A11": "QA",
        "A12": "Backend/API",
        "A13": "Teacher console",
        "A14": "Parent console",
        "A15": "Adaptive engine",
        "A16": "Research",
        "A17": "Gamification",
        "A18": "Content QA",
        "A19": "API env",
        "A20": "Games",
        "A21": "Content/RAG",
        "A22": "Release",
        "A23": "Integration",
        "A24": "Illustration",
        "A25": "Git hygiene",
    }
    status_rows = [row if len(row) == 6 else (row[0], roles[row[0]], *row[1:]) for row in status_rows]
    add_table(doc, ["Agent", "工作流", "状态", "完成/进展", "阻塞/风险", "检查"], status_rows, [760, 1120, 860, 2880, 2580, 1160], body_size=6.9)

    add_heading(doc, "已完成工作", 1)
    for item in [
        "A04/A11：完成 Bug 202 `/about` logged-out public practice showcase 修复；source regression red-first/green 18/18，浏览器验证 mission section 无 progressbar、无 In progress/Round progress。",
        "A11：完成 20260704 deliverable bug report 6 项核验；仅 `/about` progress state 保持为 immediate real bug，其余多数为 not reproduced/not current/inconclusive。",
        "A11：完成 Bug 167-202 核验与 fix verification；20 个真实/当前问题本地修复验证通过，focused regression 109/109、root type-check 通过、local browser smoke 通过。",
        "A06：完成多轮 Manim v2 source-architecture/owner-gate/source mismatch/final closure guard 传播，确保缺失 A11/A18/A22 owner reports 时不能误报 complete。",
        "A25：完成 dirty-worktree governance update；owner package blocker records 10/10 valid、pending 0；A12/A15 recommended worktree routing修复；dirty-map current。",
        "A10：完成本次 report evidence collection 与 fresh checks；未编辑 feature code。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "进行中工作", 1)
    for item in [
        "A25 closure loop：slice/extract complete，validate active；merge/cleanup blocked，等待 owner authorization 与 compose deletion confirmation。",
        "A06 Manim v2：source architecture 已可交接，但仍需 A11 browser visual/interaction report、A18 teaching-quality final decisions、A22 clean release gate。",
        "A11/A04 bug fixes：当前是 dirty-root local verification，需要 clean release slice 或 reviewed staging path 才能转为生产发布证据。",
        "A03/A04/A07/A12/A13/A15/A18/A20 owner packages：多数 isolated package checks 仍受 shared schema、package baseline、storage/API exports 或 content promotion gates 阻塞。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "阻塞项", 1)
    blocker_rows = [
        ("A25/A22", "Release source clean blocked", "A22 release-source clean=no；4206 expanded status entries；root main不能作为发布源。"),
        ("A25", "Owner authorization blocked", "71 pending canonical authorization rows；0 valid authorization rows；0 cleanup-authorized rows；0 executable rows。"),
        ("A25", "Validation hold", "等待 owner 确认 `/Users/dongpinhu/.config/superpowers/worktrees/MAIS-MVP/A10-A22-A08-A12-A06-compose-20260628` exact deletion 已完成。"),
        ("A06/A11/A18/A22", "Manim final gate blocked", "缺 A11 browser visual/interaction report、A18 final teaching decisions、A22 clean release gate。"),
        ("A08/A12/A10/A22", "Shared contract/package baseline drift", "多个 owner worktree type-check/test 失败来自 shared types、userStore/API exports、missing dependencies/package baseline。"),
    ]
    add_table(doc, ["Owner", "阻塞", "证据/影响"], blocker_rows, [1300, 2300, 5760], header_fill=RISK_FILL, body_size=8.6)

    add_heading(doc, "风险", 1)
    for item in [
        "Dirty-root blast radius 很大：当前 1455 collapsed status entries、4206 expanded entries；任何单点修复都可能混入 unrelated owner/agent changes。",
        "A11/A04 fixes 虽有 strong local evidence，但仍非 clean-release proof；若直接从 dirty root 发布，会违反 A22 release-source policy。",
        "A25 reports 48,689,219,667 residual generated-artifact bytes across 2 cleanup targets；未获 owner exact approval前不能清理。",
        "A06 Manim v2 证据链很强，但仍是 owner-gate open 状态；若缺 A11/A18/A22 reports 就声称完成，会产生验收风险。",
        "Build 未在本次 report run 中执行；当前可报告的是 type-check、analytics、dirty-map、diff-check，而非 production build readiness。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "测试与构建状态", 1)
    checks = [
        ("Fresh A25 dirty map", "`npm run release:dirty-map -- --assert-current --max-age-minutes 240`", "PASS；latest map current；4206 expanded entries。"),
        ("Fresh root type-check", "`npm run type-check -- --pretty false`", "PASS。"),
        ("Fresh analytics", "`npm run test:analytics`", "PASS；27/27。"),
        ("Fresh whitespace check", "`git diff --check`", "PASS；无输出。"),
        ("A04 focused regression", "`npx tsx --test tests/e2e/reported-bug-source-regressions.test.ts`", "PASS；18/18；local `/about` browser proof passed。"),
        ("A11 bugfix verification", "Focused source/unit + local browser smoke", "PASS；109/109；root type-check passed；browser smoke passed。"),
        ("A06 Manim evidence", "Multiple focused/adjacent/Manim suites", "PASS in A06 slices；latest full Manim suite up to 1628/1628；owner gates still open。"),
        ("Build", "`npm run build`", "Not run；report-only run in dirty root，A22 clean release/build gate should run from clean source path。"),
    ]
    add_table(doc, ["检查", "命令/范围", "结果"], checks, [1700, 3600, 4060], header_fill=OK_FILL, body_size=8.4)

    add_heading(doc, "变更文件", 1)
    add_para(doc, "窗口内文件 mtime 统计：741 files excluding `.git`, `node_modules`, `.next`, `.tmp`, `.vercel`; top areas were `coordination` 644, `components` 78, `app` 3, `lib` 3, `tests` 2. Current git status remains much larger than the reporting-window mtime slice.")
    changed_rows = [
        ("Committed", ".github/workflows/ci.yml", "`e909992b0 ci: make backup snapshot checks push-safe`。"),
        ("A04/A11", "components/practice/PersonalizedPracticeMissionShowcase.tsx; tests/e2e/reported-bug-source-regressions.test.ts", "Public practice mission preview no longer looks like logged-out user progress。"),
        ("A11", "components/learning/LearningRoadmap.tsx; app/teacher/assignments/page.tsx; components/teacher/TeacherEmptyWorkspace.tsx; components/lesson/WorkedExampleIllustration.tsx; components/lesson/workedExampleIllustrationMetadata.ts; components/visualizations/ConfiguredVisualizationLab.tsx; lib/server/userStore/studentActivityPersistence.ts", "20 real/current bugs locally verified as fixed; needs clean-slice review。"),
        ("A06", "components/visualizations/three/manim/* plus A06 report/log", "Manim v2 owner-gate/source mismatch/final-closure guard propagation。"),
        ("A25", "coordination/release-intake/*; coordination/session-logs/2026-07-04-A25-dirty-worktree-governance-status.md", "Dirty-worktree closure loop, owner authorization, blocker-report, and release-source evidence updates。"),
        ("Current status", "components 560; coordination 447; lib 132; app 110; data 81; tests 61; scripts 34", "Current dirty status totals: 1455 entries = 391 modified, 1 deleted, 1063 untracked。"),
    ]
    add_table(doc, ["类别", "文件/区域", "说明"], changed_rows, [1400, 4300, 3660], body_size=8.0)

    add_heading(doc, "明日优先级", 1)
    for item in [
        "A25/A10/A22：先处理 owner input readiness。让 Dr. Peter Hu 对 71 条 canonical authorization rows 批准/拒绝/暂缓，并确认 compose worktree deletion 状态。",
        "A22/A11：把 A11/A04 已验证 bug fixes 打成 clean reviewed release slice；只从 clean worktree/clean clone/pruned staging 跑 build/regression。",
        "A06/A11/A18/A22：补齐 Manim v2 三份 canonical owner reports，再让 A06 final objective audit intake 消费它们。",
        "A08/A12/A10/A22：优先处理 shared schema/storage/API/package-baseline drift，因为它阻塞 A03/A04/A07/A13/A15/A18/A20 owner package closure。",
        "A18/A21/A23/A24：继续保持 candidate-to-live gate，不要把 content/RAG/illustration packages 直接推进 live source。",
    ]:
        add_numbered(doc, item)

    add_heading(doc, "需要 Dr. Peter Hu 决策", 1)
    decision_rows = [
        ("1", "确认 A10-A22-A08-A12-A06 compose worktree exact deletion 是否完成，以解除 A25 validation hold。"),
        ("2", "对 71 条 canonical authorization rows 做 approve/reject/defer；其中包括 7 条 Wave01 package-resync、24 条 owner-package、38 条 physical-lifecycle、2 条 A22 generated-artifact residual cleanup rows。"),
        ("3", "指定是否将 A11/A04 本地已验证 bug fixes 进入 clean release slice；若是，要求 A22 使用 clean/pruned release path 并让 A11跑 targeted regression。"),
        ("4", "要求 A11、A18、A22分别补交 A06 Manim v2 canonical reports：browser visual/interaction、teaching-quality final decisions、clean release gate。"),
        ("5", "决定是否允许 A22 generated-artifact residual cleanup 进入 exact approval 流程；未批准前保持 0 executable cleanup rows。"),
    ]
    add_table(doc, ["序号", "决策事项"], decision_rows, [700, 8660], header_fill=PENDING_FILL, body_size=8.8)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
