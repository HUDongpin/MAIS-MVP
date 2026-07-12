# -*- coding: utf-8 -*-
"""Build the 2026-06-12 MAIS-MVP president report DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "coordination/reports/2026-06-12-president-report.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "555555"
HEADER_FILL = "F2F4F7"
LIGHT_FILL = "F7F9FC"
RISK_FILL = "FFF2CC"
GOOD_FILL = "E2F0D9"


def dxa(inches: float) -> int:
    return int(round(inches * 1440))


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Calibri"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 20, True, DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(subtitle)
    set_run_font(r, 10.5, False, MUTED)


def add_para(doc: Document, text: str = "", bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, 11, True)
        rest = text[len(bold_prefix) :]
        if rest:
            r = p.add_run(rest)
            set_run_font(r, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, 11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    width_dxa = [dxa(width) for width in widths]

    for row in table.rows:
        for cell, width_in, width_twips in zip(row.cells, widths, width_dxa):
            cell.width = Inches(width_in)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in width_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.2, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 9.2,
    status_fill_column: int | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], HEADER_FILL)
        set_cell_text(table.rows[0].cells[i], header, True, font_size, DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            if status_fill_column is not None and i == status_fill_column:
                if "通过" in text or "READY" in text or "完成" in text:
                    shade_cell(row.cells[i], GOOD_FILL)
                elif "失败" in text or "阻止" in text or "风险" in text:
                    shade_cell(row.cells[i], RISK_FILL)
            set_cell_text(row.cells[i], text, False, font_size)
    set_table_widths(table, widths)
    doc.add_paragraph()


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, val in rows:
        row = table.add_row()
        shade_cell(row.cells[0], LIGHT_FILL)
        set_cell_text(row.cells[0], key, True, 9.5, DARK_BLUE)
        set_cell_text(row.cells[1], val, False, 9.5)
    set_table_widths(table, [1.55, 4.95])
    doc.add_paragraph()


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("MAIS-MVP President Report | 2026-06-12")
    set_run_font(r, 8.5, False, MUTED)


def main() -> None:
    doc = Document()
    style_doc(doc)
    doc.core_properties.title = "MAIS-MVP 2026-06-12 President Report"
    doc.core_properties.subject = "Daily bilingual president report for Dr. Peter Hu"
    doc.core_properties.author = "S10 / Codex automation"

    add_title(
        doc,
        "MAIS-MVP 8 AM President Report",
        "Daily bilingual coordination report for Dr. Peter Hu",
    )
    add_kv_table(
        doc,
        [
            ("Report date", "2026-06-12"),
            ("Report time", "08:00 Asia/Hong_Kong; generated after fresh checks on 2026-06-12 afternoon HKT"),
            ("Project", "MAIS-MVP"),
            ("Reporting session", "S10 / Automation mais-mvp-9-am-president-report"),
            ("Audience", "Dr. Peter Hu"),
            ("Reporting window", "2026-06-11 08:00 to 2026-06-12 08:00 Asia/Hong_Kong"),
        ],
    )

    doc.add_heading("Chinese Executive Summary / 中文执行摘要", level=1)
    add_bullets(
        doc,
        [
            "本窗口不是空窗期：S01 完成首页与 About CTA checkmark 细节统一，S04 修复 Practice Arena 两个学生端问题，S05 修复 Lesson 星球入口重复/缓慢动画，S09 修正首页英文 headline，S13 将教师 mastery target 变成可交给 S15 的 adaptive signal，S22 完成两次生产部署。",
            "S22 在 2026-06-11 12:38 HKT 与 16:04 HKT 两次 owner-requested current-working-tree production publish 均达到 Vercel READY；`www.mais.hk`、`mais.hk`、login、student roadmap 等 smoke URL 返回 HTTP 200。",
            "今日 S10 复跑的本地基线为绿色：`npm run type-check` 通过，`npm run test:analytics` 23/23 通过，`npm run release:preflight -- --json` 通过，`npm run build` 通过并生成 128 个 static pages。",
            "最大发布风险仍是 dirty root：`release:root-deploy-preflight` 保护性失败，报告 3,435 status entries、260 tracked modified、4 tracked deleted、3,171 untracked status entries、32,509 untracked files。后续应继续使用 S22 pruned staging 或 clean reviewed release worktree。",
            "最需要 Dr. Peter Hu 关注的是：继续批准 Adventure UI 作为 Practice Arena 正式方向后的 S04/S11 回归节奏；是否允许继续 dirty-root emergency deploy；以及 S13 teacher mastery target 进入 S15 adaptive recommendation 的优先级。",
        ],
    )

    doc.add_heading("English Executive Summary", level=1)
    add_bullets(
        doc,
        [
            "The 2026-06-11 08:00 to 2026-06-12 08:00 HKT window had fresh work across S01, S04, S05, S09, S10, S13, and S22.",
            "Product work was narrow but useful: home/About CTA polish, Practice Arena collision and topic-selection fixes, faster Lesson planet transitions, homepage copy correction, and teacher mastery targets converted into an adaptive-signal handoff.",
            "S22 completed two owner-requested production publishes during the window. Vercel reported READY and user-facing smoke checks returned HTTP 200.",
            "Fresh local checks are green: type-check passed, analytics/adaptive tests passed 23/23, release preflight passed, and the Next production build passed with 128 static pages.",
            "The main management issue remains release control. Direct root deploy is still blocked by the dirty-root guard and should not be used unless Dr. Peter Hu explicitly approves an emergency exception.",
        ],
    )

    doc.add_heading("报告窗口摘要", level=1)
    add_table(
        doc,
        ["Item", "Conclusion"],
        [
            ["报告窗口", "2026-06-11 08:00 至 2026-06-12 08:00 Asia/Hong_Kong"],
            ["活跃 session logs", "S01、S04、S05、S09、S10、S13、S22"],
            ["新 blocker 文件", "未发现 fresh blocker file；风险主要来自各会话日志和 release guard"],
            ["新 decision / integration", "S10 记录 Practice Arena Adventure UI 正式方向，并写入 S04/S11 handoff"],
            ["生产部署", "S22 两次生产发布：12:38 HKT 与 16:04 HKT，均 READY，custom domains smoke 通过"],
            ["No assigned work in this reporting window", "No；本窗口有明确 fresh work"],
        ],
        [1.6, 4.9],
    )

    doc.add_heading("整体项目进度", level=1)
    add_bullets(
        doc,
        [
            "学生端 polish 继续推进：Home/About CTA 视觉统一，Practice Arena hero collision 消除，Topic Mission 点击现在会切换到对应 topic question path。",
            "Lesson 体验改善：跨星球进入不再叠加 source-page full-screen intro、route-level loading 和 destination intro；StrictMode replay 下也不会重新打开 Galaxy directory。",
            "教师端与 adaptive 的接口更清晰：S13 已把 manual mastery target 转成 UI-derived gap/priority signal，并给 S15 留出字段契约。",
            "生产发布能力仍可用，但当前靠 S22 pruned staging 控制风险；根目录本身仍不适合直接发布。",
            "构建结果提示性能关注点：student lesson/roadmap/visualization first-load JS 仍约 2.8 MB，需要后续 S22/S01/S03/S06 视发布目标评估分包或懒加载。",
        ],
    )

    doc.add_heading("S01-S25 session 状态表", level=1)
    session_rows = [
        ["S01", "完成", "Home 与 About CTA checkmark 改成圆角 SVG；桌面/移动截图确认 label fits", "无", "type-check passed；Playwright home/about desktop/mobile passed"],
        ["S02", "无新日志", "无 fresh dashboard evidence", "等待 dashboard/progress 任务", "未运行会话级检查"],
        ["S03", "无新日志", "无 fresh roadmap evidence", "等待 learning-path/roadmap 任务", "未运行会话级检查"],
        ["S04", "完成", "修复 Practice Arena title/star collision；Topic Mission 点击现在进入 topic-scoped free selection", "未新增 E2E；tests/e2e 仍属 S11", "type-check passed；Playwright/API repro passed"],
        ["S05", "完成", "Lesson planet entry 去重并加速；修复 StrictMode replay 导致 Galaxy reopen", "无", "type-check passed；4 条 Playwright route/direct/same-planet regression passed"],
        ["S06", "无新日志", "无 fresh visualization evidence", "等待 visualization lab 任务", "未运行会话级检查"],
        ["S07", "无新日志", "无 fresh AI Tutor evidence", "等待 tutor/provider 任务", "未运行会话级检查"],
        ["S08", "无新日志", "无 fresh shared-state evidence", "等待 analytics/shared types 任务", "未运行会话级检查"],
        ["S09", "完成", "Homepage headline 从 fine 改为 fun", "只做 targeted copy verification", "rg verification passed；未跑 broader tests"],
        ["S10", "完成", "完成 2026-06-11 report；记录 Practice Adventure UI direction 与 S04/S11 handoff；本次生成 2026-06-12 report", "无", "本次 fresh baseline checks completed"],
        ["S11", "无新日志", "无 fresh QA log", "需接手 Practice Arena Adventure UI targeted regression", "未运行 S11 full E2E"],
        ["S12", "无新日志", "无 fresh backend/API evidence", "Forum/teacher/adaptive 后续 API 若进入实现需协调", "未运行会话级检查"],
        ["S13", "完成", "Teacher mastery target 现在显示 bounded target、gap、priority、adaptive handoff attributes", "S15 仍需 server-side consumption", "type-check passed；desktop/mobile Playwright passed"],
        ["S14", "无新日志", "无 fresh parent-console evidence", "等待 parent workflow 任务", "未运行会话级检查"],
        ["S15", "无新日志", "无 fresh adaptive-engine evidence", "需决定是否优先消费 S13 teacher-priority signal", "未运行会话级检查"],
        ["S16", "无新日志", "无 fresh research evidence", "等待 learning-science/research 任务", "未运行会话级检查"],
        ["S17", "无新日志", "无 fresh gamification evidence", "若 Practice reward/streak 改动需协调 S04/S20", "未运行会话级检查"],
        ["S18", "无新日志", "无 fresh content QA evidence", "等待 content QA/signoff 任务", "未运行会话级检查"],
        ["S19", "无新日志", "无 fresh env/provider evidence", "等待 env parity/live provider smoke 任务", "未运行会话级检查"],
        ["S20", "无新日志", "无 fresh game-based learning evidence", "Practice game handoff 若改 semantics 需协调", "未运行会话级检查"],
        ["S21", "无新日志", "无 fresh content pipeline evidence", "等待 RAG/content package 任务", "未运行会话级检查"],
        ["S22", "完成", "两次 owner-requested production publish；custom domains HTTP 200", "dirty-root residual risk；未跑 broad production Playwright", "preflight/env/build/deploy/smoke passed；root guard remains blocked"],
        ["S23", "无新日志", "无 fresh integration evidence", "Practice Adventure handoff 已由 S10 写入，后续可接 release intake", "未运行会话级检查"],
        ["S24", "无新日志", "无 fresh exact-layer evidence", "等待 illustration exact-layer 任务", "未运行会话级检查"],
        ["S25", "无新日志", "无 fresh git hygiene evidence", "需要 dirty-tree ownership map/release slicing", "S10 only counted status；no Git mutation"],
    ]
    add_table(doc, ["Session", "Status", "Window result", "Risk / next", "Checks"], session_rows, [0.55, 0.8, 2.25, 1.55, 1.35], 8.0, 1)

    doc.add_heading("已完成工作", level=1)
    add_bullets(
        doc,
        [
            "S01：`components/home/PedaNovaHomeHero.tsx` 与 `components/home/HeroSection.tsx` 的 CTA checkmark 已统一为 round-cap/round-join SVG。",
            "S04：`components/practice/PracticeAdventureArenaShell.tsx` 和 preview hero 改为正常 flex row，移动端 title 与 star badge 不再碰撞。",
            "S04：`app/practice/page.tsx` 修复 topic card 只改视觉状态的问题；点击 Topic Mission 后请求 `/api/questions?grade=S3&topicId=quadratic-patterns...` 并显示 topic question `q5`。",
            "S05：`components/lesson/LessonView.tsx` 的 planet-entry flow 去除重复 intro 和 StrictMode replay 问题。",
            "S09：`components/home/PedaNovaHomeHero.tsx` 英文 headline 改为 `Math learning should be / fun and personalized!`。",
            "S13：`components/teacher/TeacherManagementViews.tsx` 新增 mastery target gap、priority tier、status 和 `data-adaptive-target-*` handoff attributes。",
            "S22：12:38 HKT 与 16:04 HKT 两次 production deploy 均完成，`www.mais.hk` 与 `mais.hk` smoke 通过。",
        ],
    )

    doc.add_heading("进行中工作", level=1)
    add_bullets(
        doc,
        [
            "Practice Arena：Adventure UI 已由 owner approved 为正式方向；下一步 S04 深化 lower controls/empty states，S11 做 desktop/mobile regression。",
            "Teacher adaptive signal：S13 完成 UI-derived handoff，S15 仍需决定并实现 server-side recommendation consumption。",
            "Release hygiene：S22 pruned staging 仍是可用生产路径；S25 需要把 dirty tree 切成可 review 的 release slices。",
            "Performance watch：build 中 student lesson/roadmap/visualization first-load bundle 接近 2.8 MB，后续应评估 code splitting。",
        ],
    )

    doc.add_heading("Blockers / 阻塞", level=1)
    add_table(
        doc,
        ["Priority", "Blocker", "Impact", "Owner / next action"],
        [
            ["P1", "Direct root deploy blocked by dirty worktree", "不能从当前根目录直接安全部署；需要 S22 pruned staging 或 clean release worktree", "Dr. Hu 确认继续禁用 direct root deploy；S22/S25 执行 release slicing"],
            ["P1", "Practice Arena no broad S11 regression after fresh fixes", "学生端核心练习页已有修复，但 release confidence 依赖 S11 targeted suite", "S11 跑 `/practice` desktop/mobile/topic/free-selection/game handoff checks"],
            ["P2", "Teacher mastery target not yet consumed by adaptive engine", "教师目标目前是 UI/handoff signal，不会自动改变 S15 recommendation", "Dr. Hu 决定优先级；S15 实现 bounded server-side consumption"],
            ["P2", "No broad production Playwright after S22 deployments", "生产 domain HTTP 200，但未覆盖 auth/storage/practice/lesson/teacher flows", "S11/S22 排 targeted production smoke"],
        ],
        [0.65, 2.0, 2.05, 1.8],
        8.6,
    )

    doc.add_heading("Risks / 风险", level=1)
    add_bullets(
        doc,
        [
            "Dirty root risk：当前 `release:root-deploy-preflight` 报 3,435 status entries、260 tracked modified、4 tracked deleted、3,171 untracked status entries、32,509 untracked files。",
            "Release drift risk：S22 本窗口生产发布来自 owner-requested dirty working tree，即使 staging 已排除 forbidden paths，也需要后续 clean review。",
            "Regression risk：S01/S04/S05/S13 都有 UI/flow 改动，虽然局部检查通过，但 full-product E2E 未在本窗口完成。",
            "Local artifact risk：`.local/hk-math-db.sqlite*` 与 `output/playwright/` 在窗口内有验证产物，应继续保持在非发布路径。",
            "Performance risk：Next build 显示部分 student surfaces first-load JS 较大，可能影响移动端首屏体验。",
        ],
    )

    doc.add_heading("Test / Build Status / 测试与构建状态", level=1)
    add_table(
        doc,
        ["Check", "Result", "Evidence"],
        [
            ["npm run type-check", "通过", "tsc --noEmit --incremental false passed"],
            ["npm run test:analytics", "通过", "23 tests passed；includes adaptiveLearning and learningAnalytics"],
            ["npm run release:preflight -- --json", "通过", "freeBytes 80,864,952,320；isolated `.tmp` E2E/staging roots configured"],
            ["npm run build", "通过", "Next.js 15.5.15 build passed；128 static pages generated"],
            ["npm run release:root-deploy-preflight -- --json", "保护性失败", "Dirty root blocked；3435 status entries；use S22 pruned staging or clean worktree"],
            ["S22 production smoke", "通过", "Two in-window deployments READY；custom domains and selected routes returned HTTP 200"],
            ["Full S11/S22 broad E2E", "未运行", "本报告未运行全产品 Playwright；建议作为明日优先项"],
        ],
        [2.0, 1.2, 3.3],
        8.0,
        1,
    )

    doc.add_heading("Changed Files / 变更文件", level=1)
    add_para(doc, "窗口内按 mtime 扫描到 44 个 touched files/artifacts；下面列出管理相关和产品源文件，省略部分截图细节。")
    add_table(
        doc,
        ["Category", "Files / directories", "Notes"],
        [
            ["Home", "`components/home/PedaNovaHomeHero.tsx`; `components/home/HeroSection.tsx`", "CTA checkmark 与 headline copy"],
            ["Practice", "`app/practice/page.tsx`; `app/practice/adventure-ui-preview/page.tsx`; `components/practice/PracticeAdventureArenaShell.tsx`", "hero collision 与 Topic Mission behavior"],
            ["Lesson", "`components/lesson/LessonView.tsx`", "planet-entry loading/StrictMode fix"],
            ["Teacher", "`components/teacher/TeacherManagementViews.tsx`", "mastery target adaptive-signal handoff"],
            ["Coordination", "`coordination/decisions/2026-06-11-practice-adventure-ui-direction.md`; `coordination/integration/2026-06-11-practice-adventure-ui-s04-s11-handoff.md`; session logs", "owner decision and S04/S11 handoff"],
            ["Reports", "`coordination/reports/2026-06-11-president-report.docx`; render PNG/PDF artifacts; design-memory README", "previous report and Practice Arena design memory updated inside current window"],
            ["Verification", "`output/playwright/...`; `.local/hk-math-db.sqlite*`", "browser evidence/local test DB artifacts; not feature source"],
            ["Current dirty tree", "3,435 status entries from root-deploy preflight", "Large pre-existing dirty tree remains outside the 44-file mtime window"],
        ],
        [1.0, 3.2, 2.3],
        8.4,
    )

    doc.add_heading("Tomorrow Priorities / 明日优先事项", level=1)
    add_bullets(
        doc,
        [
            "S11：跑 Practice Arena Adventure UI targeted regression，覆盖 mobile/desktop、topic selection、free-selection answer path、adaptive default path、game handoff。",
            "S22/S25：继续 dirty-tree release slicing；把 production deploy path 固定为 pruned staging 或 clean worktree，避免未经 owner 批准的 dirty-root override。",
            "S15/S13：决定 teacher mastery target adaptive signal 的消费范围、权重上限和 regression cases。",
            "S01/S09：快速确认 Home/About headline/CTA 在 production domain 上的视觉与 copy 一致。",
            "S05/S11：为 Lesson planet-entry fix 增加或记录 targeted E2E coverage，防止再次出现 double loader/Galaxy reopen。",
        ],
    )

    doc.add_heading("Owner Decisions Needed / 需要 Dr. Peter Hu 决策", level=1)
    add_table(
        doc,
        ["Decision", "Recommendation", "Reason"],
        [
            ["Practice Arena Adventure UI", "保持已批准方向，并授权 S04/S11 做下一轮 integration/regression", "S10 已记录 decision；S04 已修复关键 interaction bug"],
            ["Dirty-root production deployment", "只在明确 owner emergency approval 下允许 `MAIS_ALLOW_DIRTY_ROOT_DEPLOY=1`", "本窗口 deploy 成功，但 dirty-root guard 仍正确阻止 direct deploy"],
            ["Teacher target -> Adaptive engine", "决定是否列为 S15 明日/本周优先项", "S13 已交出可消费字段契约；现在需要 adaptive semantics"],
            ["Full production regression", "授权 S11/S22 跑 focused production smoke 或 full release matrix", "当前只有 curl smoke；未覆盖真实学习/教师/家长 flows"],
            ["Performance follow-up", "决定 student lesson/roadmap/visualization first-load bundle 是否进入 S22 performance task", "build 显示约 2.8 MB first-load JS，可能影响移动端体验"],
        ],
        [2.05, 2.05, 2.4],
        8.6,
    )

    add_footer(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
