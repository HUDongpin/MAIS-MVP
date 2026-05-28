from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from build_adaptive_learning_apa7_docx import (
    add_caption,
    add_formula,
    add_heading,
    add_note,
    add_page_number,
    add_paragraph,
    set_cell_margins,
    set_cell_shading,
    set_paragraph_spacing,
    set_run_font,
    set_table_borders,
    set_table_width,
    set_update_fields_on_open,
)


ROOT = Path("/Users/dongpinhu/Desktop/MAIS-MVP")
PAPER_DIR = Path("/Users/dongpinhu/Desktop/MAIS Research Proposals/Paper")
ZH_OUT = PAPER_DIR / "2026-05-17-MAIS-v2-v3-adaptive-learning-APA7-ZH.docx"
EN_OUT = PAPER_DIR / "2026-05-17-MAIS-v2-v3-adaptive-learning-APA7-EN.docx"


REFERENCES = [
    "Ainsworth, S. (2006). DeFT: A conceptual framework for considering learning with multiple representations. Learning and Instruction, 16(3), 183-198. https://doi.org/10.1016/j.learninstruc.2006.03.001",
    "Asai, A., Wu, Z., Wang, Y., Sil, A., & Hajishirzi, H. (2024). Self-RAG: Learning to retrieve, generate, and critique through self-reflection. In International Conference on Learning Representations. https://proceedings.iclr.cc/paper_files/paper/2024/hash/25f7be9694d7b32d5cc670927b8091e1-Abstract-Conference.html",
    "Athey, S., & Imbens, G. (2016). Recursive partitioning for heterogeneous causal effects. Proceedings of the National Academy of Sciences, 113(27), 7353-7360. https://doi.org/10.1073/pnas.1510489113",
    "Baker, R. S., & Hawn, A. (2022). Algorithmic bias in education. International Journal of Artificial Intelligence in Education, 32(4), 1052-1092. https://doi.org/10.1007/s40593-021-00285-9",
    "Bloom, B. S. (1984). The 2 sigma problem: The search for methods of group instruction as effective as one-to-one tutoring. Educational Researcher, 13(6), 4-16. https://doi.org/10.3102/0013189X013006004",
    "Brown, T. B., Mann, B., Ryder, N., Subbiah, M., Kaplan, J., Dhariwal, P., Neelakantan, A., Shyam, P., Sastry, G., Askell, A., Agarwal, S., Herbert-Voss, A., Krueger, G., Henighan, T., Child, R., Ramesh, A., Ziegler, D. M., Wu, J., Winter, C., ... Amodei, D. (2020). Language models are few-shot learners. In Advances in Neural Information Processing Systems, 33, 1877-1901. https://proceedings.neurips.cc/paper/2020/hash/1457c0d6bfcb4967418bfb8ac142f64a-Abstract.html",
    "Cepeda, N. J., Pashler, H., Vul, E., Wixted, J. T., & Rohrer, D. (2006). Distributed practice in verbal recall tasks: A review and quantitative synthesis. Psychological Bulletin, 132(3), 354-380. https://doi.org/10.1037/0033-2909.132.3.354",
    "Corbett, A. T., & Anderson, J. R. (1995). Knowledge tracing: Modeling the acquisition of procedural knowledge. User Modeling and User-Adapted Interaction, 4(4), 253-278. https://doi.org/10.1007/BF01099821",
    "Desmarais, M. C., & Baker, R. S. J. d. (2012). A review of recent advances in learner and skill modeling in intelligent learning environments. User Modeling and User-Adapted Interaction, 22(1-2), 9-38. https://doi.org/10.1007/s11257-011-9106-8",
    "Dwork, C. (2006). Differential privacy. In M. Bugliesi, B. Preneel, V. Sassone, & I. Wegener (Eds.), Automata, Languages and Programming (pp. 1-12). Springer. https://doi.org/10.1007/11787006_1",
    "Gervet, T., Koedinger, K., Schneider, J., & Mitchell, T. (2020). When is deep learning the best approach to knowledge tracing? Journal of Educational Data Mining, 12(3), 31-54. https://doi.org/10.5281/zenodo.4143614",
    "Hogan, A., Blomqvist, E., Cochez, M., d'Amato, C., de Melo, G., Gutierrez, C., Kirrane, S., Labra Gayo, J. E., Navigli, R., Neumaier, S., Ngomo, A.-C. N., Polleres, A., Rashid, S. M., Rula, A., Schmelzeisen, L., Sequeda, J., Staab, S., & Zimmermann, A. (2021). Knowledge graphs. ACM Computing Surveys, 54(4), Article 71. https://doi.org/10.1145/3447772",
    "Holstein, K., McLaren, B. M., & Aleven, V. (2019). Co-designing a real-time classroom orchestration tool to support teacher-AI complementarity. Journal of Learning Analytics, 6(2), 27-52. https://doi.org/10.18608/jla.2019.62.3",
    "Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., Xu, Y., Ishii, E., Bang, Y., Madotto, A., & Fung, P. (2023). Survey of hallucination in natural language generation. ACM Computing Surveys, 55(12), Article 248. https://doi.org/10.1145/3571730",
    "Kasneci, E., Sessler, K., Kuchemann, S., Bannert, M., Dementieva, D., Fischer, F., Gasser, U., Groh, G., Gunnemann, S., Hullermeier, E., Krusche, S., Kutyniok, G., Michaeli, T., Nerdel, C., Pfeffer, J., Poquet, O., Sailer, M., Schmidt, A., Seidel, T., ... Kasneci, G. (2023). ChatGPT for good? On opportunities and challenges of large language models for education. Learning and Individual Differences, 103, Article 102274. https://doi.org/10.1016/j.lindif.2023.102274",
    "Khajah, M. M., Lindsey, R. V., & Mozer, M. C. (2016). How deep is knowledge tracing? In Proceedings of the 9th International Conference on Educational Data Mining (pp. 94-101). https://educationaldatamining.org/EDM2016/proceedings/paper_122.pdf",
    "Koedinger, K. R., Corbett, A. T., & Perfetti, C. (2012). The knowledge-learning-instruction framework: Bridging the science-practice chasm to enhance robust student learning. Cognitive Science, 36(5), 757-798. https://doi.org/10.1111/j.1551-6709.2012.01245.x",
    "Kulik, J. A., & Fletcher, J. D. (2016). Effectiveness of intelligent tutoring systems: A meta-analytic review. Review of Educational Research, 86(1), 42-78. https://doi.org/10.3102/0034654315581420",
    "Lee, J., & Yeung, D.-Y. (2019). Knowledge query network for knowledge tracing: How knowledge interacts with skills. In Proceedings of the Ninth International Conference on Learning Analytics & Knowledge (pp. 491-500). https://doi.org/10.1145/3303772.3303786",
    "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Kuttler, H., Lewis, M., Yih, W.-t., Rocktaschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. In Advances in Neural Information Processing Systems, 33, 9459-9474. https://proceedings.neurips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html",
    "Ma, W., Adesope, O. O., Nesbit, J. C., & Liu, Q. (2014). Intelligent tutoring systems and learning outcomes: A meta-analysis. Journal of Educational Psychology, 106(4), 901-918. https://doi.org/10.1037/a0037123",
    "Macina, J., Daheim, N., Chowdhury, S., Sinha, T., Kapur, M., Gurevych, I., Sachan, M., & Nouri, E. (2023). MathDial: A dialogue tutoring dataset with rich pedagogical properties grounded in math reasoning problems. In Findings of the Association for Computational Linguistics: EMNLP 2023 (pp. 5602-5621). https://doi.org/10.18653/v1/2023.findings-emnlp.372",
    "Maynez, J., Narayan, S., Bohnet, B., & McDonald, R. (2020). On faithfulness and factuality in abstractive summarization. In Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics (pp. 1906-1919). https://doi.org/10.18653/v1/2020.acl-main.173",
    "Ouyang, L., Wu, J., Jiang, X., Almeida, D., Wainwright, C. L., Mishkin, P., Zhang, C., Agarwal, S., Slama, K., Ray, A., Schulman, J., Hilton, J., Kelton, F., Miller, L., Simens, M., Askell, A., Welinder, P., Christiano, P., Leike, J., & Lowe, R. (2022). Training language models to follow instructions with human feedback. In Advances in Neural Information Processing Systems, 35, 27730-27744. https://proceedings.neurips.cc/paper_files/paper/2022/hash/b1efde53be364a73914f58805a001731-Abstract-Conference.html",
    "Pardos, Z. A., & Heffernan, N. T. (2010). Modeling individualization in a Bayesian networks implementation of knowledge tracing. In P. De Bra, A. Kobsa, & D. Chin (Eds.), User Modeling, Adaptation, and Personalization (pp. 255-266). Springer. https://doi.org/10.1007/978-3-642-13470-8_24",
    "Pardos, Z. A., & Heffernan, N. T. (2011). KT-IDEM: Introducing item difficulty to the knowledge tracing model. In G. Biswas, S. Bull, J. Kay, & A. Mitrovic (Eds.), Artificial Intelligence in Education (pp. 243-254). Springer. https://doi.org/10.1007/978-3-642-22362-4_21",
    "Pavlik, P. I., Jr., Cen, H., & Koedinger, K. R. (2009). Performance factors analysis: A new alternative to knowledge tracing. In V. Dimitrova, R. Mizoguchi, B. du Boulay, & A. Graesser (Eds.), Artificial Intelligence in Education (pp. 531-538). IOS Press. https://doi.org/10.3233/978-1-60750-028-5-531",
    "Pelanek, R. (2017). Bayesian knowledge tracing, logistic models, and beyond: An overview of learner modeling techniques. User Modeling and User-Adapted Interaction, 27(3-5), 313-350. https://doi.org/10.1007/s11257-017-9193-2",
    "Piech, C., Spencer, J., Huang, J., Ganguli, S., Sahami, M., Guibas, L. J., & Sohl-Dickstein, J. (2015). Deep knowledge tracing. In Advances in Neural Information Processing Systems, 28. https://papers.nips.cc/paper/5654-deep-knowledge-tracing",
    "Ritter, S., Yudelson, M., Fancsali, S. E., & Berman, S. R. (2016). How mastery learning works at scale. In Proceedings of the Third (2016) ACM Conference on Learning @ Scale (pp. 71-79). https://doi.org/10.1145/2876034.2876039",
    "Roediger, H. L., III, & Karpicke, J. D. (2006). Test-enhanced learning: Taking memory tests improves long-term retention. Psychological Science, 17(3), 249-255. https://doi.org/10.1111/j.1467-9280.2006.01693.x",
    "Shute, V. J. (2008). Focus on formative feedback. Review of Educational Research, 78(1), 153-189. https://doi.org/10.3102/0034654307313795",
    "Sweller, J., van Merrienboer, J. J. G., & Paas, F. (2019). Cognitive architecture and instructional design: 20 years later. Educational Psychology Review, 31(2), 261-292. https://doi.org/10.1007/s10648-019-09465-5",
    "Tack, A., & Piech, C. (2022). The AI teacher test: Measuring the pedagogical ability of Blender and GPT-3 in educational dialogues. In Proceedings of the 15th International Conference on Educational Data Mining. https://educationaldatamining.org/edm2022/proceedings/2022.EDM-short-papers.54/index.html",
    "VanLehn, K. (2011). The relative effectiveness of human tutoring, intelligent tutoring systems, and other tutoring systems. Educational Psychologist, 46(4), 197-221. https://doi.org/10.1080/00461520.2011.611369",
    "Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., Gomez, A. N., Kaiser, L., & Polosukhin, I. (2017). Attention is all you need. In Advances in Neural Information Processing Systems, 30. https://papers.nips.cc/paper/7181-attention-is-all-you-need",
    "Wager, S., & Athey, S. (2018). Estimation and inference of heterogeneous treatment effects using random forests. Journal of the American Statistical Association, 113(523), 1228-1242. https://doi.org/10.1080/01621459.2017.1319839",
    "Weidinger, L., Uesato, J., Rauh, M., Griffin, C., Huang, P.-S., Mellor, J., Glaese, A., Cheng, M., Balle, B., Kasirzadeh, A., Biles, C., Brown, S., Kenton, Z., Hawkins, W., Stepleton, T., Birhane, A., Haas, J., Rimell, L., Hendricks, L. A., ... Gabriel, I. (2022). Taxonomy of risks posed by language models. In 2022 ACM Conference on Fairness, Accountability, and Transparency (pp. 214-229). https://doi.org/10.1145/3531146.3533088",
    "Wei, J., Wang, X., Schuurmans, D., Bosma, M., Xia, F., Chi, E. H., Le, Q. V., Zhou, D., & Narang, S. (2022). Chain-of-thought prompting elicits reasoning in large language models. In Advances in Neural Information Processing Systems, 35, 24824-24837. https://proceedings.neurips.cc/paper_files/paper/2022/hash/9d5609613524ecf4f15af0f7b31abca4-Abstract-Conference.html",
    "Wood, D., Bruner, J. S., & Ross, G. (1976). The role of tutoring in problem solving. Journal of Child Psychology and Psychiatry, 17(2), 89-100. https://doi.org/10.1111/j.1469-7610.1976.tb00381.x",
    "Yan, L., Sha, L., Zhao, L., Li, Y., Martinez-Maldonado, R., Chen, G., Li, X., Jin, Y., & Gasevic, D. (2024). Practical and ethical challenges of large language models in education: A systematic scoping review. British Journal of Educational Technology, 55(1), 90-112. https://doi.org/10.1111/bjet.13370",
    "Yang, Q., Liu, Y., Chen, T., & Tong, Y. (2019). Federated machine learning: Concept and applications. ACM Transactions on Intelligent Systems and Technology, 10(2), Article 12. https://doi.org/10.1145/3298981",
    "Yudelson, M. V., Koedinger, K. R., & Gordon, G. J. (2013). Individualized Bayesian knowledge tracing models. In H. C. Lane, K. Yacef, J. Mostow, & P. Pavlik (Eds.), Artificial Intelligence in Education (pp. 171-180). Springer. https://doi.org/10.1007/978-3-642-39112-5_18",
    "Zhang, J., Shi, X., King, I., & Yeung, D.-Y. (2017). Dynamic key-value memory networks for knowledge tracing. In Proceedings of the 26th International Conference on World Wide Web (pp. 765-774). https://doi.org/10.1145/3038912.3052580",
    "Zimmerman, B. J. (2002). Becoming a self-regulated learner: An overview. Theory Into Practice, 41(2), 64-70. https://doi.org/10.1207/s15430421tip4102_2",
]


FORMULAS = [
    (
        r"C_t = \{c_1, c_2, \ldots, c_k\}, \quad c_i = (a_i, k_i, q_i, g_i, e_i)",
        "1",
    ),
    (
        r"\begin{aligned} d_t ={}& \operatorname*{arg\,max}_{c_i \in C_t} \operatorname{baseScore}(c_i \mid s_t) \\ &\text{subject to } \operatorname{hardGuard}(c_i)=1 \end{aligned}",
        "2",
    ),
    (
        r"r_t = \operatorname*{arg\,max}_{c_i \in C_t} s_{\phi}(c_i \mid x_t, C_t, \pi, e_t)",
        "3",
    ),
    (
        r"\begin{aligned} V(r_t) ={}& \mathbb{1}[\operatorname{candidateId}(r_t) \in C_t] \cdot \mathbb{1}[\operatorname{questionIds}(r_t) \subseteq Q(c_i)] \\ &\cdot \mathbb{1}[\operatorname{guardrails}(r_t)=\operatorname{satisfied}] \cdot \mathbb{1}[\operatorname{bilingualAudit}(r_t)=1] \end{aligned}",
        "4",
    ),
    (
        r"D_t = \begin{cases} r_t, & \text{if providerReady} \land \text{validJSON} \land V(r_t)=1 \\ d_t, & \text{otherwise} \end{cases}",
        "5",
    ),
    (
        r"\operatorname{conf}_{AI} = \min\{1,\;0.36E + 0.30G_a + 0.22M + 0.12Q\}",
        "6",
    ),
    (
        r"P(L_{t+1}=1) = P(L_t=1 \mid X_t) + [1-P(L_t=1 \mid X_t)]T",
        "7",
    ),
    (
        r"\begin{aligned} G_E=(&V_{KC}\cup V_{lesson}\cup V_{item}\cup V_{mis}\cup V_{visual},\\ &E_{pre}\cup E_{assess}\cup E_{remediate}\cup E_{visualize}) \end{aligned}",
        "8",
    ),
    (
        r"R_{pre}(k)=\prod_{p\in Pre(k)} \mathbb{1}[P(L_p=1)\ge \tau_p]",
        "9",
    ),
    (
        r"\begin{aligned} S_{evid}(c)={}&\frac{1}{|\operatorname{Claims}(c)|}\sum_{q\in \operatorname{Claims}(c)} \\ &\max_{e\in E_c}\operatorname{sim}(q,e) \end{aligned}",
        "10",
    ),
    (
        r"\begin{aligned} c_t^{*}={}&\operatorname*{arg\,max}_{c_i\in C_t} U(c_i\mid s_t,G_E,E_c) \\ &\text{subject to } R_{pre}(k_i)=1,\;S_{evid}(c_i)\ge\tau_e,\;V(c_i)=1 \end{aligned}",
        "11",
    ),
    (
        r"U(c)=w_m\Delta M(c)+w_r Ret(c)+w_f Fair(c)-w_l Load(c)-w_{\rho}Risk(c)-w_{\kappa}Cost(c)",
        "12",
    ),
    (
        r"D_t^{v3}=\begin{cases} c_t^{*}, & \text{if evidenceSupported}\land\text{validated}\land\text{teacherPolicyAllowed}\\ d_t, & \text{otherwise} \end{cases}",
        "13",
    ),
    (
        r"\tau(x)=\mathbb{E}[Y(1)-Y(0)\mid X=x]",
        "14",
    ),
    (
        r"\begin{aligned} \min_{\theta}\sum_{m=1}^{M}\frac{n_m}{n}\mathcal{L}_m(\theta) \quad \text{subject to}\\ \text{privacy, fairness, latency, and audit constraints} \end{aligned}",
        "15",
    ),
]


ZH = {
    "lang": "zh",
    "out": ZH_OUT,
    "running": "MAIS v2/v3 自适应学习对比报告",
    "title": "MAIS v2/v3 自适应学习对比报告",
    "subtitle": "从受防护 LLM 重排到可验证教学决策系统",
    "meta": [
        ("项目", "MAIS：面向香港 P1-S6 学生的双语数学学习平台"),
        ("报告日期", "2026-05-17"),
        ("报告类型", "学术式技术与教育分析报告"),
        ("引用格式", "APA 7th author-date；References 仅收录学术期刊或学术会议论文"),
        ("公式格式", "所有显示公式均为标准 LaTeX 文本块，不使用 Word 原生 equation 对象"),
    ],
    "core_title": "核心判断",
    "core_text": "v2 是 BKT guardrails 之上的受防护 LLM-assisted reranking；v3 则把 BKT、教育知识图谱、RAG、验证服务、审计日志和教师在环控制整合为可验证教学决策架构。两者不是简单强弱之分，而是从“候选内重排”走向“证据支持、可审计、可回滚的教学决策系统”。",
    "abstract_title": "摘要",
    "abstract": [
        "本报告比较 MAIS v2 与拟议 v3 自适应学习架构。v2 保留 Bayesian Knowledge Tracing（BKT）和确定性规则作为候选生成器，只允许 LLM 在候选集合内重排、生成学习理由和教师审计备注；当 provider 失败、JSON 无效或校验器拒绝输出时，系统回退到确定性 v1/v2 基座。v3 则进一步把 BKT safety floor、教育知识图谱、RAG 证据 grounding、validation services、audit logging 和 teacher-in-the-loop 融合为一套可验证教学决策系统。",
        "文献综述显示，BKT 与知识追踪研究为可解释学生建模提供了稳定基础，ITS meta-analysis 支持自适应辅导系统的学习成效；LLM 与 instruction tuning 研究解释了语言模型为何适合候选解释、反馈生成和排序，但教育综述、幻觉研究和语言模型风险分类也说明 LLM 不能被设定为未受约束的教学权威。v3 的教育价值来自把 mastery learning、cognitive load、formative feedback、spacing/retrieval practice、self-regulated learning、multiple representations 与 teacher-AI complementarity 转化为可审计的系统约束。",
        "本报告的结论是：v2 适合成为高级 AI-assisted 层，提升个性化解释、双语反馈和教师审计；v3 则应作为下一阶段核心架构，把推荐从“选下一步”升级为“在证据、先备、风险、成本和教师治理约束下做可验证教学决策”。",
    ],
    "toc": "目录",
    "toc_note": "注：如目录页码未自动显示，请在 Word 或 LibreOffice 中更新域。正文结构已使用真实 Heading 样式。",
}


EN = {
    "lang": "en",
    "out": EN_OUT,
    "running": "MAIS v2/v3 Adaptive Learning Comparative Report",
    "title": "MAIS v2/v3 Adaptive Learning Comparative Report",
    "subtitle": "From Guarded LLM Reranking to Verifiable Instructional Decision-Making",
    "meta": [
        ("Project", "MAIS: A bilingual mathematics learning platform for Hong Kong P1-S6 students"),
        ("Report date", "2026-05-17"),
        ("Report type", "Academic technical and educational analysis report"),
        ("Citation style", "APA 7th author-date; References include only academic journals or conference papers"),
        ("Formula format", "All display formulas are standard LaTeX text blocks, not Word native equation objects"),
    ],
    "core_title": "Core Judgment",
    "core_text": "v2 is guarded LLM-assisted reranking on top of BKT guardrails. v3 turns BKT, an educational knowledge graph, RAG, validation services, audit logging, and teacher-in-the-loop control into a verifiable instructional decision architecture. The key distinction is not model size; it is the move from candidate reranking to evidence-supported, auditable, and rollback-capable instructional decision-making.",
    "abstract_title": "Abstract",
    "abstract": [
        "This report compares MAIS v2 with the proposed v3 adaptive learning architecture. v2 preserves Bayesian Knowledge Tracing (BKT) and deterministic rules as the candidate generator, while allowing the LLM only to rerank candidates, produce learner-facing rationales, and write teacher audit notes. If the provider fails, JSON is invalid, or the validator rejects an output, the system falls back to the deterministic baseline. v3 extends this architecture by combining the BKT safety floor, an educational knowledge graph, retrieval-augmented generation (RAG) evidence grounding, validation services, audit logging, and teacher-in-the-loop control.",
        "The literature indicates that BKT and knowledge tracing provide an interpretable foundation for learner modeling, and intelligent tutoring system meta-analyses support the learning value of adaptive tutoring. Research on Transformers, large language models, and instruction tuning explains why LLMs can support explanation, feedback, and candidate ranking, but education reviews, hallucination research, and language-model risk taxonomies show that LLMs should not become unrestricted instructional authorities. v3 is justified by translating mastery learning, cognitive load theory, formative feedback, spacing and retrieval practice, self-regulated learning, multiple representations, and teacher-AI complementarity into auditable system constraints.",
        "The central conclusion is that v2 should be positioned as a premium AI-assisted layer for richer explanations, bilingual feedback, and teacher auditability. v3 should become the next core architecture, upgrading the system from recommending a next step to making verifiable instructional decisions under evidence, prerequisite, risk, cost, and teacher-governance constraints.",
    ],
    "toc": "Table of Contents",
    "toc_note": "Note. If page numbers do not update automatically, update fields in Microsoft Word or LibreOffice. The body uses real Word Heading styles.",
}


def add_toc(paragraph, lang: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Microsoft Word 或 LibreOffice 中自动更新。" if lang == "zh" else "The table of contents will update automatically in Microsoft Word or LibreOffice."
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run._r.append(placeholder)
    run._r.append(fld_char3)
    set_run_font(run, size=11)


def add_bullet(doc: Document, text: str, lang: str = "zh") -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    set_paragraph_spacing(p, after=4, line=1.2)
    r = p.add_run(text)
    set_run_font(r, size=10.5, east_asia="Songti SC" if lang == "zh" else "Times New Roman")


def add_numbered(doc: Document, text: str, lang: str = "zh") -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.38)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_paragraph_spacing(p, after=4, line=1.2)
    r = p.add_run(text)
    set_run_font(r, size=10.5, east_asia="Songti SC" if lang == "zh" else "Times New Roman")


def add_paras(doc: Document, paras: list[str]) -> None:
    for text in paras:
        add_paragraph(doc, text, line=1.28)


def add_title_page(doc: Document, cfg: dict) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=36, after=10, line=1.15)
    r = p.add_run(cfg["title"])
    set_run_font(r, size=22 if cfg["lang"] == "zh" else 21, bold=True, color="0B2545", east_asia="Heiti SC" if cfg["lang"] == "zh" else "Times New Roman")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=24, line=1.2)
    r = p.add_run(cfg["subtitle"])
    set_run_font(r, size=14, italic=True, color="1F4D78", east_asia="Heiti SC" if cfg["lang"] == "zh" else "Times New Roman")

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in cfg["meta"]:
        row = table.add_row()
        for i, value in enumerate((k, v)):
            cell = row.cells[i]
            set_cell_margins(cell, top=110, bottom=110, start=140, end=140)
            if i == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=10.4, bold=(i == 0), color="0B2545" if i == 0 else "000000", east_asia="Heiti SC" if i == 0 and cfg["lang"] == "zh" else "Songti SC")
    set_table_borders(table)
    set_table_width(table, [4.1, 12.25])

    add_note(doc, cfg["core_title"], cfg["core_text"])


def add_reference_style(doc: Document) -> None:
    styles = doc.styles
    if "APA Reference" not in styles:
        ref_style = styles.add_style("APA Reference", 1)
    else:
        ref_style = styles["APA Reference"]
    ref_style.base_style = styles["Normal"]
    ref_style.font.size = Pt(9.4)
    ref_style.paragraph_format.left_indent = Inches(0.5)
    ref_style.paragraph_format.first_line_indent = Inches(-0.5)
    ref_style.paragraph_format.space_after = Pt(5)
    ref_style.paragraph_format.line_spacing = 1.12


def setup_doc(cfg: dict) -> Document:
    doc = Document()
    set_update_fields_on_open(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.color.rgb = RGBColor.from_string("1F4D78")
    add_reference_style(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = header.add_run(cfg["running"])
    set_run_font(hrun, size=9, color="606A78", east_asia="Heiti SC" if cfg["lang"] == "zh" else "Times New Roman")
    add_page_number(section.footer.paragraphs[0])
    return doc


def add_formula_by_id(doc: Document, idx: int) -> None:
    formula, number = FORMULAS[idx - 1]
    add_formula(doc, formula, number)


def add_v2_candidate_table(doc: Document, lang: str) -> None:
    if lang == "zh":
        add_caption(doc, "表 1\nv2 受防护 LLM 自适应流程")
        headers = ["环节", "输入", "允许的 LLM 行为", "硬约束与回退"]
        rows = [
            ("候选生成", "BKT 状态、先备关系、review/repair 规则、题库映射", "无；LLM 不参与候选生成", "候选必须来自确定性引擎，不能凭空生成题目或跳过先备"),
            ("候选重排", "BKT 候选集、近期表现、错题摘要、候选 evidence", "只在候选内排序，生成 learnerReason 与 teacherAuditNote", "candidateId 必须存在，questionIds 必须属于被选候选"),
            ("验证服务", "LLM JSON、候选元数据、guard flags", "无；本地 validator 做确定性检查", "格式错误、越权、缺字段或 guard 失败即拒绝"),
            ("最终决策", "validator 结果、provider 状态、fallback 策略", "无；系统选择 LLM 结果或确定性回退", "失败时返回 d_t，保证学习路径不中断"),
        ]
    else:
        add_caption(doc, "Table 1\nGuarded LLM adaptation flow in v2")
        headers = ["Stage", "Inputs", "Allowed LLM behavior", "Hard constraints and fallback"]
        rows = [
            ("Candidate generation", "BKT state, prerequisites, review/repair rules, item mapping", "None; the LLM does not generate candidates", "Candidates must come from the deterministic engine; no invented items or skipped prerequisites"),
            ("Candidate reranking", "BKT candidate set, recent performance, mistake summary, candidate evidence", "Rank only within candidates; generate learnerReason and teacherAuditNote", "candidateId must exist; questionIds must belong to the selected candidate"),
            ("Validation service", "LLM JSON, candidate metadata, guard flags", "None; the local validator performs deterministic checks", "Malformed output, boundary violation, missing fields, or guard failure triggers rejection"),
            ("Final decision", "Validator result, provider status, fallback policy", "None; the system chooses the LLM output or fallback", "Failure returns d_t so the learning path continues"),
        ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell, top=95, bottom=95, start=120, end=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9.0, bold=True, color="0B2545", east_asia="Heiti SC" if lang == "zh" else "Times New Roman")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i], top=90, bottom=90, start=120, end=120)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=8.8, bold=(i == 0), east_asia="Songti SC" if lang == "zh" else "Times New Roman")
    set_table_borders(table)
    set_table_width(table, [3.0, 4.8, 4.4, 4.3])
    add_paragraph(doc, "", after=4)


def add_v3_architecture_table(doc: Document, lang: str) -> None:
    if lang == "zh":
        add_caption(doc, "表 2\nv3 自适应技术模块与教育功能")
        headers = ["模块", "功能", "主要证据来源", "教育价值"]
        rows = [
            ("BKT safety floor", "更新 KC 掌握概率，控制 review、repair、practice、lesson、challenge 边界", "作答结果、错题、attempt、streak、review due", "保持可解释、低风险、可回退的学生模型"),
            ("教育知识图谱", "连接 KC、课程目标、题目、lesson、可视化、错因和先备关系", "课程目标、题库元数据、教师标注、错因标签", "避免跳过先备，支持跨模块补救与可视化推荐"),
            ("RAG evidence grounding", "为候选解释检索证据包，限制 unsupported claims", "lesson 片段、题目映射、知识图谱边、教师备注", "降低幻觉，使学生和教师理由可核验"),
            ("Validation services", "检查 candidateId、questionIds、guard flags、证据支持和双语字段", "候选集、policy、schema、audit rules", "把 LLM 输出放进可审计安全边界"),
            ("Teacher-in-the-loop", "教师确认、覆盖、批量布置、标注理由并写回评估", "教师操作日志、override 原因、班级层数据", "形成 teacher-AI complementarity，而不是替代教师"),
        ]
    else:
        add_caption(doc, "Table 2\nv3 adaptive technology modules and educational functions")
        headers = ["Module", "Function", "Primary evidence sources", "Educational value"]
        rows = [
            ("BKT safety floor", "Updates KC mastery and controls review, repair, practice, lesson, and challenge boundaries", "Responses, mistakes, attempts, streaks, review due dates", "Maintains an interpretable, low-risk, fallback-ready learner model"),
            ("Educational knowledge graph", "Connects KCs, curriculum objectives, items, lessons, visualizations, misconceptions, and prerequisites", "Curriculum goals, item metadata, teacher labels, misconception tags", "Prevents prerequisite skips and supports cross-module repair and visualization choices"),
            ("RAG evidence grounding", "Retrieves evidence packs for candidate explanations and limits unsupported claims", "Lesson snippets, item mappings, graph edges, teacher notes", "Reduces hallucination and makes rationales verifiable"),
            ("Validation services", "Checks candidateId, questionIds, guard flags, evidence support, and bilingual fields", "Candidate set, policy, schema, audit rules", "Keeps LLM output inside auditable safety boundaries"),
            ("Teacher-in-the-loop", "Enables teacher approval, override, batch assignment, rationale labels, and evaluation feedback", "Teacher action logs, override reasons, class-level data", "Creates teacher-AI complementarity rather than teacher replacement"),
        ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell, top=95, bottom=95, start=120, end=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=8.9, bold=True, color="0B2545", east_asia="Heiti SC" if lang == "zh" else "Times New Roman")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i], top=90, bottom=90, start=120, end=120)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=8.45, bold=(i == 0), east_asia="Songti SC" if lang == "zh" else "Times New Roman")
    set_table_borders(table)
    set_table_width(table, [3.0, 5.0, 4.45, 4.05])
    add_paragraph(doc, "", after=4)


def add_comparison_table(doc: Document, lang: str) -> None:
    if lang == "zh":
        add_caption(doc, "表 3\nv2 与 v3 自适应学习架构核心比较")
        headers = ["维度", "v2：受防护 LLM 重排", "v3：可验证教学决策系统"]
        rows = [
            ("决策权", "BKT 生成候选；LLM 在候选内重排；validator 通过才采用", "BKT 候选、KG 先备、RAG 证据、validator 和教师策略共同约束"),
            ("学生模型", "BKT mastery、streak、review due、错题摘要", "BKT safety floor + 知识图谱状态 + retrieval evidence + 可选 KT 特征"),
            ("LLM 角色", "解释、审计备注、候选排序；不直接生成教学边界", "证据内解释、候选比较、风险摘要；仍不能越过 hard guards"),
            ("RAG 角色", "可选或轻量证据提示", "核心证据层，限制 unsupported claims 并支持引用追踪"),
            ("验证/回退", "schema + candidate validation + fallback to d_t", "多门验证：先备、证据、风险、成本、公平性、教师策略；可回滚"),
            ("教师控制", "查看 teacherAuditNote 和 signalsUsed", "审批、覆盖、批量布置、override learning、班级治理 dashboard"),
            ("软件复杂度", "中等：provider、prompt、cache、rate limit、JSON validator", "高：event ingestion、KG/RAG、validation service、audit console、experiment tracking、monitoring"),
            ("成本与风险", "token 成本、延迟、幻觉、隐私和服务可用性", "更高工程成本，但通过 evidence grounding 和治理降低教育风险"),
            ("教育价值", "更自然的个性化反馈和双语解释", "可验证的 mastery repair、spaced review、visual module selection 和教师批准推荐"),
        ]
    else:
        add_caption(doc, "Table 3\nCore comparison of v2 and v3 adaptive learning architectures")
        headers = ["Dimension", "v2: Guarded LLM reranking", "v3: Verifiable instructional decision system"]
        rows = [
            ("Decision authority", "BKT generates candidates; the LLM reranks within candidates; validator approval is required", "BKT candidates, KG prerequisites, RAG evidence, validators, and teacher policy jointly constrain decisions"),
            ("Learner model", "BKT mastery, streaks, review due dates, mistake summaries", "BKT safety floor plus graph state, retrieval evidence, and optional KT features"),
            ("LLM role", "Explanation, audit notes, and candidate ranking; no direct instructional boundary setting", "Evidence-bound explanation, candidate comparison, and risk summary; still cannot bypass hard guards"),
            ("RAG role", "Optional or lightweight evidence prompt", "Core evidence layer that limits unsupported claims and supports citation traceability"),
            ("Validation/fallback", "Schema and candidate validation plus fallback to d_t", "Multiple gates: prerequisite, evidence, risk, cost, fairness, teacher policy, and rollback"),
            ("Teacher control", "Inspection of teacherAuditNote and signalsUsed", "Approval, override, batch assignment, override learning, and class-level governance dashboard"),
            ("Software complexity", "Moderate: provider, prompt, cache, rate limit, JSON validator", "High: event ingestion, KG/RAG, validation service, audit console, experiment tracking, monitoring"),
            ("Cost and risk", "Token cost, latency, hallucination, privacy, and service availability", "Higher engineering cost, but evidence grounding and governance lower educational risk"),
            ("Educational value", "More natural personalized feedback and bilingual explanation", "Verifiable mastery repair, spaced review, visual-module selection, and teacher-approved recommendations"),
        ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell, top=95, bottom=95, start=120, end=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9.0, bold=True, color="0B2545", east_asia="Heiti SC" if lang == "zh" else "Times New Roman")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i], top=86, bottom=86, start=120, end=120)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=8.55, bold=(i == 0), east_asia="Songti SC" if lang == "zh" else "Times New Roman")
    set_table_borders(table)
    set_table_width(table, [3.0, 6.65, 6.85])
    note = "注：v2 与 v3 的差异主要在决策架构和证据治理，不在是否调用更大的 LLM。" if lang == "zh" else "Note. The v2/v3 distinction is primarily about decision architecture and evidence governance, not merely about using a larger LLM."
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=8, line=1.12)
    r = p.add_run(note)
    set_run_font(r, size=9.2, italic=True, color="555555", east_asia="Songti SC" if lang == "zh" else "Times New Roman")


def add_zh_body(doc: Document) -> None:
    add_heading(doc, "Section 1. v2：有 LLM 自适应", level=1)
    add_heading(doc, "1.1 概念介绍：BKT 候选集之上的 guarded LLM reranking", level=2)
    add_paras(
        doc,
        [
            "在 MAIS 当前 v2 定义中，LLM 自适应不是让语言模型自由决定课程路径，而是把 LLM 放在 BKT-generated candidate set 之后，作为受限制的候选重排器、解释生成器和教师审计助手。BKT 与规则引擎先依据掌握概率、先备知识、错题记录、review due 和题库映射生成候选动作；LLM 只能在这些候选内部排序，不能生成新题、跳过先备知识、改变掌握概率或取消 due-review/repair guard。这一结构保留了 BKT 的可解释安全边界，同时利用 LLM 的自然语言综合能力。",
            "v2 的工程流可以概括为四步：第一，BKT/规则引擎生成候选；第二，LLM 接收候选摘要、近期表现和 policy 说明后输出排序、learnerReason、teacherAuditNote 和 signalsUsed；第三，本地 validator 对 candidateId、questionIds、schema、guard flags 和双语字段进行确定性校验；第四，若校验失败或 provider 不可用，系统返回确定性决策 d_t。这个流程体现了教育 AI 的基本原则：LLM 可以辅助解释和排序，但不能成为不受约束的教学权威。",
            "LLM 的技术基础来自 Transformer attention 架构、大规模预训练、few-shot learning、instruction tuning/RLHF 和 chain-of-thought prompting（Brown et al., 2020; Ouyang et al., 2022; Vaswani et al., 2017; Wei et al., 2022）。这些研究说明，大模型可以整合上下文、生成可读解释并对复杂输入做条件化推断。然而，教育场景要求的不是语言流畅，而是可证据化、可审计、可回退的教学合理性。",
        ],
    )
    add_v2_candidate_table(doc, "zh")

    add_heading(doc, "1.2 文献综述：LLM 教育应用、AI tutor 评估与风险治理", level=2)
    add_paras(
        doc,
        [
            "LLM 教育研究已经形成清晰共识：大模型在个性化反馈、学习解释、教师备课、对话辅导和内容改写方面具有潜力，但其教学使用必须受到准确性、隐私、公平性和教师角色约束。Kasneci 等（2023）从学习支持和挑战两个方面讨论了 LLM 的教育机会；Yan 等（2024）系统综述了 LLM 教育应用中的实践与伦理难题。这些综述支持 v2 的定位：LLM 应当增强学习反馈和教师可见性，而不是直接替代学生模型或教师判断。",
            "对 AI tutor 的评估不能只看语言是否自然。Tack 和 Piech（2022）提出 AI Teacher Test，强调模型是否能表现出合适的教学行为；Macina 等（2023）的 MathDial 数据集将数学推理问题与具有教学属性的对话相结合，说明数学 tutor 需要面向解释、追问、提示和错误修复进行评价。对 MAIS v2 来说，LLM 输出应当被评估为“是否在候选约束内产生教育上合适的解释与排序”，而不是“是否生成了漂亮文字”。",
            "风险研究则解释了为什么 v2 必须采用 validation 和 fallback。Ji 等（2023）指出自然语言生成系统存在 hallucination；Maynez 等（2020）从摘要生成角度区分 factuality 与 faithfulness；Weidinger 等（2022）建立了语言模型风险分类，覆盖偏差、误用、隐私和社会影响。教育领域还要特别注意算法偏差和群体影响，Baker 和 Hawn（2022）对教育算法偏差的综述说明，模型与推荐系统可能在不同学生群体之间产生系统性差异。因此，v2 的 LLM confidence 不应来自模型自报，而应来自本地 evidence、guardrail consistency 和 validation log。",
            "v2 同时继承智能导师系统与知识追踪传统。Corbett 和 Anderson（1995）提出 BKT，用概率方式追踪程序性知识习得；Pardos 和 Heffernan（2010）与 Yudelson 等（2013）把 BKT 推向个性化参数估计；Pardos 和 Heffernan（2011）进一步把题目难度纳入知识追踪。ITS meta-analysis 也表明，高质量自适应辅导系统可以改善学习结果（Kulik & Fletcher, 2016; Ma et al., 2014; VanLehn, 2011）。这些文献共同支持 v2 的混合定位：BKT/ITS 负责可信决策底座，LLM 负责解释和排序增强。",
        ],
    )

    add_heading(doc, "1.3 数学公式：候选集、受限重排、校验与回退", level=2)
    add_paras(
        doc,
        [
            "令 C_t 表示时刻 t 由 BKT 和确定性 policy 生成的候选集合。每个候选 c_i 包含 action、knowledge component、questionIds、guard flags 与 evidence summary。v2 的关键是限制 LLM 的动作空间：它只能在 C_t 内进行排序和解释。",
        ],
    )
    for idx in range(1, 7):
        add_formula_by_id(doc, idx)
        if idx == 1:
            add_paragraph(doc, "候选集公式强调，LLM 面对的不是开放课程空间，而是已经通过 BKT 与规则生成的教学动作集合。")
        elif idx == 2:
            add_paragraph(doc, "d_t 是确定性基线，任何 LLM 输出都必须能回退到这一安全决策。")
        elif idx == 3:
            add_paragraph(doc, "s_phi 表示 LLM 辅助打分函数，但最大化范围仍被 C_t 限制。")
        elif idx == 4:
            add_paragraph(doc, "V(r_t) 是本地验证门，强调 candidateId、questionIds、guardrails 与双语审计字段都必须满足。")
        elif idx == 5:
            add_paragraph(doc, "最终决策是明确的 fallback rule：provider 或校验失败不应中断教学，而应返回确定性推荐。")
        elif idx == 6:
            add_paragraph(doc, "本地 AI confidence 由证据深度 E、guard 一致性 G_a、mastery margin M 和 question integrity Q 构成，用来约束而不是放大 LLM 自信。")

    add_heading(doc, "1.4 教育中的应用：双语解释、教师审计与题组排序", level=2)
    add_paras(
        doc,
        [
            "v2 的第一项教育价值是把概率状态转化为学生能理解的反馈。BKT 可以判断某个知识点掌握概率不足，但 LLM 可以把这个判断改写为双语学习理由，例如解释为什么应先修补分数化简、再进入比例应用。Shute（2008）关于 formative feedback 的综述说明，反馈应清楚、具体、面向任务，并帮助学习者调整下一步行动；v2 的 learnerReason 正是把底层模型状态转化为这种可行动反馈。",
            "第二项价值是教师审计。teacherAuditNote、signalsUsed 和 fallback log 能帮助教师判断系统推荐是否合理。Holstein 等（2019）关于 teacher-AI complementarity 的研究提示，AI 工具应增强教师对课堂状态的可见性，而不是把教师排除在决策之外。因此，v2 的教师端不应只显示“AI 推荐”，而应显示候选、证据、被拒绝原因和确定性 baseline。",
            "第三项价值是题组排序。v2 可以在 review、repair、practice 等候选之间考虑近期错题分布、学习节奏和题型多样性，从而使练习序列更贴近学生当前困难。与此同时，v2 也有明确限制：不能让 LLM 生成新题或给出答案，不能绕过 due review，不能在证据不足时声称学生已经掌握。这样的限制把 v2 定位为高级 AI-assisted adaptive layer，而不是替代 BKT 的新权威。",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Section 2. v3：涉及的自适应技术", level=1)
    add_heading(doc, "2.1 概念介绍：从候选重排到可验证教学决策", level=2)
    add_paras(
        doc,
        [
            "v3 的核心不是“使用更强 LLM”，而是把自适应学习升级为 verifiable instructional decision-making。该架构保留 BKT 作为 interpretable safety floor，同时增加教育知识图谱、RAG evidence grounding、validation services、audit logging、teacher-in-the-loop 和回滚治理。LLM 在 v3 中仍然不是教学边界的制定者，而是证据内解释、候选比较和风险摘要的助手。",
            "v3 的候选决策过程可以描述为：学习事件进入 event ingestion；BKT 更新每个知识组件的掌握概率；教育知识图谱判断先备 readiness、课程目标、题目映射、lesson 与可视化模块关系；RAG 从课程、lesson、题库元数据、错因标签和教师备注检索证据包；LLM 只在证据包和候选集内部进行解释与排序；validator 执行先备、证据、风险、成本和教师策略校验；教师可在 audit console 中确认、覆盖或批量布置推荐。",
            "这一设计与知识追踪和学生建模文献一致。Desmarais 和 Baker（2012）以及 Pelanek（2017）强调，learner modeling 不只是预测答题正确率，也要支持教学解释和系统行动。Koedinger 等（2012）的 KLI framework 进一步指出，教学系统需要连接知识、学习过程和教学行动。v3 正是把这种连接写入系统架构。",
        ],
    )
    add_v3_architecture_table(doc, "zh")

    add_heading(doc, "2.2 文献综述：KT、KG/RAG 与教育理论支撑", level=2)
    add_paras(
        doc,
        [
            "知识追踪研究为 v3 的 learner-state modeling 提供基础。经典 BKT 可解释且适合安全边界（Corbett & Anderson, 1995），个性化 BKT 能处理学生差异（Pardos & Heffernan, 2010; Yudelson et al., 2013），KT-IDEM 和 Performance Factors Analysis 进一步引入题目难度与成功/失败历史（Pardos & Heffernan, 2011; Pavlik et al., 2009）。深度知识追踪、DKVMN 和 knowledge query network 展示了更强的序列建模能力，但也带来可解释性与校准问题（Gervet et al., 2020; Khajah et al., 2016; Lee & Yeung, 2019; Piech et al., 2015; Zhang et al., 2017）。因此，v3 可以把深度 KT 作为离线排序信号，但仍应让 BKT 作为安全底座。",
            "RAG 与知识图谱为 v3 的 evidence grounding 提供方法。Lewis 等（2020）提出 retrieval-augmented generation，使生成过程结合外部非参数记忆；Asai 等（2024）的 Self-RAG 将检索、生成与自我批判结合起来。Hogan 等（2021）对 knowledge graphs 的综述说明，图结构适合表达实体、关系、约束和可查询知识。对 MAIS 来说，教育知识图谱可以连接知识组件、先备关系、题目、lesson、可视化模块和错因标签；RAG 则把这些关系转化为 LLM 解释必须依赖的 evidence pack。",
            "v3 的教育理论基础包括 mastery learning、认知负荷、形成性反馈、间隔与提取练习、自我调节学习、多重表征和支架式教学。Bloom（1984）的 two-sigma problem 支持高质量个别化辅导的目标；Sweller 等（2019）说明教学设计必须控制认知负荷；Shute（2008）强调反馈应可行动；Cepeda 等（2006）和 Roediger 与 Karpicke（2006）分别支持间隔练习与 retrieval practice；Zimmerman（2002）强调学习者自我调节；Ainsworth（2006）提出多重表征学习框架；Wood 等（1976）关于 scaffolding 的经典研究支持逐步撤除支持。v3 可以把这些理论转化为多目标 utility 与验证门。",
            "教师在环不是附加功能，而是 v3 的治理结构。Holstein 等（2019）证明 teacher-AI complementarity 可以通过课堂编排工具实现，Ritter 等（2016）也展示 mastery learning at scale 需要系统化的补救和进阶机制。v3 的 teacher console 应记录 override、approval、reason labels 和班级层趋势，使教师反馈成为系统改进数据。",
        ],
    )

    add_heading(doc, "2.3 数学公式：BKT、知识图谱、证据支持与多目标决策", level=2)
    add_paragraph(doc, "v3 首先保留 BKT 更新作为安全底座。公式 (7) 表示每次证据更新后，模型再加入学习转移概率 T。")
    add_formula_by_id(doc, 7)
    add_paragraph(doc, "教育知识图谱把课程、题库、lesson、可视化和错因标签纳入同一关系结构。")
    add_formula_by_id(doc, 8)
    add_paragraph(doc, "先备 readiness 可以作为 hard gate：若任一关键先备未达阈值，系统不能推荐挑战或新知识推进。")
    add_formula_by_id(doc, 9)
    add_paragraph(doc, "RAG evidence support 用检索证据与候选声明之间的相似或匹配关系度量解释是否有支撑。")
    add_formula_by_id(doc, 10)
    add_paragraph(doc, "在 v3 中，最终候选不是单纯由 LLM 选择，而是在 BKT 候选集、知识图谱、证据支持和验证门约束下最大化多目标效用。")
    add_formula_by_id(doc, 11)
    add_formula_by_id(doc, 12)
    add_paragraph(doc, "效用函数中的正向项包括 mastery gain、retention 和 fairness，惩罚项包括 cognitive load、risk 与 cost。若证据、验证或教师策略失败，系统回退到确定性 baseline。")
    add_formula_by_id(doc, 13)

    add_heading(doc, "2.4 教育中的应用：mastery repair、spaced review 与可视化推荐", level=2)
    add_paras(
        doc,
        [
            "v3 可以把 mastery repair 做得更可信：当学生在二次函数顶点、截距或对称轴上反复出错时，BKT 检测低掌握概率，知识图谱定位相关先备，RAG 找到对应 lesson、worked example 和错因说明，LLM 生成证据支持的双语解释，教师端显示候选与可覆盖选项。学生看到的是清晰下一步，教师看到的是可审计证据链。",
            "v3 也能支持 spaced review 和 retrieval practice。系统不只根据“今天是否做错”推荐练习，还可根据 review due、记忆保持风险和题型分布安排复习。Cepeda 等（2006）与 Roediger 和 Karpicke（2006）的研究说明，间隔与提取练习对长期保持有重要作用；v3 可以把这些理论变成 retention utility 和 review scheduler，而不是口头原则。",
            "在可视化数学学习中，v3 的知识图谱能把抽象技能连接到 visualization modules。若学生对函数图像、斜率、截距或概率模拟缺乏直观理解，系统可以推荐相应互动可视化，而不是只继续推纸笔题。Ainsworth（2006）的多重表征框架说明，不同表征可以互补、限制误解并促进抽象理解；这为 MAIS 的 visualization lab 与 adaptive engine 结合提供了理论依据。",
            "最后，v3 把教师批准推荐作为常规模式而非异常模式。教师可以对系统推荐进行确认、调整或批量布置，override data 反过来用于评估 candidate generator 和 utility weights。这样，v3 不让 LLM 直接越过教学边界，而是把它放进可验证、可审计、可回滚的教学决策系统。",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Section 3. v2 与 v3 综合比较", level=1)
    add_comparison_table(doc, "zh")
    add_heading(doc, "3.1 关键差异", level=2)
    add_paras(
        doc,
        [
            "v2 与 v3 的最大差异是系统权力结构。v2 解决的是“如何安全使用 LLM 来改进候选排序和解释”；v3 解决的是“如何让每个教学决策都能被证据、先备、风险、教师策略和回退机制验证”。因此，v2 的主要成功标准是 LLM 输出是否合规、是否改善解释质量和教师可读性；v3 的主要成功标准则扩展到证据覆盖、unsupported claim rate、teacher override rate、公平性差距、成本控制和长期学习收益。",
            "从教育价值看，v2 更像一个高级表达层。它把 BKT 的概率与规则翻译成学生和教师能理解的自然语言，并在安全候选内选择更合适的下一步。v3 则把教育理论、课程结构和教师治理写入决策架构，使系统能解释为什么推荐某个 lesson、为什么先修补某个先备、为什么选择可视化而不是继续练题，以及为什么在证据不足时应回退。",
            "从风险治理看，v2 通过 validation/fallback 约束 LLM；v3 通过 evidence grounding、知识图谱、teacher-in-the-loop、监控和回滚建立完整治理链。若 MAIS 要进入学校级部署，v3 的价值更高，因为学校不仅关心推荐是否聪明，也关心推荐是否可审计、可解释、可控、可复盘。",
        ],
    )
    add_heading(doc, "3.2 推荐定位", level=2)
    add_numbered(doc, "v2 定位为 AI-assisted adaptive reranking：提升双语反馈、个性化解释、教师审计和候选排序质量。")
    add_numbered(doc, "v3 定位为 verifiable instructional decision architecture：把 BKT、KG、RAG、验证、审计和教师控制作为统一系统。")
    add_numbered(doc, "v2 可先用于高级版或教师试点；v3 应作为下一代学校部署基础，尤其适合需要审计、合规和教师治理的场景。")
    add_numbered(doc, "两者都不应允许 LLM 直接生成教学边界；任何 LLM 贡献都必须能被候选集、证据、validator 和 fallback 约束。")

    doc.add_page_break()
    add_heading(doc, "Section 4. v4 版本展望：证据验证、因果、多模态与隐私保护生态", level=1)
    add_heading(doc, "4.1 开发定位", level=2)
    add_paras(
        doc,
        [
            "如果 v3 的目标是“可验证教学决策系统”，v4 应进一步发展为 evidence-validated, causal, multimodal, privacy-preserving adaptive ecosystem。也就是说，v4 不只是更强的模型，而是能持续回答：某个推荐是否真的提升学习收益、对哪些学生有效、是否造成群体不公平、是否保护学生隐私、能否跨班级和学校泛化。",
            "v4 仍不应让 LLM 直接越过教学边界。相反，v4 应在 v3 的治理框架上增强因果评估、长期学习收益监测、多模态学习分析、隐私保护协作训练和学校级治理 dashboard。LLM 继续作为解释、证据摘要和教师协作界面，而不是最终教学权威。",
        ],
    )
    add_heading(doc, "4.2 方法路线", level=2)
    add_paragraph(doc, "第一，v4 可引入 causal inference 与 uplift modeling，评估某类推荐对不同学生的异质性效果。Athey 和 Imbens（2016）以及 Wager 和 Athey（2018）的研究为 heterogeneous treatment effects 提供方法基础。")
    add_formula_by_id(doc, 14)
    add_paragraph(doc, "第二，v4 可发展多模态 learning analytics，把作答日志、可视化互动、手写过程、停留时间、提示使用和教师标注合并为更丰富的学习证据。这里的关键不是收集更多数据，而是把数据纳入明确的同意、最小化、审计和教学目的限制。")
    add_paragraph(doc, "第三，v4 可探索 federated 或 privacy-preserving analytics。Yang 等（2019）概述了 federated machine learning，Dwork（2006）奠定了 differential privacy 的形式化基础。对 K-12 教育系统来说，这些方法可帮助不同学校共享模型改进而不直接集中敏感学生数据。")
    add_formula_by_id(doc, 15)
    add_paragraph(doc, "第四，v4 需要 longitudinal mastery modeling，把短期答题正确率扩展到长期保持、迁移能力和概念稳定性。系统应追踪学生在数周或数月尺度上的 mastery repair、review retention、跨题型迁移和教师确认结果。")
    add_heading(doc, "4.3 学校部署与治理", level=2)
    add_paras(
        doc,
        [
            "v4 的学校部署应采用 co-design loop：教师、课程负责人、家长代表和技术团队共同定义推荐边界、审计指标、隐私规则和失败处理流程。系统 dashboard 应显示班级趋势、群体公平性、teacher override、fallback rate、cost per learner、unsupported claims 和 learning gain，而不只显示使用量。",
            "A/B 测试和准实验应分层进行：先在 shadow mode 比较 v3/v4 建议与教师判断，再在 teacher-approved mode 中观察学习收益，最后才有限自动化。任何自动化推荐都应保留 rollback governance，使学校可以按班级、年级、题型或模块关闭高风险功能。",
            "v4 的成功标准不是“AI 更像教师”，而是“教师和系统共同做出更可靠的教学决策”。这一方向与 Bloom（1984）的高质量个别化辅导目标一致，也与 Holstein 等（2019）的 teacher-AI complementarity 相容。",
        ],
    )
    add_heading(doc, "引用与资料使用说明", level=1)
    add_paragraph(doc, "本报告使用附件 Research Paper_MAIS_v3_adaptive_engine.docx 和 MAIS-MVP 本地项目背景作为 provided background，但正式 References 仅收录学术期刊或学术会议论文。内部报告、网页、课程指南、MCP 文档、厂商文档、博客、白皮书、书籍和 preprint-only 来源不进入 References。")


def add_en_body(doc: Document) -> None:
    add_heading(doc, "Section 1. v2: LLM-Adaptive Learning", level=1)
    add_heading(doc, "1.1 Concept: Guarded LLM reranking over a BKT candidate set", level=2)
    add_paras(
        doc,
        [
            "In the current MAIS v2 definition, LLM adaptation does not mean allowing a language model to freely decide the instructional path. The LLM sits after the BKT-generated candidate set as a constrained candidate reranker, explanation generator, and teacher-audit assistant. BKT and deterministic rules first generate candidate actions from mastery probability, prerequisites, mistakes, review due dates, and item mapping. The LLM may rank only within these candidates; it may not invent new items, skip prerequisites, change mastery probabilities, or cancel due-review and repair guards.",
            "The v2 engineering flow has four steps. First, the deterministic adaptive core generates candidates. Second, the LLM receives candidate summaries, recent performance, and policy instructions, then returns a ranking, learnerReason, teacherAuditNote, and signalsUsed. Third, a local validator deterministically checks candidateId, questionIds, schema validity, guard flags, and bilingual fields. Fourth, if validation fails or the provider is unavailable, the system returns the deterministic decision d_t. This makes v2 an LLM-assisted system, not an LLM-authoritative system.",
            "The technical basis for this capability comes from Transformer attention, large-scale pretraining, few-shot learning, instruction tuning/RLHF, and chain-of-thought prompting (Brown et al., 2020; Ouyang et al., 2022; Vaswani et al., 2017; Wei et al., 2022). These studies explain why large models can synthesize context, produce readable explanations, and condition on complex input. However, educational deployment requires evidence, auditability, and safe fallback rather than fluency alone.",
        ],
    )
    add_v2_candidate_table(doc, "en")

    add_heading(doc, "1.2 Literature Review: LLMs in education, AI tutor evaluation, and risk governance", level=2)
    add_paras(
        doc,
        [
            "A consistent finding in the LLM-in-education literature is that large models can support personalized feedback, learning explanations, teacher preparation, dialogue tutoring, and content transformation, but their use must be constrained by accuracy, privacy, fairness, and teacher role design. Kasneci et al. (2023) discuss educational opportunities and challenges of LLMs, while Yan et al. (2024) systematically review practical and ethical issues. These reviews support the v2 design position: the LLM should improve feedback and teacher visibility, not replace the learner model or teacher judgment.",
            "AI tutor evaluation should not be reduced to naturalness of language. Tack and Piech (2022) propose the AI Teacher Test to measure pedagogical behavior, and Macina et al. (2023) introduce MathDial as a mathematics tutoring dataset with rich pedagogical properties. For MAIS v2, the LLM output should be evaluated by whether it produces an educationally appropriate explanation and ranking within candidate constraints, not by whether the prose sounds teacher-like.",
            "Risk research explains why v2 requires validation and fallback. Ji et al. (2023) review hallucination in natural language generation; Maynez et al. (2020) analyze factuality and faithfulness; Weidinger et al. (2022) propose a taxonomy of language-model risks including bias, misuse, privacy, and social harms. Educational systems must also address algorithmic bias, and Baker and Hawn (2022) show how educational algorithms can affect groups differently. Therefore, v2 AI confidence should not be a model self-report; it should be computed from local evidence, guardrail consistency, and validator logs.",
            "v2 also inherits the tradition of intelligent tutoring systems and knowledge tracing. Corbett and Anderson (1995) introduced BKT to trace procedural knowledge; Pardos and Heffernan (2010) and Yudelson et al. (2013) individualized BKT parameters; Pardos and Heffernan (2011) added item difficulty. Meta-analyses show that intelligent tutoring systems can improve learning outcomes (Kulik & Fletcher, 2016; Ma et al., 2014; VanLehn, 2011). Together, these sources support a hybrid architecture: BKT/ITS methods provide the trusted decision floor, while the LLM enhances explanation and candidate ranking.",
        ],
    )

    add_heading(doc, "1.3 Mathematical Formulation: Candidate set, constrained reranking, validation, and fallback", level=2)
    add_paragraph(doc, "Let C_t be the candidate set generated at time t by BKT and deterministic policy. Each candidate c_i includes an action, a knowledge component, questionIds, guard flags, and an evidence summary. The key v2 constraint is that the LLM's action space is limited to C_t.")
    for idx in range(1, 7):
        add_formula_by_id(doc, idx)
        if idx == 1:
            add_paragraph(doc, "The candidate-set definition makes clear that the LLM receives a bounded instructional space, not an open curriculum space.")
        elif idx == 2:
            add_paragraph(doc, "d_t is the deterministic baseline to which every LLM-assisted decision must be able to fall back.")
        elif idx == 3:
            add_paragraph(doc, "s_phi denotes the LLM-assisted scoring function, but the maximization domain remains C_t.")
        elif idx == 4:
            add_paragraph(doc, "V(r_t) is the local validation gate for candidate identity, item containment, guardrail satisfaction, and bilingual audit fields.")
        elif idx == 5:
            add_paragraph(doc, "The final decision is an explicit fallback rule: provider failure or validation failure returns the deterministic recommendation.")
        elif idx == 6:
            add_paragraph(doc, "Local AI confidence is composed from evidence depth E, guard agreement G_a, mastery margin M, and question integrity Q; it constrains rather than amplifies model self-confidence.")

    add_heading(doc, "1.4 Educational Applications: Bilingual feedback, teacher audit, and item-set ordering", level=2)
    add_paras(
        doc,
        [
            "The first educational value of v2 is converting probabilistic state into feedback students can act on. BKT can identify that mastery is below threshold; the LLM can translate that state into a bilingual rationale explaining why fraction simplification should be repaired before ratio applications. Shute's (2008) review of formative feedback emphasizes that feedback should be clear, specific, task-focused, and useful for next action. v2 learnerReason turns model state into such actionable feedback.",
            "The second value is teacher auditability. teacherAuditNote, signalsUsed, and fallback logs help teachers judge whether the recommendation is reasonable. Holstein et al. (2019) show that teacher-AI complementarity depends on tools that improve teachers' visibility into classroom state rather than excluding them from decisions. Accordingly, the v2 teacher interface should show candidates, evidence, rejection reasons, and the deterministic baseline.",
            "The third value is item-set ordering. v2 can compare review, repair, and practice candidates using recent mistake patterns, learning rhythm, and item diversity. At the same time, it must not allow the LLM to generate new items, reveal answers, bypass due review, or claim mastery when evidence is insufficient. This positions v2 as a premium AI-assisted adaptive layer, not a replacement for BKT.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Section 2. v3: Adaptive Technologies Involved", level=1)
    add_heading(doc, "2.1 Concept: From candidate reranking to verifiable instructional decision-making", level=2)
    add_paras(
        doc,
        [
            "The core of v3 is not a stronger LLM. It is a shift toward verifiable instructional decision-making. The architecture preserves BKT as an interpretable safety floor and adds an educational knowledge graph, RAG evidence grounding, validation services, audit logging, teacher-in-the-loop control, and rollback governance. In v3 the LLM remains an assistant for evidence-bound explanation, candidate comparison, and risk summarization rather than an authority that sets instructional boundaries.",
            "The v3 decision process can be described as follows: learning events enter event ingestion; BKT updates mastery for each knowledge component; the educational knowledge graph evaluates prerequisites, curriculum objectives, item mapping, lessons, and visualization relations; RAG retrieves evidence packs from curriculum, lessons, item metadata, misconception tags, and teacher notes; the LLM explains and compares only within the evidence pack and candidate set; validators check prerequisites, evidence, risk, cost, and teacher policy; and teachers can approve, override, or batch assign recommendations in an audit console.",
            "This design is consistent with the knowledge tracing and learner modeling literature. Desmarais and Baker (2012) and Pelanek (2017) emphasize that learner modeling should support instructional action, not merely next-response prediction. Koedinger et al. (2012) argue that instructional systems should connect knowledge, learning processes, and instructional actions. v3 encodes that connection into the software architecture.",
        ],
    )
    add_v3_architecture_table(doc, "en")

    add_heading(doc, "2.2 Literature Review: Knowledge tracing, KG/RAG, and educational theory", level=2)
    add_paras(
        doc,
        [
            "Knowledge tracing research provides the foundation for v3 learner-state modeling. Classical BKT is interpretable and appropriate as a safety boundary (Corbett & Anderson, 1995). Individualized BKT handles student differences (Pardos & Heffernan, 2010; Yudelson et al., 2013). KT-IDEM and Performance Factors Analysis add item difficulty and histories of success or failure (Pardos & Heffernan, 2011; Pavlik et al., 2009). Deep Knowledge Tracing, DKVMN, and Knowledge Query Networks improve sequence modeling but raise explainability and calibration questions (Gervet et al., 2020; Khajah et al., 2016; Lee & Yeung, 2019; Piech et al., 2015; Zhang et al., 2017). v3 can use deep KT as an offline ranking feature while preserving BKT as the safety floor.",
            "RAG and knowledge graphs provide v3's evidence-grounding methods. Lewis et al. (2020) propose retrieval-augmented generation, which combines parametric models with non-parametric memory. Asai et al. (2024) extend this direction through retrieval, generation, and self-critique. Hogan et al. (2021) review knowledge graphs as structures for representing entities, relations, constraints, and queryable knowledge. For MAIS, the educational knowledge graph can connect knowledge components, prerequisites, items, lessons, visualization modules, and misconception tags, while RAG converts those relations into the evidence pack that LLM explanations must use.",
            "The educational theory base for v3 includes mastery learning, cognitive load theory, formative feedback, spacing and retrieval practice, self-regulated learning, multiple representations, and scaffolding. Bloom's (1984) two-sigma problem motivates high-quality individualized support; Sweller et al. (2019) show that instruction should manage cognitive load; Shute (2008) emphasizes actionable feedback; Cepeda et al. (2006) and Roediger and Karpicke (2006) support spaced practice and retrieval practice; Zimmerman (2002) highlights self-regulated learning; Ainsworth (2006) frames learning with multiple representations; and Wood et al. (1976) provide the classic scaffolding account. v3 can translate these theories into utility terms and validation gates.",
            "Teacher-in-the-loop is not an add-on; it is part of v3 governance. Holstein et al. (2019) show how classroom orchestration tools can support teacher-AI complementarity, and Ritter et al. (2016) demonstrate that mastery learning at scale requires systematic repair and progression mechanisms. A v3 teacher console should record overrides, approvals, reason labels, and class-level trends so that teacher feedback becomes model-improvement data.",
        ],
    )

    add_heading(doc, "2.3 Mathematical Formulation: BKT, knowledge graph, evidence support, and multi-objective decision-making", level=2)
    add_paragraph(doc, "v3 first preserves the BKT update as the safety floor. Formula (7) shows how the model incorporates the learning transition after evidence has been observed.")
    add_formula_by_id(doc, 7)
    add_paragraph(doc, "The educational knowledge graph brings curriculum, items, lessons, visualizations, and misconception tags into one relation structure.")
    add_formula_by_id(doc, 8)
    add_paragraph(doc, "Prerequisite readiness can be implemented as a hard gate: if any key prerequisite is below threshold, challenge or new-topic advancement is blocked.")
    add_formula_by_id(doc, 9)
    add_paragraph(doc, "RAG evidence support measures whether retrieved evidence supports the claims made for a candidate.")
    add_formula_by_id(doc, 10)
    add_paragraph(doc, "The final v3 candidate is not chosen by the LLM alone. It maximizes utility under the BKT candidate set, knowledge graph constraints, evidence support, and validation gates.")
    add_formula_by_id(doc, 11)
    add_formula_by_id(doc, 12)
    add_paragraph(doc, "The utility function rewards mastery gain, retention, and fairness while penalizing cognitive load, risk, and cost. If evidence, validation, or teacher policy fails, the system falls back to the deterministic baseline.")
    add_formula_by_id(doc, 13)

    add_heading(doc, "2.4 Educational Applications: Mastery repair, spaced review, and visualization recommendation", level=2)
    add_paras(
        doc,
        [
            "v3 can make mastery repair more trustworthy. When a learner repeatedly struggles with vertex, intercept, or axis-of-symmetry tasks in quadratic functions, BKT detects low mastery, the knowledge graph identifies related prerequisites, RAG retrieves the relevant lesson, worked example, and misconception evidence, the LLM produces evidence-supported bilingual explanations, and the teacher view shows the candidate and override option. The student sees a clear next step; the teacher sees an auditable evidence chain.",
            "v3 can also support spaced review and retrieval practice. The system recommends practice not only because a learner made an error today, but because review is due, retention risk is high, or item-type coverage is incomplete. Research on distributed practice and retrieval practice supports this design (Cepeda et al., 2006; Roediger & Karpicke, 2006). v3 turns those principles into retention utility and a review scheduler.",
            "For visual mathematics learning, the knowledge graph can connect abstract skills to visualization modules. If a learner lacks intuition about function graphs, slope, intercepts, or probability simulation, the system can recommend an interactive visualization instead of simply assigning more paper-style items. Ainsworth's (2006) multiple-representation framework supports the idea that representations can complement one another, constrain misconceptions, and promote abstraction.",
            "Finally, v3 treats teacher-approved recommendation as a normal mode rather than an exception. Teachers can confirm, adjust, or batch assign recommendations, and override data can improve the candidate generator and utility weights. Thus v3 keeps the LLM inside a verifiable, auditable, and rollback-capable instructional decision system.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Section 3. Synthesis of v2 and v3", level=1)
    add_comparison_table(doc, "en")
    add_heading(doc, "3.1 Core Differences", level=2)
    add_paras(
        doc,
        [
            "The largest difference between v2 and v3 is the distribution of authority. v2 solves the problem of how to use an LLM safely to improve candidate ranking and explanation. v3 solves the larger problem of making every instructional decision verifiable through evidence, prerequisites, risk, teacher policy, and fallback. Accordingly, v2 success is measured by compliant LLM outputs, better explanations, and teacher readability; v3 success expands to evidence coverage, unsupported claim rate, teacher override rate, fairness gaps, cost control, and long-term learning gain.",
            "Educationally, v2 functions as an advanced expression layer. It translates BKT probabilities and rules into language that students and teachers can understand, then chooses among safe candidates. v3 writes educational theory, curriculum structure, and teacher governance into the decision architecture. It can explain why the system recommends a lesson, why prerequisite repair comes first, why a visualization is preferable to another item set, and why evidence insufficiency should trigger fallback.",
            "From a risk-governance perspective, v2 constrains the LLM through validation and fallback. v3 creates a fuller governance chain through evidence grounding, knowledge graph constraints, teacher-in-the-loop control, monitoring, and rollback. For school-level deployment, v3 is more valuable because schools need recommendations that are auditable, explainable, controllable, and reviewable, not merely intelligent-looking.",
        ],
    )
    add_heading(doc, "3.2 Recommended Positioning", level=2)
    add_numbered(doc, "Position v2 as AI-assisted adaptive reranking: richer bilingual feedback, personalized explanation, teacher audit, and candidate ranking.", "en")
    add_numbered(doc, "Position v3 as a verifiable instructional decision architecture: BKT, KG, RAG, validation, audit, and teacher control in one system.", "en")
    add_numbered(doc, "Use v2 first for premium or teacher-pilot features; use v3 as the next school-deployment foundation where auditability and governance matter.", "en")
    add_numbered(doc, "Do not allow either version to let the LLM directly set instructional boundaries; every LLM contribution must be constrained by candidates, evidence, validators, and fallback.", "en")

    doc.add_page_break()
    add_heading(doc, "Section 4. Outlook for v4: Evidence-Validated, Causal, Multimodal, and Privacy-Preserving Ecosystem", level=1)
    add_heading(doc, "4.1 Development Positioning", level=2)
    add_paras(
        doc,
        [
            "If v3 aims to be a verifiable instructional decision system, v4 should evolve into an evidence-validated, causal, multimodal, and privacy-preserving adaptive ecosystem. The system should be able to answer whether a recommendation actually improves learning, for which learners it works, whether it introduces inequity, whether student privacy is protected, and whether the pattern generalizes across classrooms and schools.",
            "v4 should still not allow the LLM to bypass instructional boundaries. Instead, it should build on v3 governance by adding causal evaluation, long-term learning-gain monitoring, multimodal learning analytics, privacy-preserving collaborative training, and school-level governance dashboards. The LLM remains an interface for explanation, evidence summarization, and teacher collaboration rather than the final instructional authority.",
        ],
    )
    add_heading(doc, "4.2 Methodological Roadmap", level=2)
    add_paragraph(doc, "First, v4 can introduce causal inference and uplift modeling to estimate heterogeneous effects of recommendations. Athey and Imbens (2016) and Wager and Athey (2018) provide methodological foundations for heterogeneous treatment effects.")
    add_formula_by_id(doc, 14)
    add_paragraph(doc, "Second, v4 can develop multimodal learning analytics by combining response logs, visualization interactions, handwriting process data, dwell time, hint use, and teacher labels. The point is not to collect more data for its own sake, but to place data under clear consent, minimization, audit, and instructional-purpose limits.")
    add_paragraph(doc, "Third, v4 can explore federated or privacy-preserving analytics. Yang et al. (2019) survey federated machine learning, and Dwork (2006) formalizes differential privacy. In K-12 education, these methods can help schools share model improvements without centralizing sensitive student records.")
    add_formula_by_id(doc, 15)
    add_paragraph(doc, "Fourth, v4 should extend short-term correctness into longitudinal mastery modeling, tracking mastery repair, retention, transfer across item types, and teacher-confirmed conceptual stability over weeks or months.")
    add_heading(doc, "4.3 School Deployment and Governance", level=2)
    add_paras(
        doc,
        [
            "v4 school deployment should use a co-design loop in which teachers, curriculum leaders, parent representatives, and technical teams jointly define recommendation boundaries, audit indicators, privacy rules, and failure-handling procedures. Dashboards should show class trends, fairness gaps, teacher overrides, fallback rates, cost per learner, unsupported claims, and learning gain rather than only usage volume.",
            "A/B tests and quasi-experiments should progress in layers: shadow mode to compare v3/v4 recommendations against teacher judgment, teacher-approved mode to observe learning impact, and only then limited automatic mode. Every automatic recommendation should retain rollback governance so schools can disable high-risk functions by class, grade, item type, or module.",
            "The v4 success criterion is not whether AI sounds more like a teacher. It is whether teachers and the system make more reliable instructional decisions together. This direction is consistent with Bloom's (1984) goal of high-quality individualized tutoring and Holstein et al.'s (2019) account of teacher-AI complementarity.",
        ],
    )
    add_heading(doc, "Citation and Source-Use Note", level=1)
    add_paragraph(doc, "This report uses the attached Research Paper_MAIS_v3_adaptive_engine.docx and the local MAIS-MVP project context as provided background. Formal References include only academic journals or conference papers. Internal reports, web pages, curriculum guides, MCP documentation, vendor documents, blogs, white papers, books, and preprint-only sources are excluded from References.")


def add_front_matter(doc: Document, cfg: dict) -> None:
    add_title_page(doc, cfg)
    doc.add_page_break()
    add_heading(doc, cfg["abstract_title"], level=1)
    add_paras(doc, cfg["abstract"])
    add_heading(doc, cfg["toc"], level=1)
    add_toc(doc.add_paragraph(), cfg["lang"])
    add_paragraph(doc, cfg["toc_note"], after=10, line=1.2)
    doc.add_page_break()


def add_references(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "References", level=1)
    for ref in REFERENCES:
        p = doc.add_paragraph(style="APA Reference")
        r = p.add_run(ref)
        set_run_font(r, size=9.1, east_asia="Times New Roman", latin="Times New Roman")


def build_doc(cfg: dict) -> Path:
    doc = setup_doc(cfg)
    add_front_matter(doc, cfg)
    if cfg["lang"] == "zh":
        add_zh_body(doc)
    else:
        add_en_body(doc)
    add_references(doc)
    cfg["out"].parent.mkdir(parents=True, exist_ok=True)
    doc.save(cfg["out"])
    return cfg["out"]


def main() -> None:
    for cfg in (ZH, EN):
        path = build_doc(cfg)
        print(path)


if __name__ == "__main__":
    main()
