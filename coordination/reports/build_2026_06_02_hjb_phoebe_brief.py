from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-06-02-沪教版RAG与AI内容进度简报-Phoebe.docx")

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

FONT_BODY = "Microsoft YaHei"
FONT_HEADING = "Microsoft YaHei"
FONT_EN = "Calibri"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8F1FA"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E2EC"
GREEN_FILL = "EAF7EF"
GREEN_TEXT = "166534"
YELLOW_FILL = "FFF7E0"
YELLOW_TEXT = "8A6100"
RED_FILL = "FDECEC"
RED_TEXT = "A51D2A"
WHITE = "FFFFFF"


def set_run_font(run, name=FONT_BODY, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D0D5DD", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_row_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_para(doc, text="", size=11, color="000000", bold=False, italic=False, align=None, before=0, after=6, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, before, after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        set_paragraph_spacing(p, before=16, after=8, line=1.10)
        size, color = 16, BLUE
    elif level == 2:
        set_paragraph_spacing(p, before=12, after=6, line=1.10)
        size, color = 13, BLUE
    else:
        set_paragraph_spacing(p, before=8, after=4, line=1.10)
        size, color = 12, DARK_BLUE
    run = p.add_run(text)
    set_run_font(run, name=FONT_HEADING, size=size, color=color, bold=True)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_paragraph_spacing(p, before=0, after=5, line=1.167)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_paragraph_spacing(p, before=0, after=5, line=1.167)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE, title_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA - TABLE_INDENT_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    for border_cell in table._cells:
        set_cell_border(border_cell, color="B8CCE4", size="8")
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=1, after=3, line=1.12)
    r = p.add_run(title)
    set_run_font(r, size=11, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=1, line=1.16)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    add_para(doc, "", after=4)


def add_status_pill(cell, label, kind):
    fills = {"green": GREEN_FILL, "yellow": YELLOW_FILL, "red": RED_FILL, "blue": LIGHT_BLUE}
    colors = {"green": GREEN_TEXT, "yellow": YELLOW_TEXT, "red": RED_TEXT, "blue": DARK_BLUE}
    shade_cell(cell, fills[kind])
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=0, line=1.05)
    run = p.add_run(label)
    set_run_font(run, size=9.5, color=colors[kind], bold=True)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_row_cant_split(table.rows[0])
    set_row_repeat_header(table.rows[0])
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        shade_cell(hdr[idx], header_fill)
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, before=0, after=0, line=1.08)
        run = p.add_run(text)
        set_run_font(run, size=font_size, color=INK, bold=True)
    for row in rows:
        table_row = table.add_row()
        set_row_cant_split(table_row)
        cells = table_row.cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx != 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=0, after=0, line=1.12)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color="1D2939")
            if idx == 0:
                run.bold = True
    add_para(doc, "", after=6)
    return table


def set_doc_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, margin, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in (
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ):
        style = styles[name]
        style.font.name = FONT_HEADING
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEADING)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_running_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(hp, before=0, after=0, line=1)
    run = hp.add_run("MAIS 沪教版内容进度简报 | Phoebe 审阅版")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(fp, before=0, after=0, line=1)
    run = fp.add_run("内部课程 QA 文件")
    set_run_font(run, size=8.5, color=MUTED)


def add_cover(doc):
    add_para(doc, "内部课程 QA 简报", size=10.5, color=BLUE, bold=True, after=8)
    title = doc.add_paragraph()
    set_paragraph_spacing(title, before=0, after=4, line=1.0)
    r = title.add_run("沪教版数学 RAG 与 AI 内容进度简报")
    set_run_font(r, name=FONT_HEADING, size=24, color=INK, bold=True)
    sub = doc.add_paragraph()
    set_paragraph_spacing(sub, before=0, after=16, line=1.15)
    r = sub.add_run("给中国数学课程专家 Phoebe 的审阅版")
    set_run_font(r, size=13.5, color="475467")

    rows = [
        ("收件人", "Phoebe｜中国数学课程专家"),
        ("报告日期", "2026年6月2日"),
        ("项目范围", "MAIS 沪教版小学、初中、高中数学内容管线"),
        ("文档定位", "进度判断、专家审阅入口、风险与下一步建议"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1600, 7600])
    for i, (label, value) in enumerate(rows):
        shade_cell(table.cell(i, 0), LIGHT_GRAY)
        p = table.cell(i, 0).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.08)
        run = p.add_run(label)
        set_run_font(run, size=9.5, color=INK, bold=True)
        p = table.cell(i, 1).paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.12)
        run = p.add_run(value)
        set_run_font(run, size=10.5, color="1D2939")

    add_para(doc, "", after=8)
    add_callout(
        doc,
        "一句话判断",
        "沪教版 P1-S6 的安全 RAG 约束、AI 题库和主要 Lesson 内容已基本成形；题目 QA 是当前最成熟的一环。专家审阅的重点应放在课程顺序、年级适配、教材风格、中文术语自然度、以及小学插图和高中教师签核的补齐。",
        fill=LIGHT_BLUE,
    )


def build_doc():
    doc = Document()
    set_doc_styles(doc)
    add_running_header_footer(doc)
    doc.core_properties.title = "沪教版数学 RAG 与 AI 内容进度简报"
    doc.core_properties.subject = "MAIS MAINLAND_HJB curriculum QA progress brief"
    doc.core_properties.author = "MAIS 内容与课程 QA 协调"

    add_cover(doc)

    add_heading(doc, "一、阅读口径", 1)
    add_callout(
        doc,
        "重要口径",
        "这里的 RAG 不是教材原文全文库，而是安全抽象层：只保留教材/测评的章节结构、知识点、能力标签、常见误区、题型模式和生成约束；不保留教材原文、练习题原题、答案、解法、页码、图片、OCR 文本或 embedding。",
        fill="F6FAFF",
    )
    add_para(
        doc,
        "本简报按五个内容生产环节汇总：课程标准与 RAG 约束、AI 生成教材、AI 生成题目、教材 QA、题目 QA。状态来自项目内 S18/S05/S10 等会话记录、content-qa 产物和当前数据文件，不代表已经完成正式对外发布公告。",
        size=10.8,
    )

    doc.add_page_break()
    add_heading(doc, "二、总览结论", 1)
    rows = [
        ("课程标准 / RAG", "安全抽象层已覆盖 P1-S6", "绿灯", "小学 70 教材卡 + 83 测评卡；初中 22 教材卡 + 6 测评卡 + 26 试卷模式卡；高中 21 教材卡 + 63 测评/考试模式卡。"),
        ("AI 生成教材", "小学、初中、高中均有 Lesson 形态", "黄绿", "初中 22 个 lesson 已 approved/integrated；高中 21 个 lesson 自动 QA 通过；小学为基于 topic/question 的 lesson seed，需要专家确认教学完整度。"),
        ("AI 生成题目", "当前主线 4,500 题", "绿灯", "小学 1,500 题、初中 1,500 题、高中 V2 1,500 题，均为沪教版专属题库主线。"),
        ("教材 QA", "自动 QA 已有，人工签核未完全闭环", "黄灯", "高中记录显示 21/21 通过自动验证，但独立教师签核 pending；初中已有 approved 状态，小学缺独立教材 QA 报告。"),
        ("题目 QA", "三段题库均已达到绿色或准绿色", "绿灯", "小学 DeepSeek V4 Pro 最终 1500 pass；初中 1500 pass 且 150 抽样 approved；高中 V2 1500/1500 answer-matched、0 P0/P1/P2。"),
        ("插图 / 视觉资产", "小学题目插图仍未生成", "红黄", "已统计 607 题必须配图、538 题强烈建议配图，但 GPT Image2 生成受 API key / budget gate 限制。"),
    ]
    table = add_table(doc, ["模块", "当前进度", "状态", "Phoebe 需要关注的含义"], rows, [1550, 2050, 1000, 4760], font_size=9.2)
    for row in table.rows[1:]:
        status = row.cells[2].text
        row.cells[2].text = ""
        add_status_pill(row.cells[2], status, "green" if status == "绿灯" else "yellow" if status == "黄灯" or status == "黄绿" else "red")

    doc.add_page_break()
    add_heading(doc, "三、按学段进度", 1)
    stage_rows = [
        ("小学 P1-P6", "70 教材安全卡；83 测评安全卡。P1-P6 均已覆盖，但 P3 上册标注含“老课本暂用”口径。", "1,500 题；P1-P6 各 250。已生产集成到沪教版小学 Lesson 与 Practice Arena。", "最终 DeepSeek V4 Pro QA：1500 pass、0 warn、0 fail；发布结论记录为 owner-approved production launch。", "需重点审年级适配、P3 版本口径、低年级题目图像化需求。"),
        ("初中 S1-S3", "22 教材卡；6 个 S1 测评卡；26 个 S1-S3 试卷模式卡；共享中考模式层。", "22 个 generated lessons；1,500 题，S1/S2/S3 各 500。", "确定性审计 1500/1500 pass；0 duplicate；150 行样本 approved；生产集成批准。", "可进入专家抽样审题；另需确认部分几何判定题是否需增加条件限定。"),
        ("高中 S4-S6", "21 个高中教材安全卡；63 个沪教版高中测评/考试模式卡；5 本 owner-provided PDF 只保留安全抽象。", "21 个 generated lessons；当前主线题库为 V2 1,500 题，S4/S5/S6 各 500。", "V2 质量审计：1500/1500 可解且答案匹配；0 P0/P1/P2；370/370 manual queue reviewed。", "不要用早期 V4 blocked 口径作为当前状态；高中 Lesson 还需独立教师签核。"),
    ]
    add_table(doc, ["学段", "RAG / 课程约束", "AI 教材与题目", "QA 状态", "专家审阅重点"], stage_rows, [1120, 2320, 2230, 1940, 1750], font_size=8.8)

    doc.add_page_break()
    add_heading(doc, "四、题目 QA 细节", 1)
    add_para(doc, "题目管线当前是沪教版内容中最接近可规模化审阅的一环。以下状态建议作为 Phoebe 进行数学课程抽审的入口。")
    qa_rows = [
        ("小学", "1,500", "P1-P6 各 250", "DeepSeek 初审发现 55 风险行；S18 修复 62 行后，最终 DeepSeek V4 Pro 1500 pass、0 warn/fail。", "已获 owner approval 并生产集成。"),
        ("初中", "1,500", "S1-S3 各 500", "201/201 targeted repair 完成；确定性审计 1500 pass、0 fail、0 duplicate；150 抽样 approved。", "production-integration-approved。"),
        ("高中 V2", "1,500", "S4-S6 各 500", "质量审计 1500/1500 answer-matched，0 P0/P1/P2，370/370 manual queue reviewed。", "质量绿；正式对外发布仍应走 owner/S11 app/browser/build gates。"),
    ]
    add_table(doc, ["范围", "题量", "分布", "QA 证据", "当前判断"], qa_rows, [900, 900, 1450, 4300, 1810], font_size=9)

    add_heading(doc, "五、教材 QA 与课程专家审阅切入点", 1)
    add_callout(
        doc,
        "建议 Phoebe 先看什么",
        "不要从 4,500 题全量逐题审起。建议先用课程专家视角确认“沪教版课程顺序、年级定位、知识点表达、课堂教学完整度、考试风格”是否成立，再抽样进入题目数学正确性。",
        fill=YELLOW_FILL,
        title_color=YELLOW_TEXT,
    )
    add_bullet(doc, "先审课程结构：小学 P1-P6、初中 S1-S3、高中 S4-S6 的年级/学期/章节映射是否符合沪教版教学实际。")
    add_bullet(doc, "再审 Lesson 形态：初中 22 个 generated lessons、高中 21 个 generated lessons、小学 lesson seed 是否达到“可给学生学”的颗粒度。")
    add_bullet(doc, "再审题目抽样：建议按低小、中高小、初中、高中四个桶抽样，每桶覆盖 Foundation / Core / Challenge / Exam。")
    add_bullet(doc, "最后审发布风险：尤其是小学插图、几何图形/钟表/统计图等视觉题，以及高中教师指南和独立签核。")

    doc.add_page_break()
    add_heading(doc, "六、当前风险与待决事项", 1)
    risk_rows = [
        ("小学插图缺口", "1,145/1,500 题建议配图，其中 607 题必须配图。当前仅有需求清单，没有生成图片。", "确认是否先启动 A_required 607 题的图片生产；数学标签/坐标/图形需程序化叠加。"),
        ("高中教材签核", "高中 21 个 lesson 自动验证通过，但 QA acceptance 仍显示 independent teacher signoff pending。", "可把 Phoebe 的高中教材抽审作为教师签核的一部分。"),
        ("小学教材形态", "小学 lesson 由 topic/question seed 生成，缺少像高中那样完整的独立教材 QA 报告。", "建议抽 6 个代表单元做完整教材 QA：P1、P2、P3、P4、P5、P6 各 1 个。"),
        ("候选版本混淆", "高中历史上有 V1/V2/V3/V4 多个候选；当前主线是 V2，V4 早期 blocked 结论不可再当最终口径。", "对外与专家沟通统一称“高中当前主线 V2 题库”。"),
        ("工程验证补齐", "题目 QA 与内容 QA 已较强，但正式发布还应补 type-check/build/browser smoke 与全量 MVP gate 的最新记录。", "发布前由 S11/S10 统一跑 release gate。"),
    ]
    add_table(doc, ["风险/缺口", "现象", "建议动作"], risk_rows, [1700, 4100, 3560], font_size=9)

    add_heading(doc, "七、给 Phoebe 的建议审阅清单", 1)
    add_bullet(doc, "课程适配：确认沪教版章节顺序、学期归属、P3 旧课本暂用口径、高中 S4-S6 对应是否合理。")
    add_bullet(doc, "数学严谨性：优先抽查 Challenge / Exam 难度题、几何判定题、概率统计题、函数与解析几何题。")
    add_bullet(doc, "语言与术语：检查简体中文是否符合大陆课堂表达，避免港式/英式残留，避免“AI生成感”的模板化表达。")
    add_bullet(doc, "教学完整度：判断 lesson 是否包含足够的概念引入、例题、误区、练习和迁移，而不只是题目包装。")
    add_bullet(doc, "源距离与原创性：确认内容没有复刻教材/试卷原题结构，尤其是测评模式卡驱动的题目。")
    add_bullet(doc, "视觉题规范：小学配图应由程序化数学图层控制关键数字、坐标、标签和图形，图片模型只负责情境背景。")

    doc.add_page_break()
    add_heading(doc, "八、证据索引", 1)
    add_para(doc, "以下为项目内可追溯证据，供需要时由团队打开核对。")
    evidence_rows = [
        ("小学题目发布结论", "coordination/content-qa/mainland-hjb-primary-generated-bank-v1-1500/s18-promotability-decision.md"),
        ("初中题目发布结论", "coordination/content-qa/mainland-hjb-junior-generated-bank-v2-1500/s18-promotability-decision.md"),
        ("高中 V2 题目质量审计", "coordination/content-qa/mainland-hjb-high-generated-bank-v2/quality-audit.md"),
        ("高中教材 QA acceptance", "coordination/content-qa/mainland-hjb-high-lessons-v1/qa-acceptance-report.md"),
        ("小学插图需求统计", "coordination/content-qa/hjb-primary-illustration-audit/summary.md"),
        ("高中安全 RAG crosswalk", "coordination/content-qa/mainland-hjb-high-textbooks-v1/crosswalk.md"),
    ]
    add_table(doc, ["证据", "路径"], evidence_rows, [2400, 6960], font_size=8.8)

    add_heading(doc, "九、建议的下一步", 1)
    add_bullet(doc, "安排 Phoebe 先做 1 小时口径审阅：确认课程映射、教材形态和抽样策略。")
    add_bullet(doc, "准备专家抽样包：小学 30 题、初中 30 题、高中 40 题，加 6 个 lesson 样本。")
    add_bullet(doc, "若 Phoebe 认可方向，再启动小学 A_required 插图生产与数学图层规范。")
    add_bullet(doc, "发布前由 S11/S10 做一次统一 release gate：type-check、build、浏览器 smoke、Practice/Lesson 专项检查。")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_doc()
    print(out)
