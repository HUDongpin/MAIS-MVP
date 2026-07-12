from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-10-president-report.docx")
REPORT_DATE = "2026-07-10"
WINDOW = "2026-07-09 08:00 to 2026-07-10 08:00 Asia/Hong_Kong"
FORMAT_VERSION = "President Report Business Brief v1"

# Bilingual rendering override: this installed sans font reliably covers both
# Latin and Simplified Chinese in Word and the required LibreOffice QA pass.
FONT_LATIN = "Hiragino Sans GB"
FONT_CJK = "Hiragino Sans GB"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().get_or_add_rFonts()
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        size, color = 16, BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        size, color = 13, BLUE
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        size, color = 12, DEEP_BLUE
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.2, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header)
        set_run_font(r, size=9.1, bold=True)
        set_cell_shading(cell, header_fill)
    row_fills = row_fills or {}
    for row_index, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        fill = row_fills.get(row_index)
        for idx, value in enumerate(row):
            cell = table_row.cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx in (0, 1) and len(str(value)) < 24:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=body_size)
            if fill:
                set_cell_shading(cell, fill)
    set_table_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=10.1)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=10.1)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    hr = header.add_run("MAIS-MVP | Daily coordination report | A01-A25")
    set_run_font(hr, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    fr = footer.add_run(f"{REPORT_DATE} | {FORMAT_VERSION} | Page ")
    set_run_font(fr, size=8.5, color=MUTED)
    add_page_field(footer)


def add_title_block(doc):
    prepared = datetime.now(ZoneInfo("Asia/Hong_Kong")).strftime("%Y-%m-%d %H:%M HKT")
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", REPORT_DATE),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Prepared during run", prepared),
        ("Reporting window", WINDOW),
        ("UTC filter window", "2026-07-09 00:00 to 2026-07-10 00:00 UTC"),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Format", FORMAT_VERSION),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL, body_size=8.7)


def agent_rows():
    no_log = "窗口内无 fresh session log。"
    statuses = {
        "A01": ("mtime-only", f"{no_log} app/login/page.tsx 与 LanguageToggle.tsx 在窗口内变化。", "需 A01/A09 确认来源、范围与可见行为。", "无 A01 targeted check。"),
        "A02": ("无 fresh 日志", no_log, "未发现 dashboard/progress 责任域窗口内 mtime 证据。", "无 A02 check。"),
        "A03": ("无 fresh 日志", no_log, "未发现 roadmap/topic 责任域窗口内 mtime 证据。", "无 A03 check。"),
        "A04": ("mtime-only", f"{no_log} app/practice/page.tsx 与 questions 路径有变化。", "需与 A12/A18/A21 区分 API、题库与 live integration 责任。", "无 A04 targeted check。"),
        "A05": ("mtime-only", f"{no_log} lessons 数据与 162 个 lesson illustration assets 变化。", "不得视为 live-ready；需 A18/A21/A24 provenance、QA 与 promotion gate。", "无 A05 lesson rerun。"),
        "A06": ("mtime-only", f"{no_log} 4 个 three/manim runtime 文件变化。", "需 A06 isolated validation；A22 top candidate type-check 仍红。", "无 A06 visualization test。"),
        "A07": ("无 fresh 日志", no_log, "无 AI Tutor provider 行为或 credential placement 证据。", "无 live provider check。"),
        "A08": ("mtime-only / 红门", f"{no_log} AppProviders.tsx 变化；fresh analytics 为绿。", "Fresh root type-check 的 QuestionDiagram narrowing 需 A08/A11 路由。", "Analytics 27/27；root type-check 红。"),
        "A09": ("mtime-only", f"{no_log} LanguageToggle.tsx 变化。", "需 copy/accessibility owner 确认，与 A01 协调。", "无 A09 check。"),
        "A10": ("Active", "窗口内生成 2026-07-09 president report；本次生成 2026-07-10 report。", "仅 reporting/coordination；未改 feature code。", "DOCX build、render、结构检查。"),
        "A11": ("mtime-only / 红门", f"{no_log} 1 个 lesson E2E 文件变化。", "Fresh root type-check 另在 practice-bank-solvability.spec.ts 报错；需 regression owner 处理。", "未运行 Playwright。"),
        "A12": ("mtime-only / work order", f"{no_log} API、questionStore、userStore 等 9 个 backend 文件变化。", "A22 work order 包含 A12；需要 handoff、backend tests 与 build evidence。", "Fresh root type-check 红；无 backend suite。"),
        "A13": ("mtime-only / work order", f"{no_log} teacher console 约 9 个文件变化。", "A22 work order 包含 A13；需 targeted teacher checks 与 session log。", "无 A13 check。"),
        "A14": ("无 fresh 日志", no_log, "未发现 parent console 责任域窗口内 mtime 证据。", "无 A14 check。"),
        "A15": ("无 fresh 日志", no_log, "无 fresh adaptive engine 行为变更；只消费 analytics 结果。", "Analytics 27/27。"),
        "A16": ("未归属报告", f"{no_log} Technical-Review 下 3 个报告文件变化。", "需确认作者、结论状态与是否进入正式决策链。", "文档未单独复核。"),
        "A17": ("无 fresh 日志", no_log, "无 reward economy / motivation 责任域窗口内证据。", "无 A17 check。"),
        "A18": ("mtime-only / QA 缺口", f"{no_log} 内容、lesson illustrations 与 3 个候选题包变化。", "尚无 A18 independent QA / acceptance log；不得 promotion。", "无 A18 QA rerun。"),
        "A19": ("无 fresh 日志", no_log, "无 API env、Vercel variables 或 provider readiness 证据。", "无 live env check。"),
        "A20": ("无 fresh 日志", no_log, "无 game loop / game data 责任域窗口内 mtime 证据。", "无 A20 game check。"),
        "A21": ("mtime-only / candidate", f"{no_log} 3 个 generated question packs 与 illustration assets 变化。", "需 A18 signoff、A23 promotion plan、A04/A05 live owner、A11/A22 gates。", "无 package-local A21 check。"),
        "A22": ("Evidence active", "无 direct session log；A25 evidence 记录 top 与 fallback candidate validation。", "Top candidate 红；fallback codex/s22-release-hygiene-2026-06-15 type-check/build 绿但未选。", "Fallback type-check/build pass；top smoke 2/2、TS 360 红。"),
        "A23": ("无 fresh 日志", no_log, "无 candidate-to-live promotion record；内容链不得绕过。", "无 A23 check。"),
        "A24": ("mtime-only / exact-layer 缺口", f"{no_log} 162 个 lesson illustration assets 与 illustration data 变化。", "需 deterministic exact-layer/provenance handoff 与 A18 final QA。", "无 A24 validator。"),
        "A25": ("Active", "2 个 fresh logs；747 个 release-intake 文件；已输出 remediation routing、owner work orders 与 fallback validation evidence。", "Validate/merge 仍 blocked；42 canonical rows pending，cleanup/executable rows 为 0。", "In-window gates passed；fresh dirty-map 7,211；no staged。"),
    }
    return [(f"A{i:02d}", *statuses[f"A{i:02d}"]) for i in range(1, 26)]


def build_doc():
    doc = Document()
    setup_document(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "窗口内有 fresh AI activity：2 个 session logs（A10、A25）与 747 个 A25/A22 release-intake files；不是空窗口。"),
            ("2", "A22 证据出现可行替代路线：top candidate codex/A22-us-region-alignment 仍有 360 条 TS error lines 且 build 需刷新；fallback codex/s22-release-hygiene-2026-06-15 已在 clean validation 中 type-check/build 通过，但尚未被选为 release source。"),
            ("3", "窗口内另有 202 个非 coordination 文件出现 mtime 变化，包括 teacher/backend、lessons、visualization、题包与 162 个 lesson illustration assets；没有相应 fresh owner logs，因此只能标为 mtime-only，不能当作完成或签核。"),
            ("4", "Fresh post-window gates：A25 dirty-map 刷新通过，root 现有 7,211 expanded entries、0 staged；analytics 27/27 通过；root type-check 因 QuestionDiagram narrowing 在 2 个测试文件出现 22 条 diagnostics 而失败。"),
            ("5", "Dr. Peter Hu 的最小决策集：保持 no-dirty-root-deploy；决定 A22 修 top candidate 还是评审绿色 fallback；处理 validation hold/source staleness；要求 mtime-only 变更由 owning agents 补 log、checks 与 promotion evidence。"),
        ],
        [700, 8660],
        body_size=8.6,
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "The window contains fresh AI activity: two session logs (A10 and A25) and 747 A25/A22 release-intake files. The no-work fallback does not apply."),
            ("2", "A22 evidence now offers a viable alternative path. The top candidate still has 360 TypeScript error lines and needs a build refresh; fallback codex/s22-release-hygiene-2026-06-15 passed clean type-check and build validation but has not been selected as the release source."),
            ("3", "Another 202 non-coordination files changed by mtime, including teacher/backend, lessons, visualization, question packages, and 162 lesson illustration assets. Without fresh owner logs, these are evidence-only changes, not accepted completion."),
            ("4", "Fresh post-window gates show 7,211 expanded dirty entries and zero staged entries. Analytics passed 27/27. Root type-check failed with 22 diagnostics in two test files because QuestionDiagram union members were accessed without narrowing."),
            ("5", "Dr. Peter Hu should keep the dirty root out of deployment, choose top-candidate repair versus fallback review, resolve the validation-hold/source-staleness path, and require owner logs and checks for the unlogged file clusters."),
        ],
        [700, 8660],
        body_size=8.6,
    )

    add_heading(doc, "报告窗口摘要（中文）", 1)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ("窗口", "2026-07-09 08:00 至 2026-07-10 08:00 Asia/Hong_Kong。"),
            ("Fresh session logs", "2 个：coordination/session-logs/2026-07-09-A10.md 与 coordination/session-logs/2026-07-07-A25.md。"),
            ("Blockers / decisions / integration / content QA", "四个目录在窗口内均无新增文件。"),
            ("Reports", "3 个：2026-07-09-president-report.docx 与 2 个 builder cache 文件。"),
            ("Release intake", "747 个文件；A25/A22 形成 owner work orders、clean-source queue、fallback validation 与 validate/merge evidence。"),
            ("非 coordination mtime 证据", "202 个文件：162 public assets、11 components、7 data、7 app、6 lib、3 Technical-Review、1 tests，加 5 个 root/local 文件。"),
            ("空窗口 fallback", "不适用；本窗口存在 A10/A25 日志与 A25/A22 evidence activity。"),
            ("Post-window fresh checks", "Dirty-map、analytics、no-staged、git diff --check 通过；root type-check 失败。"),
        ],
        [2450, 6910],
        body_size=8.3,
    )

    add_heading(doc, "整体项目进展（中文）", 1)
    add_bullet(doc, "A25-owned release intake 已把 A22 top candidate 的 360 条 TS errors 路由给 A06/A08/A10/A12/A13，并生成 5 份 owner work orders；这些 work orders 不授权直接修改 candidate。")
    add_bullet(doc, "A22-owned fallback candidate codex/s22-release-hygiene-2026-06-15 在 clean snapshot 中 type-check 与 build 均通过，tracked status 前后为 0；但 releaseSourceSelected=false，不能视为发布决定。")
    add_bullet(doc, "A10-owned daily reporting 连续性保持；本次沿用 President Report Business Brief v1，且不改 feature code。")
    add_bullet(doc, "内容/lesson asset、teacher/backend 与 visualization 变更缺少同窗口 owner logs；项目进度必须分开记录“文件变化”和“已验证完成”。")
    add_bullet(doc, "Root main 继续是 inventory/reporting source；fresh dirty inventory 已增至 7,211 expanded entries，不具备 release eligibility。")

    add_heading(doc, "A01-A25 会话状态表（中文）", 1)
    add_table(
        doc,
        ["ID", "状态", "窗口内活动", "当前重点 / 阻塞", "Checks"],
        agent_rows(),
        [620, 1180, 3000, 3020, 1540],
        body_size=7.0,
        row_fills={9: OK_FILL, 21: PENDING_FILL, 24: OK_FILL},
    )

    add_heading(doc, "已完成工作（中文）", 1)
    add_bullet(doc, "A10 在窗口内完成 2026-07-09 president report DOCX 与结构验证；本次继续固定 A01-A25 全表和双语摘要格式。")
    add_bullet(doc, "A25 在 fresh log 中完成 top-candidate typecheck remediation routing、currentness stale-loop closeout 与 5-owner work-order packaging；相关 evidence gates 均保持 fail-closed。")
    add_bullet(doc, "A22 fallback candidate snapshot 记录 type-check 通过、build 通过、tracked mutation 未发生；该结果是候选验证完成，不等于 release-source selection。")
    add_bullet(doc, "A25/A22 clean-source validation queue 已明确：top candidate gates red；8 个 fallback 中 1 个绿色，另外 7 个 sweep validation failed。")
    add_bullet(doc, "Fresh analytics deterministic suite 27/27 通过；fresh no-staged gate 与 git diff --check 通过。")

    add_heading(doc, "进行中工作（中文）", 1)
    add_bullet(doc, "A25 closure 仍在 validate；42 个 canonical authorization rows 待处理，当前 focus row 为 1，cleanupAuthorizedRows=0，executableRows=0。")
    add_bullet(doc, "A22 需在修复 top candidate 与评审绿色 fallback 之间形成明确 clean-source selection review；目前 releaseSourceSelected=false。")
    add_bullet(doc, "A08/A11 需在 isolated scope 中处理 QuestionDiagram union narrowing 与相关 test typing；本次报告不实施修复。")
    add_bullet(doc, "A12/A13 需为 backend/teacher mtime-only 变更补 session log、targeted tests 与 build evidence。")
    add_bullet(doc, "A05/A18/A21/A24 需为 lesson data、3 个 generated packs 与 162 个 illustration assets 完成 provenance、QA、exact-layer 与 promotion handoff。")

    add_heading(doc, "阻塞项（中文）", 1)
    add_table(
        doc,
        ["Blocker", "Evidence", "Owner / next action"],
        [
            ("Dirty root 不可发布", "Fresh map: 1,499 collapsed；7,211 expanded；401 modified；1 deleted；6,809 untracked files；0 staged。", "A22/A25/A10：只从 clean worktree/clone/reviewed slice 发布。"),
            ("Fresh root type-check 红", "22 条 diagnostics，位于 lib/mvpReadiness.test.ts 与 tests/e2e/practice-bank-solvability.spec.ts；TS2339/TS7006。", "A08/A11：在 isolated worktree 修 narrowing 与 test types，随后复跑 type-check。"),
            ("Top candidate 红", "codex/A22-us-region-alignment：focused smoke 2/2；type-check 360 error lines；build refresh required。", "A22 + A06/A08/A10/A12/A13：修复或停止投资。"),
            ("Fallback 未选", "codex/s22-release-hygiene-2026-06-15 type-check/build 绿，但 releaseSourceSelected=false。", "Owner/A22：先做 selection review，再决定是否成为 release source。"),
            ("Validation hold / source stale", "Gate 为 blocked-source-stale；worktree 仍 registered，路径存在，记录 896 status entries。", "Owner/A25/A22：确认 prior confirmation 与当前 worktree 状态，禁止直接 cleanup。"),
            ("未归属 mtime 变更", "202 个非 coordination 文件变化但无 fresh owner logs。", "相关 A01/A05/A06/A08/A11/A12/A13/A18/A21/A24 owners 补 provenance 与 checks。"),
        ],
        [1900, 4200, 3260],
        body_size=7.5,
        row_fills={0: RISK_FILL, 1: RISK_FILL, 2: RISK_FILL, 4: PENDING_FILL, 5: PENDING_FILL},
    )

    add_heading(doc, "风险（中文）", 1)
    add_bullet(doc, "Release risk：dirty root 与绿色 fallback 的证据层不同；不能用 root analytics 通过替代 candidate build/release proof。")
    add_bullet(doc, "Provenance risk：mtime-only 文件可能来自用户或其他 agent；没有 log、diff review 与 checks 时不能归属、合并或 promotion。")
    add_bullet(doc, "Content risk：3 个 generated packs 与 162 个 illustration assets 若绕过 A18 -> A24 -> A23 -> A04/A05/A11/A22 链，会把候选内容误当 live content。")
    add_bullet(doc, "Coordination risk：A22 top candidate、绿色 fallback、42-row authorization backlog 与 validation hold 同时存在，容易让单项绿灯被误读为整体 ready。")
    add_bullet(doc, "Generated-state risk：post-window residual evidence 检测到 .next/.tmp active writers；当前不应授权 cleanup apply。")

    add_heading(doc, "测试 / 构建状态（中文）", 1)
    add_table(
        doc,
        ["Check", "Result", "Interpretation"],
        [
            ("Fresh npm run type-check", "FAIL：22 diagnostics；2 个测试文件；TS2339/TS7006。", "Root 当前不再 type-green；需要 A08/A11 owner fix。"),
            ("Fresh npm run test:analytics", "PASS：27 tests passed，0 failed。", "A08/A15 deterministic analytics/adaptive checks 绿色。"),
            ("Fresh A25 dirty-map", "PASS：7,211 expanded entries。", "作为 post-window 最新 inventory baseline；不代表 release ready。"),
            ("Fresh no-staged gate", "PASS：stagedEntryCount=0。", "报告运行未 staging。"),
            ("Fresh git diff --check", "PASS：无输出。", "未发现 whitespace diagnostics。"),
            ("A25 in-window currentness", "PASS：session log 记录 796/796 与 153/153 gates，failures 0。", "Release-intake evidence 当时内部一致。"),
            ("A22 top candidate", "Smoke 2/2 PASS；type-check FAIL 360；build refresh required。", "不能 promotion。"),
            ("A22 fallback candidate", "Type-check PASS；build PASS；tracked mutation=false。", "可进入 clean-source selection review，但尚未被选。"),
            ("Fresh npm run build / Playwright by A10", "Not run。", "A10 report-only；root type-check 已红，dirty-root build/E2E 不能证明 release readiness。"),
        ],
        [2500, 2850, 4010],
        body_size=7.5,
        row_fills={0: RISK_FILL, 1: OK_FILL, 6: RISK_FILL, 7: OK_FILL},
    )

    add_heading(doc, "变更文件（中文）", 1)
    add_table(
        doc,
        ["Category", "Count / examples", "Meaning"],
        [
            ("窗口扫描总量", "954 files（排除 .git、node_modules、.next、.tmp、coverage、Playwright outputs）。", "Mtime 证据；不等于 Git ownership 或 completion。"),
            ("coordination", "752：747 release-intake、2 session logs、3 report artifacts。", "主要 fresh logged activity 为 A10/A25/A22 evidence control plane。"),
            ("public assets", "162：mainland-hjb-primary 与 mainland-pep-junior lesson illustration concept/worked-example PNG。", "A05/A18/A21/A24 provenance 与 QA gate 未在窗口日志中出现。"),
            ("app/components", "18：login、practice、teacher、AppProviders、LanguageToggle、4 个 three/manim files。", "跨 A01/A04/A06/A08/A09/A12/A13；需 owner handoff。"),
            ("data/lib/tests", "14：lessons、illustration data、3 generated packs、questionStore/userStore、MVP/teacher tests、1 lesson E2E。", "跨 A05/A08/A11/A12/A18/A21/A24；当前 root type-check 红。"),
            ("Technical / root reports", "6：3 Technical-Review files + 3 root DOCX reports。", "无 fresh owner log；需确认 status、source 与决策用途。"),
            ("Local-only", "2：.local/hk-math-db.sqlite、.DS_Store。", "不得进入 release source 或业务成果统计。"),
            ("Fresh current owner buckets", "A25 5,287；A10 486；A06 460；A12 178；A24 168；A05 129；A21 76。", "Dirty inventory 仍需按 owner/slice review；不得整体 merge。"),
        ],
        [1900, 4460, 3000],
        body_size=7.5,
    )

    add_heading(doc, "明日优先事项（中文）", 1)
    add_numbered(doc, "A22 完成 clean-source selection review：明确修 top candidate，或把 codex/s22-release-hygiene-2026-06-15 提交为候选 release source；两条路线不可同时含糊推进。")
    add_numbered(doc, "A08/A11 在 isolated worktree 修复 QuestionDiagram narrowing 与 test typing，并复跑 root/package type-check。")
    add_numbered(doc, "A12/A13 对 backend/teacher mtime-only cluster 补 session log、targeted backend/teacher tests 与 build evidence。")
    add_numbered(doc, "A05/A18/A21/A24 对 lesson data、3 个 generated packs 与 162 assets 完成 provenance、independent QA、exact-layer 与 promotion packet。")
    add_numbered(doc, "A25 在 owner decision 后刷新 42-row authorization backlog、validation hold 与 validate-to-merge gates；继续保持 0 cleanup/executable rows。")

    add_heading(doc, "需要 Owner 决策（中文）", 1)
    add_bullet(doc, "请 Dr. Peter Hu 确认：dirty root 继续只作 inventory/reporting source，不作为 production deploy source。")
    add_bullet(doc, "请决定 A22 路线：继续修 codex/A22-us-region-alignment，或优先评审已 type-check/build 通过的 codex/s22-release-hygiene-2026-06-15。")
    add_bullet(doc, "请确认 validation-hold prior confirmation 与当前事实是否一致：worktree 仍 registered、路径仍存在且有 896 status entries；在一致性恢复前不授权 cleanup。")
    add_bullet(doc, "请批准或暂缓当前 1 个 focus authorization row，并指定 42 个 pending canonical rows 的下一批最小范围。")
    add_bullet(doc, "请要求 A01/A05/A06/A08/A11/A12/A13/A18/A21/A24 对 202 个 mtime-only 文件分别补 owner log、checks 和 keep/quarantine/promotion 决定。")
    add_bullet(doc, "请保持 generated-artifact cleanup 暂缓：post-window .next/.tmp 有 active writers，cleanup apply 当前不安全且未授权。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
