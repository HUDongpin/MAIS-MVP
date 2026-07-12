from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-06-06-president-report.docx")


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "111827"
MUTED = "4B5563"
LIGHT = "F2F4F7"
CALLOUT = "F8FAFC"
BORDER = "C9D1D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    font = run.font
    font.name = "Calibri"
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_p(doc: Document, text: str = "", style: str | None = None, bold=False, color=None) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_font(run, bold=bold, color=color)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_font(run, bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_font(run, bold=True, color=INK)
        set_cell_shading(hdr[i], LIGHT)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_font(run, size=9.2)
    set_table_width(table, widths)
    doc.add_paragraph()


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, before, after, color in [
        (1, 16, 16, 8, BLUE),
        (2, 13, 12, 6, BLUE),
        (3, 12, 8, 4, DARK_BLUE),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("MAIS-MVP Daily President Report | S10 Automation")
    set_font(run, size=8.5, color=MUTED)


def build() -> None:
    doc = Document()
    configure_doc(doc)
    doc.core_properties.title = "MAIS-MVP President Report - 2026-06-06"
    doc.core_properties.subject = "Daily AI coordination report"
    doc.core_properties.author = "S10 Codex automation"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("MAIS-MVP 每日总裁报告 / Daily President Report")
    set_font(run, size=20, bold=True, color=INK)
    add_p(doc, "Report date: 2026-06-06 08:00 HKT | Reporting window: 2026-06-05 08:00 to 2026-06-06 08:00 Asia/Hong_Kong", color=MUTED)
    add_p(doc, "Prepared for: Dr. Peter Hu | Prepared by: S10 Tooling, Docs, and Executive Reporting Automation", color=MUTED)

    add_heading(doc, "中文执行摘要", 1)
    add_bullets(doc, [
        "本窗口不是空窗期：S04/S11 完成 Arkansas G6-G12 题库 live adapter 与回归；S04/S23 完成 California Math Practice Beta 3000 题本地集成、发布检查，并由 S23 推广到 mais.hk 生产环境。",
        "S19 已清除两个关键生产环境门：DeepSeek 文本 AI 在生产环境为 live/configured；Neon Postgres 已连接 Production/Preview，注册/登录持久化 smoke 通过。",
        "S07 解决 AI Tutor 语音失败后无法重试、Nova 语音回声被麦克风收录两个问题，并新增 Nova Lens 企业级全局选择/治理/API 能力，目标检查与 build 通过。",
        "S18 将 Mainland BNU junior lessons 从 77 个 P0/P1 DeepSeek 问题降到 0，并给出 approved-for-lesson-integration-review；仍需 S05/S11/S23 才能进入 live lessons。",
        "S11 产出 Parent Console 10 个 P1/P2 缺陷报告和修复回归门；Student 与 handwriting board 审计仍在进行中，未见最终报告。",
        "主要风险是 dirty worktree 极大，生产推广使用 pruned staging snapshot 而不是 clean Git commit；必须尽快由 S25/S10/S22 做发布切片和回滚证据整理。",
    ])

    add_heading(doc, "English Executive Summary", 1)
    add_bullets(doc, [
        "The day delivered visible product progress: California Math Practice Beta is now promoted to production, and Arkansas G6-G12 live-bank integration passed targeted regression.",
        "Production platform readiness improved materially: S19 configured DeepSeek text AI and durable Neon Postgres for Production/Preview, then passed registration/login persistence smoke.",
        "AI Tutor reliability advanced through voice retry and echo fixes; Nova Lens enterprise agent surfaces, policy routes, audit history, and focused tests were added.",
        "Content/research output was heavy but mostly candidate-gated: BNU junior lessons passed final S18 QA for integration review; Florida/New York textbook packages and illustration batches remain review-only.",
        "Release risk remains high because the tree has thousands of dirty/generated files and the production promotion was not tied to a clean commit or PR slice.",
    ])

    add_heading(doc, "报告窗口摘要", 1)
    add_bullets(doc, [
        "窗口：2026-06-05 08:00 至 2026-06-06 08:00（香港时间）。",
        "新鲜 session logs：S04、S05、S07、S10、S11、S16、S18、S19、S21、S22、S23。",
        "新鲜 blocker records：S18 AR live import authorization、S19 Postgres/Neon 条款与 URL（已解决）、S04 California preview deploy timeout（之后由 S23 staging/production promotion 绕过）。",
        "本次 S10 新鲜检查：`npm run type-check` 通过；`npm run test:analytics` 通过 23/23。",
    ])

    add_heading(doc, "整体项目进展", 1)
    add_bullets(doc, [
        "生产端：mais.hk 已服务含 California Math Practice Beta 文案的 promoted deployment；生产 JS chunk scan 显示 `California Math Practice Beta` 存在，旧 `California Curriculum` 不存在。",
        "平台端：DeepSeek 文本 AI、图像、语音在生产状态均 configured；Postgres 持久化配置已覆盖 Production/Preview。",
        "学习内容端：California 3000 题、Arkansas K-G12 3000 题、BNU primary 已有 integration evidence；BNU junior lessons 已过 S18 final QA，但尚未 live integration。",
        "质量端：TypeScript、analytics/adaptive、question-bank、MVP、局部 Playwright、默认 build 多项检查在窗口内通过；Parent Console 仍有 10 个 P1/P2 待修。",
        "发布治理端：S22 清理默认 `.next` stale build failure 与 Next runtime re-export warnings；但 Git 工作区仍极脏，需要 release intake。",
    ])

    session_rows = [
        ["S01", "无新鲜日志", "未见窗口内独立 S01 交付。", "保持 app shell 稳定；避免与 S23/S07 全局 UI 改动冲突。"],
        ["S02", "无新鲜日志", "Dashboard/adaptive surface 被 S23 间接触及。", "如要开放 P1 loop，需 S02 审 UI 与 fallback。"],
        ["S03", "无新鲜日志", "Roadmap/curriculum 结构无独立 handoff。", "等待内容集成任务。"],
        ["S04", "完成/部分阻塞", "AR G6-G12 live adapter + regression green；CA 3000 题本地集成、build/E2E green；preview deploy timeout。", "确认生产后监控；避免把 beta 称为完整课程。"],
        ["S05", "完成", "California textbook type mismatch 已不复现；type-check 与 question-bank rerun 通过。", "接手 BNU junior lesson integration review。"],
        ["S06", "无新鲜日志", "Visualization 无独立交付。", "等待后续 QA 或 value checks。"],
        ["S07", "完成", "AI Tutor voice retry、voice echo suppression、Nova Lens enterprise agents/API/governance 已完成。", "需 live provider QA 与更广 E2E release gate。"],
        ["S08", "无新鲜日志", "Shared types/user-store 被 S07/S23 间接扩展。", "如继续 P1 world-model loop，需 S08 contract review。"],
        ["S09", "无新鲜日志", "无新鲜 copy/i18n handoff；多个 candidate packages 需要后续 S09。", "审 California/Florida/NY/BNU public copy。"],
        ["S10", "完成", "前日报告、README/AGENTS/API-key memory、S23-S25 协调扩展、automation rename；本报告生成中。", "继续 release notes 与 President Report。"],
        ["S11", "进行中", "AR regression green；Parent Console 10 个 P1/P2 报告与 regression gate；Student/handwriting audit 开始。", "等待 S14/S12 修复授权；完成 student/handwriting final reports。"],
        ["S12", "无新鲜日志", "Parent/API/storage 多处被 S19/S23/S11 触及或点名。", "处理 parent API bugs 与 production auth transport flakiness。"],
        ["S13", "无新鲜日志", "Teacher review/operations 被 S07/S23/S22 触及。", "审 teacher operations/governance UX 与 release scope。"],
        ["S14", "无新鲜日志", "Parent Console 被 S11 审计并发现 10 个 P1/P2。", "需要 owner 授权修复 parent UI/workflow。"],
        ["S15", "无新鲜日志", "Adaptive guards/API 被 S23 触及；analytics/adaptive tests green。", "审 P1 loop/adaptive semantics。"],
        ["S16", "完成研究切片", "World-model literature/feasibility、P0 fraction simulator、P1 shadow demo synthetic evidence 完成。", "不得进入 real-student；下一步需 owner 分配实现 owners。"],
        ["S17", "无新鲜日志", "Gamification/reward economy 无独立新交付。", "如 games/rewards上线需重新协调。"],
        ["S18", "完成/待集成", "AR QA handoff；TX high illustration QA；BNU junior lesson final QA approved for integration review；P0 simulator S18 review complete。", "S05/S11/S23 集成；S18 可做 P2 polish。"],
        ["S19", "完成", "DeepSeek Production configured；Neon Postgres Production/Preview configured；production persistence smoke passed。", "需 owner 提供/批准 admin smoke 方式。"],
        ["S20", "无新鲜日志", "Game surfaces 无独立新交付。", "等待 game-specific release task。"],
        ["S21", "高产/候选态", "Florida HS textbooks、NY HS/MS textbooks、TX/BNU illustration batches、exact-layer gates；多项 remain candidate-only。", "S18/S09/S05/S11 review before live。"],
        ["S22", "完成", "默认 build stale `.next` 问题清理；runtime re-export warnings 消除；build/type/analytics/E2E spot checks green。", "继续 release hygiene 与 clean deployment path。"],
        ["S23", "完成/高风险", "P1 platform loop；BNU primary copy；California beta preview + production promotion ready/verified。", "需要 S25/S22 整理 dirty snapshot 与 release record。"],
        ["S24", "无新鲜日志", "Exact-layer lead 无独立日志；S21/S18 多次点名 exact-layer gate。", "接手 deterministic overlays when assigned。"],
        ["S25", "无新鲜日志", "Git hygiene lead 新增后暂无独立 intake。", "应尽快做 dirty-tree inventory 与 PR slice plan。"],
    ]
    add_heading(doc, "S01-S25 Session 状态表", 1)
    add_table(doc, ["Session", "状态", "窗口证据", "下一步"], session_rows, [720, 1180, 4460, 3000])

    add_heading(doc, "已完成工作", 1)
    add_bullets(doc, [
        "S04/S11：Arkansas G6-G12 候选包 live adapter 集成并通过 `npm run type-check`、`npm run test:question-bank` 69/69、`npm run test:mvp` 29/29。",
        "S23：California Math Practice Beta promoted to production；production chunks 已验证新文案存在、旧文案不存在；UI login smoke 对 Student Shirleen 成功。",
        "S19：Vercel Production/Preview 已有 encrypted `POSTGRES_URL` 与 `HK_MATH_STORAGE_PROVIDER=postgres`；Neon resource connected；registration/login persistence smoke 通过。",
        "S07：AI Tutor microphone retry/echo fixes；Nova Lens runs/policy API、global overlay、teacher operations governance view、targeted E2E/API tests 与 build 通过。",
        "S18：BNU junior lessons final full DeepSeek QA 315 sections，0 P0/P1；manual acceptance artifact 和 final decision 已完成。",
        "S22：默认 `.next` build blocker 与 runtime re-export warnings 清理，`rm -rf .next && npm run build` 通过。",
    ])

    add_heading(doc, "进行中工作", 1)
    add_bullets(doc, [
        "S11 Parent Console 修复回归门等待 S14/S12 实施；Student P0/P1/P2 和 handwriting board 审计仍在日志阶段。",
        "S21 Texas K-G5 illustration Batch 008 因 built-in `image_gen` TooManyRequests 暂停，G3-G5 Unit 2 Lesson 3 未完成。",
        "S16 learner-world-model P0/P1 仅为 offline/synthetic research evidence；live DeepSeek/Qwen 输出与真实学生闭环均未启动。",
        "BNU junior lessons 已通过 S18 QA，但 live lesson integration、route regression、promotion planning 仍需 S05/S11/S23。",
    ])

    add_heading(doc, "阻塞项", 1)
    add_bullets(doc, [
        "S04 California preview deploy blocker：direct Vercel CLI 上传 timeout，archive 模式无 URL/record；S23 后续用 pruned staging preview + promote 绕过，但根因仍需 S22/S19 关注。",
        "S18 AR live-import blocker 已被后续 S04/S11 授权工作解决；blocker record 仍作为窗口证据保留。",
        "S19 Postgres terms/URL blockers 已解决；唯一剩余验证缺口是 authenticated admin `/api/admin/storage/health` 未跑。",
        "S21 image generation rate limit 阻塞部分 Texas/BNU illustration continuation。",
    ])

    add_heading(doc, "风险", 1)
    add_bullets(doc, [
        "发布风险最高：`git status --short` 约 3,227 条，mtime scan 约 25,062 个窗口内文件；production promotion 来自 pruned staging snapshot，不是 clean commit。",
        "Parent Console 仍有 10 个 confirmed P1/P2；其中 child context、notices/messages/thread URL state 等问题可能影响家长体验。",
        "Production auth/API transport 有零星 `socket hang up` / empty reply 证据，虽 UI login smoke 成功，但建议 S12/S19 查日志。",
        "大量 candidate textbook/illustration/content 包尚未 S18/S09/S05/S11 完整门控，不能直接宣传为 live curriculum。",
        "BNU junior final QA 有 25 个 P2 warning 被接受为非阻塞，但进入广泛公开前仍建议 polish backlog。",
    ])

    add_heading(doc, "测试 / Build 状态", 1)
    test_rows = [
        ["Fresh S10", "`npm run type-check`", "通过", "当前 dirty tree 通过。"],
        ["Fresh S10", "`npm run test:analytics`", "通过 23/23", "learningAnalytics + adaptiveLearning tests。"],
        ["S04/S11", "`npm run test:question-bank`", "通过 69/69", "Arkansas/California/NC/FL live route/question coverage。"],
        ["S04", "`npm run build`", "通过", "California beta local production build。"],
        ["S22", "`rm -rf .next && npm run build`", "通过", "默认 `.next` stale state 已清理。"],
        ["S23", "California focused Playwright", "通过 1/1", "Teacher/student assignment flow。"],
        ["S11", "Parent E2E/stress", "部分失败", "入口/login reliability 与 10 个 P1/P2 产品缺陷。"],
    ]
    add_table(doc, ["来源", "检查", "状态", "说明"], test_rows, [1000, 2740, 1200, 4420])

    add_heading(doc, "变更文件", 1)
    add_bullets(doc, [
        "Git dirty summary：约 3,227 条 tracked/untracked status entries；其中包含大量 generated content、public illustration assets、coordination artifacts 和 feature files。",
        "窗口 mtime scan：coordination/research/content/report 约 704 个文件；全树排除 `node_modules`、`.next`、`.tmp` 后约 25,062 个文件。",
        "关键 live/source areas：`data/questions.ts`、`data/topics.ts`、US CA/AR adapters、`lib/questionBankSolvability.ts`、`lib/fullQuestionBankSolvability.test.ts`。",
        "关键 AI/platform areas：`components/ai/AITutorProvider.tsx`、`components/ai/NovaLensGlobalOverlay.tsx`、`app/api/nova-lens/`、`app/api/admin/nova-lens/`、`lib/server/userStore.ts`、`types/index.ts`。",
        "关键 release/teacher/parent areas：pilot loop API、teacher review parent draft route、parent notices view、teacher route wrappers、teacher operations/governance views。",
        "关键 coordination areas：`coordination/content-qa/`、`coordination/research/`、`coordination/reports/`、`coordination/session-logs/`。",
    ])

    add_heading(doc, "明日优先事项", 1)
    add_bullets(doc, [
        "S25/S10/S22：把当前 dirty snapshot 做 release intake、owner map、PR/commit slice recommendation，明确 California production deployment 的可回滚来源。",
        "S14/S12/S11：修复 Parent Console 10 个 P1/P2，并让 S11 添加/跑对应 regression gate。",
        "S05/S11/S23：启动 BNU junior lessons integration review，保持 S18 final QA 边界，不跳过 25 P2 backlog 的公开前处理判断。",
        "S07/S11/S13/S12：给 Nova Lens 做 release-readiness review，尤其 governance policy、teacher operations UX、provider live behavior、audit redaction。",
        "S19/S12：补 authenticated admin storage health 验证方案，查 production auth/API transport flakiness。",
        "S18/S21/S09：安排 Florida/NY textbooks 与 Texas/BNU illustrations 的独立 curriculum/source/copy/accessibility review。",
    ])

    add_heading(doc, "需要 Dr. Peter Hu 决策 / 提醒", 1)
    add_bullets(doc, [
        "是否接受 California Math Practice Beta 当前生产推广的 dirty-staging 风险，或要求 S25/S22 先产出 clean release slice 再继续后续上线。",
        "是否授权 S14/S12 立即修复 Parent Console 10 个 P1/P2，并授权 S11 写入/运行对应 regression tests。",
        "是否授权 S05/S11/S23 将 S18-approved BNU junior lessons 推进 live integration review。",
        "是否提供现有 admin production session/credential，或批准临时 admin smoke account，用于 authenticated storage health 检查。",
        "是否启动 live DeepSeek/Qwen learner-world-model smoke；若启动，需 S07/S19 secret-safe runtime handling，且禁止真实学生使用。",
        "是否安排 S18/S09 对 Florida/NY textbook packages 与 Texas/BNU illustration batches 做正式发布前 review。",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
