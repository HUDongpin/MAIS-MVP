# -*- coding: utf-8 -*-
"""Build the 2026-06-10 MAIS-MVP president report DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "coordination" / "reports" / "2026-06-10-president-report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                set_cell_width(row.cells[idx], width)
                row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)


def style_cell_text(cell, bold: bool = False, color: str | None = None) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(8.5)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Aptos Display"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(31, 41, 55)


def add_para(doc: Document, text: str = "", bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix) :])
        runs = [r1, r2]
    else:
        runs = [p.add_run(text)]
    for run in runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9.5)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.2)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        set_cell_shading(hdr[i], "E5E7EB")
        style_cell_text(hdr[i], bold=True)
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, value in enumerate(row_data):
            row_cells[i].text = value
            style_cell_text(row_cells[i])
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(9.5)
    for style_name in ["List Bullet", "List Number"]:
        styles[style_name].font.name = "Aptos"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[style_name].font.size = Pt(9.2)


def build() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MAIS-MVP President Report")
    r.bold = True
    r.font.name = "Aptos Display"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(17, 24, 39)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Daily AI Coordination Report for Dr. Peter Hu")
    sr.font.name = "Aptos"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor(75, 85, 99)

    metadata = [
        ["Report date", "2026-06-10"],
        ["Report time", "8:00 AM Asia/Hong_Kong"],
        ["Reporting window", "2026-06-09 08:00 to 2026-06-10 08:00 Asia/Hong_Kong"],
        ["Project", "MAIS-MVP"],
        ["Reporting session", "S10 recurring president-report automation"],
        ["Audience", "Dr. Peter Hu"],
    ]
    add_table(doc, ["Field", "Value"], metadata, [2200, 7160])

    add_heading(doc, "中文 Executive Summary", 1)
    for item in [
        "报告窗口内有新进展，不是空窗期：S01 完成论坛首屏重构，S04 推进 Practice Arena 游戏化视觉方向并保存三个设计方案，S09 修正美国课程注册/登录年级命名，S10 完成昨日报告收尾并运行今日基线检查。",
        "当前产品健康度为“可继续推进但需强发布控制”：`type-check`、analytics、release preflight、production build 全部通过；直接从当前根目录部署仍被 3,400 条 dirty status 保护性拦截。",
        "最大管理风险仍是脏工作树和未完成的全产品回归。今日 UI 变化集中在学生论坛、练习入口、登录/注册年级标签，需要 S11/S22 尽快补充回归确认。",
        "最需要 Dr. Peter Hu 决策的是：论坛可见发帖/回复工作台是否保持移除并改为新学生友好流程；Practice Arena 当前视觉方向是否继续深化；是否以 S22 clean/pruned staging 作为唯一发布路径直到根目录清理完毕。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "English Executive Summary", 1)
    for item in [
        "The window had fresh work: S01 rebuilt the forum first screen, S04 advanced the Practice Arena adventure UI and preserved three design directions, S09 corrected US grade labeling on registration/login, and S10 closed the prior report cycle.",
        "Fresh baseline checks are green: type-check passed, analytics tests passed 23/23, release preflight passed with 79.3 GB free, and the production build passed with 128 static pages.",
        "The main release risk remains unchanged: direct root deploy is blocked by the dirty worktree, now 3,400 status entries. Use the reviewed S22 staging path or a clean release worktree.",
        "The decisions needed today are visual/product decisions for forum posting and Practice Arena, plus release-control direction before any production promotion.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "报告窗口摘要", 1)
    summary_rows = [
        ["报告窗口", "2026-06-09 08:00 至 2026-06-10 08:00 HKT"],
        ["新鲜日志", "S01、S04、S09、S10"],
        ["新鲜 blocker", "S01 旧 type-check/build blocker 已解决；无新的未解决 S01 blocker"],
        ["No assigned work in this reporting window", "No"],
        ["协调重点", "论坛首屏、Practice Arena 游戏化视觉、美国课程年级命名、发布基线检查"],
        ["跨会话依赖", "S01/S12 论坛发帖与 roster 后端；S04/S11/S22 练习页回归；S09 后续全站美国标签审计；S22/S25 脏根目录发布治理"],
    ]
    add_table(doc, ["项目", "结论"], summary_rows, [2600, 6760])

    add_heading(doc, "整体项目进度", 1)
    for item in [
        "学生端体验继续向更清晰的年龄段差异化推进：论坛根据 primary/K/访客与 junior/high 呈现不同视觉路线，Practice Arena 进入更强游戏化的入口与任务面板。",
        "美国课程本地化细节提升：注册和登录的 US curriculum 年级显示从香港 P/S 命名切换为 K、G1-G12 与 Elementary/Middle/High 分组。",
        "工程基线短期稳定：今日 S10 复跑的 TypeScript、analytics、release preflight 和 Next build 均通过。",
        "发布成熟度仍受 dirty tree、全产品 E2E 缺口、视觉验收与跨会话 owner 决策影响。今天应优先做 release slicing 和 targeted regression，而不是继续扩大功能面。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "S01-S25 Session Status", 1)
    session_rows = [
        ["S01", "活跃-完成", "论坛首屏重构为 Math Park / Adventure Island 两套体验；移除下方 thread/reply/workbench 区域", "若需要发帖/回复，需设计新 modal 或独立流程", "旧 practice type-check blocker 已解决；primary class posting 仍是后端 roster/enrollment 后续项", "type-check、forum browser smoke、class-forum Playwright 3 passed / 3 skipped"],
        ["S02", "无新日志", "无 fresh dashboard evidence", "等待 dashboard/adaptive UI 后续任务", "无新 blocker", "未运行会话级检查"],
        ["S03", "无新日志", "无 fresh roadmap evidence", "等待课程路线任务", "无新 blocker", "未运行会话级检查"],
        ["S04", "活跃-进行中", "保存 Scheme A/B/C 设计记忆；完成 `/practice` adventure shell、topic mission、lower control polish", "继续视觉验收；QuestionCard 内部与 streak 数据源仍需后续", "`Current Streak` 是占位；全产品回归未跑", "多次 type-check 通过；desktop/mobile Playwright 无横向溢出"],
        ["S05", "无新日志", "无 fresh lesson evidence", "等待 lesson/content integration 任务", "无新 blocker", "未运行会话级检查"],
        ["S06", "无新日志", "无 fresh visualization evidence", "等待 visualization 任务", "无新 blocker", "未运行会话级检查"],
        ["S07", "无新日志", "无 fresh AI Tutor evidence", "等待 provider/tutor 任务", "无新 blocker", "未运行会话级检查"],
        ["S08", "无新日志", "无 fresh shared-state evidence", "等待 shared analytics/type task", "无新 blocker", "未运行会话级检查"],
        ["S09", "活跃-完成", "注册/登录 US 年级标签改为 K/G1-G12 与美制学段分组", "可继续审计其它 US curriculum surfaces", "in-app browser 拒绝 localhost；用本地 render/source 检查替代", "type-check 通过；register/login label checks 通过"],
        ["S10", "活跃-完成", "完成 2026-06-09 report 收尾；生成今日 2026-06-10 report；运行 fresh baseline", "报告后更新 automation memory", "DOCX render QA 取决于本地工具；根目录 deploy guard 仍失败", "type-check、analytics、preflight、build 通过；root deploy preflight 保护性失败"],
        ["S11", "无新日志", "无 fresh QA/release log", "需要对 forum/practice/login/register 做 targeted regression", "全产品 E2E 未在本窗口由 S10 运行", "未运行 S11 suite"],
        ["S12", "无新日志", "无 fresh backend/API evidence", "论坛 posting/roster/enrollment 后续需 S12 协调", "无新 blocker", "未运行会话级检查"],
        ["S13", "无新日志", "无 fresh teacher-console evidence", "等待 teacher workflow 任务", "无新 blocker", "未运行会话级检查"],
        ["S14", "无新日志", "无 fresh parent-console evidence", "等待 parent workflow 任务", "无新 blocker", "未运行会话级检查"],
        ["S15", "无新日志", "无 fresh adaptive-engine evidence", "等待 adaptive regression/US content priority", "无新 blocker", "未运行会话级检查"],
        ["S16", "无新日志", "无 fresh research evidence", "等待 research/evaluation 任务", "无新 blocker", "未运行会话级检查"],
        ["S17", "无新日志", "无 fresh gamification evidence", "Practice streak/reward source 若推进需协调", "无新 blocker", "未运行会话级检查"],
        ["S18", "无新日志", "无 fresh curriculum QA evidence", "等待内容签核/QA 任务", "无新 blocker", "未运行会话级检查"],
        ["S19", "无新日志", "无 fresh env/provider evidence", "等待 env parity 或 live provider smoke 任务", "无新 blocker", "未运行会话级检查"],
        ["S20", "无新日志", "无 fresh game-based learning evidence", "Practice/game unlock 深化需协调", "无新 blocker", "未运行会话级检查"],
        ["S21", "无新日志", "无 fresh content pipeline evidence", "等待 content/RAG package task", "无新 blocker", "未运行会话级检查"],
        ["S22", "无新日志", "无 fresh S22 log；S10 复查 release guard", "需要 clean/pruned staging 发布路径维护", "直接根目录部署被 dirty worktree 阻止", "root deploy preflight failed by design"],
        ["S23", "无新日志", "无 fresh integration evidence", "等待 promotion/readiness task", "无新 blocker", "未运行会话级检查"],
        ["S24", "无新日志", "无 fresh exact-layer evidence", "等待 illustration exact-layer task", "无新 blocker", "未运行会话级检查"],
        ["S25", "无新日志", "无 fresh git hygiene log", "需要 dirty-tree ownership map / release slicing", "3400 status entries remain", "S10 only counted status; no Git mutation"],
    ]
    add_table(doc, ["Session", "状态", "已完成", "进行中", "Blockers", "Checks"], session_rows, [750, 1050, 2450, 1850, 1900, 1360])

    add_heading(doc, "已完成工作", 1)
    for item in [
        "S01：论坛首屏重构完成，primary/K/访客使用 Math Park，junior/high 使用 Adventure Island，并保留 search/filter/pagination/API-backed browsing。",
        "S01：按 owner 后续要求移除页面下方 selected discussion、thread replies、forum inbox、class snapshot、composer 和 moderation workbench。",
        "S04：Practice Arena 三个视觉方案已保存到 `coordination/reports/2026-06-10-practice-arena-design-memory/`，作为后续 MAIS 设计记忆。",
        "S04：将 approved adventure UI 方向接入真实 `/practice`，grade chips/topic mission cards 连接现有 filter state，并把下方 controls 与 QuestionPager 改成同一视觉语言。",
        "S09：美国课程注册和登录的年级显示完成本地化，使用 Elementary/Middle/High 与 K/G1-G12，不改变内部 GradeId。",
        "S10：今日 fresh checks 完成，DOCX 报告生成流程建立。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "进行中工作", 1)
    for item in [
        "Practice Arena 仍需要 owner 视觉验收；若通过，应继续把 `PracticeQuestionCard` 内部、streak 数据源和 reward 状态接入同一 adventure language。",
        "论坛可见发帖/回复入口被移除后，若产品仍需要学生投稿，应设计新的 modal 或独立 flow，不能简单恢复旧 workbench。",
        "US curriculum 标签逻辑已覆盖 register/login，但其它 US-facing surfaces 仍建议由 S09 做一次 copy/accessibility audit。",
        "发布侧仍需 S22/S25 把 dirty root 切片、映射 owner，并保持 clean/pruned staging 作为可控发布路径。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Blockers", 1)
    for item in [
        "直接根目录部署仍被保护性拦截：3,400 status entries，257 tracked modified，4 tracked deleted，3,139 untracked status entries，32,405 untracked files。",
        "全产品 E2E/release regression 本窗口未运行；论坛与 Practice Arena 都有可见 UI 变化，不能仅凭 type-check/build 视为 release complete。",
        "论坛发帖/回复工作流的可见入口已按 owner 要求移除；是否恢复为新流程需要 owner 产品决策。",
        "S01 旧 blocker 已解决，不再阻塞：`components/practice/PracticeArenaIslandHero.tsx` 的 out-of-scope type-check 问题不在当前工作树复现。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "风险", 1)
    for item in [
        "脏工作树过大，使任何直接部署、review、回滚和责任归属都存在高风险。",
        "S04 的 `Current Streak` 是视觉占位，若对学生展示为真实指标会造成信任风险；上线前需接入 durable source 或改文案。",
        "S01 移除了旧下方工作台，降低视觉复杂度，但同时减少了 visible posting/reply affordance；需要明确目标用户是否仍应在首屏发帖。",
        "多处 UI 变化依赖 Playwright/smoke 局部验证，缺少 S11 统一回归矩阵。",
        "窗口内存在本地 SQLite、Playwright 输出、截图和临时 tsconfig 文件变动；这些应继续保持为本地/证据产物，不进入 release slice。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "测试 / Build 状态", 1)
    check_rows = [
        ["npm run type-check", "通过", "S10 fresh run；tsc --noEmit --incremental false 无错误"],
        ["npm run test:analytics", "通过", "23/23 passed；项目脚本清理 `.tmp`"],
        ["npm run release:preflight -- --json", "通过", "freeBytes 79,330,250,752；超过 20 GB 门槛"],
        ["npm run build", "通过", "Next.js build 成功；生成 128 static pages"],
        ["npm run release:root-deploy-preflight -- --json", "保护性失败", "dirty worktree 阻止直接 root deploy；应使用 S22 staging 或 clean worktree"],
        ["S01 class-forum Playwright", "通过/跳过符合预期", "3 passed, 3 skipped by intentional project split；desktop/mobile smoke 通过"],
        ["S04 Practice Playwright checks", "通过", "desktop/mobile 无横向溢出；topic mission interaction 更新 existing select"],
        ["全产品 E2E", "未运行", "应由 S11/S22 在 release gate 中安排"],
    ]
    add_table(doc, ["检查", "状态", "证据 / 备注"], check_rows, [2500, 1400, 5460])

    add_heading(doc, "Changed Files", 1)
    file_rows = [
        ["S09 copy/i18n", "`app/register/page.tsx`; `app/login/page.tsx`; `components/ui/GradeSelector.tsx`; `lib/i18n.ts`"],
        ["S01 forum UI/assets", "`components/forum/ForumWorkspace.tsx`; `components/layout/Navbar.tsx`; `public/forum-assets/*.png`; `design-qa.md`"],
        ["S04 Practice Arena", "`app/practice/page.tsx`; `app/practice/adventure-ui-preview/page.tsx`; `app/practice/adventure-ui-preview/assets/math-adventure-island-map.png`; `components/practice/PracticeAdventureArenaShell.tsx`; `components/practice/assets/math-adventure-island-map.png`"],
        ["Coordination/reporting", "`coordination/session-logs/2026-06-10-S01.md`; `2026-06-10-S04.md`; `2026-06-09-S09.md`; `2026-06-09-S10.md`; `coordination/blockers/2026-06-10-S01-forum-verification-practice-typecheck.md`; `coordination/reports/2026-06-10-practice-arena-design-memory/*`; current report DOCX/builder"],
        ["Generated/local evidence", "`.playwright-cli/*`; `output/playwright/forum-*.png`; `.local/hk-math-db.sqlite*`; `next-env.d.ts`; temporary Playwright tsconfig. Treat as evidence/local artifacts, not product release changes."],
    ]
    add_table(doc, ["类别", "窗口内观察到的文件"], file_rows, [2100, 7260])

    add_heading(doc, "明日优先级", 1)
    for item in [
        "S11/S22：针对 forum、practice、register、login 跑 targeted regression，并决定是否需要 full release matrix。",
        "S04：让 owner 视觉验收 live `/practice`，确认是否继续深化 question card、streak source、reward/game unlock 状态。",
        "S01/S12：确认论坛是否需要新 posting/reply modal；若需要，先定义 backend roster/enrollment 边界。",
        "S22/S25：继续 dirty-tree ownership map 与 release slicing；禁止从 dirty root 直接 deploy。",
        "S09：审计其它 US curriculum surfaces，避免残留 P/S 或香港学段命名。",
    ]:
        add_number(doc, item)

    add_heading(doc, "Dr. Peter Hu 需要决策", 1)
    for item in [
        "论坛：当前可见发帖/回复工作台已移除。是否接受“只浏览/筛选/讨论列表”的首屏，还是要求 S01/S12 设计新的学生友好发帖 modal/flow？",
        "Practice Arena：是否批准当前 Math Adventure Island live 方向继续深化，还是要求停留在设计记忆/preview 状态？",
        "发布：在 dirty root 清理完成前，是否明确要求所有 production/preview 发布只能使用 S22 pruned staging 或 clean reviewed worktree？",
        "质量门：是否要求今天补跑 S11/S22 full-product E2E，还是先跑 forum/practice/login/register targeted matrix？",
        "US curriculum：是否授权 S09 扩展美国年级标签审计到 dashboard、practice、lesson、parent/teacher console 中所有 US-facing surfaces？",
    ]:
        add_bullet(doc, item)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
