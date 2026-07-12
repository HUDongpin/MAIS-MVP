from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-07-president-report.docx")

FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.4, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, bold=True)
        set_cell_shading(cell, header_fill)
    row_fills = row_fills or {}
    for row_index, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        fill = row_fills.get(row_index)
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=body_size)
            if fill:
                set_cell_shading(cells[idx], fill)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DEEP_BLUE),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    prepared = datetime.now().strftime("%Y-%m-%d %H:%M HKT")
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-07"),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Prepared during run", prepared),
        ("Reporting window", "2026-07-06 08:00 to 2026-07-07 08:00 Asia/Hong_Kong"),
        ("UTC filter window", "2026-07-06 00:00 to 2026-07-07 00:00 UTC"),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL, body_size=8.8)


def agent_rows():
    inactive = "No fresh in-window session log; represented only by A25 owner-package/readiness evidence where applicable."
    statuses = {
        "A01": ("A25-consumed", "No direct A01 log. A25 recorded owner-package authorization for app shell final-state evidence.", "A01 package still not merged; release-source and lifecycle gates remain blocked.", "No A01-owned check run."),
        "A02": ("A25-consumed", "No direct A02 log. A25 recorded dashboard owner-package authorization evidence.", "Dashboard package remains part of dirty-root closure, not release-ready.", "No A02-owned check run."),
        "A03": ("A25-consumed", "No direct A03 log. A25 recorded curriculum-roadmap authorization in first five-row batch.", "Package closure still depends on global validation and clean source.", "No A03-owned check run."),
        "A04": ("A25-consumed", "No direct A04 log. A25 recorded practice authorization in first five-row batch.", "Practice package remains not merged and not deployable.", "No A04-owned check run."),
        "A05": ("A25-consumed", "No direct A05 log. A25 recorded lesson authorization in first five-row batch.", "Lesson package remains behind package/worktree gates.", "No A05-owned check run."),
        "A06": ("No fresh log", inactive, "Visualization package remains dirty and owner-routed by A25.", "No A06-owned check run."),
        "A07": ("Idempotent recheck", "No direct A07 log. A25 rechecked prior A07 authorization as already recorded.", "AI tutor package remains in dirty-root closure; no provider behavior changed.", "No live provider check."),
        "A08": ("Idempotent recheck", "No direct A08 log. A25 rechecked prior A08 authorization as already recorded.", "Shared type/state package still part of blocked closure.", "No A08-owned check run."),
        "A09": ("Pending owner input", "No direct A09 log. A25 current focus row for copy/i18n/accessibility remains pending.", "Needs explicit owner authorization text if final state is reviewed commit.", "No A09-owned check run."),
        "A10": ("Active", "Verified Technical-Review advisory materials and generated Codex verification DOCX outputs.", "Report/document scope only; no feature implementation.", "DOCX structural/render QA passed for Technical-Review outputs."),
        "A11": ("No fresh log", inactive, "Regression evidence remains represented through A25/A22 package gates; no current release-ready regression package.", "No A11-owned Playwright run."),
        "A12": ("No fresh log", inactive, "Backend/API package remains owner-routed by A25 evidence; shared/backend closure still blocked by package/worktree gates.", "No A12-owned backend check."),
        "A13": ("A25-consumed", "No direct A13 log. A25 recorded teacher-console owner-package authorization.", "Teacher package remains unmerged and non-deployable.", "No A13-owned check run."),
        "A14": ("Idempotent recheck", "No direct A14 log. A25 rechecked prior parent-console authorization.", "Parent package still not selected as release source.", "No A14-owned check run."),
        "A15": ("Idempotent recheck", "No direct A15 log. A25 rechecked prior adaptive-engine authorization.", "Adaptive package still part of blocked closure.", "No A15-owned check run."),
        "A16": ("No fresh log", inactive, "No fresh A16 repo log in this window; no research implementation action.", "No A16 check run."),
        "A17": ("Pending owner input", "No direct A17 log. A25 current focus row for gamification/motivation remains pending.", "Needs explicit owner authorization text.", "No A17-owned check run."),
        "A18": ("A25-consumed", "No direct A18 log. A25 recorded A18/A21 curriculum/content authorization.", "Content package remains unmerged and still gated by release/QA.", "No A18-owned QA check."),
        "A19": ("No fresh log", inactive, "No API environment work in this window; no provider/env action.", "No live env check."),
        "A20": ("A25-consumed", "No direct A20 log. A25 recorded game-based-learning owner-package authorization.", "Game package remains in dirty closure; no release selection.", "No A20-owned game check."),
        "A21": ("A25-consumed", "No direct A21 log. A25 recorded content/RAG owner-package authorization in first batch.", "Content/RAG package remains gated by QA/integration/release.", "No A21 package-local check."),
        "A22": ("Evidence only", "Logged two-row cleanup instruction evidence; A25 added clean-source candidate evidence.", "Release source still blocked; top clean candidate type-check failed.", "Focused smoke 2/2 passed; candidate type-check failed with 652 TS error lines."),
        "A23": ("Pending owner input", "No direct A23 log. A25 current focus row for integration/promotion remains pending.", "Needs explicit owner authorization text.", "No A23-owned check run."),
        "A24": ("Idempotent recheck", "No direct A24 log. A25 rechecked prior exact-layer authorization.", "Exact-layer package remains part of blocked closure.", "No A24-owned check run."),
        "A25": ("Active", "Built release-intake guardrails, owner authorization packets, backlog queue, validate-to-merge exit criteria, and A22 clean-source candidate evidence.", "Closure remains validate/blocked; no cleanup, deploy, staging, merge, or executable rows.", "In-window aggregate reached 128/128 before report; fresh post-map aggregate now stale 11/128."),
    }
    rows = []
    for i in range(1, 26):
        agent = f"A{i:02d}"
        rows.append((agent, *statuses[agent]))
    return rows


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "报告窗口内存在 fresh AI session activity，集中在 A10-owned Technical-Review verification、A22-owned cleanup evidence，以及 A25-owned git hygiene/release intake。"),
            ("2", "A10 完成 Technical-Review advisory verification，生成 Codex verification report、tracked DOCX 与 clean DOCX，并完成 16-page render QA；没有改 feature code。"),
            ("3", "A25 将 owner authorization、backlog queue、validate-to-merge exit criteria、A22 clean-source candidate review、focused smoke 与 type-check evidence 推进，但所有结果仍为 evidence-only。"),
            ("4", "发布仍 blocked：fresh A22 release-source clean gate failed，root `main` 当前 5,548 expanded dirty entries；A22 top clean candidate focused smoke passed 2/2，但 candidate type-check failed with 652 TypeScript error lines。"),
            ("5", "当前最关键 owner 决策：授权或拒绝 A09/A17/A23 当前 focus rows，解除 compose-worktree validation hold，选择 clean-source release path，且继续禁止 dirty-root deploy。"),
        ],
        [700, 8660],
        body_size=8.7,
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "Fresh activity exists in this window across A10 Technical-Review verification, A22 cleanup evidence, and A25 git hygiene/release intake."),
            ("2", "A10 produced verified advisory DOCX outputs and a Codex verification report without touching feature code."),
            ("3", "A25 deepened the release-intake control plane: authorization packets, backlog queue, validate-to-merge exit criteria, A22 clean-source candidate review, focused smoke, and candidate type-check evidence."),
            ("4", "Release remains blocked. The dirty root has 5,548 expanded status entries; A22 release-source clean fails. The top clean candidate smoke passed 2/2, but candidate type-check failed with 652 TypeScript error lines."),
            ("5", "Owner attention should go to A09/A17/A23 focus-row authorization, the compose-worktree validation hold, clean-source release selection, and maintaining the no-dirty-root-deploy rule."),
        ],
        [700, 8660],
        body_size=8.7,
    )

    add_heading(doc, "报告窗口摘要（中文）", 1)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ("窗口", "2026-07-06 08:00 至 2026-07-07 08:00 Asia/Hong_Kong。"),
            ("Fresh session logs", "`coordination/session-logs/2026-07-06-A10-technical-review-verification.md`、`2026-07-06-A22.md`、`2026-07-06-A25.md`。"),
            ("Blockers / decisions", "`coordination/blockers/` 与 `coordination/decisions/` 在窗口内没有新增文件；阻塞项来自 A25/A22 release-intake gates。"),
            ("Changed files", "窗口内 1,586 个 filesystem-changed files（排除 `.git`、`node_modules`、`.next`、`.tmp`）：1,558 个在 `coordination/release-intake/`，12 个在 `coordination/reports/`，3 个 session logs，Technical-Review 10 个文件，2 个 scripts，另有 `.DS_Store`。"),
            ("Fresh report checks", "A25 dirty-map currentness passed after refresh；no-staged passed；root type-check passed；analytics tests passed 27/27；A22 release-source clean failed；A25 aggregate currentness failed 11/128 because downstream evidence became stale against the fresh map。"),
        ],
        [2450, 6910],
        body_size=8.5,
    )

    add_heading(doc, "整体项目进展（中文）", 1)
    add_bullet(doc, "A10-owned Technical-Review verification advanced project governance: stale advisory claims were corrected and clean/tracked DOCX deliverables were produced with render QA.")
    add_bullet(doc, "A25-owned release intake advanced from package listing into stronger decision plumbing: focus batches, backlog equations, owner-input packets, validate-to-merge exit criteria, and A22 clean-source candidate gates.")
    add_bullet(doc, "A22-owned release path is clearer but still blocked: top clean candidate `codex/A22-us-region-alignment` has git/archive evidence and smoke pass, but type-check is red and no release source is selected.")
    add_bullet(doc, "A01-A24 product/content owners mostly had no direct new logs; their status is represented through A25 owner-package authorization and readiness evidence.")

    add_heading(doc, "A01-A25 会话状态表（中文）", 1)
    add_table(
        doc,
        ["ID", "状态", "窗口内活动", "当前重点 / 阻塞", "Checks"],
        agent_rows(),
        [720, 1100, 3080, 2940, 1520],
        body_size=6.7,
        row_fills={9: OK_FILL, 21: RISK_FILL, 24: PENDING_FILL},
    )

    add_heading(doc, "已完成工作（中文）", 1)
    add_bullet(doc, "A10 verified Technical-Review advisory material against repo evidence and live public metadata, then produced `Technical-Review/Codex-Verification-Report.md`, tracked DOCX, and clean DOCX.")
    add_bullet(doc, "A10 rendered both Technical-Review DOCX outputs to 16 PNG pages plus PDF and visually inspected the clean/tracked contact sheets and representative dense/redline pages.")
    add_bullet(doc, "A22 recorded two-row cleanup execution-instruction evidence without running cleanup apply or deleting generated directories; deployment remained blocked.")
    add_bullet(doc, "A25 recorded multiple owner-package authorization rows, integrated authorization backlog queue, validate-to-merge exit criteria, transition forecast, and A22 clean-source candidate review/focused-smoke/type-check evidence.")
    add_bullet(doc, "A25 maintained non-destructive boundaries across evidence work: no staging, commit, merge, push, reset, cleanup apply, worktree deletion, branch deletion, or deploy.")

    add_heading(doc, "进行中工作（中文）", 1)
    add_bullet(doc, "A25 closure remains in `validate`: validateExitReady=false, handoffStatus=blocked-before-merge, pendingCanonicalAuthorizationRows=42, focusBatchPendingRows=3, deferredPhysicalLifecycleRows=38.")
    add_bullet(doc, "A25 current focus rows are `a09-copy-i18n-accessibility`, `a17-gamification-and-motivation`, and `a23-integration-and-promotion-lead`; they wait for explicit owner authorization.")
    add_bullet(doc, "A22 clean-source selection remains in progress: 9 clean worktree candidates exist, but the top candidate is promotion-blocked and type-check-red.")
    add_bullet(doc, "A25/A22 evidence must be refreshed after any new dirty-map change before using currentness claims for merge/release planning.")

    add_heading(doc, "阻塞项（中文）", 1)
    add_table(
        doc,
        ["Blocker", "Evidence", "Owner / next action"],
        [
            ("Release source dirty", "Fresh `assert-release-source-clean.mjs` failed; root `main` has 5,548 expanded dirty entries.", "A22/A25/A10: only clean worktree, clean clone, reviewed clean slice, or owner-approved pruned staging can release."),
            ("A25 aggregate stale after fresh map", "`assert-dirty-worktree-remediation-current.mjs --json` failed: 11/128 current checks passed.", "A25: refresh downstream release-intake evidence after the 08:03 dirty map if it will be used for decisions."),
            ("Top clean candidate type-check red", "`codex/A22-us-region-alignment` type-check failed with exit 2 and 652 TypeScript error lines.", "A22 plus owning sessions: fix candidate or select another clean-source path before promotion."),
            ("Owner focus rows pending", "A09/A17/A23 focus rows are waiting for owner authorization; backlog remains 42 pending rows.", "Owner: approve/reject with selectedFinalState and explicit negative boundaries."),
            ("Validation hold and physical lifecycle", "A25 evidence still reports validation hold and 38 deferred physical-lifecycle rows.", "Owner/A25: confirm compose-worktree hold and decide physical lifecycle closure sequence."),
        ],
        [2100, 3700, 3560],
        body_size=7.4,
        row_fills={0: RISK_FILL, 1: RISK_FILL, 2: RISK_FILL, 3: PENDING_FILL, 4: PENDING_FILL},
    )

    add_heading(doc, "风险（中文）", 1)
    add_bullet(doc, "Release risk: dirty-root deployment would mix 5,548 expanded entries across many owners and violate A22 release-source policy.")
    add_bullet(doc, "Evidence-staleness risk: A25 generated a fresh dirty map during the report run; downstream packets now need regeneration before they can be treated as current.")
    add_bullet(doc, "Clean-candidate risk: `codex/A22-us-region-alignment` has smoke evidence but fails type-check; promotion without type/build/regression closure would be premature.")
    add_bullet(doc, "Scope risk: the report run must not repair feature code; red package/worktree gates need routed owner sessions, not A10 report edits.")
    add_bullet(doc, "Generated-artifact risk: cleanup remains evidence-only unless owner gives exact execution instructions and A22 confirms no traces/reports/logs/evidence need preservation.")

    add_heading(doc, "测试 / 构建状态（中文）", 1)
    add_table(
        doc,
        ["Check", "Result", "Interpretation"],
        [
            ("`npm run release:dirty-map -- --reason ...`", "Pass; wrote `2026-07-07-A25-dirty-tree-map-20260707T000303Z.md`; 5,548 expanded entries.", "Fresh A25 baseline for this report."),
            ("`node coordination/release-intake/assert-no-staged-changes.mjs`", "Pass; staged entries 0.", "No staged Git mutation."),
            ("`npm run release:dirty-map -- --assert-current --max-age-minutes 120`", "Pass; current at 5,548 expanded entries.", "Dirty-map itself is fresh."),
            ("`node coordination/release-intake/assert-release-source-clean.mjs`", "Fail; root release source dirty.", "A22 release remains blocked."),
            ("`node coordination/release-intake/assert-dirty-worktree-remediation-current.mjs --json`", "Fail; 11/128 current checks passed.", "Downstream A25/A22 packets stale against fresh map."),
            ("`npm run type-check -- --pretty false`", "Pass.", "Dirty-root TypeScript currently green; not release evidence."),
            ("`npm run test:analytics`", "Pass; 27/27.", "Learning analytics/adaptive unit slice green."),
            ("`git diff --check`", "Pass; no diagnostics.", "No whitespace-conflict diagnostics."),
            ("`npm run build` / Playwright", "Not run by report owner.", "Dirty-root build/E2E would not satisfy A22 clean-source release rule."),
        ],
        [3350, 2250, 3760],
        body_size=7.4,
        row_fills={0: OK_FILL, 1: OK_FILL, 2: OK_FILL, 3: RISK_FILL, 4: RISK_FILL, 5: OK_FILL, 6: OK_FILL, 7: OK_FILL, 8: PENDING_FILL},
    )

    add_heading(doc, "变更文件（中文）", 1)
    add_table(
        doc,
        ["Category", "Count / examples", "Meaning"],
        [
            ("Session logs", "3 files: A10 Technical-Review verification, A22 cleanup evidence, A25 release-intake.", "Fresh direct work evidence."),
            ("Release intake", "1,558 files under `coordination/release-intake/`; top-level subset includes 637 owner authorization/intake, 238 closure/validation gates, 78 A22 release/cleanup files, 73 owner mapping/pathspec files, 10 dirty maps.", "Dominant A25/A22 governance artifact set."),
            ("Technical-Review", "10 files including source advisory docs, workbook, Codex verification report, tracked DOCX, clean DOCX, and local lock/history files.", "A10 advisory verification deliverables and source materials."),
            ("Reports", "12 files under `coordination/reports/`, mainly previous report DOCX/render artifacts and builder.", "Prior report deliverable evidence within the mtime window."),
            ("Scripts", "`scripts/refresh-dirty-tree-map.mjs`, `scripts/cleanup-generated-artifacts.mjs`.", "A25/A22 tooling changed in release-intake scope."),
            ("Local metadata", "`.DS_Store`.", "Local metadata; not product evidence."),
        ],
        [1880, 4460, 3020],
        body_size=7.3,
    )

    add_heading(doc, "明日优先事项（中文）", 1)
    add_numbered(doc, "A25 refresh downstream release-intake evidence against the 5,548-entry dirty map before any merge/release planning.")
    add_numbered(doc, "Owner decide A09/A17/A23 current focus rows or explicitly defer them with blocker text.")
    add_numbered(doc, "A22 fix or replace the top clean candidate after the 652-line type-check failure; do not select it for release until type/build/regression gates are green.")
    add_numbered(doc, "A25/A22 keep no-staged and no-dirty-root-deploy evidence fresh; release source must be clean or reviewed/pruned.")
    add_numbered(doc, "A10/A25 keep package closure routed to owning sessions; avoid report-run feature fixes.")

    add_heading(doc, "需要 Owner 决策（中文）", 1)
    add_bullet(doc, "请 Dr. Peter Hu 决定是否授权当前三个 owner-package focus rows：`a09-copy-i18n-accessibility`、`a17-gamification-and-motivation`、`a23-integration-and-promotion-lead`。")
    add_bullet(doc, "请确认 compose-worktree validation hold 是否可以解除，尤其是 exact deletion / lifecycle closure 的状态。")
    add_bullet(doc, "请决定 A22 clean-source release path：继续修 `codex/A22-us-region-alignment`，选择其他 clean candidate，或暂停 promotion。")
    add_bullet(doc, "请确认任何 generated-artifact cleanup 必须有 exact command、exact target、evidence-preservation check 与 separate execution instruction。")
    add_bullet(doc, "请继续明确：dirty root 只可作为 inventory/reporting source，不可作为 production deploy source。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
