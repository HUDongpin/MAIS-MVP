from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-03-president-report.docx")

BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_para(doc, text="", style=None, size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = ""
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.5, bold=True, color=BLACK)
        set_cell_shading(hdr[idx], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=body_size, color=BLACK)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DEEP_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-03"),
        ("Reporting window", "2026-07-02 08:00 to 2026-07-03 08:00 Asia/Hong_Kong"),
        ("Prepared at", "2026-07-03 12:07 HKT"),
        ("Prepared by", "A10 reporting automation; source evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL)


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "A22 已在 2026-07-03 00:27 HKT 通过 Vercel 将当前最新 MAIS-MVP 发布到 `www.mais.hk` 与 `www.mais.ac`；Vercel inspect 为 Ready，双域名 HTTP/API smoke 与 A11 production auth/API preflight 通过。"),
            ("2", "本次发布不是从 clean Git commit 发布，而是 owner-requested current latest dirty-root state 的 A22 pruned staging package；A22 未直接从 dirty root 发布，且 staging forbidden-path scan 为 0。"),
            ("3", "A11 在窗口内完成两个 DOCX bug report 的核验与一轮局部修复：20260702 报告 26 项中 7 项为真、2 项 partial；20260703 报告 6 项中 3 项为真/部分为真。"),
            ("4", "A05、A07、A09、A13、A20 等 owner surfaces 有局部修复或路由：US-CA elementary completion panel removal、Adventure Island unlock、teacher assignment filters、lesson audio fallback、option capitalization、worked-example concrete model。"),
            ("5", "A25 fresh dirty-map 通过，当前 expanded status entries 为 3777；owner package approval requests 24 项、physical lifecycle approval requests 34 项均 current，但 aggregate remediation gate 仍红，root 仍禁止作为常规发布源。"),
        ],
        [700, 8660],
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "A22 completed a production deployment to `www.mais.hk` and `www.mais.ac` from a pruned staging package; Vercel inspect is Ready and safe live smoke checks passed."),
            ("2", "This was not a clean-commit release. The source remains a heavily dirty integration inventory, so future production work still needs A22 clean/pruned release evidence."),
            ("3", "A11 verified two external bug-report DOCX files. Several claims were confirmed, several were downgraded or not reproduced, and seven real issues received local fixes with focused tests."),
            ("4", "A06 continued Manim v2 closure/evidence propagation and supplied limited build-blocker fixes used by A22, but A18 teaching acceptance, A11 broad browser evidence, A22 clean release evidence, and final 4/4 audit remain open."),
            ("5", "A25 fresh release intake shows 3,777 expanded dirty entries, 24 owner-package approvals, and 34 physical-lifecycle approvals. No staging or cleanup is authorized by those artifacts alone."),
        ],
        [700, 8660],
    )

    add_heading(doc, "汇报窗口摘要", 1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("窗口", "2026-07-02 08:00 至 2026-07-03 08:00 Asia/Hong_Kong。报告生成时间为 2026-07-03 12:07 HKT，使用窗口内证据与报告时 fresh checks。"),
            ("活跃 session logs", "A22 production deploy、A11 bug verification/fix、A06 Manim closure、A05 lesson panel removal、A10 reporting、A25 release-intake evidence。"),
            ("专门 blocker/decision 文件", "`coordination/blockers/` 与 `coordination/decisions/` 在窗口内无新增文件；阻塞状态主要来自 session logs 与 A25 release-intake artifacts。"),
            ("报告时 fresh checks", "A25 dirty-map passed；owner-package approval requests current；physical-lifecycle approval requests current；no-staged gate passed；aggregate remediation currentness failed due stale downstream readiness/decision packets。"),
        ],
        [1800, 7560],
    )

    add_heading(doc, "整体项目进展", 1)
    add_para(doc, "项目在产品可用性与发布工程上都有实质推进：生产域名已更新，若只看 `www.mais.hk` 与 `www.mais.ac` 的基础访问、登录页、匿名 `/api/me`、公共题库、受保护 lesson redirect 与 AI Tutor status，A22/A11 smoke 均显示正常。")
    add_para(doc, "同时，代码库仍处于强 dirty-root 状态。A25 fresh dirty-map 显示 1438 collapsed status entries、3777 expanded status entries、389 tracked modified、1 tracked deleted、3387 untracked files。当前最重要的管理动作不是继续扩大功能面，而是按 A01-A25 owner package 关停、提交、归档或写明 blocker。")
    add_para(doc, "A11 bug verification 将外部 bug 报告拆成可执行路由：roadmap blank rendering、Nova Tutor voice endpoint、teacher assignment filter、Adventure Island progression、lesson audio responsiveness、P5 worked-example visual 等需要按 owner surface 清理。A06 Manim 继续强化 closure evidence，但尚未达到最终验收。")

    add_heading(doc, "A01-A25 会话状态表", 1)
    status_rows = [
        ("A01", "App shell", "无 fresh log；被 A11 路由", "About logged-out mission state/copy 由 A11 路由给 A01/A09；A25 package仍 blocked。"),
        ("A02", "Dashboard", "无 fresh log；A02/A15 package blocked", "A11 20260702 对 Personalized Learning 多数未复现；dashboard/adaptive closure仍需 owner package处理。"),
        ("A03", "Roadmap", "无 fresh log；有新 A11 路由", "20260703 report 确认 primary/student roadmap blank/skeleton 风险；需 A03/A22 定位。"),
        ("A04", "Practice", "被 A11/A20 消费", "Adventure Island progression 与 practice handoff 相关；A20 已做局部修复，A04仍需 clean-slice配合。"),
        ("A05", "Lesson", "活跃/局部完成", "移除 US-CA elementary completion panel；参与 lesson audio、option text、worked-example visual fixes。"),
        ("A06", "Visualization", "高活跃；未最终验收", "Manim evidence propagation 与 A22 build-blocker fixes完成；A18/A11/A22/final audit gate仍开。"),
        ("A07", "AI tutor", "无 fresh log；被 A11 路由", "Nova Tutor voice live synthesis失败；A11另做 bounded lesson-audio fallback，需 A07/A12/A19/A22处理voice endpoint。"),
        ("A08", "State/analytics", "无 fresh log；shared contract blocked", "A25 aggregate仍显示 Wave 02 readiness stale/blocked。"),
        ("A09", "Copy/i18n/a11y", "无 fresh log；参与局部修复", "A11/A05 option capitalization 涉及 copy display；A25 package仍 blocked。"),
        ("A10", "Tooling/report", "活跃", "生成本报告并消费 A01-A25 evidence；无 feature code edits。"),
        ("A11", "QA", "高活跃", "完成 20260701、20260702、20260703 bug-report verification；局部 fix pass 24/24 focused tests passed。"),
        ("A12", "Backend/API", "支持性活跃", "A22 production smoke 消费 API evidence；A11/A05/A20 修复涉及 auth/API/persistence boundaries。"),
        ("A13", "Teacher console", "被 A11 fix pass 激活", "Bugs 138-140 teacher assignment filters fixed locally；仍需 owner package closure。"),
        ("A14", "Parent console", "无 fresh log", "A13/A14 console closure仍 dirty-open decision。"),
        ("A15", "Adaptive", "无 fresh log", "A02/A15 dashboard-adaptive package仍在 dirty lifecycle 中。"),
        ("A16", "Research", "无 fresh log；相对 ready", "A25 owner matrix中仍是最轻的 research evidence row；无新研究任务。"),
        ("A17", "Motivation", "无 fresh log", "A17/A20 game/motivation closure仍 blocked。"),
        ("A18", "Curriculum QA", "无 fresh log；A06 gate consumer", "A06 final teaching acceptance仍需 A18；content package仍 blocked。"),
        ("A19", "API env", "支持性证据", "A22 env preflight confirms required production variable names present 15/15；未记录 secret values。"),
        ("A20", "Games", "活跃/局部完成", "Bug 127 Adventure Island locked-state fix；renderer failure不再 demote verified unlock。"),
        ("A21", "Content/RAG", "无 fresh log", "A18/A21 content evidence package仍需签核/closure。"),
        ("A22", "Release engineering", "完成生产发布", "Pruned staging package deploy Ready；root仍不合格为 clean release source。"),
        ("A23", "Integration", "无 fresh log", "candidate-to-live promotion仍需 A23计划；本窗口无新集成决定。"),
        ("A24", "Exact layer", "无 fresh log", "content/exact-layer package仍在 A18/A21/A23/A24 chain中。"),
        ("A25", "Git hygiene/release intake", "高活跃；仍 blocked", "fresh dirty-map 3777；24 owner-package approvals、34 physical lifecycle approvals current；aggregate remediation red。"),
    ]
    add_table(doc, ["Agent", "职责", "窗口状态", "证据/下一步"], status_rows, [900, 1700, 2400, 4360], body_size=8.7)

    add_heading(doc, "已完成工作", 1)
    add_table(
        doc,
        ["Owner", "完成项", "验证"],
        [
            ("A22/A25/A19/A06/A11", "发布当前最新站点到 `www.mais.hk` 与 `www.mais.ac`，使用 pruned staging package，不直接从 dirty root deploy。", "release:build-gate passed；222 static pages；staging 2187 files/215MB；forbidden path count 0；Vercel Ready；双域名 smoke passed；A11 production preflight 1/1 passed。"),
            ("A11", "核验 20260702 bug report 26 项，并对 7 个 real bugs 做局部 fix pass。", "fix verification `npx tsx --test ...` passed 24/24；source regressions 17/17；visualization diagnostics 43/43。"),
            ("A11", "核验 20260703 bug report 6 个 unique claims。", "3 项 real/部分为真：About logged-out state、roadmap blank/primary map blank、Nova Tutor voice；Lessons P0 outage 与 visualization clipping 未按原文复现。"),
            ("A05", "移除 screenshot 指向的 US-CA elementary lesson completion/quick self-check panel。", "lessonCompletionChecklist test passed；41 seed helper sweep 0 failures；local DOM smoke 200 and panel texts absent。"),
            ("A06", "多轮 Manim v2 review-slice/evidence propagation 与 A22 build-blocker type fixes。", "多组 focused/adjacent tests green；full Manim directory suite reported 1525/1525 或 1529/1529；A22 consumed three A06 build-blocker fixes。"),
            ("A25", "刷新 release-intake：dirty-map、owner package approval requests、physical lifecycle approval requests。", "fresh dirty-map passed at 3777 entries；owner-package current gate passed 24 requests；physical lifecycle current gate passed 34 requests；no staged entries。"),
        ],
        [1500, 4300, 3560],
    )

    add_heading(doc, "进行中工作", 1)
    add_table(
        doc,
        ["Owner", "状态"],
        [
            ("A22/A25/A10", "生产站点已上线，但 source-control closure 未完成；后续发布必须继续使用 clean worktree/clean clone/reviewed clean slice/pruned staging。"),
            ("A11/A03/A07/A12/A19/A22", "20260703 report 暴露的 roadmap blank rendering 与 Nova Tutor voice endpoint 仍需 owner sessions 处理。"),
            ("A11/A13/A20/A07/A05/A09", "七个 real bug 的局部 fixes 已在 dirty root 验证，但需要切成 owner packages 与 A22 clean/pruned release path 后才能声明生产推广。"),
            ("A06/A18/A11/A22", "Manim v2 reusable skill/source-architecture work继续推进，但整体 completion 仍需 A18 teaching、A11 browser、A22 clean release、final 4/4 audit。"),
            ("A25", "Aggregate remediation gate 仍需刷新下游 readiness/decision packets；当前 red 的主因多为 stale relative to fresh dirty map，而不是所有检查都重新失败。"),
        ],
        [1900, 7460],
    )

    add_heading(doc, "阻塞项", 1)
    add_table(
        doc,
        ["阻塞", "影响", "下一步"],
        [
            ("Dirty root", "fresh A25 dirty-map 为 3777 expanded entries；root 仍不可作为常规 production source。", "继续执行 A22 clean/pruned source rule；先关 owner packages。"),
            ("A25 aggregate remediation gate red", "69 current checks 中 40 passed、29 failed；失败多为 Wave 02-06 readiness、owner decision packets、A22 residual evidence stale。", "用 refresh runner 更新下游 artifacts，或按 owner package 写 blocker/closure report。"),
            ("Worktree lifecycle", "37 worktrees；29 dirty open decisions；5 clean-diverged open decisions；34 physical lifecycle approvals needed。", "Owner 必须逐项选择 final state；approval artifacts 不等于执行授权。"),
            ("Confirmed live/product bugs", "Roadmap blank rendering 与 Nova Tutor voice endpoint 仍影响用户路径；teacher/practice/lesson fixes尚未 clean-slice发布。", "A03/A07/A12/A19/A22 与对应 owner sessions 接手。"),
            ("A06 final acceptance", "Manim local/test evidence强，但未通过最终接受链。", "A18/A11/A22 提交真实 owner-gate evidence 后再做 final 4/4 objective audit。"),
        ],
        [2300, 3550, 3510],
    )

    add_heading(doc, "风险", 1)
    add_table(
        doc,
        ["风险", "说明", "控制"],
        [
            ("可追溯性风险", "本次 production deploy 来自 pruned staging current latest dirty state，不是 clean commit。", "把本次 deploy 视作 owner-requested exception path；下一步补 clean package/review evidence。"),
            ("验证边界风险", "A11/A05/A06 多为 focused/source/local checks；dirty-root `npm run type-check` 仍不适合当全局成功信号。", "报告只引用局部证据；A22 release claims 只引用 pruned staging/build/smoke。"),
            ("下游 readiness stale", "fresh dirty-map 后，A25多个 downstream readiness/decision packets stale；aggregate gate red。", "先刷新 A25 downstream artifacts，再做 owner decisions。"),
            ("空间/构建风险", "窗口内曾出现 ENOSPC；报告时 df 显示约 46GiB available。", "A22 cleanup-generated-artifacts 仍需 dry-run then exact owner authorization。"),
            ("未发布修复风险", "A11 fix pass解决的 bugs 尚未 clean-source生产发布。", "把 fixes 切进 owner packages 并走 A22 clean/pruned release。"),
        ],
        [2300, 3900, 3160],
    )

    add_heading(doc, "测试 / 构建状态", 1)
    add_table(
        doc,
        ["Scope", "Result"],
        [
            ("Fresh A25 dirty-map", "Passed；`coordination/release-intake/2026-07-03-A25-dirty-tree-map-20260703T040704Z.md`；expanded entries 3777。"),
            ("A25 owner/physical gates", "owner-package approval requests current passed：24 requests；physical lifecycle approval requests current passed：34 requests；no-staged gate passed。"),
            ("A25 aggregate remediation", "Failed：40/69 current checks passed；failures mainly stale downstream readiness/decision packets relative to fresh dirty map。"),
            ("A22 production deploy", "`npm run release:build-gate ...` passed；222 static pages；staging forbidden-path count 0；Vercel Ready；custom-domain safe HTTP/API probes passed；A11 production preflight passed 1/1。"),
            ("A11 bug fix pass", "`npx tsx --test ...` passed 24/24；source regressions passed 17/17；visualization diagnostics passed 43/43；`npm run type-check` still red only in pre-existing A06 visualization/manim review-package tests per A11 report。"),
            ("A05 lesson panel removal", "focused helper/test and local DOM smoke passed；no production deploy claimed。"),
            ("A06 Manim", "focused RED/GREEN and adjacent suites passed across several slices; full Manim directory suite reported 1525/1525 or 1529/1529。"),
            ("Broad build/type-check", "Not rerun as a broad success gate from dirty root for this report；A22 production claim uses build-gate/pruned staging/live smoke instead。"),
        ],
        [2700, 6660],
    )

    add_heading(doc, "变更文件", 1)
    add_table(
        doc,
        ["类别", "代表文件/范围"],
        [
            ("窗口内 mtime summary", "831 report-facing files changed excluding `.git`, `node_modules`, `.next`, `.tmp`, `.local`, and known secret files：coordination 755、components 61、app 5、tests 2、scripts 2、config/root 6。"),
            ("A22 deploy evidence", "`coordination/session-logs/2026-07-02-A22-mais-domains-production-deploy.md`; `coordination/reports/2026-07-03-A22-mais-domains-production-deploy.md`; A25 dirty-map/staging evidence。"),
            ("A11 verification/fix evidence", "`coordination/reports/2026-07-03-A11-*`; screenshots under `coordination/reports/screenshots/2026-07-03-A11-deliverable-bug-verification/`; tests `tests/e2e/reported-bug-source-regressions.test.ts`。"),
            ("A11/A13/A20/A07/A05/A09 touched source", "`app/practice/page.tsx`, `app/teacher/assignments/page.tsx`, `components/gamification/AdventureIslandGame.tsx`, `components/lesson/LessonView.tsx`, `components/lesson/WorkedExampleIllustration.tsx`, `components/practice/PracticeQuestionCard.tsx`。"),
            ("A05 lesson panel", "`components/lesson/lessonCompletionChecklist.ts`, `components/lesson/lessonCompletionChecklist.test.ts`, `components/lesson/LessonView.tsx`。"),
            ("A06 Manim/Visualization", "`components/visualizations/three/manim/*`; examples include `mathSceneV2FinalObjectiveAuditRecordIntake.ts`, `mathSceneV2FinalObjectiveAuditRequestPacket.ts`, `mathSceneV2OwnerGateRerunIntake.ts`, `mathSceneRunFromBeat.ts`。"),
            ("A25 release intake", "`coordination/release-intake/2026-07-03-A25-*`, `latest-A25-*`, owner pathspecs, approval requests, lifecycle and remediation currentness artifacts。"),
        ],
        [2300, 7060],
    )

    add_heading(doc, "明日优先级", 1)
    add_table(
        doc,
        ["优先级", "事项"],
        [
            ("P0", "A22/A25/A10：不要把 dirty root 当作常规 release source；若需新发布，先定义 clean/pruned source evidence。"),
            ("P0", "A25：刷新 stale downstream readiness/decision packets，使 aggregate remediation gate 从 stale-red 回到可决策状态。"),
            ("P1", "A03/A22：优先修复/核验 roadmap blank rendering，尤其 `/student/roadmap/primary` 全图 blank 问题。"),
            ("P1", "A07/A12/A19/A22：处理 Nova Tutor voice endpoint/provider path，A19 只负责 redacted env readiness。"),
            ("P1", "A11/A13/A20/A07/A05/A09：把 dirty-root local fixes 切成 owner packages，决定哪些需要生产推广。"),
            ("P1", "A06/A18/A11/A22：完成 Manim v2 final gate evidence，不扩大到 course-wide lab generation。"),
            ("P2", "A22：在 owner 明确授权后继续 generated-artifact dry-run/apply cleanup；保留 Playwright/report evidence。"),
        ],
        [900, 8460],
    )

    add_heading(doc, "需要 Owner 决策", 1)
    add_table(
        doc,
        ["Decision", "Owner action needed"],
        [
            ("Production source posture", "确认本次 A22 pruned-staging production deploy 的后续治理：补 clean release slice、保留为 exception evidence，或要求回滚/再发布。"),
            ("Roadmap/voice priority", "决定 A03 roadmap blank rendering 与 A07 Nova Tutor voice endpoint 的优先级和 owner assignment。"),
            ("Bug-fix promotion", "决定 A11 local fixes 中哪些进入下一轮 A22 clean/pruned production release。"),
            ("Owner package approvals", "按 A25 current owner-package approval requests 选择 final state，尤其 A25/A22/A06/A12/A11/A10/A05/A18-A21。"),
            ("Physical lifecycle approvals", "按 A25 current physical lifecycle requests 选择 root/main 与 33 个 worktree/branch final states；未授权前不可 cleanup。"),
            ("A06 final acceptance", "确认 A18 teaching review、A11 browser scope、A22 clean release evidence、final 4/4 objective audit 的完成路径。"),
        ],
        [2700, 6660],
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("MAIS-MVP daily report - 2026-07-03 - A10/A25 evidence based")
    set_run_font(r, size=9, color=MUTED)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
