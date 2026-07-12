from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/dongpinhu/Desktop/MAIS-MVP")
OUT = ROOT / "coordination/reports/2026-06-19-president-report.docx"

REPORT_DATE = "2026-06-19"
RUNTIME = "2026-06-19 11:25-12:10 HKT"
WINDOW = "2026-06-18 08:00 至 2026-06-19 08:00 Asia/Hong_Kong"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "PingFang SC")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                cell.width = Inches(widths[idx] / 1440)
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(widths[idx]))
                tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_paragraph(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("MAIS-MVP President Report")
    set_run_font(run, 20, True, "0B2545")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    meta = (
        f"Report date: {REPORT_DATE} | Report time: 8:00 AM Asia/Hong_Kong | "
        "Reporting session: S10 | Audience: Dr. Peter Hu"
    )
    run = p.add_run(meta)
    set_run_font(run, 9.5, False, "555555")


def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, 15, True, "2E74B5")
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, 12, True, "1F4D78")
    return p


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, 10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, 10)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, 10)
    return p


def add_table(doc, headers, rows, widths, font_size=8.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_shading(hdr[i], "F2F4F7")
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, font_size, True, "0B2545")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, font_size)
    set_table_width(table, widths)
    return table


def build_doc():
    doc = Document()
    doc.core_properties.title = "MAIS-MVP President Report 2026-06-19"
    doc.core_properties.author = "S10 Codex reporting automation"
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("MAIS-MVP | S10 daily president report | 2026-06-19")
    set_run_font(run, 8.5, False, "555555")

    add_title(doc)
    add_table(
        doc,
        ["字段", "内容"],
        [
            ["项目", "MAIS-MVP bilingual math learning platform"],
            ["报告窗口", WINDOW],
            ["运行时间", RUNTIME],
            ["正式范围", "S01-S25 coordination model; no feature-code edits by S10"],
            ["核心证据", "session logs, blockers, release intake, S11/S19/S22 reports, git/disk checks"],
        ],
        [2100, 7140],
        9,
    )

    add_h1(doc, "中文 Executive Summary")
    add_bullet(doc, "本窗口有明确推进，不属于“无任务窗口”。S19 完成 LRS 本地与生产连接，生产 /api/lrs/status、/api/lrs/smoke 与 /api/learning-events 均已红acted 通过。")
    add_bullet(doc, "S22 修复教师备课 isolated-app 门禁漂移，教师端本地桌面门禁最终为 27 passed / 1 skipped，移动教师 smoke 1 passed。")
    add_bullet(doc, "S11 学生端 E2E 仍为红灯：注册/登录、Lesson、Practice、移动 Visualization、California、Mainland PEP/HJB、DeepSeek/SimpleTex 前置阻断仍需分 owner 修复。")
    add_bullet(doc, "S10 fresh checks 显示 root type-check 与 analytics/adaptive tests 已恢复通过；但执行 analytics 脚本按 package 定义清除了 .tmp，本报告已把该副作用列为 S22/S25 需要注意的证据状态变化。")
    add_bullet(doc, "发布风险从磁盘不足转为 dirty-root/release-slice 风险：release guard 仍阻止 root deploy；S10 probe 在加入本报告 artifacts 前报告 1711 status entries。")

    add_h1(doc, "English Executive Summary")
    add_bullet(doc, "This was an active reporting window. S19 connected local and production MAIS-MVP to the Learning Record Store with redacted credential handling and verified status, smoke write, and authenticated learning-events ingestion.")
    add_bullet(doc, "S22 stabilized the teacher prep isolated-app harness. The teacher local desktop gate ended at 27 passed / 1 skipped, and the mobile teacher smoke passed.")
    add_bullet(doc, "The student E2E release gate remains red across registration, lessons, Practice, mobile Visualization, California, Mainland PEP/HJB, and pre-provider DeepSeek/SimpleTex blockers.")
    add_bullet(doc, "Fresh S10 checks passed: root type-check and analytics/adaptive tests. The analytics script removed .tmp as part of its package-defined cleanup, so historical local E2E artifacts should be treated as cleared.")
    add_bullet(doc, "Direct production release from the dirty root remains blocked. Disk is now sufficient, but S25/S22 still need a clean or pruned release slice.")

    add_h1(doc, "报告窗口摘要")
    add_body(doc, f"报告窗口：{WINDOW}。本次报告覆盖 S01-S25 正式会话范围；窗口内活跃证据主要来自 S02、S04、S05、S06、S07、S09、S10、S11、S12、S13、S15、S18、S19、S22、S23、S24、S25。")
    add_body(doc, "No assigned work in this reporting window: No. 本窗口内存在教师端、学生端、LRS、Visualization release、STEMBENCH 报告/技能蒸馏等多条工作线。")
    add_body(doc, "S10 未修改功能代码；写入范围限制在 coordination/reporting artifacts 和本报告生成脚本。")

    add_h1(doc, "整体项目进展")
    add_body(doc, "产品面上，教师端本地 E2E 门禁大幅好转，LRS/xAPI 生产接入完成，Visualization release 证据链继续收敛。质量面上，学生端 release gate 仍红，需要先修复核心 Practice、Lesson、注册/语言选择、课程资产与多地区路线漂移。发布面上，磁盘容量已恢复，但 dirty-root 仍阻止直接部署。")
    add_body(doc, "当前更适合按 release slice 推进：先让 S25 明确候选文件边界，再由 S22/S11/S12 对该 slice 做构建、回归、生产/存储 parity，而不是从根目录直接发布。")

    add_h1(doc, "S01-S25 会话状态表")
    session_rows = [
        ["S01", "无新日志", "注册/登录入口受 S11/S09 测试反馈影响。", "等待 selector/UI 基线复核。"],
        ["S02", "活跃", "AdaptiveKnowledgeGalaxy landmark 补齐；dashboard/California gate 改善。", "需重跑 student-frontend。"],
        ["S03", "无新日志", "HJB/PEP roadmap 入口和语言/demo selector 仍在学生门禁失败簇中。", "等待 S11/S09 复核。"],
        ["S04", "活跃", "Practice free-selection unlock 与 grade sync 修复。", "需重跑 practice-pager/OCR。"],
        ["S05", "活跃", "Lesson H1 使用 API lesson title。", "需重跑 lesson/student gates。"],
        ["S06", "活跃", "Visualization S11 regression evidence 可用；strict verifier 仍未 ready。", "等待 S12/S22/S25 parity。"],
        ["S07", "活跃", "AI tutor status 返回 redacted provider/model metadata。", "DeepSeek matrix 待 broader rerun。"],
        ["S08", "无新日志", "无窗口内直接证据。", "关注 shared schema/type drift。"],
        ["S09", "活跃", "Register labels/IDs 更明确。", "重跑 app-shell/student smoke。"],
        ["S10", "活跃", "生成本报告；完成 STEMBENCH math/non-math DOCX、source distillation、global skill install。", "保持 report-only 边界。"],
        ["S11", "活跃", "Teacher gate 几乎全绿；student gate 红并形成中文失败报告。", "优先学生端分簇修复后重跑。"],
        ["S12", "活跃", "Visualization API negative/idempotency coverage 添加；login selected-grade 修复。", "backend spec 待 rerun；LRS consumed by S19。"],
        ["S13", "活跃", "Teacher console product flow 被 S11 gate 覆盖通过。", "关注 California assignment rerun。"],
        ["S14", "无新日志", "无窗口内 parent console 证据。", "无直接动作。"],
        ["S15", "活跃", "未改 adaptive engine；consumed Practice unlock semantics。", "adaptive Practice/dashboard rerun。"],
        ["S16", "无新日志", "无窗口内 research 证据。", "无直接动作。"],
        ["S17", "无新日志", "无窗口内 reward-economy 证据。", "无直接动作。"],
        ["S18", "活跃", "Mainland PEP High topic metadata 补齐。", "插图替代/课程 QA 待定。"],
        ["S19", "活跃", "LRS local/production env and connectivity verified; Preview unavailable。", "维护 redacted inventory。"],
        ["S20", "无新日志", "无窗口内 game-based learning 证据。", "无直接动作。"],
        ["S21", "无新日志", "Mainland PEP High/US content assets 被 S11 failures 牵涉。", "待 owner 分配 pipeline/asset work。"],
        ["S22", "活跃", "Teacher prep isolated harness fixed; release guard now dirty-root blocked。", "准备 clean/pruned slice gate。"],
        ["S23", "活跃", "支持 Mainland PEP metadata/withdrawn illustration integration state。", "重跑 roadmap/lesson gates。"],
        ["S24", "活跃", "保持 withdrawn illustration zero-image assertions；未恢复资产。", "等 owner-approved exact-layer replacements。"],
        ["S25", "活跃", "Visualization release slice inventory and recovery plan written。", "需 owner 选择 slice/路径。"],
    ]
    add_table(doc, ["会话", "状态", "窗口内结论", "下一步"], session_rows, [700, 1050, 4550, 3060], 7.6)

    add_h1(doc, "已完成工作")
    for item in [
        "S19：完成 LRS/xAPI server client、/api/lrs/status、/api/lrs/smoke、/api/learning-events ingestion；local 与 production smoke 均 accepted=1。",
        "S22/S11：教师端本地门禁最终 27 passed / 1 skipped；mobile teacher smoke 1 passed；isolated prep 原失败归因并修复为测试/harness 漂移。",
        "S04/S05/S07/S09/S12/S18/S23/S24：为学生门禁修复或重基线提供多项小范围支持，包括 Practice、Lesson H1、AI status metadata、register labels、selected-grade、Mainland PEP metadata、withdrawn illustration state。",
        "S10：完成 STEMBENCH math SVG DOCX、non-math/deep distillation、global skill install；本报告 DOCX 生成并执行 fresh lightweight checks。",
        "S10 fresh checks：root type-check passed；analytics/adaptive tests 23/23 passed。",
    ]:
        add_bullet(doc, item)

    add_h1(doc, "进行中工作")
    for item in [
        "S11 学生端 release gate：仍为红灯，需按失败簇分配 owner 后重跑 core/mobile/multicourse/live-provider 前置检查。",
        "S06 Visualization release：local/static regression 继续推进，但 strict verifier 仍需 S12/S22/S25 外部证据和 clean release/parity。",
        "S18/S21/S23/S24 Mainland PEP High 插图与内容链：withdrawn high-school illustrations 需要 owner-approved replacement exact-layer assets 后再恢复断言。",
        "S25/S22 release hygiene：磁盘已经恢复，但 direct root deploy 仍被 dirty-root release guard 阻止。",
    ]:
        add_bullet(doc, item)

    add_h1(doc, "阻塞")
    for item in [
        "Direct root deploy blocked：`node scripts/release-env-guard.mjs root-deploy --json` 失败，原因是 worktree dirty；S10 probe 在加入本报告 artifacts 前报告 1711 status entries。",
        "学生端 E2E gate red：Practice combobox/unlock、lesson titles、registration/demo selectors、mobile Visualization、California assignment、Mainland roadmap/PEP High assets 仍阻断。",
        "S06 strict completion not ready：S12/S22 production/storage parity 和 S25 release-slice 决策仍缺。",
        "Vercel Preview LRS env 未完成：Vercel project 没有 connected Git repository；Production 已配置并验证。",
        "历史 `.tmp` Playwright artifact cache 已被 `npm run test:analytics` package script 清除；如需要旧 trace/video，只能依赖已复制到 coordination 的报告证据。",
    ]:
        add_bullet(doc, item)

    add_h1(doc, "风险")
    for item in [
        "Dirty tree 范围很大：S10 probe 在加入本报告 artifacts 前报告 `git status --short` 1711 entries；`git status --porcelain -uall` 曾统计 3017 file entries，public/coordination/app/test 都有大量未整理内容。",
        "LRS Production 已临时 overlay 发布成功，但源改动仍在 dirty tree 中；后续必须由 owner 决定如何 review/stage/commit，不能让临时状态长期漂移。",
        "Student gate failures 混合真实产品问题与测试基线漂移；需要 S11 分簇 rerun 防止误判 provider 或课程质量。",
        "Mainland PEP High illustration deletion/withdrawn state 若未明确产品决策，会持续阻断 lesson release assertions。",
        "S10 运行 analytics 脚本带来的 `.tmp` 清理改变了 S22/S25 早前磁盘证据状态；报告以当前状态为准。",
    ]:
        add_bullet(doc, item)

    add_h1(doc, "测试/构建状态")
    check_rows = [
        ["S10 fresh", "npm run type-check -- --pretty false", "Passed", "Root TypeScript gate now clean in this run."],
        ["S10 fresh", "npm run test:analytics", "Passed", "23 tests passed; script removed `.tmp` by package definition."],
        ["S22/S11", "Teacher desktop gate", "Passed with skip", "27 passed / 1 skipped after S22 fix."],
        ["S22/S11", "Teacher mobile smoke", "Passed", "1 passed."],
        ["S11", "Student core desktop", "Red", "6 passed / 16 failed."],
        ["S11", "Student mobile smoke", "Red", "1 failed / 8 skipped."],
        ["S11", "Student multicourse", "Red", "10 passed / 11 failed."],
        ["S11/S07/S19", "DeepSeek live text", "Inconclusive", "Blocked before provider calls in earlier run; status endpoint later passed."],
        ["S11/S19/S12", "SimpleTex/live OCR", "Inconclusive", "Practice combobox blocker prevented provider rows."],
        ["S22/S25", "release-env-guard root-deploy", "Blocked", "Dirty worktree; probe reported 1711 status entries before report artifacts; disk no longer blocker."],
        ["S10", "npm run build", "Not run", "Report automation avoided broad build/deploy work; S22 owns release gate."],
    ]
    add_table(doc, ["Owner", "Check", "Status", "Evidence"], check_rows, [1100, 2700, 1200, 4360], 8)

    add_h1(doc, "变更文件")
    add_body(doc, "窗口内 dirty-file mtime sample before this report: 205 entries. Largest buckets: HU-XIANGEN-STEM-BENCH 113, coordination 48, tests 22, app 10, components 6, lib 2, data 1, .env.local.example 1, skill files 2.")
    changed_rows = [
        ["S19", ".env.local.example; app/api/learning-events/route.ts; app/api/lrs/status/route.ts; app/api/lrs/smoke/route.ts; lib/server/lrsClient.ts; lib/server/lrsClient.test.ts", "LRS local/production integration; real secrets not recorded."],
        ["S22/S11", "tests/e2e/isolated-app.ts; tests/e2e/teacher-prep-toolchain.spec.ts; tests/e2e/teacher-console-button-matrix.spec.ts", "Teacher prep isolated harness and teacher gate drift."],
        ["S11", "tests/e2e/app-shell-auth.spec.ts; student-smoke; student-frontend; practice-pager; student-button-dropdown-matrix; mainland/california specs; helpers.ts", "Student gate repair/rebaseline work."],
        ["S04/S05/S07/S09/S12/S18", "app/practice/page.tsx; components/lesson/LessonView.tsx; app/api/ai-tutor/status/route.ts; app/register/page.tsx; app/api/auth/login/route.ts; data/mainlandPepHighTopics.ts", "Targeted product/test support for S11 gates."],
        ["S10", "HU-XIANGEN-STEM-BENCH/; hu-xiangen-*-visualization-svg-skill.md; coordination/reports/2026-06-18-stembench-math-svg-implementation-report.docx", "STEMBENCH source distillation and DOCX report."],
        ["S25/S22", "coordination/release-intake/2026-06-18-S25-visualization-release-slice-inventory.md; coordination/blockers/2026-06-18-S22-S25-production-release-space-dirty-tree.md", "Release-slice and dirty-root blocker evidence."],
    ]
    add_table(doc, ["Owner", "Representative files", "Meaning"], changed_rows, [950, 5400, 3010], 7.5)

    add_h1(doc, "明日优先级")
    for item in [
        "S25/S22：基于 dirty-root probe 和 S25 inventory 选择 clean/pruned release slice；禁止 dirty-root deploy。",
        "S11/S04/S05/S09/S02/S15：先修复并重跑学生 core Practice、registration、lesson、dashboard gates。",
        "S18/S21/S23/S24：决定 Mainland PEP High withdrawn illustration replacement/zero-image policy，并给 S11 可测 contract。",
        "S06/S12/S22：在 slice 与 parity 准备好后重跑 Visualization backend/API、production route/storage parity、strict verifier。",
        "S19/S12/S22：保持 LRS Production 监控，整理 reviewable source diff；Preview env 等 connected Git 决策。",
    ]:
        add_number(doc, item)

    add_h1(doc, "需要 Dr. Peter Hu 决策")
    for item in [
        "是否批准 S25/S22 走 clean/pruned release worktree 或 staging path；不建议 dirty-root deploy override。",
        "学生端红灯修复优先级：先 Practice/registration/lesson，还是先 Mainland/California 课程链？",
        "Mainland PEP High 已撤回插图：是否要求 S24 先补 exact-layer replacement，或在短期 release 中继续 zero-image contract？",
        "LRS：Production 已通；是否需要补 Vercel Preview 连接 Git 或接受 Production-only 验证路径？",
        "E2E artifact retention：是否调整 `npm run test:analytics` 的 `.tmp` 清理行为，避免 future evidence cache 被轻量测试清除？",
    ]:
        add_bullet(doc, item)

    add_h1(doc, "证据与边界说明")
    add_body(doc, "S10 读取 AGENTS.md、automation memory、coordination/session-logs、coordination/blockers、coordination/release-intake、coordination/reports，并运行 git/disk/type-check/analytics/release-guard probes。未复制、打印、总结或记录任何真实 credential value。")
    add_body(doc, "本报告是 coordination artifact，不代表 feature-code approval、Git staging/commit、branch、push、deploy approval 或 production write approval。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_doc()
    print(path)
