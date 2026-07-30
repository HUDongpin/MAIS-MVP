from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-27-president-report.docx")
REPORT_DATE = "2026-07-27"
WINDOW = "2026-07-26 08:00 至 2026-07-27 08:00 Asia/Shanghai（UTC+8）"
FORMAT = "President Report Business Brief v1"

FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"
RISK_FILL = "FCE4D6"


def set_run_font(run, size=11, bold=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    language = rpr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        rpr.append(language)
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "zh-CN")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_widths(table, widths):
    if sum(widths) != 9360:
        raise ValueError("All report tables must use 9,360 DXA.")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag, value in (("tblW", "9360"), ("tblInd", "120")):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_pr.append(node)
        node.set(qn("w:type"), "dxa")
        node.set(qn("w:w"), value)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)


def keep_row(row, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if header:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)


def add_table(doc, headers, rows, widths, body_size=7.4, row_fills=None, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_widths(table, widths)
    keep_row(table.rows[0], header=True)
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(value), size=8.5, bold=True)
        shade(cell, header_fill)
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        keep_row(row)
        fill = (row_fills or {}).get(row_index)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i < 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(str(value)), size=body_size)
            if fill:
                shade(cell, fill)
    set_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def para(doc, text="", size=10, bold=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), size=15, bold=True, color=BLUE)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    set_run_font(p.add_run(text), size=9.4)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    for kind, text in (("begin", None), ("instrText", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
        if kind in ("begin", "separate", "end"):
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText" if kind == "instrText" else "w:t")
            if kind == "instrText":
                node.set(qn("xml:space"), "preserve")
            node.text = text
        run._r.append(node)


def setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for name in ("Normal", "Heading 1"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("MAIS-MVP | Daily coordination report | A01-A25"), size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run(f"{REPORT_DATE} | {FORMAT} | Page "), size=8.5, color=MUTED)
    add_page_field(footer)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)
    doc.core_properties.title = "MAIS-MVP President Report — 2026-07-27"
    doc.core_properties.subject = "Daily bilingual coordination report for Dr. Peter Hu"
    doc.core_properties.author = "A10 reporting automation"
    doc.core_properties.keywords = "MAIS-MVP, A01-A25, president report"


def agent_rows():
    rows = {
        "A01": ("无新正式交接", "未见 app shell / auth 的独立 handoff。", "不以本窗口 onboarding/layout 变更代替入口验收。"),
        "A02": ("Git 提交证据", "Dashboard next-step 与任务优先级面板；tour 触及 adaptive UI。", "补 dashboard / adaptive 消费面回归。"),
        "A03": ("无新正式交接", "未见 roadmap、grade/topic 结构记录。", "保持课程结构验收边界。"),
        "A04": ("Git 提交证据", "US-CA item QA、ten-frame figures、pager 重构和 Practice tests。", "补 Practice handoff 与 targeted E2E。"),
        "A05": ("Git 提交证据", "lesson menu、pager、state reset、FigureScroll、loading boundary 修复。", "补 lesson route / content acceptance。"),
        "A06": ("Git 提交证据", "11 个 signature lab math audit 均有证明；Ellipse hover 修复。", "clean source 跑浏览器回归。"),
        "A07": ("Git 提交证据", "Tutor moderation、key-gated US provider / Mathpix、US RAG grounding。", "补 provider live / safety handoff，不披露 secrets。"),
        "A08": ("Git 提交证据", "shared types、analytics runner 与 coherence data 随变更更新。", "A08 核对共享语义与回归范围。"),
        "A09": ("无新正式交接", "未见 i18n 或 accessibility 独立审查记录。", "后续覆盖 tour、lesson 与 teacher 新文案。"),
        "A10": ("已完成", "前一窗口总统报告、A10 session log、A22/A25 协调证据已提交。", "本报告保持 08:00 cutoff 与证据分层。"),
        "A11": ("门禁待补", "CI 新增 component / source-regression gates，修复 rotted lesson tests。", "补 post-PR #69–#80 targeted regression 结论。"),
        "A12": ("Git 提交证据", "teacher assessment / assignment persistence 与 server-side teacher operations。", "补 API contract / persistence 验收。"),
        "A13": ("Git 提交证据", "Teacher Overview learning-insights charts 与 CA G1 demo classroom seed。", "补 teacher workflow / data-state handoff。"),
        "A14": ("无新正式交接", "未见 Parent Console 变更或验收。", "不推断 parent acceptance。"),
        "A15": ("Git 提交证据", "CCSS prerequisite DAG、跨年级 backtracking、diagnostic placement。", "补 adaptive semantics / persistence / regression handoff。"),
        "A16": ("无新正式交接", "未见 research / learning-science 新记录。", "保持现状。"),
        "A17": ("边界待补", "lesson-practice 接入 Practice Island stars；无 A17 sign-off。", "确认 reward economy 边界与回归。"),
        "A18": ("Git 提交证据", "US math item audit / CA repair；TEKS exception decision 已记录。", "仍需独立 QA acceptance。"),
        "A19": ("Git 提交证据", ".env.local.example 的 provider variable 更新；无真实环境验证。", "不触碰或记录 secret；补 redacted parity evidence。"),
        "A20": ("无新正式交接", "未见 game-based learning 新 handoff。", "保持 A17/A20 边界。"),
        "A21": ("Git 提交证据", "US RAG grounding 与 AR/FL generated-content audit inputs。", "补 candidate package / RAG handoff。"),
        "A22": ("已部署 / 待补", "main @ 04cdfe8919 已部署并 CERTIFIED_WITH_FINDINGS。", "PR #69–#80 尚无 clean-source release proof。"),
        "A23": ("无 promotion 记录", "未见 candidate-to-live promotion decision。", "保持 A18→A23→A11/A22 gate。"),
        "A24": ("Git 提交证据", "Question figures / FigureScroll 有工程变更；无 exact-layer sign-off。", "需按 A21/A18/A24 链路验收。"),
        "A25": ("已完成", "11:47、15:19、17:47 三张 clean-source map 均为 0 dirty / 0 unmapped。", "任何新候选仍须重新 map。"),
    }
    return [(f"A{i:02d}", *rows[f"A{i:02d}"]) for i in range(1, 26)]


def build():
    doc = Document()
    setup(doc)
    prepared = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M CST")
    para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    add_table(doc, ["Field", "Value"], [
        ("Report date", REPORT_DATE),
        ("Reporting cutoff", "08:00 Asia/Shanghai (UTC+8)"),
        ("Prepared during run", prepared),
        ("Reporting window", WINDOW),
        ("Project", "MAIS-MVP"),
        ("Reporting agent", "A10 — Tooling, docs, and report coordination"),
        ("Formal scope", "A01-A25"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Format", FORMAT),
        ("Evidence basis", "AGENTS.md; coordination logs/decisions/release intake; Git committer history; safe fresh checks."),
    ], [2200, 7160], body_size=8.0, header_fill=LIGHT_BLUE_FILL)

    heading(doc, "Chinese Executive Summary")
    add_table(doc, ["序号", "重点"], [
        ("1", "项目状态为绿中带黄。按 committer timestamp，窗口内 main 记录 78 笔提交（39 merge、39 non-merge），覆盖 158 个不同路径；提交规模证明集成活动，不等同于 A01-A25 的正式验收。"),
        ("2", "产品推进集中在 US math item quality / figures、Visualization proofs、CCSS adaptive placement、Tutor safety/provider/RAG、lesson/practice/onboarding、dashboard 与 teacher workflow；所有表述均以 Git 或已存协调证据为界。"),
        ("3", "发布来源在已验证范围内为绿：A25 于 11:47、15:19、17:47（Asia/Shanghai）三次记录 0 dirty / 0 unmapped；A22 已将 main @ 04cdfe8919 部署并记录 CERTIFIED_WITH_FINDINGS。PR #69–#80 尚无同级 clean-source 发布证明。"),
        ("4", "报告时本地安全检查通过：type-check、analytics 58/58、question-bank、MVP readiness、git diff --check，以及无 staged entries。未运行 post-PR #69–#80 production build、broad Playwright 或新的 live smoke。"),
        ("5", "治理仍需收口：除 A10 报告协调外，未发现本窗口新的 feature-owner handoff、A18 content acceptance 或 A23 promotion decision；TEKS full-text exception 为本窗口唯一新正式 decision record。"),
    ], [700, 8660], body_size=8.0)

    heading(doc, "English Executive Summary")
    add_table(doc, ["No.", "Summary"], [
        ("1", "Project health is green with a yellow evidence layer. Main records 78 commits in the window by committer timestamp (39 merges and 39 non-merges) across 158 distinct paths; integration activity is not formal owner acceptance."),
        ("2", "The train covers US math-item quality and figures, visualization proofs, CCSS adaptive placement, Tutor safety/provider/RAG work, lesson/practice/onboarding, dashboard, and teacher workflows. Claims are limited to Git or recorded coordination evidence."),
        ("3", "The verified release source is green within scope: three A25 maps recorded zero dirty and zero unmapped entries, and A22 deployed main @ 04cdfe8919 as CERTIFIED_WITH_FINDINGS. PRs #69–#80 do not yet have equivalent clean-source release proof."),
        ("4", "Report-time checks passed: type-check, analytics 58/58, question-bank, MVP readiness, whitespace check, and no staged entries. No post-PR #69–#80 production build, broad Playwright, or new live smoke was run."),
        ("5", "Governance remains incomplete: other than A10 report coordination, no new feature-owner handoff, A18 content acceptance, or A23 promotion decision was found. The Texas TEKS full-text exception is the only new formal decision record."),
    ], [700, 8660], body_size=8.0)

    heading(doc, "报告窗口摘要（中文）")
    add_table(doc, ["项目", "结果"], [
        ("正式窗口", WINDOW),
        ("Git 活动", "committer 时间戳 09:36 至 19:37：78 笔（39 merge、39 non-merge）；含 32 个 PR merge 与 7 个同步/repair merge。"),
        ("代码变更量", "non-merge numstat：199 file-change entries，16,984 additions / 1,834 deletions；158 个不同路径。两笔 source commit 于 01:54 / 02:22 authored、09:36 才 committed，已按 committer 时间计入。"),
        ("新正式会话记录", "仅 A10 session log（用于前一窗口总统报告）；未见本窗口新的 feature-owner handoff。"),
        ("新 blocker / promotion", "未发现窗口内新的 blocker report 或 A23 candidate-to-live decision。"),
        ("新 decision", "Texas TEKS full-text exception：仅限 19 TAC Chapter 111；production content freeze 前须对 TEA 再核验。"),
        ("A25 发布入口", "11:47、15:19、17:47 Asia/Shanghai 三张 map 均为 0 dirty、0 unmapped；报告生成时 16:05 的 fresh map 同为 0/0，但属于 cutoff 后观察。"),
        ("A22 发布证据", "main @ 04cdfe8919 已部署到 www.mais.ac / www.mais.hk；full CI 成功，certification 为 CERTIFIED_WITH_FINDINGS。"),
    ], [2450, 6910], body_size=7.15)

    heading(doc, "整体项目进展（中文）")
    bullet(doc, "内容与 Practice：US-CA math item QA 修复进入主线，并把 item audit 扩展到 AR/FL；K-1 ten-frame figures、double-ten-frame 换行和 CI figure audit 均有提交与门禁证据。")
    bullet(doc, "课程与可视化：11 个 signature visualization benches 均补上数学 audit proof，Ellipse hover 参数修复完成；lesson menu、Lesson↔Practice Arena pager、in-round state 和 FigureScroll 亦有修复。")
    bullet(doc, "自适应与学生体验：CCSS prerequisite DAG、跨年级 backtracking、diagnostic placement、dashboard next-step / urgency order 与 student guided tour 进入集成历史。")
    bullet(doc, "Tutor、教师与平台：AI Tutor moderation、key-gated US provider / Mathpix handwriting、US-curriculum RAG grounding、teacher charts 和 CA G1 demo classroom seed 进入主线；真实 provider 与生产环境均未在本报告中暴露或假定为已验收。")
    bullet(doc, "质量与发布：CI 新增 content-safety、source-regression 和 discovery-based component test gates；A22 的已部署来源有 full CI 与认证记录。后续合并包仍须 A11/A22 在明确 clean source 上验证。")

    heading(doc, "A01-A25 会话状态表（中文）")
    status_fills = {i: PENDING_FILL for i in range(25)}
    for i in (9, 21, 24):
        status_fills[i] = OK_FILL
    status_fills[10] = RISK_FILL
    add_table(doc, ["Agent", "状态", "窗口内证据", "下一步 / 风险"], agent_rows(), [620, 1340, 3950, 3450], body_size=6.0, row_fills=status_fills)

    heading(doc, "已完成工作（中文）")
    add_table(doc, ["负责方 / 包", "完成事项", "可核验证据"], [
        ("A04 + A18", "US-CA item QA 修复、AR/FL audit 扩展、K-1 ten-frame question figures 与 figure audit。", "PR #51/#54；question-bank gate 在报告时通过。"),
        ("A06", "11 个 signature visualization math audit proofs；Ellipse hover parameter 修复。", "PR #47/#50/#52；无 A06 formal handoff。"),
        ("A05 + A04", "Lesson menu / pager / state / FigureScroll 修复；lesson practice 进入 Practice Island stars。", "PR #55/#60/#62–#66/#74；需 UI acceptance。"),
        ("A15", "CCSS coherence DAG、跨年级 prerequisite repair 与 diagnostic placement。", "PR #56/#72；analytics 58/58 通过。"),
        ("A07 + A12", "Tutor moderation、provider / handwriting route、US RAG grounding 与 server persistence。", "PR #57/#58/#76；无 fresh provider/live handoff。"),
        ("A02 + A13", "Dashboard next-step / urgency order；teacher learning-insights charts 和 CA G1 seed。", "PR #77/#78；需 workflow acceptance。"),
        ("A10 + A11 + A22 + A25", "CI gate 扩充、production certification orchestrator、三份 clean-source maps、main-sync deployment record。", "PR #68/#71/#73/#80；A22 source 仅覆盖 @04cdfe8919。"),
    ], [1820, 4700, 2840], body_size=6.65, row_fills={6: OK_FILL})

    heading(doc, "进行中工作（中文）")
    add_table(doc, ["负责方", "进行中 / 待收口事项", "安全下一步"], [
        ("A11 + A22", "PR #69–#80 在 A22 已部署来源之后合并，尚未取得等价 build、broad Playwright、staging/live evidence。", "指定 clean main/worktree，形成独立 release packet。"),
        ("A04/A05/A06/A07/A12/A13/A15", "主要功能已有提交，但 formal handoff 未记录范围、checks、未覆盖流与 rollback boundary。", "各 owner 补最小 handoff；不把 commit message 当验收。"),
        ("A18 + A21 + A23", "question/RAG/TEKS 活动没有本窗口独立 A18 acceptance、A21 candidate handoff 与 A23 promotion decision。", "继续 candidate-only，完整链路后才进入 live surface。"),
        ("A19", "provider variable / .env.local.example 变更没有真实环境 parity evidence。", "仅补 redacted inventory 与 owner-approved live smoke。"),
        ("A25 + A10", "每次新 release candidate 需要新的 dirty-map，报告时 clean root 不是发布授权。", "保持 source、commit、checks 和时间戳可追溯。"),
    ], [1750, 4700, 2910], body_size=6.75, row_fills={0: RISK_FILL, 2: PENDING_FILL})

    heading(doc, "阻塞项（中文）")
    add_table(doc, ["阻塞项", "证据", "Owner / 安全下一步"], [
        ("本窗口无新增正式 blocker report", "coordination/blockers/ 未见窗口内新文件；历史 blocker 不自动解除。", "维持历史 blocker 状态，不将其误报为本次完成。"),
        ("后续合并包无发布证明", "A22 deployment / certification 指向 main @ 04cdfe8919；PR #69–#80 随后进入 main。", "A11/A22 在 clean source 重新跑门禁；未完成前不称 release-ready。"),
        ("正式 owner handoff 缺失", "除 A10 外未见 feature-owner session handoff。", "相关 owner 记录范围、checks、未覆盖流、依赖和 rollback。"),
        ("内容推广证据缺失", "无 A18 independent acceptance、A21 package handoff 或 A23 promotion decision。", "冻结 candidate-to-live promotion。"),
    ], [2050, 4110, 3200], body_size=6.7, row_fills={1: RISK_FILL, 2: PENDING_FILL, 3: RISK_FILL})

    heading(doc, "风险（中文）")
    bullet(doc, "集成风险：单窗口 78 笔提交横跨 UI、API、server persistence、data、tests、CI 与 release tooling；TypeScript / unit-style green 不覆盖所有浏览器和生产路径。")
    bullet(doc, "时间线风险：Git 筛选采用 committer time。两笔 question-bank source commits 在 08:00 前 authored、09:36 才进入 main；报告已明确区分作者时间和集成时间。")
    bullet(doc, "生产风险：A22 的 P0 认证结果为通过，但仍保留既有 P1 findings（缺 x-content-type-options、guest 页面 401 console error）；新的合并包尚无 live-domain 检查。")
    bullet(doc, "课程与合规风险：TEKS 完整法规文本例外只限 19 TAC Chapter 111；Texas SBOE 正在 review，production content freeze 前应以 TEA 原始来源复核。")
    bullet(doc, "治理风险：没有 formal handoff 时，难以准确评估未覆盖用户流、跨包依赖和回滚边界，尤其是 Tutor、adaptive、teacher persistence 与内容推广。")

    heading(doc, "测试 / 构建状态（中文）")
    add_table(doc, ["检查", "结果", "解释"], [
        ("A25 fresh dirty-map", "PASS", "2026-07-27 16:05 Asia/Shanghai：0 dirty、0 unmapped、0 ambiguous；仅为 cutoff 后 inventory。"),
        ("npm run type-check", "PASS", "报告时当前 checkout 通过。"),
        ("npm run test:analytics", "PASS", "58/58 通过。"),
        ("npm run test:question-bank", "PASS", "报告时通过；覆盖本窗口 question / figure 相关 gate。"),
        ("npm run test:mvp", "PASS", "报告时 MVP readiness 测试通过。"),
        ("git diff --check / staged", "PASS", "未发现 whitespace errors；无 staged entries。"),
        ("A22 historical clean-source CI", "PASS / scoped", "main @ 04cdfe8919 full CI 成功；不能外推至 PR #69–#80。"),
        ("production build / broad Playwright / new live smoke", "Not run", "本报告不在 dirty/root inventory 代替 A11/A22 的 clean-source release gate。"),
    ], [2300, 1500, 5560], body_size=6.8, row_fills={0: OK_FILL, 1: OK_FILL, 2: OK_FILL, 3: OK_FILL, 4: OK_FILL, 5: OK_FILL, 6: PENDING_FILL, 7: RISK_FILL})

    heading(doc, "变更文件（中文）")
    add_table(doc, ["类别", "代表性窗口内文件", "说明"], [
        ("Practice / question quality", "data/usCaliforniaQuestions.ts；data/usCaliforniaPracticeFigures.ts；components/practice/QuestionFigure.tsx；scripts/audit-us-*-math-item-quality.mjs", "US item QA、K-1 figure 渲染与 CI audit。"),
        ("Lesson / learning", "components/lesson/LessonView.tsx；components/lesson/worlds/*；components/lesson/ccss/FigureScroll.tsx；app/*/loading.tsx", "pager、menu、state、scroll affordance 与 streaming cleanup。"),
        ("Adaptive / dashboard", "lib/adaptiveLearning.ts；lib/diagnosticPlacement.ts；data/ccssCoherenceMap.ts；app/dashboard/page.tsx", "CCSS DAG、placement 与 next-step surface。"),
        ("Tutor / server", "app/api/ai-tutor/resolve/route.ts；lib/server/llmProvider.ts；lib/server/usMathTutorStandards.ts；lib/server/userStore/*", "moderation、provider / OCR、US RAG 与 persistence。"),
        ("Teacher / onboarding", "components/teacher/*；app/student/assessments/[assessmentId]/page.tsx；components/onboarding/*", "charts、CA G1 seed 和 guided tour。"),
        ("CI / release / governance", ".github/workflows/ci.yml；package.json；scripts/prod-certification.mjs；coordination/release-intake/*；coordination/decisions/*", "quality gates、certification、clean-source evidence 与 TEKS decision。"),
    ], [1900, 4800, 2660], body_size=6.25)

    heading(doc, "明日优先事项（中文）")
    add_table(doc, ["负责方", "优先事项", "完成定义"], [
        ("A11 + A22", "为 PR #69–#80 建立干净源 release package。", "记录 commit、build、targeted / broad browser coverage、staging/live result与未覆盖项。"),
        ("A04/A05/A06/A07/A12/A13/A15", "补已合并功能的 formal owner handoff。", "包含文件范围、checks、未覆盖用户流、依赖与 rollback boundary。"),
        ("A18 + A21 + A23", "复核 question/RAG/TEKS 后续工作没有绕过 promotion gate。", "A18 acceptance、A21 handoff、A23 decision 依次可核验。"),
        ("A25 + A10", "任何新的 release candidate 先重跑 dirty-map 并维持报告证据分层。", "clean source、0 unmapped、报告时间和 commit 可追溯。"),
        ("A19", "仅在授权范围内补 redacted provider/environment parity evidence。", "不输出、保存或提交任何 credential 值。"),
    ], [1850, 4300, 3210], body_size=6.8, row_fills={0: RISK_FILL, 2: PENDING_FILL})

    heading(doc, "需要 Owner 决策（中文）")
    add_table(doc, ["决策", "请确认"], [
        ("后续发布门槛", "对 PR #69–#80，是否要求 A11 targeted + broad browser evidence、A22 clean-source build 与必要 live smoke 全部记录后，才允许 promotion / release-ready 表述？"),
        ("合并后 handoff 规则", "是否要求每个 merged A01-A25 package 在下一个 08:00 cutoff 前提交最小 handoff；未交接的工作仅标为 Git evidence？"),
        ("内容推广边界", "是否继续明确：US item / RAG / TEKS 相关提交在 A18 acceptance、A21 handoff、A23 promotion、A11/A22 gate 前，均不视为 live curriculum acceptance？"),
        ("Texas 后续范围", "是否授权 A18/A21 以已批准的 19 TAC Chapter 111 exception 开展 Phase 0 standards / crosswalk candidate work，并在 production freeze 前复核 TEA 原始版本？"),
    ], [2050, 7310], body_size=7.1, row_fills={0: PENDING_FILL, 1: PENDING_FILL, 2: PENDING_FILL, 3: PENDING_FILL})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
