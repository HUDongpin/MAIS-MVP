# -*- coding: utf-8 -*-
"""Build the 2026-06-13 MAIS-MVP president report DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "coordination/reports/2026-06-13-president-report.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "555555"
HEADER_FILL = "F2F4F7"
GOOD_FILL = "E2F0D9"
WARN_FILL = "FFF2CC"
RISK_FILL = "FCE4D6"


def dxa(inches: float) -> int:
    return int(round(inches * 1440))


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Calibri"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("MAIS-MVP President Report")
    set_run_font(r, 20, True, DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Daily AI coordination report for Dr. Peter Hu / Dr. Peter Hu 每日 AI 协作报告")
    set_run_font(r, 10.5, False, MUTED)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, 11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, 10.3)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    width_dxa = [dxa(width) for width in widths]

    for row in table.rows:
        for cell, width_in, width_twips in zip(row.cells, widths, width_dxa):
            cell.width = Inches(width_in)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(width_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in width_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 8.4, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 8.4,
    status_column: int | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], HEADER_FILL)
        set_cell_text(table.rows[0].cells[i], header, True, font_size, DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            if status_column is not None and i == status_column:
                if any(token in text for token in ["完成", "通过", "READY", "Passed"]):
                    shade_cell(row.cells[i], GOOD_FILL)
                elif any(token in text for token in ["阻塞", "失败", "Blocked", "风险"]):
                    shade_cell(row.cells[i], RISK_FILL)
                elif any(token in text for token in ["进行", "待", "需"]):
                    shade_cell(row.cells[i], WARN_FILL)
            set_cell_text(row.cells[i], text, False, font_size)
    set_table_widths(table, widths)
    doc.add_paragraph()


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in rows:
        row = table.add_row()
        shade_cell(row.cells[0], HEADER_FILL)
        set_cell_text(row.cells[0], key, True, 9.3, DARK_BLUE)
        set_cell_text(row.cells[1], value, False, 9.3)
    set_table_widths(table, [1.65, 4.85])
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    style_doc(doc)
    add_title(doc)

    add_kv_table(
        doc,
        [
            ("Report date", "2026-06-13"),
            ("Report time", "8:00 AM Asia/Hong_Kong"),
            ("Runtime", "2026-06-13 08:02-08:13 HKT"),
            ("Project", "MAIS-MVP"),
            ("Reporting session", "S10"),
            ("Audience", "Dr. Peter Hu"),
            ("Reporting window", "2026-06-12 08:00 to 2026-06-13 08:00 Asia/Hong_Kong"),
            ("Automation ID", "mais-mvp-9-am-president-report"),
        ],
    )

    doc.add_heading("中文 Executive Summary", level=1)
    add_bullets(
        doc,
        [
            "本窗口有大量实质推进：学生入口、Dashboard、Practice、Lesson、Visualization、AI Tutor、Teacher/Parent console、游戏、内容 QA、插图生产、集成和 Vercel 发布均有新证据。",
            "质量状态比昨日更好：本次 S10 新跑 `npm run type-check`、`npm run test:analytics` 和 `npm run build` 全部通过；S22 也完成多次生产发布，其中 Texas K-G5 插图页已上线。",
            "最大管理风险不是单点构建，而是工作树切片：当前 `git status --short` 有 3,651 项，含 285 modified、48 deleted、3,318 untracked。需要 S25/S22 把可发布工作拆成可审查切片。",
            "内容链条推进很快，但多数包仍是候选或 integration review，不等于 production approval。S18/S24/S23/S11/S22 的签核顺序不能跳过。",
            "今天最需要 Dr. Peter Hu 决策：是否继续允许 dirty-root guarded deploy、三档难度是否成为全项目规范、以及下一批内容/插图的生产优先级。",
        ],
    )

    doc.add_heading("English Executive Summary", level=1)
    add_bullets(
        doc,
        [
            "The reporting window shows broad progress across UI polish, student learning flows, Visualization Lab, AI Tutor, teacher/parent consoles, games, content QA, illustration production, integration, and production deployment.",
            "Fresh S10 checks are green: `npm run type-check`, `npm run test:analytics` with 23 passing tests, and `npm run build` all passed. S22 also published production updates, including the Texas K-G5 illustration route.",
            "The main release risk is still coordination hygiene, not a current compile failure: the root tree has 3,651 dirty entries and needs S25/S22 slicing before clean release claims.",
            "Content velocity is high, but many assets remain candidate-only or approved for integration review only. Production claims still require S18/S24/S23/S11/S22 gates.",
            "Owner decisions needed: dirty-root deploy policy, project-wide Low/Medium/High difficulty normalization, and which curriculum/illustration stream should receive the next production push.",
        ],
    )

    doc.add_heading("报告窗口摘要", level=1)
    add_kv_table(
        doc,
        [
            ("Window", "2026-06-12 08:00 至 2026-06-13 08:00 HKT"),
            ("Active sessions", "S01, S02, S04, S05, S06, S07, S09, S10, S12, S13, S14, S18, S20, S21, S22, S23, S24"),
            ("No fresh logs", "S03, S08, S11, S15, S16, S17, S19, S25"),
            ("Blocker", "1 个：S18 BNU primary question-bank regression 被共享 TypeScript / nested-copy 问题阻止。S10 新鲜 type-check/build 已通过，但该专项回归仍需重跑后才能关闭。"),
            ("Deployment", "S22 记录 California middle-school、English-only US view、current online production、Texas K-G5 illustration production publish；主要别名 `mais.hk` / `www.mais.hk` / `mais-mvp.vercel.app` 指向 Ready 部署。"),
            ("Current tree", "3,651 status entries: 285 modified, 48 deleted, 3,318 untracked。"),
        ],
    )

    doc.add_heading("整体项目进展", level=1)
    add_para(
        doc,
        "MAIS-MVP 在本窗口从单点页面修补推进到更完整的学习平台闭环：学生端入口和课程页面更简洁，Visualization Lab 完成 756/756 深链运行扫测，AI Tutor 的文本/图像/语音边界被验证，Teacher Console 和 Parent Console 增加了更多可执行工作流。生产环境也完成了 owner-requested guarded deploy，Texas K-G5 图例页面已上线。"
    )
    add_para(
        doc,
        "同时，内容团队完成了多条中国大陆和美国课程 QA/插图链路：HJB junior 88 个缺陷题被修复，BNU junior 12 个目标行被修复，Texas K-G5 300 张 deterministic 图例通过 S18/S09 integration-review 审查并由 S23/S22 推到生产；PEP high 旧风格已撤下并生成候选重绘包。"
    )
    add_para(
        doc,
        "当前最大瓶颈仍是 release intake：大量 session 同时产出，根目录 dirty-tree 很大。虽然新鲜 type-check / analytics / build 都通过，仍需要 S25 ownership map、S22 clean/pruned deploy hygiene、S11 targeted regression 才能把这些成果转为低风险发布切片。"
    )

    doc.add_heading("S01-S25 Session 状态表", level=1)
    session_rows = [
        ["S01", "完成", "Home/login/register/nav/language 大量 UI polish：移动 hero、课程下拉、登录/注册背景、demo account、语言菜单。", "多次 `npm run type-check` 与 Playwright/Browser desktop/mobile checks 通过。", "纳入 S25 切片；避免继续扩大同文件 diff。"],
        ["S02", "完成", "Dashboard 改为 US grade labels；移除学生 leaderboard；增加 Lesson/课时 shortcut 与中英提示。", "`npm run type-check` 与 dashboard desktop/mobile checks 通过。", "S11 可补 dashboard regression。"],
        ["S03", "无新日志", "Curriculum roadmap 无窗口内新证据。", "未运行。", "等待后续 roadmap/content assignment。"],
        ["S04", "完成", "Practice Arena 删除 `Question bank ready / Collect Gems / Earn Stars` 状态奖励条。", "`npm run type-check`、desktop/mobile `/practice` checks 通过。", "S11 可补 practice visual regression。"],
        ["S05", "完成/局部待回归", "Lesson 删除 galaxy/fullscreen/AI pill 等干扰，重排为左侧目录+右侧内容，统一卡片和 Jump button。", "多数局部 type/browser checks 通过；后续几项被 unrelated nested-copy type errors 阻止。", "S11/S22 清洁环境后做 Lesson regression。"],
        ["S06", "完成", "Visualization Lab 删除 preview 控件，修复坐标反射、deep-link、tile cap、template mapping；完成 756/756 runtime sweep。", "全量 desktop deep-link 756/756 PASS，mobile representative 31/31 PASS；报告已写。", "S18 做课程/数学语义签核。"],
        ["S07", "完成/有 follow-up", "Nova Tutor 命名、选中文本、图片附件、语音/语音状态边界验证；晚间计划 resize handle。", "API/compiler checks 0 diagnostics；lesson selection 4 passed；AI Tutor smoke 通过。", "S19 确认 live provider env；完成 resize 验证。"],
        ["S08", "无新日志", "Shared state/analytics 无窗口内新日志。", "本次 S10 fresh analytics 23/23 PASS。", "如决定三档难度规范，S08/S10 共同 owning migration。"],
        ["S09", "完成", "Texas K-G5 300 张 deterministic 图例做 copy/accessibility review。", "Package audit 300 approved-for-integration-review，0 package/row issues。", "生产 approval 仍需 owner/human gate。"],
        ["S10", "完成", "上一日报告、三档难度协调清理、AGENTS no-S26+ 边界更新；本次生成 2026-06-13 DOCX。", "`git diff --check` / rg / table scan 通过；本次 fresh checks 绿。", "继续做 daily synthesis 与 release slicing coordination。"],
        ["S11", "无新日志", "QA/release regression 无窗口内新日志。", "未运行 broad E2E。", "应优先教师/家长/Practice/Game/Question-bank targeted gates。"],
        ["S12", "完成/部分套件红", "Auth JSON guard/rate-limit/cache-control；admin storage export 红acted snapshot；tenant/revision metadata。", "Focused API/unit/type/build passed；backend Playwright 2 passed, 3 existing suite-drift failures。", "S11/S02/S15 处理 backend E2E drift；未来 durable limiter/RLS。"],
        ["S13", "进行中", "Teacher Console 补大量 enterprise loops：dashboard launch、roster slips、reports、queues、resources、review lessons、live follow-up、rewards、prep teams、analytics、lesson-kit publish、grading、mastery、inbox。", "多次 temporary type-check excluding nested copy passed；browser evidence 局部；official type-check 多次受 nested-copy 阻止。", "需 S22 dev/build isolation + S11 maintained E2E。"],
        ["S14", "完成/待复测", "Parent Console 增强 overview、child detail、reports、messages、notices。", "`npm run type-check` 与 desktop route sweep 通过；mobile overview 通过。", "Dev-server conflict 后需补 full mobile sweep。"],
        ["S15", "无新日志", "Adaptive engine 无窗口内新日志。", "本次 analytics/adaptive unit tests 23/23 PASS。", "如 teacher mastery target 进入 adaptive，需 S15 regression。"],
        ["S16", "无新日志", "Research/learning science 无窗口内新日志。", "未运行。", "可支持 adaptive/content evaluation design。"],
        ["S17", "无新日志", "Gamification economy 无窗口内新日志。", "未运行。", "S20 draft games 接入真钱包前必须协调 S17/S12。"],
        ["S18", "完成/有内容 gates", "BNU/HJB/PEP/US CA/AR/TX 多条 QA；HJB junior 88 repaired；BNU junior 12 repaired；PEP high旧图撤下重绘；Texas K-G5 integration review。", "多项 package-local checks PASS；BNU primary question-bank regression 被共享编译问题阻止。", "明确 production approval 边界；继续 S18/S24 review。"],
        ["S19", "无新日志", "API/env 无窗口内新日志。", "S22 env preflight 7/7 required production variable names present。", "Provider live checks仍按 S19/S07/S15 owner approval。"],
        ["S20", "完成/内部 draft", "Adventure/Fishing audit，Fishing aim/feedback；Math Match Quest、Mighty Tank Battle、Math Virus Blaster 内部 draft games。", "Game logic tests、targeted compiler checks、browser checks 多项通过；full checks曾受 unrelated blockers。", "不要上导航；S11/S17/S12 gates 后再生产化。"],
        ["S21", "进行中", "California/Arkansas K-G5 candidate generation；Texas K-G5 300 deterministic diagrams；PEP primary/junior v2 review packages。", "Package-local validation PASS；生产 TypeScript/build 多数不适用。", "继续 AR/CA generation；交 S18/S24 review。"],
        ["S22", "完成/风险受控", "Clean preview、California/English-only/direct production、Texas K-G5 production publish；preflights/build/smoke。", "release/env/publish preflights、type-check/build、production smoke PASS；Ready deployments。", "dirty-root deploy policy 需 owner确认；继续 clean staging。"],
        ["S23", "完成", "HJB high default治理；AR G6-G8 local illustration integration；Texas K-G5 public/live promotion。", "Audit/validation/type-check/build/browser/curl checks PASS；AR broad test blocked by unrelated MathMatchQuest error。", "继续候选到 live 的 gate coordination。"],
        ["S24", "完成/候选待审", "Texas K-G5 300 exact-layer diagrams；AR G6-G8 exact-layer repair；Texas G6-G8/HS assets staging；PEP primary v2 candidates。", "Package validators、PNG counts、dimension checks PASS。", "高阶 exact overlays 与 final approval 仍需 S18/human。"],
        ["S25", "无新日志", "Git hygiene/release intake 无窗口内新日志。", "S10 fresh dirty-tree inventory: 3,651 entries。", "最优先：ownership map + PR/commit slicing plan。"],
    ]
    add_table(doc, ["Session", "Status", "Window work", "Evidence / checks", "Next"], session_rows, [0.45, 0.75, 2.55, 1.65, 1.10], 7.2, 1)

    doc.add_heading("已完成工作", level=1)
    add_bullets(
        doc,
        [
            "学生入口和学习页：Home、login/register、dashboard、practice、lesson UI 的多处 owner-requested 清理和移动端修复完成。",
            "Visualization Lab：756 个 deep-linked labs 两轮全量 desktop runtime sweep 全部 PASS；mobile representative sweep 31/31 PASS。",
            "AI Tutor：文本选择、图片附件、语音/语音合成未配置边界、guest/auth 边界得到 targeted verification。",
            "Backend/storage：Auth API hardening、rate-limit headers、redacted admin storage snapshot、tenant/revision metadata 完成。",
            "Teacher/Parent Console：教师端闭环大量增强；家长端 reports/messages/notices/child workflows 更完整。",
            "Games：Fishing aim/feedback 修复，Adventure diagnostics 增强，三个内部 draft games 生成并隐藏在直接 route 中。",
            "Content QA：HJB junior 88 个缺陷题修复，BNU junior 12 个目标行修复，多条 mainland/US 内容和插图审查完成。",
            "Release：California/US English-only/current site/Texas K-G5 多次生产发布 Ready；本次 fresh build 也通过。",
        ],
    )

    doc.add_heading("进行中工作", level=1)
    add_bullets(
        doc,
        [
            "S13 Teacher Console 仍需清洁环境下的 end-to-end regression，不应只依赖逐页 ad hoc smoke。",
            "S21 California/Arkansas K-G5 candidate generation 还未完成全部 300-slot coverage；所有候选仍需 S18/S24/human review。",
            "PEP primary/junior v2 illustration packages 是 review-only，尚未进入 live approval list 或 student-facing integration。",
            "California high-school 20 concept redraws 只获 placement-review/candidate 背景资格，0 production-approved。",
            "PEP high 旧风格已撤下，22 个概念候选包待 S18/S24/owner approval 后再集成。",
            "S07 resize follow-up 已规划但日志中未见最终验证闭环。",
        ],
    )

    doc.add_heading("Blockers", level=1)
    add_bullets(
        doc,
        [
            "正式 blocker：`coordination/blockers/2026-06-12-S18-bnu-primary-question-bank-regression.md`。`npm run test:question-bank` 未进入 Node tests，因共享 TypeScript/schema 与 nested-copy compile issue 被阻止。",
            "S10 fresh checks 现在显示 `npm run type-check` 和 `npm run build` 通过；但 S18 blocker 的专项 question-bank regression 未重跑，不能自动视为业务质量已关闭。",
            "S13/S14/S20 多次遇到共享 `.next` / dev-server stale chunk 或端口冲突；需要 S22-style isolated build/dev output 才能作为 release evidence。",
            "S20 Fishing historical stress harness 不符合产品 eligibility 规则；不能把旧 90/90 red stress 作为产品失败证据，需 S11/S20 修 harness。",
            "内容生产的主要 blocker 是审批链：candidate-only / integration-review 不等于 production release。",
        ],
    )

    doc.add_heading("Risks", level=1)
    add_bullets(
        doc,
        [
            "Dirty tree 高风险：3,651 status entries 会增加 review、merge、deploy 回滚和 release attribution 难度。",
            "Dirty-root guarded deploy 虽然本窗口成功，但不能替代长期 clean branch / release slicing 流程。",
            "许多新游戏和 teacher workflows 还缺 maintained Playwright coverage；当前局部 smoke 不能等价为全产品回归。",
            "S12 auth rate limits 仍是 in-memory/per-instance；production abuse protection 需要 durable store 或 edge/WAF 策略。",
            "部分 student routes first-load JS 较大，build route table 显示 `/student/lessons`、roadmap、visualizations 约 2.8 MB；后续应做 performance pass。",
            "PEP primary question images 872 张被记录为禁止学生展示；若未严格保持 approval gate，存在数学不一致风险。",
        ],
    )

    doc.add_heading("Test / Build 状态", level=1)
    test_rows = [
        ["Fresh S10", "`npm run type-check`", "通过", "2026-06-13 08:05 HKT，`tsc --noEmit --incremental false` clean。"],
        ["Fresh S10", "`npm run test:analytics`", "通过", "23 tests PASS；learning analytics + adaptive tests。"],
        ["Fresh S10", "`npm run build`", "通过", "Next build 使用 `tsconfig.next.json`，137 static pages generated。"],
        ["S22", "Release/env/publish preflights", "通过", "Production variable names 7/7 present; guarded publish preflights passed。"],
        ["S22", "Production smoke", "通过", "`mais.hk` aliases、login/register/lesson/practice/visualization/API smoke returned expected 200/401。"],
        ["S06", "Visualization runtime sweeps", "通过", "Desktop 756/756 PASS；mobile representative 31/31 PASS。"],
        ["S12", "Backend broad Playwright", "部分失败", "2 passed, 3 suite-drift failures：registration grade、adaptive difficulty label、temporary password redirect。"],
        ["S18", "`npm run test:question-bank`", "阻塞", "BNU primary blocker recorded；fresh type/build green but专项回归需重跑。"],
        ["S11", "Full E2E release suite", "未运行", "本窗口无 S11 新日志；建议今日补 targeted gates。"],
    ]
    add_table(doc, ["Source", "Check", "Status", "Notes"], test_rows, [0.8, 1.75, 0.7, 3.25], 8.3, 2)

    doc.add_heading("Changed Files", level=1)
    add_para(doc, "当前 root dirty-tree 摘要：3651 status entries；285 modified、48 deleted、3318 untracked。以下为报告窗口内高价值变更分组，不展开所有 generated/candidate assets。")
    changed_rows = [
        ["Coordination/reporting", "`AGENTS.md`, `coordination/session-logs/2026-06-12-S*.md`, S18 blocker, S23/S24 integration notes, S22/S06/S20 reports。"],
        ["Student UI", "`app/page.tsx`, `app/login/page.tsx`, `app/register/page.tsx`, `app/dashboard/page.tsx`, `app/practice/page.tsx`, `components/home/`, `components/layout/`, `components/lesson/`, `components/practice/`, `components/ui/LanguageToggle.tsx`。"],
        ["Visualization", "`components/visualizations/VisualizationLabPage.tsx`, `ConfiguredVisualizationLab.tsx`, `CoordinatePlaneDemo.tsx`, `data/visualizationLabs.ts`。"],
        ["AI/backend", "`components/ai/`, `app/api/ai-tutor/`, `app/api/nova-lens/`, `app/api/auth/*`, `app/api/me/route.ts`, `lib/server/userStore.ts`, `lib/server/authRouteGuards.ts`。"],
        ["Teacher/Parent", "`app/teacher/`, `components/teacher/`, `app/parent/`, `components/parent/`。"],
        ["Games", "`components/gamification/FishingGame.tsx`, `AdventureIslandGame.tsx`, `components/games/`, `app/games/`, `data/*Battle*.ts`, `lib/*Battle*.ts`, `public/games/`。"],
        ["Content/illustrations", "`coordination/content-qa/`, `data/generated-content/`, `public/lesson-illustrations/`, `public/question-illustrations/`, selected `data/*Questions.ts` / `data/*Lessons.ts` / `lib/*test.ts` files。"],
        ["Release config", "`next.config.ts`, `tsconfig.json`, `tsconfig.next.json`, `.vercelignore`, deployment scripts/reports`。"],
    ]
    add_table(doc, ["Group", "Representative paths"], changed_rows, [1.35, 5.15], 8.3)

    doc.add_heading("明日优先级", level=1)
    add_bullets(
        doc,
        [
            "S25 先做 dirty-tree ownership map 和 release/PR slice recommendation；不要继续把所有成果堆在 root dirty tree。",
            "S22 固化 clean/pruned build 与 dev-server isolation；保留 dirty-root deploy override 只作为 owner-approved emergency path。",
            "S11 跑 targeted regression：login/register, dashboard, practice, lesson, visualization, teacher, parent, games, question-bank。",
            "S18/S24/S23 继续把 candidate-only 包推进到明确 gate：reject、integration-review、owner-approved、production-ready，不要混用状态词。",
            "S13 与 S14 在 isolated environment 下补 browser regressions，把 teacher/parent enterprise loops 转为 maintained tests。",
            "S21 继续 Arkansas/California K-G5 candidate coverage；完成后交 S18/S24，而不是直接进入 public/live routes。",
        ],
    )

    doc.add_heading("Owner Decisions Needed", level=1)
    add_bullets(
        doc,
        [
            "决定 dirty-root deploy policy：是否继续允许 `MAIS_ALLOW_DIRTY_ROOT_DEPLOY=1` guarded publish，或要求所有后续生产发布先由 S25/S22 切 clean release slice。",
            "确认难度体系方向：全项目正式迁移为 Low/Medium/High，还是保留 Foundation/Core/Challenge/Exam 为 canonical 并在 UI 派生三档。",
            "确认 Texas K-G5 已上线图例是否进入正式验收口径；若是，安排 S11/S22 做 post-production regression record。",
            "选择下一条内容生产优先级：Arkansas/California K-G5 generation、PEP primary v2、PEP junior question illustration v2、PEP high redraw、或 California high-school redraw/exact-layer。",
            "确认是否给 S11 一个专门窗口来把 S13 teacher enterprise loops、S14 parent console、S20 games 转成 maintained Playwright coverage。",
            "确认 PEP primary question images 的 872 个失败历史记录是否需要进一步生成 owner-facing archive package，或只保留 coordination record。",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
