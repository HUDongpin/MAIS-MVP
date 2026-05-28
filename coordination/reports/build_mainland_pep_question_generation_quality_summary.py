from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MARKDOWN_PATH = ROOT / "coordination/reports/2026-05-23-mainland-pep-question-generation-quality-summary.md"
DOCX_PATH = ROOT / "coordination/reports/2026-05-23-mainland-pep-question-generation-quality-summary.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_east_asia_font(run, font_name: str = "Microsoft YaHei") -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_font(paragraph, font_name: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        set_east_asia_font(run, east_asia)


def set_style_font(style, size_pt: float, color: str | None = None, bold: bool | None = None) -> None:
    font = style.font
    font.name = "Calibri"
    font.size = Pt(size_pt)
    if color:
        font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_spacing(style, before_pt: float, after_pt: float, line: float) -> None:
    fmt = style.paragraph_format
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(after_pt)
    fmt.line_spacing = line


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    set_style_font(normal, 11, "000000")
    set_spacing(normal, 0, 6, 1.10)

    title = styles["Title"]
    set_style_font(title, 20, "0B2545", True)
    set_spacing(title, 0, 8, 1.10)

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, 10.5, "555555")
    set_spacing(subtitle, 0, 10, 1.10)

    h1 = styles["Heading 1"]
    set_style_font(h1, 16, "2E74B5", True)
    set_spacing(h1, 16, 8, 1.10)

    h2 = styles["Heading 2"]
    set_style_font(h2, 13, "2E74B5", True)
    set_spacing(h2, 12, 6, 1.10)

    h3 = styles["Heading 3"]
    set_style_font(h3, 12, "1F4D78", True)
    set_spacing(h3, 8, 4, 1.10)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        set_style_font(style, 11, "000000")
        set_spacing(style, 0, 8, 1.167)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def configure_section(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "MAIS | Mainland PEP Question Generation QA"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.runs[0].font.size = Pt(8.5)
    paragraph.runs[0].font.color.rgb = RGBColor.from_string("666666")
    set_east_asia_font(paragraph.runs[0])

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = "Documentation-only report generated from existing S18 QA artifacts"
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.runs[0].font.size = Pt(8)
    footer_paragraph.runs[0].font.color.rgb = RGBColor.from_string("666666")
    set_east_asia_font(footer_paragraph.runs[0])


def add_text_with_inline_code(paragraph, text: str, font_size: float = 11.0, bold: bool = False) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[1:-1] if is_code else part)
        run.font.name = "Courier New" if is_code else "Calibri"
        run.font.size = Pt(font_size if not is_code else max(font_size - 0.5, 8))
        run.bold = bold
        if is_code:
            run.font.color.rgb = RGBColor.from_string("1F4D78")
            set_east_asia_font(run, "Microsoft YaHei")
        else:
            set_east_asia_font(run)


def split_table_line(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_line(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def table_widths(headers: list[str]) -> list[int]:
    cols = len(headers)
    joined = "|".join(headers)
    if cols == 2:
        if "优先级" in joined:
            return [1200, 8160]
        if "指标" in joined:
            return [4300, 5060]
        if "项目" in joined:
            return [4100, 5260]
        return [3600, 5760]
    if cols == 3:
        if "说明" in joined:
            return [2500, 1600, 5260]
        if "可上线建议" in joined:
            return [2200, 3500, 3660]
        if "主要风险" in joined:
            return [1600, 5900, 1860]
        return [2800, 2200, 4360]
    if cols == 4:
        return [1800, 2100, 2100, 3360]
    if cols == 6:
        return [1300, 1100, 1200, 1200, 1100, 3460]
    if cols == 7:
        return [1050, 950, 1050, 1050, 950, 950, 3360]
    base = CONTENT_WIDTH_DXA // cols
    widths = [base] * cols
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_cell_margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in CELL_MARGINS_DXA.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def configure_table_xml(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


def mark_header_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    headers = rows[0]
    widths = table_widths(headers)
    table = doc.add_table(rows=len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    configure_table_xml(table, widths)
    mark_header_repeat(table.rows[0])

    small_table = len(headers) >= 6
    font_size = 8.4 if small_table else 9.2

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_width(cell, widths[c_idx])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 and len(value) <= 24 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_text_with_inline_code(paragraph, value, font_size=font_size, bold=(r_idx == 0))
            if r_idx == 0:
                shade_cell(cell, "F2F4F7")
            elif value in {"不可上线，需人工签核", "暂不公开；人工抽样后再评估提升", "不公开", "不可接入 app；先完成 pending-s18-review"}:
                shade_cell(cell, "FFF2CC")

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def parse_markdown(doc: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            table_lines = [split_table_line(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(split_table_line(lines[i]))
                i += 1
            add_table(doc, table_lines)
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_text_with_inline_code(paragraph, line[2:], font_size=20, bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
            i += 1
            continue

        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            i += 1
            continue

        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_text_with_inline_code(paragraph, line[2:])
            i += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_text_with_inline_code(paragraph, numbered.group(2))
            i += 1
            continue

        paragraph = doc.add_paragraph()
        add_text_with_inline_code(paragraph, line)
        i += 1


def main() -> None:
    doc = Document()
    configure_section(doc)
    configure_styles(doc)
    parse_markdown(doc, MARKDOWN_PATH.read_text(encoding="utf-8"))
    doc.core_properties.title = "内地人教版数学题库生成与质量检查专项总结"
    doc.core_properties.subject = "Mainland PEP question generation QA summary"
    doc.core_properties.author = "MAIS S10"
    doc.core_properties.comments = "Generated from existing S18 QA artifacts; documentation-only report."
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
