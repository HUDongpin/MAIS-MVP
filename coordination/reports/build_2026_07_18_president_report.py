from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-18-president-report.docx")
REPORT_DATE = "2026-07-18"
WINDOW = "2026-07-17 08:00 to 2026-07-18 08:00 Asia/Taipei (UTC+8)"
UTC_WINDOW = "2026-07-17 00:00 to 2026-07-18 00:00 UTC"
FORMAT_VERSION = "President Report Business Brief v1"

# Business Brief v1 consistency overrides on the standard_business_brief preset.
FONT_LATIN = "Hiragino Sans GB"
FONT_CJK = "Hiragino Sans GB"
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().get_or_add_rFonts()
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    if sum(widths) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.1
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    size = 16 if level == 1 else 13
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.1, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header)
        set_run_font(r, size=9, bold=True)
        set_cell_shading(cell, header_fill)
    row_fills = row_fills or {}
    for row_index, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        fill = row_fills.get(row_index)
        for idx, value in enumerate(row):
            cell = table_row.cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx in (0, 1) and len(str(value)) < 24:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=body_size)
            if fill:
                set_cell_shading(cell, fill)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=10)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=10)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    for style_name in ["Normal", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("MAIS-MVP | Daily coordination report | A01-A25")
    set_run_font(hr, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run(f"{REPORT_DATE} | {FORMAT_VERSION} | Page ")
    set_run_font(fr, size=8.5, color=MUTED)
    add_page_field(footer)
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_title_block(doc):
    prepared = datetime.now(ZoneInfo("Asia/Hong_Kong")).strftime("%Y-%m-%d %H:%M HKT")
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", REPORT_DATE),
        ("Report time", "8:00 AM Asia/Hong_Kong / Asia/Taipei"),
        ("Prepared during run", prepared),
        ("Reporting window", WINDOW),
        ("UTC filter window", UTC_WINDOW),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files and all Git refs"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Format", FORMAT_VERSION),
        ("Formal main range", "3b888f43c -> 933a05db59 (63 commits in window)"),
        ("Current repository", "main at 79be1d0cbc, synchronized with origin/main; current HEAD is 20 commits after cutoff"),
        ("Cutoff note", "Formal activity stops at 2026-07-18 08:00 UTC+8; current checks and later commits are labeled generation-time/post-cutoff."),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL, body_size=8.2)


def agent_rows():
    no_log = "窗口内无 fresh owner log。"
    rows = {
        "A01": ("集成-only", "`72fcd5bd00` A01 audited snapshot，20 paths。", no_log),
        "A02": ("集成 / 旧日志", "`a436cf1d09` snapshot；`0a0f7dd165` grade persistence。", "7/13 A02 log仅为 carry-forward。"),
        "A03": ("集成-only", "`6161d0ec16` snapshot；`f1b93e2387` roadmap shell fix。", no_log),
        "A04": ("集成 / 旧日志", "snapshot、question-bank QA cherry-pick、figure spec。", "7/13 A04 log非 fresh。"),
        "A05": ("集成-only", "`175bfa973a` snapshot；`fb3b3eac27` speed visualization source。", no_log),
        "A06": ("大包集成", "459-path snapshot + lab-logo regression slice。", "无 fresh A06 log；recorded runtime-visualization blocker。"),
        "A07": ("Fresh / completed", "14:11 fresh log；classroom policy controls；AI tutor region/voice fixes。", "TDD 5/5，targeted 17/17，backend 41/41；voice仍需 broad browser gate。"),
        "A08": ("集成-only", "`0c97afd669` state/analytics snapshot。", "无 fresh log；generation-time analytics 27/27。"),
        "A09": ("集成-only", "`1662413744` copy/i18n/a11y snapshot。", no_log),
        "A10": ("集成 / report", "`ea329b9104` tooling snapshot；prior-report continuity来自 automation memory。", "7/17 DOCX当前已不在 checkout；本次为 report-only。"),
        "A11": ("集成-only / gate gap", "66 test/config paths + profile fallback test。", "无 fresh A11 full-regression handoff；recorded routing blocker。"),
        "A12": ("Fresh / disabled hold", "14:35 fresh log；fail-closed Google OAuth + callback export fix。", "13/13、56/56、type-check pass；Playwright无 final result，OAuth保持关闭。"),
        "A13": ("集成 / 旧成果", "28-path snapshot + teacher BUG-001..010 cherry-pick。", no_log),
        "A14": ("提交-only", "snapshot；avatar object-storage persistence + fallback test。", no_log),
        "A15": ("集成-only", "8-path adaptive snapshot。", "无 fresh adaptive regression；post-cutoff import gate现为 red。"),
        "A16": ("集成-only", "2-path research snapshot。", no_log),
        "A17": ("集成-only", "3-path gamification snapshot。", no_log),
        "A18": ("Fresh shared log / HOLD", "65 QA files、CA K5 template/OPC evidence。", "8 generated-content candidates待 row-level review。"),
        "A19": ("历史整合", "`.env.local.example` + 6/29 US-West evidence。", "不是 fresh env verification；Google OAuth env未配置/启用。"),
        "A20": ("提交-only", "draft game compile support + older Fishing/Adventure fixes。", no_log),
        "A21": ("Fresh shared log / HOLD", "29 candidate artifacts、268 reviewed media paths、portability fix。", "Texas仍 needs-repair / candidate-only。"),
        "A22": ("集成-only / gate gap", "10 release scripts + runtime/build configuration。", "无 fresh final build/Playwright handoff；generation-time source-clean pass。"),
        "A23": ("无 fresh session", "仅整合旧 S23 promotion records。", "最新书面结论仍是不进入 live app。"),
        "A24": ("路径整合 / no signoff", "A24-owned净路径 169；figure-spec evidence进入主线。", "无 fresh exact-layer signoff。"),
        "A25": ("Fresh / major integration", "12:59 fresh log；fingerprint、63 main commits、15 archive tags、3 pushes。", "缺最终 cleanup handoff与整合后全量 gate evidence。"),
    }
    return [(f"A{i:02d}", *rows[f"A{i:02d}"]) for i in range(1, 26)]


def build_doc():
    doc = Document()
    setup_document(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "本窗口有实质活动：fresh-log active IDs 为 A07、A12、A18、A21、A25；A25 主导 preserve-first physical cleanup，将 main 从 `3b888f43c` 推进至 `933a05db59`，窗口内 63 commits。"),
            ("2", "Main 净整合 1,498 files（1,114 added / 383 modified / 1 deleted），约 2.70M insertions。21 个 A01-A19/A21/A22 snapshot commits及其他 owner slices进入主线，但多数是历史脏工作树切片，不代表 21 个代理当日独立开发。"),
            ("3", "Fresh成果包括 A07 classroom-policy controls、A12 disabled-by-default Google OAuth、A18/A21 reviewed evidence packages，以及 A25 fingerprint/release-governance；A12 browser结果、A18/A21 live promotion与 A25 final handoff仍未闭环。"),
            ("4", "Generation-time Git hygiene明显改善：current main与 origin/main同步，仅 1 worktree；fresh dirty-map 0 entries、A22 source-clean gate pass、package gate 8/8 valid。此状态含 cutoff 后20 commits，不能替代窗口内 final release proof。"),
            ("5", "当前仍非 release-ready：type-check被 `.next/types/validator.ts` 的两个缺失 game-page引用阻断；import gate被 post-cutoff `adaptive-eval.mjs` 的 `.tmp`依赖阻断；整合后 build与 broad Playwright未运行。"),
        ],
        [700, 8660],
        body_size=8.0,
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "The window was active. Fresh logs exist for A07, A12, A18, A21, and A25. A25 led a preserve-first physical cleanup that advanced main from `3b888f43c` to `933a05db59` through 63 in-window commits."),
            ("2", "Main absorbed 1,498 net files: 1,114 added, 383 modified, and one deleted, with about 2.70 million insertions. Most Axx snapshot commits are historical dirty-worktree integrations, not 21 independent same-day development sessions."),
            ("3", "Fresh outcomes include A07 classroom-policy controls, A12 disabled-by-default Google OAuth, A18/A21 reviewed evidence packages, and A25 fingerprint/release governance. Browser, promotion, and final cleanup handoffs remain incomplete."),
            ("4", "Generation-time Git hygiene is materially better: main equals origin/main, only one worktree remains, the fresh dirty map has zero entries, the A22 source-clean gate passes, and all eight package manifests validate. This includes 20 post-cutoff commits."),
            ("5", "The project is not release-ready. Type-check is blocked by two missing game-page references in generated `.next` types; import checking is blocked by a generated `.tmp` adaptive-eval dependency; no post-integration build or broad Playwright gate is available."),
        ],
        [700, 8660],
        body_size=8.0,
    )

    add_heading(doc, "报告窗口摘要（中文）", 1)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ("正式窗口", WINDOW),
            ("Fresh logs", "A07、A12、A18/A21 shared、A25；其他 Axx 主要为 commit/path evidence 或旧日志整合。"),
            ("Main movement", "`3b888f43c` -> `933a05db59`；63 main commits；origin/main于 18:51、19:26、20:25 HKT推送。"),
            ("Net change", "1,498 files；1,114 A / 383 M / 1 D；2,695,409 insertions / 33,285 deletions。"),
            ("Evidence caveat", "全 refs 70 commits中6个是 patch-equivalent archive重复；唯一独立非-main为 `ad40fed27b` CCSS PR。"),
            ("Coordination files", "窗口内 blockers/decisions/integration净变化为0；`coordination/integration/`当前不存在。"),
            ("CCSS PR", "`origin/pr/ccss-code-compliance` `ad40fed27b` 于02:41 HKT；基于窗口终点，尚未在 cutoff main。"),
            ("Generation-time state", "HEAD `79be1d0cbc`，比 cutoff晚20 commits；main=origin/main；1 worktree；Git dirty entries 0。"),
            ("空窗口 fallback", "不适用；存在 fresh logs与主线整合，因此不使用 `No assigned work in this reporting window`。"),
        ],
        [2450, 6910],
        body_size=7.35,
    )

    add_heading(doc, "整体项目进展（中文）", 1)
    add_bullet(doc, "A25 preserve-first cleanup把原先39 registered worktrees、15,187 dirty entries、20 unique commits的复杂状态转换为可审计切片；窗口内主线完成63 commits并推送，当前只剩1个clean worktree。")
    add_bullet(doc, "A07交付 classroom AI policy：open / limited / fallback-only、teacher authorization、student strictest-policy resolution；targeted tests与type-check在A07日志中绿色。")
    add_bullet(doc, "A12交付 fail-closed Google OAuth/OIDC slice，默认关闭；Node/auth/storage tests绿色，但未取得Playwright final result，也未运行live callback。")
    add_bullet(doc, "A18/A21把reviewed QA与candidate evidence拆成独立包并移除machine-local路径；8个live-data candidates继续HOLD，未绕过A23 promotion gate。")
    add_bullet(doc, "A01-A22多领域历史切片已进入main，包括visualization、backend、lessons、tests、media、content QA与release tooling；规模大，必须用A11/A22 post-integration gates验证。")
    add_bullet(doc, "Current Git hygiene为green，但current TypeScript/import gates为red；发布判断必须以clean-source full gate为准，而不是以Git clean等同release-ready。")

    add_heading(doc, "A01-A25 会话状态表（中文）", 1)
    add_table(
        doc,
        ["ID", "状态", "窗口内活动", "当前重点 / 阻塞"],
        agent_rows(),
        [620, 1320, 3680, 3740],
        body_size=6.15,
        row_fills={5: PENDING_FILL, 6: OK_FILL, 10: RISK_FILL, 11: PENDING_FILL, 17: PENDING_FILL, 20: PENDING_FILL, 21: RISK_FILL, 22: PENDING_FILL, 23: PENDING_FILL, 24: PENDING_FILL},
    )

    add_heading(doc, "已完成工作（中文）", 1)
    add_bullet(doc, "A25：完成fingerprint工具链（tests 16/16、independent review无P0/P1/P2），保存15个archive tags，并把63个窗口提交推入main。")
    add_bullet(doc, "A25 integration：21个A01-A19/A21/A22 audited snapshot commits及A20/A24等独立切片进入主线；报告明确标注为历史切片整合。")
    add_bullet(doc, "A07：classroom-policy controls完成；TDD 5/5、targeted 17/17、related backend 41/41、imports/type-check/diff-check pass。")
    add_bullet(doc, "A12：disabled Google OAuth完成；slice 13/13、auth/storage 56/56、type-check pass；callback export另有修复。")
    add_bullet(doc, "A18/A21：A18 JSON 17/17与MJS 9/9；A21 JSON 13/13与MJS 4/4；isolated validator 41 topics、492/492 questions、0 findings。")
    add_bullet(doc, "Generation-time：fresh A25 dirty-map 0 entries；A22 source-clean pass；release package gate 8/8 valid；analytics 27/27与Git diff checks pass。")

    add_heading(doc, "进行中工作（中文）", 1)
    add_bullet(doc, "A25需补final cleanup handoff：明确worktree/branch cleanup时间线、最终指纹、preserve/archive证据、63 commits review summary与整合后full gates。")
    add_bullet(doc, "A11/A22需在clean/current approved source运行type-check、build、targeted Playwright与release regression；当前Git clean不替代这些门禁。")
    add_bullet(doc, "A20/A22需判定两个missing game-page refs是stale `.next`还是source-route regression；A15/A22需修正或隔离adaptive-eval generated import contract。")
    add_bullet(doc, "A12/A19/A11/A22：Google OAuth继续disabled，等待server-only env placement、live callback与browser regression。")
    add_bullet(doc, "A18/A21/A23：8个generated-content versions等待row-level QA和promotion approval；Texas remains needs-repair / candidate-only。")
    add_bullet(doc, "A18/A03/A05/A06/A11/A22：review窗口内未合并的CCSS PR `ad40fed27b`及其30/30 self-reported tests。")

    add_heading(doc, "阻塞项（中文）", 1)
    add_table(
        doc,
        ["Blocker", "Evidence", "Owner / next action"],
        [
            ("Post-integration full gate缺失", "1,498-file integration后无完整type-check/build/A11 regression/A22 release handoff。", "A11/A22：clean-source full gate。"),
            ("Generation-time type-check red", "`.next/types/validator.ts`引用两个不存在的game pages。", "A20/A22：isolated regenerate/route review。"),
            ("Generation-time import gate red", "`scripts/adaptive-eval.mjs` imports `.tmp/adaptive-eval/lib/adaptiveLearningEval.js`。", "A15/A22：建立stable import/build contract。"),
            ("A25 final handoff缺失", "A25 fresh log停在12:58 fingerprint；未记录final cleanup与全量gate。", "A25/A10/A22：补executive handoff。"),
            ("Google OAuth activation hold", "Default disabled；Playwright无final result；未测live callback/env。", "A12/A19/A11/A22：保持fail-closed。"),
            ("Content promotion hold", "8 generated-content candidates；Texas needs-repair；S23仍决定不进入live app。", "A18/A21/A23/A24 + live owners。"),
            ("Recorded package blockers", "Package gate仍记录A12 storage、A06 visualization、A11 routing、content chain四个blocker reports。", "对应owners：fresh revalidation/prioritization。"),
            ("CCSS PR not merged at cutoff", "`ad40fed27b`基于`933a05db59`；self-reported tsc 0 + tests 30/30。", "A18/A03/A05/A06/A11/A22 review。"),
        ],
        [1800, 4300, 3260],
        body_size=6.65,
        row_fills={0: RISK_FILL, 1: RISK_FILL, 2: RISK_FILL, 3: PENDING_FILL, 4: PENDING_FILL, 5: PENDING_FILL, 6: PENDING_FILL, 7: PENDING_FILL},
    )

    add_heading(doc, "风险（中文）", 1)
    add_bullet(doc, "Integration scale risk：1,498 files与2.70M insertions在单一窗口进入main；任何release-ready声明都需要post-integration full gates。")
    add_bullet(doc, "Attribution risk：只有A07、A12、A18、A21、A25有fresh logs；其余Axx多为historical snapshot integration，不能当作当天独立完成。")
    add_bullet(doc, "Generated-state risk：stale `.next`与`.tmp` contract可能制造false red，也可能掩盖真实route/import regression；当前证据不足以二选一。")
    add_bullet(doc, "Promotion risk：A18/A21 candidate artifacts若绕过A23/A11/A22 gates进入live surfaces，会把未修复Texas/row-level问题带入production。")
    add_bullet(doc, "Security/readiness risk：Google OAuth尚未进行live callback与A19 env verification；启用前必须保持server-only、fail-closed。")
    add_bullet(doc, "Cutoff drift risk：current HEAD比窗口终点多20 commits；generation-time green/red checks描述current state，不得倒推为窗口完成证据。")

    add_heading(doc, "测试 / 构建状态（中文）", 1)
    add_table(
        doc,
        ["Check", "Result", "Interpretation"],
        [
            ("A07 logged gates", "PASS：5/5、17/17、41/41；imports/type-check/diff-check。", "Fresh A07 slice evidence。"),
            ("A12 logged gates", "PASS：13/13、56/56、type-check；Playwright无final result。", "OAuth保持disabled；no live callback。"),
            ("A18/A21 package gates", "PASS：JSON 17/17 + 13/13；MJS 9/9 + 4/4；492/492。", "Candidate/evidence only；not live approval。"),
            ("A25 fingerprint", "PASS：16/16；independent review无P0/P1/P2；diff-check pass。", "Fresh fingerprint stage；final handoff缺失。"),
            ("CCSS PR", "Self-reported PASS：tsc 0；targeted 30/30。", "Unmerged at cutoff；needs independent review。"),
            ("Fresh dirty/source/package gates", "PASS：dirty-map 0；source-clean；8 packages valid；4 reviewed + 4 blocker reports。", "Generation-time current Git/release governance。"),
            ("Fresh analytics", "PASS：27/27。", "A08/A15 deterministic suite green。"),
            ("Fresh npm run type-check", "FAIL：2 TS2307 missing game-page refs from `.next/types/validator.ts`。", "Post-cutoff current HEAD; source vs generated drift unresolved。"),
            ("Fresh npm run check:imports", "FAIL：adaptive-eval imports generated `.tmp` JS。", "Post-cutoff current contract red。"),
            ("No-staged / diff checks", "PASS via direct Git checks；legacy assert script absent。", "No staged content；script contract changed。"),
            ("Build / broad Playwright", "Not run。", "A10 report-only；must run under A11/A22 clean-source control。"),
        ],
        [2400, 3900, 3060],
        body_size=6.35,
        row_fills={0: OK_FILL, 1: PENDING_FILL, 2: OK_FILL, 3: OK_FILL, 4: PENDING_FILL, 5: OK_FILL, 6: OK_FILL, 7: RISK_FILL, 8: RISK_FILL, 9: OK_FILL, 10: PENDING_FILL},
    )

    add_heading(doc, "变更文件（中文）", 1)
    add_table(
        doc,
        ["Category", "Count / examples", "Meaning"],
        [
            ("Window net change", "1,498 files：1,114 A / 383 M / 1 D；2,695,409 insertions / 33,285 deletions。", "`3b888f43c` -> `933a05db59`。"),
            ("Top-level concentration", "components 556、public 262、lib 196、coordination 159、app 127、data 85、tests 65、scripts 34。", "Broad cross-surface integration。"),
            ("Largest subareas", "visualizations 456、lesson illustrations 165、content QA 138、lib/server 129、audio 76、E2E 65、app/api 58。", "A06/A05/A18/A12/A11 high review load。"),
            ("Owner-pathspec counts", "A06 462、A12 183、A24 169、A18 138、A05 129、A21 77、A11 66、A04 61。", "Ownership estimate；not same-day session count。"),
            ("Fresh logs", "A07、A12、A18/A21、A25。", "Only these qualify as fresh owner-log activity。"),
            ("Shared/high-risk files", "types、AppProviders、i18n、package/lock、next/tsconfig、Playwright、userStore/LLM、questions/topics/grades、65 E2E files。", "Require owner-routed regression。"),
            ("CCSS candidate", "Remote `ad40fed27b`，not in cutoff main。", "Review candidate, not completed main work。"),
            ("Post-cutoff excluded", "20 main commits between cutoff and current `79be1d0cbc`。", "Current-state context only。"),
            ("This run outputs", "`build_2026_07_18_president_report.py`、DOCX、A10 session log、automation memory。", "Report-only after cutoff；not feature code。"),
        ],
        [2000, 4600, 2760],
        body_size=6.45,
    )

    add_heading(doc, "明日优先事项（中文）", 1)
    add_numbered(doc, "A11/A22在clean/current approved source运行type-check、build、targeted Playwright与release matrix，形成post-integration handoff。")
    add_numbered(doc, "A20/A22隔离诊断两个missing game-page refs；A15/A22修正adaptive-eval `.tmp` import contract；不得用清理动作掩盖source regression。")
    add_numbered(doc, "A25补final physical-cleanup handoff：final fingerprint、archive tags、worktree/branch disposition、63-commit review与no-destructive-loss proof。")
    add_numbered(doc, "A12/A19/A11/A22保持Google OAuth disabled，完成env readiness、live callback与browser regression后再提交enablement decision。")
    add_numbered(doc, "A18/A21/A23逐项处理8个generated-content HOLD candidates与Texas needs-repair；未经A23/A11/A22不得live promotion。")
    add_numbered(doc, "A18/A03/A05/A06/A11/A22 review `ad40fed27b` CCSS PR的curriculum scope、route impact与independent tests。")

    add_heading(doc, "需要 Owner 决策（中文）", 1)
    add_bullet(doc, "请 Dr. Peter Hu 确认：当前main Git clean但不等于release-ready；必须先取得A11/A22 post-integration full gate。")
    add_bullet(doc, "请指定A25 final cleanup handoff与63-commit small-slice review为最高release-governance优先级，或明确可接受的剩余审计风险。")
    add_bullet(doc, "请决定Google OAuth enablement时间线；在A19 env与A11/A22 live/browser evidence齐备前继续保持disabled。")
    add_bullet(doc, "请选择首个A23 content/RAG promotion package；8个HOLD candidates与Texas needs-repair不得默认进入live。")
    add_bullet(doc, "请决定是否把`ad40fed27b` CCSS PR列为下一批review candidate，并指定A18/A03/A05/A06/A11/A22联合签核。")
    add_bullet(doc, "请要求A20/A22与A15/A22先解决current type/import red gates，再讨论任何production deploy。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
