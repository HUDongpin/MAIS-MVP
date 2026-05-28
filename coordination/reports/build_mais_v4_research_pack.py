from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_adaptive_learning_apa7_docx import (
    add_caption,
    add_formula,
    add_heading,
    add_note,
    add_page_number,
    add_paragraph,
    set_cell_margins,
    set_cell_shading,
    set_paragraph_spacing,
    set_run_font,
    set_table_borders,
    set_table_width,
    set_update_fields_on_open,
)


PAPER_DIR = Path("/Users/dongpinhu/Desktop/MAIS Research Proposals/Paper")

OUTS = {
    "research": PAPER_DIR / "2026-05-17-MAIS-v4-algorithm-research-report-APA7-ZH.docx",
    "spec": PAPER_DIR / "2026-05-17-MAIS-v4-algorithm-design-spec-ZH.docx",
    "roadmap": PAPER_DIR / "2026-05-17-MAIS-v4-engineering-roadmap-ZH.docx",
    "pilot": PAPER_DIR / "2026-05-17-MAIS-v4-evaluation-school-pilot-plan-ZH.docx",
}


REFERENCES = [
    "Ainsworth, S. (2006). DeFT: A conceptual framework for considering learning with multiple representations. Learning and Instruction, 16(3), 183-198. https://doi.org/10.1016/j.learninstruc.2006.03.001",
    "Asai, A., Wu, Z., Wang, Y., Sil, A., & Hajishirzi, H. (2024). Self-RAG: Learning to retrieve, generate, and critique through self-reflection. In International Conference on Learning Representations. https://proceedings.iclr.cc/paper_files/paper/2024/hash/25f7be9694d7b32d5cc670927b8091e1-Abstract-Conference.html",
    "Athey, S., & Imbens, G. (2016). Recursive partitioning for heterogeneous causal effects. Proceedings of the National Academy of Sciences, 113(27), 7353-7360. https://doi.org/10.1073/pnas.1510489113",
    "Baker, R. S., Hutt, S., Brooks, C. A., Srivastava, N., & Mills, C. (2024). Open science and educational data mining: Which practices matter most? In B. Paassen & C. D. Epp (Eds.), Proceedings of the 17th International Conference on Educational Data Mining (pp. 279-287). International Educational Data Mining Society. https://doi.org/10.5281/zenodo.12729816",
    "Baker, R. S., & Hawn, A. (2022). Algorithmic bias in education. International Journal of Artificial Intelligence in Education, 32(4), 1052-1092. https://doi.org/10.1007/s40593-021-00285-9",
    "Bloom, B. S. (1984). The 2 sigma problem: The search for methods of group instruction as effective as one-to-one tutoring. Educational Researcher, 13(6), 4-16. https://doi.org/10.3102/0013189X013006004",
    "Cepeda, N. J., Pashler, H., Vul, E., Wixted, J. T., & Rohrer, D. (2006). Distributed practice in verbal recall tasks: A review and quantitative synthesis. Psychological Bulletin, 132(3), 354-380. https://doi.org/10.1037/0033-2909.132.3.354",
    "Christie, S. T., Cook, C., & Rafferty, A. N. (2024). Uncertainty-preserving deep knowledge tracing with state-space models. In Proceedings of the 17th International Conference on Educational Data Mining. https://educationaldatamining.org/edm2024/proceedings/2024.EDM-posters.108/index.html",
    "Corbett, A. T., & Anderson, J. R. (1995). Knowledge tracing: Modeling the acquisition of procedural knowledge. User Modeling and User-Adapted Interaction, 4(4), 253-278. https://doi.org/10.1007/BF01099821",
    "Dwork, C. (2006). Differential privacy. In M. Bugliesi, B. Preneel, V. Sassone, & I. Wegener (Eds.), Automata, Languages and Programming (pp. 1-12). Springer. https://doi.org/10.1007/11787006_1",
    "Gervet, T., Koedinger, K., Schneider, J., & Mitchell, T. (2020). When is deep learning the best approach to knowledge tracing? Journal of Educational Data Mining, 12(3), 31-54. https://doi.org/10.5281/zenodo.4143614",
    "Holstein, K., McLaren, B. M., & Aleven, V. (2019). Co-designing a real-time classroom orchestration tool to support teacher-AI complementarity. Journal of Learning Analytics, 6(2), 27-52. https://doi.org/10.18608/jla.2019.62.3",
    "Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., Xu, Y., Ishii, E., Bang, Y., Madotto, A., & Fung, P. (2023). Survey of hallucination in natural language generation. ACM Computing Surveys, 55(12), Article 248. https://doi.org/10.1145/3571730",
    "Kasneci, E., Sessler, K., Kuchemann, S., Bannert, M., Dementieva, D., Fischer, F., Gasser, U., Groh, G., Gunnemann, S., Hullermeier, E., Krusche, S., Kutyniok, G., Michaeli, T., Nerdel, C., Pfeffer, J., Poquet, O., Sailer, M., Schmidt, A., Seidel, T., ... Kasneci, G. (2023). ChatGPT for good? On opportunities and challenges of large language models for education. Learning and Individual Differences, 103, Article 102274. https://doi.org/10.1016/j.lindif.2023.102274",
    "Lee, J., & Yeung, D.-Y. (2019). Knowledge query network for knowledge tracing: How knowledge interacts with skills. In Proceedings of the Ninth International Conference on Learning Analytics & Knowledge (pp. 491-500). https://doi.org/10.1145/3303772.3303786",
    "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Kuttler, H., Lewis, M., Yih, W.-t., Rocktaschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. In Advances in Neural Information Processing Systems, 33, 9459-9474. https://proceedings.neurips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html",
    "Li, X., Yan, L., Zhao, L., Martinez-Maldonado, R., & Gasevic, D. (2023). CVPE: A computer vision approach for scalable and privacy-preserving socio-spatial, multimodal learning analytics. In LAK 2023 Conference Proceedings: The Thirteenth International Conference on Learning Analytics & Knowledge (pp. 175-185). Association for Computing Machinery. https://doi.org/10.1145/3576050.3576145",
    "Martinez Beltran, E. T., Sanchez Sanchez, P. M., Lopez Bernal, S., Bovet, G., Gil Perez, M., Martinez Perez, G., & Huertas Celdran, A. (2023). Fedstellar: A platform for training models in a privacy-preserving and decentralized fashion. In Proceedings of the Thirty-Second International Joint Conference on Artificial Intelligence (pp. 7154-7157). https://doi.org/10.24963/ijcai.2023/838",
    "Pardos, Z. A., & Heffernan, N. T. (2010). Modeling individualization in a Bayesian networks implementation of knowledge tracing. In P. De Bra, A. Kobsa, & D. Chin (Eds.), User Modeling, Adaptation, and Personalization (pp. 255-266). Springer. https://doi.org/10.1007/978-3-642-13470-8_24",
    "Pham, D. M., Vanacore, K. P., Sales, A. C., & Gagnon-Bartsch, J. A. (2024). LOOL: Towards personalization with flexible & robust estimation of heterogeneous treatment effects. In B. Paassen & C. D. Epp (Eds.), Proceedings of the 17th International Conference on Educational Data Mining (pp. 376-384). International Educational Data Mining Society. https://doi.org/10.5281/zenodo.12729840",
    "Piech, C., Spencer, J., Huang, J., Ganguli, S., Sahami, M., Guibas, L. J., & Sohl-Dickstein, J. (2015). Deep knowledge tracing. In Advances in Neural Information Processing Systems, 28. https://papers.nips.cc/paper/5654-deep-knowledge-tracing",
    "Rachatasumrit, N., Carvalho, P. F., & Koedinger, K. R. (2024). Beyond accuracy: Embracing meaningful parameters in educational data mining. In B. Paassen & C. D. Epp (Eds.), Proceedings of the 17th International Conference on Educational Data Mining. https://educationaldatamining.org/edm2024/proceedings/2024.EDM-long-papers.17/",
    "Roediger, H. L., III, & Karpicke, J. D. (2006). Test-enhanced learning: Taking memory tests improves long-term retention. Psychological Science, 17(3), 249-255. https://doi.org/10.1111/j.1467-9280.2006.01693.x",
    "Sharma, P., Stewart, A. E. B., Li, Q., Ravichander, K., & Walker, E. (2024). Building learner activity models from log data using sequence mapping and hidden Markov models. In Proceedings of the 17th International Conference on Educational Data Mining. https://educationaldatamining.org/edm2024/proceedings/2024.EDM-short-papers.60/index.html",
    "Shute, V. J. (2008). Focus on formative feedback. Review of Educational Research, 78(1), 153-189. https://doi.org/10.3102/0034654307313795",
    "So, J., Ali, R. E., Guler, B., Jiao, J., & Avestimehr, A. S. (2023). Securing secure aggregation: Mitigating multi-round privacy leakage in federated learning. Proceedings of the AAAI Conference on Artificial Intelligence, 37(8), 9864-9873. https://doi.org/10.1609/aaai.v37i8.26177",
    "Sweller, J., van Merrienboer, J. J. G., & Paas, F. (2019). Cognitive architecture and instructional design: 20 years later. Educational Psychology Review, 31(2), 261-292. https://doi.org/10.1007/s10648-019-09465-5",
    "Wager, S., & Athey, S. (2018). Estimation and inference of heterogeneous treatment effects using random forests. Journal of the American Statistical Association, 113(523), 1228-1242. https://doi.org/10.1080/01621459.2017.1319839",
    "Weidinger, L., Uesato, J., Rauh, M., Griffin, C., Huang, P.-S., Mellor, J., Glaese, A., Cheng, M., Balle, B., Kasirzadeh, A., Biles, C., Brown, S., Kenton, Z., Hawkins, W., Stepleton, T., Birhane, A., Haas, J., Rimell, L., Hendricks, L. A., ... Gabriel, I. (2022). Taxonomy of risks posed by language models. In 2022 ACM Conference on Fairness, Accountability, and Transparency (pp. 214-229). https://doi.org/10.1145/3531146.3533088",
    "Wood, D., Bruner, J. S., & Ross, G. (1976). The role of tutoring in problem solving. Journal of Child Psychology and Psychiatry, 17(2), 89-100. https://doi.org/10.1111/j.1469-7610.1976.tb00381.x",
    "Yan, L., Sha, L., Zhao, L., Li, Y., Martinez-Maldonado, R., Chen, G., Li, X., Jin, Y., & Gasevic, D. (2024). Practical and ethical challenges of large language models in education: A systematic scoping review. British Journal of Educational Technology, 55(1), 90-112. https://doi.org/10.1111/bjet.13370",
    "Yang, Q., Liu, Y., Chen, T., & Tong, Y. (2019). Federated machine learning: Concept and applications. ACM Transactions on Intelligent Systems and Technology, 10(2), Article 12. https://doi.org/10.1145/3298981",
    "Zhang, J., Shi, X., King, I., & Yeung, D.-Y. (2017). Dynamic key-value memory networks for knowledge tracing. In Proceedings of the 26th International Conference on World Wide Web (pp. 765-774). https://doi.org/10.1145/3038912.3052580",
    "Zimmerman, B. J. (2002). Becoming a self-regulated learner: An overview. Theory Into Practice, 41(2), 64-70. https://doi.org/10.1207/s15430421tip4102_2",
]


FORMULAS = [
    (r"\tau(x)=\mathbb{E}[Y(1)-Y(0)\mid X=x]", "1"),
    (r"\hat{u}(c,x)=\hat{\tau}(c,x)\cdot q_{evidence}(c)-\lambda_r Risk(c)-\lambda_f FairGap(c)-\lambda_k Cost(c)", "2"),
    (r"z_t=f_{\theta}(x^{response}_t,x^{handwriting}_t,x^{visual}_t,x^{time}_t,x^{hint}_t,x^{teacher}_t)", "3"),
    (r"\hat{M}_{k,t+h}=g_{\psi}(M_{k,t},z_{1:t},Review_{k,t},Repair_{k,t})", "4"),
    (r"\min_{\theta}\sum_{s=1}^{S}\frac{n_s}{n}\mathcal{L}_s(\theta)\quad\text{without centralizing raw student records}", "5"),
    (r"\Pr[\mathcal{M}(D)\in O]\le e^{\epsilon}\Pr[\mathcal{M}(D')\in O]+\delta", "6"),
    (r"\Delta_{fair}(g)=\left|\operatorname{nLG}_{g}-\operatorname{nLG}_{all}\right|+\left|\operatorname{HarmRate}_{g}-\operatorname{HarmRate}_{all}\right|", "7"),
    (r"D_t^{v4}=\begin{cases} c^{*},& V_{v3}(c^{*})=1\land q_{causal}\ge\tau_c\land \Delta_{fair}\le\tau_f\land Risk\le\tau_r\\ D_t^{v3},&\text{otherwise}\end{cases}", "8"),
]


def setup_doc(title: str, subtitle: str, doc_type: str) -> Document:
    doc = Document()
    set_update_fields_on_open(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.color.rgb = RGBColor.from_string("1F4D78")
    ref_style = styles.add_style("APA Reference", 1)
    ref_style.base_style = styles["Normal"]
    ref_style.font.size = Pt(9.1)
    ref_style.paragraph_format.left_indent = Inches(0.5)
    ref_style.paragraph_format.first_line_indent = Inches(-0.5)
    ref_style.paragraph_format.space_after = Pt(5)
    ref_style.paragraph_format.line_spacing = 1.12

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(header.add_run("MAIS v4 算法研究与开发包"), size=9, color="606A78", east_asia="Heiti SC")
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=34, after=8, line=1.15)
    set_run_font(p.add_run(title), size=21, bold=True, color="0B2545", east_asia="Heiti SC")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=24, line=1.2)
    set_run_font(p.add_run(subtitle), size=13.5, italic=True, color="1F4D78", east_asia="Heiti SC")

    meta = [
        ("项目", "MAIS：P1-S6 双语数学自适应学习系统"),
        ("文件类型", doc_type),
        ("日期", "2026-05-17"),
        ("定位", "v4：evidence-validated, causal, multimodal, privacy-preserving adaptive ecosystem"),
        ("边界", "LLM 不直接越过教学边界；最终决策受 BKT/KG/RAG/validation/teacher/privacy/fairness 约束"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta:
        row = table.add_row()
        for idx, value in enumerate((k, v)):
            cell = row.cells[idx]
            set_cell_margins(cell, top=110, bottom=110, start=140, end=140)
            if idx == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            set_run_font(p.add_run(value), size=10.2, bold=(idx == 0), color="0B2545" if idx == 0 else "000000", east_asia="Heiti SC" if idx == 0 else "Songti SC")
    set_table_borders(table)
    set_table_width(table, [3.6, 12.8])
    add_paragraph(doc, "")
    return doc


def add_paras(doc: Document, paras: list[str]) -> None:
    for para in paras:
        add_paragraph(doc, para, line=1.28)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        set_paragraph_spacing(p, after=4, line=1.2)
        set_run_font(p.add_run(item), size=10.4)


def add_numbered_items(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.38)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        set_paragraph_spacing(p, after=4, line=1.2)
        set_run_font(p.add_run(item), size=10.4)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[tuple[str, ...]], widths: list[float]) -> None:
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell, top=95, bottom=95, start=120, end=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(h), size=8.9, bold=True, color="0B2545", east_asia="Heiti SC")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i], top=85, bottom=85, start=120, end=120)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=8.45, bold=(i == 0))
    set_table_borders(table)
    set_table_width(table, widths)
    add_paragraph(doc, "", after=4)


def add_references(doc: Document, refs: list[str]) -> None:
    doc.add_page_break()
    add_heading(doc, "References", level=1)
    for ref in refs:
        p = doc.add_paragraph(style="APA Reference")
        set_run_font(p.add_run(ref), size=9.1, east_asia="Times New Roman", latin="Times New Roman")


def add_formulas(doc: Document, start: int = 1, end: int = 8) -> None:
    for formula, num in FORMULAS[start - 1 : end]:
        add_formula(doc, formula, num)


def build_research_report() -> Path:
    doc = setup_doc(
        "MAIS v4 算法研究报告",
        "证据验证、因果、多模态、隐私保护的自适应学习生态",
        "学术研究报告 / APA 7th",
    )
    add_note(doc, "核心论点", "v4 不应被定义为更强 LLM 版本，而应定义为 v3 治理框架上的因果化、长期化、多模态化和隐私保护化升级。它回答的不只是“下一步推荐什么”，而是“该推荐是否真的提升长期学习、对谁有效、是否公平、是否可跨校泛化，以及是否在隐私约束内可治理”。")
    add_heading(doc, "摘要", level=1)
    add_paras(doc, [
        "本报告提出 MAIS v4 的算法研究与开发路线。v4 建立在 v3 的 BKT safety floor、educational knowledge graph、RAG evidence grounding、validation services、audit logging 和 teacher-in-the-loop 之上，新增 causal/uplift modeling、multimodal learner-state modeling、longitudinal mastery modeling、privacy-preserving cross-school analytics 和 fairness monitoring。",
        "研究核心是从预测正确率转向因果学习收益。v4 需要估计某个推荐对某类学生的增量效果，监测不同学生群体的收益差距，并在隐私保护条件下评估跨校泛化。LLM 的角色继续被限制为证据摘要、解释生成、教师协作和候选比较，不能直接生成教学边界或越过 v3 validation gates。",
    ])
    add_heading(doc, "1. 研究问题与定位", level=1)
    add_numbered_items(doc, [
        "RQ1：哪些推荐真正提升 learning gain、retention 和 transfer，而不只是提高短期答题正确率？",
        "RQ2：推荐效果是否因年级、语言、先备知识、学校、知识点和学习历史不同而显著变化？",
        "RQ3：作答、手写、可视化操作、停留时间、hint 使用和教师标注等多模态信号是否能提升 learner-state modeling？",
        "RQ4：如何在不集中敏感学生数据的前提下进行跨校模型改进、公平性监测和泛化评估？",
        "RQ5：教师如何参与 v4 推荐确认、覆盖、因果评估和 school-level governance dashboard？",
    ])
    add_heading(doc, "2. 文献综述", level=1)
    add_paras(doc, [
        "因果与异质性效果研究为 v4 的 personalization 提供核心方法。Athey 和 Imbens（2016）以及 Wager 和 Athey（2018）提供了 heterogeneous treatment effects 的统计学习基础；EDM 2024 的 LOOL 研究进一步把 CATE 估计用于教育个性化分配（Pham et al., 2024）。这意味着 v4 的目标不是让系统更会预测，而是知道哪种教学动作对哪类学生产生增量收益。",
        "多模态学习分析为 v4 提供更丰富的学习证据。LAK 2023 的 CVPE 研究展示了可扩展且隐私保护的 socio-spatial multimodal learning analytics（Li et al., 2023）；EDM 2024 的 learner activity modeling 研究展示了如何从日志序列建模学习状态（Sharma et al., 2024）。这些研究支持 MAIS 将 handwriting、visualization interaction、dwell time 和 hint traces 纳入 v4，但必须先定义 consent、minimization 与 retention policy。",
        "RAG 与 self-reflective retrieval 继续是 v4 的证据边界。Lewis 等（2020）提出 retrieval-augmented generation，Asai 等（2024）进一步强调检索、生成和自我批判的结合。v4 不把这些方法用作开放问答，而是用作 causal explanation 与 teacher audit note 的证据门。",
        "隐私保护和跨校泛化需要 federated / privacy-preserving analytics。Dwork（2006）给出 differential privacy 的形式化基础；Yang 等（2019）系统化介绍 federated machine learning；Fedstellar 展示了去中心化联邦训练平台（Martinez Beltran et al., 2023）；So 等（2023）提醒 secure aggregation 在多轮联邦学习中仍可能出现隐私泄漏。因此 v4 不应把 federated learning 当作天然隐私充分条件，而应设计 privacy budget、secure aggregation audit 和学校级退出机制。",
        "长期掌握建模需要同时吸收可解释 KT 与深度 KT。BKT 是可审计的安全底座（Corbett & Anderson, 1995; Pardos & Heffernan, 2010）；DKT、DKVMN、KQN 和不确定性保持模型提供更强序列建模能力（Christie et al., 2024; Lee & Yeung, 2019; Piech et al., 2015; Zhang et al., 2017）。Gervet 等（2020）和 Rachatasumrit 等（2024）提醒，预测更准不等于参数更有教育意义。v4 因此应把长期 mastery、retention 和 transfer 作为主要结果。",
        "治理研究要求 v4 保持教师与学校控制。算法偏差研究说明教育推荐可能扩大群体差异（Baker & Hawn, 2022）；teacher-AI complementarity 研究强调教师应参与 AI 工具设计和课堂编排（Holstein et al., 2019）。LLM 风险与教育综述进一步说明，大模型适合解释和反馈，但不能成为自治教学权威（Ji et al., 2023; Kasneci et al., 2023; Weidinger et al., 2022; Yan et al., 2024）。",
        "教育理论为 v4 的目标函数提供约束。Bloom（1984）支持个别化辅导目标，Shute（2008）支持 formative feedback，Cepeda 等（2006）和 Roediger 与 Karpicke（2006）支持 spacing/retrieval practice，Ainsworth（2006）支持多重表征，Sweller 等（2019）支持 cognitive load，Zimmerman（2002）支持自我调节学习，Wood 等（1976）支持 scaffolding。v4 应把这些理论变为 utility、risk、load、retention 和 teacher-approval gates。",
    ])
    add_heading(doc, "3. v4 算法模型", level=1)
    add_paragraph(doc, "v4 adaptive decision model 接收 v3 candidate set、BKT mastery、KG prerequisites、RAG evidence、teacher policy、multimodal traces、school context 以及 privacy/fairness constraints，输出 validated recommendation、causal/uplift estimate、longitudinal mastery impact estimate、risk/fairness/cost score 和 teacher-facing audit note。")
    add_formulas(doc, 1, 8)
    add_table(doc, "表 1\nv4 算法层与功能", ["层", "输入", "输出", "约束"], [
        ("Causal uplift", "候选、学习者特征、结果窗口", "CATE/uplift、因果信心", "证据不足则 exploratory，不进入 automatic mode"),
        ("Multimodal learner state", "作答、手写、可视化、时间、hint、教师标注", "融合状态 z_t", "consent、minimization、purpose limitation"),
        ("Longitudinal mastery", "BKT、review、repair、历史 mastery trace", "长期保持与迁移预测", "不可替代 BKT safety floor"),
        ("Privacy-preserving cross-school", "学校本地模型/摘要", "跨校泛化信号", "不集中原始学生记录，跟踪 privacy budget"),
        ("Fairness monitoring", "subgroup metrics、harm rate、fallback/override", "fairness gap 与治理告警", "超阈值回退 teacher-approved mode"),
    ], [3.1, 4.6, 4.0, 4.8])
    add_heading(doc, "4. 开发路线与评估", level=1)
    add_paras(doc, [
        "v4 应按五个阶段推进：research prototype、shadow mode、teacher-approved pilot、limited automatic mode 和 cross-school generalization。每个阶段都必须有明确 stop gate。若 causal estimate 不可用，回退 v3；若 multimodal signal 缺失，使用 text/log-only learner state；若 privacy budget 耗尽，停止跨校更新；若 fairness gap 超阈值，关闭相关推荐策略。",
        "评估需要三类证据：离线模型证据、教师审查证据和课堂学习证据。离线层关注 uplift calibration、counterfactual sanity checks、retention prediction 和 fairness gap；教师层关注 recommendation appropriateness、evidence sufficiency、risk clarity 和 override reason；课堂层关注 normalized learning gain、repair speed、review retention、transfer performance、student engagement 和 teacher workload。",
    ])
    add_references(doc, REFERENCES)
    OUTS["research"].parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTS["research"])
    return OUTS["research"]


def build_spec() -> Path:
    refs = [r for r in REFERENCES if r.startswith(("Athey", "Wager", "Pham", "Li,", "Corbett", "Pardos", "Christie", "Yang", "Dwork", "Baker, R. S., & Hawn", "Holstein", "Lewis", "Asai"))]
    doc = setup_doc("MAIS v4 算法设计规格", "接口、数据对象、公式与失败模式", "算法设计规格")
    add_heading(doc, "1. 系统边界", level=1)
    add_bullets(doc, [
        "v4 只能评估和增强 v3 已生成并验证过的候选；不能由 LLM 直接生成教学边界。",
        "BKT safety floor 继承自知识追踪传统，负责保留可解释 mastery、review 和 repair 边界（Corbett & Anderson, 1995）。",
        "RAG evidence pack 继承 v3 的检索增强生成边界，用于限制候选解释和教师审计备注（Lewis et al., 2020）。",
        "所有 causal/uplift 估计必须带 evidence quality 与 outcome window；低证据推荐只允许进入 shadow 或 exploratory 状态。",
        "多模态数据默认只保存 feature summary；原始图像、音视频或手写过程数据必须有 retention policy 和学校审批。",
    ])
    add_heading(doc, "2. 输入/输出接口草案", level=1)
    add_table(doc, "表 1\nv4 未来接口草案", ["接口", "输入", "输出", "说明"], [
        ("getV4CandidateEvaluation", "candidateId, learnerState, evidencePack", "V4DecisionRecord", "评估候选的 uplift、risk、fairness、cost"),
        ("estimateInstructionalUplift", "candidateId, learnerProfile, context", "CATE estimate", "估计推荐对该类学生的增量收益"),
        ("logMultimodalLearningEvent", "event", "eventId", "只记录允许的数据摘要和 privacy class"),
        ("getLongitudinalMasteryTrace", "studentId, skillId, window", "mastery trajectory", "查看长期掌握、保持和迁移"),
        ("getSchoolGovernanceMetrics", "scope, timeWindow", "GovernanceMetricSnapshot", "学校/班级/年级治理指标"),
    ], [3.8, 4.4, 3.8, 4.4])
    add_heading(doc, "3. 核心数据对象", level=1)
    add_table(doc, "表 2\nv4 数据对象", ["对象", "关键字段", "用途"], [
        ("V4DecisionRecord", "candidate, baseline, evidence, uplift, risk, fairness, cost, teacherAction, fallbackReason", "完整审计一个 v4 推荐"),
        ("MultimodalLearningEvent", "eventType, timestamp, skill, privacyClass, featureSummary, rawRetentionPolicy", "记录多模态学习证据"),
        ("GovernanceMetricSnapshot", "learningGain, retention, override, fallback, fairnessGap, privacyBudget, cost", "学校级治理 dashboard"),
        ("OutcomeWindow", "preState, postState, delay, transferItems, reviewDue", "定义因果/长期结果窗口"),
    ], [3.8, 8.6, 4.0])
    add_heading(doc, "4. 算法公式", level=1)
    add_formulas(doc, 1, 8)
    add_heading(doc, "5. 失败模式", level=1)
    add_bullets(doc, [
        "causal estimate unavailable：回退 v3，并记录 `fallbackReason=causal_unavailable`。",
        "multimodal signal missing：使用 text/log-only learner state，不降低现有 v3 能力。",
        "privacy budget exhausted：停止跨校更新，只保留本校本地评估。",
        "fairness gap exceeds threshold：禁用受影响策略并要求教师审批。",
        "teacher repeated rejection：标记 candidate policy review，不自动调高该策略权重。",
    ])
    add_paragraph(doc, "规格依据 v3 的 BKT/KG/RAG/validation/teacher governance，并引入 causal/uplift、multimodal、privacy-preserving 和 fairness 文献支撑（Asai et al., 2024; Athey & Imbens, 2016; Baker & Hawn, 2022; Christie et al., 2024; Dwork, 2006; Holstein et al., 2019; Li et al., 2023; Pardos & Heffernan, 2010; Pham et al., 2024; Wager & Athey, 2018; Yang et al., 2019）。")
    add_references(doc, refs)
    doc.save(OUTS["spec"])
    return OUTS["spec"]


def build_roadmap() -> Path:
    refs = [r for r in REFERENCES if r.startswith(("Pham", "Baker, R. S., Hutt", "Li,", "Martinez", "So,", "Holstein", "Baker, R. S., & Hawn", "Dwork", "Yang"))]
    doc = setup_doc("MAIS v4 工程路线图", "从研究原型到跨校泛化", "工程路线图")
    add_heading(doc, "1. 阶段计划", level=1)
    add_table(doc, "表 1\nv4 五阶段路线图", ["阶段", "工程任务", "交付物", "进入下一阶段条件"], [
        ("Phase 1 Research prototype", "data dictionary、模拟/历史数据实验、uplift 与长期 mastery notebook", "research memo、algorithm spec、offline baseline", "uplift calibration 与 schema audit 通过"),
        ("Phase 2 Shadow mode", "后台生成 v4 评估但不影响学生路径", "shadow dashboard、teacher comparison report", "teacher agreement 与 unsupported causal claim rate 达标"),
        ("Phase 3 Teacher-approved pilot", "教师端显示证据、uplift、风险、公平性提示", "approval workflow、override taxonomy", "教师接受度与学习收益初步达标"),
        ("Phase 4 Limited automatic", "仅低风险高证据场景自动化", "policy gate、rollback switch、incident playbook", "无重大 fairness/risk/privacy 告警"),
        ("Phase 5 Cross-school generalization", "隐私保护跨校评估与联邦/DP 实验", "school governance dashboard、privacy report", "跨校泛化收益超过治理成本"),
    ], [3.3, 5.0, 4.0, 4.2])
    add_heading(doc, "2. 工程模块", level=1)
    add_bullets(doc, [
        "Event ingestion：统一 learning event、multimodal event、teacher override 与 outcome window。",
        "Feature store：保存可审计 feature summary，不默认保存敏感原始数据。",
        "Causal evaluation service：输出 CATE/uplift、confidence、适用学生群体和证据质量。",
        "Longitudinal mastery service：追踪 repair speed、retention stability、transfer performance。",
        "Governance dashboard：显示 learning gain、fairness gap、fallback、override、privacy budget、cost。",
        "Rollback governance：按学校、班级、年级、知识点、题型或策略关闭 v4。",
    ])
    add_heading(doc, "3. 依赖与风险", level=1)
    add_paras(doc, [
        "v4 的主要依赖是 v3 的候选生成、RAG evidence pack、validation gates、teacher audit console 和日志体系。如果 v3 的 evidence IDs、candidate IDs 或 teacher override records 不稳定，v4 的因果评估和长期追踪会失去审计基础。",
        "工程风险包括数据稀疏、跨校分布偏移、隐私预算管理复杂、教师负担上升和 dashboard 误读。路线图必须保留 shadow mode、teacher-approved mode 和 limited automatic mode 的分层发布策略。Open science 与教育数据挖掘实践也要求报告数据、代码、假设和负结果的透明度（Baker et al., 2024）。",
    ])
    add_paragraph(doc, "路线图依据 causal personalization、privacy-preserving analytics、multimodal learning analytics、teacher-AI complementarity 和 algorithmic bias 文献（Baker & Hawn, 2022; Dwork, 2006; Holstein et al., 2019; Li et al., 2023; Martinez Beltran et al., 2023; Pham et al., 2024; So et al., 2023; Yang et al., 2019）。")
    add_references(doc, refs)
    doc.save(OUTS["roadmap"])
    return OUTS["roadmap"]


def build_pilot_plan() -> Path:
    refs = [r for r in REFERENCES if r.startswith(("Athey", "Wager", "Pham", "Baker, R. S., & Hawn", "Holstein", "Shute", "Cepeda", "Roediger", "Bloom", "Dwork", "So,", "Li,"))]
    doc = setup_doc("MAIS v4 评估与学校试点方案", "离线评估、教师审查、课堂试点、隐私与公平性 QA", "评估与学校试点方案")
    add_heading(doc, "1. 试点设计", level=1)
    add_table(doc, "表 1\nv4 评估层级", ["层级", "目标", "指标", "停止条件"], [
        ("Offline evaluation", "验证 uplift、retention、fairness 是否可估计", "uplift calibration、counterfactual sanity、retention prediction、fairness gap", "估计不稳定或 subgroup harm gap 超阈值"),
        ("Teacher review", "验证推荐是否可解释、可采纳", "appropriateness、evidence sufficiency、risk clarity、override reason", "教师连续拒绝或证据不足"),
        ("Shadow mode", "不影响学生路径地比较 v4 与 v3/教师判断", "teacher agreement、unsupported causal claim rate、latency、cost", "unsupported causal claim rate 超阈值"),
        ("Teacher-approved pilot", "教师审批后推送推荐", "approval rate、modification rate、repair speed、teacher workload", "教师负担或风险告警过高"),
        ("Classroom pilot", "评估真实学习效果", "normalized learning gain、review retention、transfer、engagement", "公平性、隐私或学习收益未达门槛"),
    ], [3.0, 4.6, 5.2, 3.7])
    add_heading(doc, "2. 结果窗口", level=1)
    add_bullets(doc, [
        "短期：推荐后 1-3 次练习的正确率、hint 使用、repair completion。",
        "中期：7-14 天 review retention、同知识点不同题型迁移。",
        "长期：4-8 周 mastery stability、跨单元 transfer、教师确认的概念稳定性。",
        "公平性：按年级、语言、学校、能力段、知识点监测 learning gain 与 harm rate。",
    ])
    add_heading(doc, "3. 隐私与伦理", level=1)
    add_bullets(doc, [
        "试点前完成学校审批、家长/学生同意、数据最小化和退出机制。",
        "跨校分析不得集中原始学生记录；必要时只交换聚合指标、差分隐私输出或联邦模型更新。",
        "多模态数据必须默认脱敏或只保存 feature summary；原始影像/手写轨迹保留需单独审批。",
        "所有自动推荐必须有 rollback switch；教师保留最终教学权威。",
    ])
    add_heading(doc, "4. 接受标准", level=1)
    add_bullets(doc, [
        "v4 相比 v3 在至少一个主要教育指标上有稳定提升，且未增加 harm gap。",
        "教师认为推荐理由、风险提示和证据链可理解、可审计、可覆盖。",
        "privacy budget、fallback rate、latency、cost per learner 均在学校可接受范围内。",
        "所有高风险推荐仍停留在 teacher-approved mode，不进入 automatic mode。",
    ])
    add_paragraph(doc, "本试点方案以高质量个别化辅导、形成性反馈、间隔/提取练习、因果异质性效果、教师-AI 互补、公平性和隐私保护研究为依据（Athey & Imbens, 2016; Baker & Hawn, 2022; Bloom, 1984; Cepeda et al., 2006; Dwork, 2006; Holstein et al., 2019; Li et al., 2023; Pham et al., 2024; Roediger & Karpicke, 2006; Shute, 2008; So et al., 2023; Wager & Athey, 2018）。")
    add_references(doc, refs)
    doc.save(OUTS["pilot"])
    return OUTS["pilot"]


def main() -> None:
    PAPER_DIR.mkdir(parents=True, exist_ok=True)
    paths = [build_research_report(), build_spec(), build_roadmap(), build_pilot_plan()]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
