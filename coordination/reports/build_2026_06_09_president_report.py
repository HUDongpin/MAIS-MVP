from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "coordination/reports/2026-06-09-president-report.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "555555"
HEADER_FILL = "F2F4F7"
LIGHT_FILL = "F7F9FC"
RISK_FILL = "FFF2CC"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Calibri"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 20, True, DARK_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(subtitle)
    set_run_font(r, 10.5, False, MUTED)


def add_para(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, 11, True)
        rest = text[len(bold_prefix) :]
        if rest:
            r = p.add_run(rest)
            set_run_font(r, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, 11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.2, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size, bold, color)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 9.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, h in enumerate(headers):
        shade_cell(table.rows[0].cells[i], HEADER_FILL)
        set_cell_text(table.rows[0].cells[i], h, True, font_size, DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            set_cell_text(row.cells[i], text, False, font_size)
    set_table_widths(table, widths)
    doc.add_paragraph()


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, val in rows:
        row = table.add_row()
        shade_cell(row.cells[0], LIGHT_FILL)
        set_cell_text(row.cells[0], key, True, 9.5, DARK_BLUE)
        set_cell_text(row.cells[1], val, False, 9.5)
    set_table_widths(table, [1.55, 4.95])
    doc.add_paragraph()


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("MAIS-MVP President Report | 2026-06-09")
    set_run_font(r, 8.5, False, MUTED)


def main() -> None:
    doc = Document()
    style_doc(doc)
    doc.core_properties.title = "MAIS-MVP 2026-06-09 President Report"
    doc.core_properties.subject = "Daily bilingual president report for Dr. Peter Hu"
    doc.core_properties.author = "S10 / Codex automation"

    add_title(
        doc,
        "MAIS-MVP 8 AM President Report",
        "Daily bilingual coordination report for Dr. Peter Hu",
    )

    add_kv_table(
        doc,
        [
            ("Report date", "2026-06-09"),
            ("Report time", "08:00 Asia/Hong_Kong; generated after fresh checks at 08:03-08:38 HKT"),
            ("Project", "MAIS-MVP"),
            ("Reporting session", "S10 / Automation mais-mvp-9-am-president-report"),
            ("Audience", "Dr. Peter Hu"),
            ("Reporting window", "2026-06-08 08:00 to 2026-06-09 08:00 Asia/Hong_Kong"),
        ],
    )

    doc.add_heading("中文 Executive Summary", level=1)
    add_bullets(
        doc,
        [
            "本窗口有明确推进，不是空窗期。活跃证据覆盖 S01、S04、S05、S06、S07、S09、S10、S11、S12、S22；其余会话无 fresh log。",
            "产品层面完成多项面向学生和教师的可见改动：Adaptive Learning 改为 Personalized Learning，AI Tutor/Nova Lens 统一为 Nova Tutor，Lesson 页加入 Core concept AI audio guide，Practice Arena 提速并修复 4/5 正确后的游戏入口。",
            "后端和发布层面推进明显：S12 完成 forum 生产级搜索、分页、附件、通知、审核队列、审计与桌面/移动 E2E；S22 在 owner 明确要求和 dirty-root override 下完成 production deployment，`https://mais.hk` 与 `https://www.mais.hk` HTTP smoke 均为 200。",
            "今天 fresh baseline 为绿色：`npm run type-check` 通过；`npm run test:analytics` 23/23 通过；`npm run release:preflight -- --json` 通过；`npm run build` 通过并生成 127 个静态页面。",
            "最大风险仍是 release hygiene：根目录仍有 3,387 个 git status entries，direct root deploy guard 继续保护性失败；production 已从极脏工作树通过 S22 staging wrapper 发布，必须尽快做 S25 release slicing、S11/S22 全量回归和 rollback 判断。",
        ],
    )

    doc.add_heading("English Executive Summary", level=1)
    add_bullets(
        doc,
        [
            "The reporting window had substantial fresh work across product, API, QA, and release engineering; it was not an inactive window.",
            "User-facing progress included Personalized Learning naming/route migration, Nova Tutor naming and behavior fixes, Lesson audio guide/player work, Practice Arena speed and unlock fixes, and visualization math typography cleanup.",
            "Backend/release progress was material: the class forum gained production-grade API/UI capabilities and targeted desktop/mobile E2E evidence, while S22 completed a Vercel production deployment to `mais.hk` after explicit owner approval for the dirty-worktree staging path.",
            "Fresh local checks are green: type-check passed, analytics/adaptive tests passed 23/23, release preflight passed, and production build passed with 127 static pages.",
            "The release risk is still high. The root worktree remains too dirty for direct deploy, full-product E2E has not been rerun after cleanup, and Mainland PEP High mobile lesson smoke still has a horizontal-overflow issue.",
        ],
    )

    doc.add_heading("报告窗口摘要", level=1)
    add_kv_table(
        doc,
        [
            ("窗口", "2026-06-08 08:00 至 2026-06-09 08:00 HKT"),
            ("活跃会话", "S01, S04, S05, S06, S07, S09, S10, S11, S12, S22"),
            ("无 fresh log 会话", "S02, S03, S08, S13, S14, S15, S16, S17, S18, S19, S20, S21, S23, S24, S25"),
            ("No assigned work in this reporting window", "No"),
            ("协调重点", "命名迁移跨 S01/S04/S09/S07/S05；Forum/API 与 release checks 跨 S12/S11/S22；生产部署由 S22 使用 pruned staging path 完成。"),
            ("新 blocker/风险证据", "S12 记录 forum full-E2E disk gate；S11 记录 PEP High mobile overflow；S22 记录 dirty-root production deployment residual risk。"),
        ],
    )

    doc.add_heading("整体项目进展", level=1)
    add_bullets(
        doc,
        [
            "学生端从 Adaptive 命名转向 Personalized/Nova Tutor，并继续改善 lesson/practice 的首屏体验、音频辅助和移动展示质量。",
            "Practice Arena 首屏 catalog payload 从 986 full questions/约 510 KB 改为 topic summary/约 7.4 KB；实际题目仍按需加载。",
            "Forum 从 MVP 形态推进到更完整的 class discussion platform，包含搜索、分页、附件、通知、teacher moderation、review queue 和 audit trail。",
            "生产站点 `mais.hk` 已发布新版本并返回 200；但发布源仍未整理为 clean reviewed release worktree。",
            "QA 状态比昨日更强，但还不能称为全量 release green：全量 E2E 未跑，PEP High mobile route smoke 未关闭，dirty-root guard 仍失败。",
        ],
    )

    doc.add_heading("S01-S25 会话状态表", level=1)
    session_rows = [
        ["S01", "完成", "Hero 和导航把 Adaptive/Adaptive Learning 改为 Personalized/Personalized Learning；新增 `/personalized-learning`，旧路由 redirect。", "确认新 URL 与旧兼容路径；继续清理历史文档/视频文案。", "type, diff, smoke"],
        ["S02", "无 fresh log", "Dashboard lead 本窗口无新交付。", "沿用已有 dashboard/adaptive UI 风险。", "未跑"],
        ["S03", "无 fresh log", "Curriculum roadmap 本窗口无新交付。", "无新增阻塞。", "未跑"],
        ["S04", "完成", "Practice copy 改 Personalized；首屏 catalog summary API；动态加载重组件；4/5 正确立即显示 Adventure CTA。", "早期 build/typing 问题已被后续 fresh build/type-check 覆盖；需生产 smoke。", "type, diff, PW"],
        ["S05", "完成", "Lesson 删除 extension/metadata badge；Galaxy copy/loading 修正；Core concept AI audio guide；audio control cleanup；title fitting。", "音频真实播放仍需用户浏览器手动确认；TTS 依赖 DashScope/Qwen env。", "type, local smoke"],
        ["S06", "完成", "Quadratic explorer equation/readout 改用显式 inline math，消除混合字体。", "无新增阻塞。", "type, PW"],
        ["S07", "完成", "Nova Lens/Tutor 根据用户输入语言回复；Voice input 先用 browser SpeechRecognition，再 fallback Qwen ASR。", "不同浏览器语音支持差异仍需生产监测。", "type, build, PW, ASR"],
        ["S08", "无 fresh log", "State/analytics 本窗口无新 session。", "Fresh analytics/adaptive tests 由 S10 跑通过。", "S10 covered"],
        ["S09", "完成", "AI Tutor/Nova Lens 统一为 Nova Tutor；dashboard Galaxy prompt；assignment 多语言字段与重复记录清理。", "本地 `.local` state 被清理；生产历史 assignment 仍需迁移策略。", "type, API smoke"],
        ["S10", "完成", "生成 2026-06-08 president report；本次生成 2026-06-09 report 并跑 fresh checks。", "DOCX render QA 取决于本机 LibreOffice/pdf2image。", "type, tests, build"],
        ["S11", "部分完成", "PEP Junior smoke 因 analytics 404 修复后 6/6 通过；PEP High API/desktop 通过。", "PEP High mobile `pep-high-s4-quadratic-inequalities` 横向溢出 49 px，不能 green。", "PW, type"],
        ["S12", "完成/待全量", "Forum 生产级功能完成；login authenticated `/login` redirect 修复；forum desktop/mobile targeted E2E 3 pass/3 skip。", "Full `npm run test:e2e` 未跑；早期 disk blocker 已缓解但 gate 未重跑。", "type, build, forum PW"],
        ["S13", "无 fresh log", "Teacher console 本窗口无独立新 log。", "Teacher-facing assignment/forum 变更需 S13 后续 review。", "未跑"],
        ["S14", "无 fresh log", "Parent console 本窗口无新 log。", "无新增阻塞。", "未跑"],
        ["S15", "无 fresh log", "Adaptive engine 本窗口无新 session。", "Personalized naming未改 API/engine 语义；S15 regression 仍需守住。", "S10 tests"],
        ["S16", "无 fresh log", "Research 本窗口无新 log。", "无新增阻塞。", "未跑"],
        ["S17", "无 fresh log", "Gamification/reward 本窗口无新 log。", "Practice game unlock 触及 game CTA，需要 S17/S20 后续观察。", "未跑"],
        ["S18", "无 fresh log", "Content QA 本窗口无新 log。", "Content signoff backlog 仍沿用历史状态。", "未跑"],
        ["S19", "无 fresh log", "API/env 本窗口无新 log。", "Production env status依赖 S22 deploy evidence；未读取 secret。", "未跑"],
        ["S20", "无 fresh log", "Game-based learning 本窗口无新 log。", "Adventure CTA flow 需要 game owner后续确认。", "未跑"],
        ["S21", "无 fresh log", "Content pipeline/RAG 本窗口无新 log。", "无新增 candidate package evidence。", "未跑"],
        ["S22", "完成/高风险", "Preview deploy ready；owner-approved production deploy ready；`mais.hk` 与 `www.mais.hk` smoke 200。", "Production 来自极脏工作树 staging wrapper；raw root deploy guard 失败。", "preflight, build, deploy"],
        ["S23", "无 fresh log", "Integration/promotion 本窗口无新 log。", "需要接手 release slicing 与 live promotion清单。", "未跑"],
        ["S24", "无 fresh log", "Exact-layer 本窗口无新 log。", "无新增阻塞。", "未跑"],
        ["S25", "无 fresh log", "Git hygiene 本窗口无新 log。", "今天最需要 S25：3387 dirty entries 的 release slicing。", "未跑"],
    ]
    add_table(doc, ["Session", "状态", "已完成/证据", "风险/下一步", "Checks"], session_rows, [0.45, 0.7, 2.55, 2.0, 0.8], 7.8)

    doc.add_heading("已完成工作", level=1)
    add_bullets(
        doc,
        [
            "S01/S04/S09/S07/S05 完成大范围但有边界的命名与体验修正：Personalized Learning、Nova Tutor、Lesson audio、Practice unlock 与 dashboard copy。",
            "S04 完成 Practice Arena 首屏性能优化，并用 Playwright 验证首屏只拉 topic-catalog，选择 Core 后才拉 full questions。",
            "S05 完成 Lesson 页多项 UI/UX 改动与 AI audio guide；本地 API/browser smoke 确认 `/api/lesson-audio` 返回 200 和 audio blob。",
            "S12 完成 forum search/pagination/attachments/notifications/moderation/review/audit，并跑过 targeted desktop/mobile forum E2E。",
            "S22 完成 Preview 和 Production deploy；production custom domains 返回 HTTP 200。",
            "S10 fresh checks 全绿：type-check、analytics/adaptive tests、release preflight、build。",
        ],
    )

    doc.add_heading("进行中工作", level=1)
    add_bullets(
        doc,
        [
            "Mainland PEP High lesson mobile overflow：S11 已定位到 `components/lesson/LessonGalaxyDirectory.tsx` 一类移动布局问题，需修复后重跑 high/junior smoke。",
            "Full-product E2E：S12/S10 清理 `.tmp` 后磁盘已恢复，但 `npm run test:e2e` 还没有重新跑。",
            "Release slicing：production 已发布，但 dirty-tree 尚未切成可 review、可回滚的 release units。",
            "Production post-deploy verification：custom domain 200 已确认，但新功能的 authenticated role flows、forum、Nova Tutor voice/audio、Practice unlock 仍需生产 smoke。",
        ],
    )

    doc.add_heading("阻塞事项", level=1)
    add_bullets(
        doc,
        [
            "Direct root deploy blocker：`npm run release:root-deploy-preflight -- --json` 保护性失败，当前 3,387 status entries、257 tracked modified、4 tracked deleted、3,126 untracked status entries、32,384 untracked files。",
            "PEP High mobile lesson smoke blocker：`pep-high-s4-quadratic-inequalities` mobile overflow 49 px，release smoke 不能标绿。",
            "Full E2E evidence gap：`.tmp` cleanup 后 release preflight 通过，但 full `npm run test:e2e` 未重跑。",
            "S12 disk blocker 状态：原 blocker 因 `.tmp` 约 91-92 GB 导致 e2e guard 失败；S10 `test:analytics` 脚本清理后目前已缓解，但 blocker 需要用正式 rerun 关闭。",
        ],
    )

    doc.add_heading("风险", level=1)
    add_bullets(
        doc,
        [
            "生产发布风险：S22 根据 owner 明确要求使用 dirty-root override 和 pruned staging wrapper 发布；这不是 clean reviewed release worktree。",
            "交叉范围风险：`app/layout.tsx`、`lib/server/userStore.ts`、`components/ai/AITutorProvider.tsx`、`components/lesson/LessonView.tsx` 被多个会话触及，后续必须做 ownership review。",
            "本地状态风险：S09 修改 `.local/hk-math-db.sqlite` 清理重复 assignment；这不是生产数据迁移。",
            "Provider 风险：Lesson audio 与 Nova voice 依赖 DashScope/Qwen/浏览器语音能力，需监控 latency、quota、unsupported-browser fallback。",
            "测试噪声风险：Vercel analytics local 404 已修复给 PEP Junior smoke，但此类 app-shell/harness 差异可能继续影响 route smoke 判断。",
        ],
    )

    doc.add_heading("测试与构建状态", level=1)
    test_rows = [
        ["Fresh S10", "`npm run type-check`", "Passed", "当前工作树 TypeScript 通过。"],
        ["Fresh S10", "`npm run test:analytics`", "Passed 23/23", "覆盖 learningAnalytics 与 adaptiveLearning tests；脚本清理 `.tmp`。"],
        ["Fresh S10", "`npm run release:preflight -- --json`", "Passed", "freeBytes 78,438,686,720；当前 `df` 约 73 GiB available。"],
        ["Fresh S10", "`npm run build`", "Passed", "Next build 成功，生成 127 static pages。"],
        ["Fresh S10", "`npm run release:root-deploy-preflight -- --json`", "Failed as expected", "Direct root deploy blocked by dirty worktree。"],
        ["S11", "PEP Junior lesson smoke", "Passed 6/6 after fix", "Analytics local 404 gate 已修。"],
        ["S11", "PEP High lesson smoke", "Partial", "API/desktop pass；mobile overflow fail。"],
        ["S12", "Forum targeted E2E", "Passed targeted", "Desktop API/UI + mobile smoke: 3 passed, 3 skipped。"],
        ["S22", "Production deploy", "Ready", "`mais.hk` 与 `www.mais.hk` HTTP 200。"],
    ]
    add_table(doc, ["来源", "检查", "结果", "说明"], test_rows, [0.75, 2.0, 1.05, 2.7], 8.6)

    doc.add_heading("变更文件", level=1)
    add_para(
        doc,
        "mtime 扫描口径：2026-06-08 08:00 至 2026-06-09 08:00 HKT，排除 `.git`、`node_modules`、`.next`、`.tmp`、真实 env/secret 文件。共 124 个文件在窗口内有非生成/非 secret mtime 变化。",
    )
    changed_rows = [
        ["tests/e2e", "57", "多项 E2E 规格 mtime 变化，S09/S11/S12/S22 相关断言和 harness 覆盖最明显。"],
        ["components", "18", "Hero/Navbar/Lesson/AI/Forum/Practice/Visualization 等用户界面改动。"],
        ["app", "16", "Personalized route、login redirect、questions API、lesson audio API、layout analytics gate、teacher assignment API。"],
        ["coordination", "13", "S01/S04/S05/S06/S07/S09/S10/S11/S12/S22 logs 与 2026-06-08 president report。"],
        ["lib/data/types", "6", "`lib/i18n.ts`、`lib/server/userStore.ts`、HK Chinese glossary/exceptions、`types/index.ts`。"],
        [".local/public/video-plan/other", "14", "本地 sqlite state、`.DS_Store`、question-illustration `.DS_Store`、video-plan 文案引用。"],
    ]
    add_table(doc, ["区域", "数量", "代表变化"], changed_rows, [1.25, 0.55, 4.7], 8.8)

    doc.add_heading("明日优先级", level=1)
    add_bullets(
        doc,
        [
            "1. S11/S05 修复 PEP High mobile horizontal overflow，并重跑 high/junior lesson route smoke。",
            "2. S11/S22 在当前磁盘恢复后跑 full `npm run test:e2e`，并把结果写入 release-readiness evidence。",
            "3. S25 做 dirty-tree release intake：按 owner/session ownership 切分 3,387 entries，给出可 review/rollback 的 release slice。",
            "4. S22/S12/S07/S05 做 production smoke：登录/角色跳转、forum、Nova Tutor language/voice、Lesson audio、Practice game unlock。",
            "5. S19/S22 核对 production env parity，只记录变量名和红acted status，不打印 secret。",
        ],
    )

    doc.add_heading("需要 Dr. Peter Hu 决策", level=1)
    add_bullets(
        doc,
        [
            "是否接受当前 production deployment 为今天基线，还是要求 rollback/freeze 直到 S25 切片和 full E2E 完成。",
            "是否立即授权 S25 做 release slicing，并指定是否允许把相关 dirty changes 分批提交/PR（当前 automation 未执行任何 git staging/commit）。",
            "是否把 PEP High mobile overflow 设为今日 P1 launch blocker，并指定 owner：S05 Lesson UI 还是 S11 QA 协调修复。",
            "是否确认 Personalized Learning 与 Nova Tutor 作为正式对外命名，并授权 S09/S10 清理 README、video-plan、旧报告以外的当前文案。",
            "是否要求 S12 forum 在 production 上做 teacher/student authenticated smoke 后再对外开放。",
        ],
    )

    add_footer(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
