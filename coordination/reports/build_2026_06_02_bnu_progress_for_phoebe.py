from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "coordination" / "reports" / "2026-06-02-mainland-bnu-rag-content-status-for-phoebe.docx"

FONT_EAST_ASIA = "Microsoft YaHei"
FONT_LATIN = "Calibri"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B677A"
GRID = "D7DEE8"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
GREEN_FILL = "E2F0D9"
GREEN_TEXT = "27632A"
YELLOW_FILL = "FFF4CE"
YELLOW_TEXT = "7A5A00"
RED_FILL = "FDE7E9"
RED_TEXT = "9B1C1C"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None):
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold


def set_para(paragraph, before: float = 0, after: float = 6, line: float = 1.1):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_text(paragraph, text: str, size: float = 11, color: str = INK, bold: bool = False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return run


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, size: float = 10.3, color: str = INK, bold: bool = False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.15)
    if align is not None:
        p.alignment = align
    add_text(p, text, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_geometry(table, widths_in: list[float], indent_dxa: int = 120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), GRID)

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [("top", 80), ("bottom", 80), ("left", 120), ("right", 120)]:
        el = margins.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            margins.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)

    widths_dxa = [round(w * 1440) for w in widths_in]
    correction = 9360 - sum(widths_dxa)
    widths_dxa[-1] += correction
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_rule(paragraph, color: str = BLUE, size: str = "8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def set_cell_status(cell, label: str, status: str):
    fill, color = {
        "green": (GREEN_FILL, GREEN_TEXT),
        "yellow": (YELLOW_FILL, YELLOW_TEXT),
        "red": (RED_FILL, RED_TEXT),
    }[status]
    set_cell_shading(cell, fill)
    set_cell_text(cell, label, size=10, color=color, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return p


def add_body(doc, text: str, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    set_para(p, after=6, line=1.1)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, bold=True)
        add_text(p, text[len(bold_prefix):])
    else:
        add_text(p, text)
    return p


def add_callout(doc, title: str, body: str, fill: str = LIGHT_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p1 = cell.paragraphs[0]
    set_para(p1, after=2, line=1.12)
    add_text(p1, title, size=11.5, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    set_para(p2, after=0, line=1.16)
    add_text(p2, body, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_status_table(doc):
    data = [
        ("课程标准 / RAG", "绿灯", "已完成", "小学、初中、高中安全 RAG 与测评模式层已建；北师大定向 RAG 回归 75/75 通过。"),
        ("AI 生成题目", "绿灯", "已上线", "P1-P6 3000 题、S1-S3 1500 题、S4-S6 1500 题，总计 6000 题已接入题库。"),
        ("QA 题目", "绿灯", "已通过", "题库回归 67/67 通过；初中 153 个 DeepSeek 问题行已修复，高中使用已批准题目包。"),
        ("AI 生成教材", "黄灯", "已生成初中草案", "初中 6 册、35 单元、105 课、6 个 DOCX 已生成；仍为 review-only，未接入公开课页。"),
        ("QA 教材", "黄灯", "未完成", "本地结构校验与 preflight 通过；DeepSeek smoke/full QA、人工审核、DOCX 渲染 QA 尚未完成。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [1.42, 0.74, 1.04, 4.3])
    headers = ["模块", "状态", "进度", "课程专家需要知道的边界"]
    for idx, header in enumerate(headers):
        set_cell_shading(table.cell(0, idx), HEADER_FILL)
        set_cell_text(table.cell(0, idx), header, size=10.2, color=DARK_BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(table.rows[0])
    for module, signal, progress, note in data:
        cells = table.add_row().cells
        set_cell_text(cells[0], module, bold=True)
        set_cell_status(cells[1], signal, "green" if signal == "绿灯" else "yellow")
        set_cell_text(cells[2], progress, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[3], note, size=10.1)
    return table


def add_inventory_table(doc):
    rows = [
        ("小学 P1-P6", "题目", "3000", "已接入", "v1 + v2 各 1500；每年级 500 题。"),
        ("初中 S1-S3", "题目", "1500", "已接入", "每年级 500 题；153 个 DeepSeek issue 行已修复/裁决。"),
        ("高中 S4-S6", "题目", "1500", "已接入", "使用已批准题目包 `question-pack.approved.json`，不是早期未通过的原始候选包。"),
        ("初中 S1-S3", "教材", "105 课 / 6 册", "review-only", "已生成 DOCX；未完成 DeepSeek full QA、人工审核和公开课页接入。"),
        ("小学 P1-P6", "课页", "lesson seeds", "可用", "已有小学 lesson seeds 与练习 checkpoints；不是完整富教材正文包。"),
        ("高中 S4-S6", "教材", "未见完整包", "待规划", "当前证据显示题库可用；完整高中教材正文包尚未形成。"),
    ]
    table = doc.add_table(rows=1, cols=5)
    set_table_geometry(table, [1.15, 0.76, 1.12, 1.0, 4.22])
    for i, h in enumerate(["范围", "内容", "数量", "状态", "说明"]):
        set_cell_shading(table.cell(0, i), HEADER_FILL)
        set_cell_text(table.cell(0, i), h, size=10.2, color=DARK_BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(table.rows[0])
    for scope, content, qty, status, note in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], scope, bold=True)
        set_cell_text(cells[1], content, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], qty, align=WD_ALIGN_PARAGRAPH.CENTER)
        if status in ["已接入", "可用"]:
            set_cell_status(cells[3], status, "green")
        elif status == "待规划":
            set_cell_status(cells[3], status, "red")
        else:
            set_cell_status(cells[3], status, "yellow")
        set_cell_text(cells[4], note, size=10.1)


def add_action_table(doc):
    rows = [
        ("1", "先审初中教材样本", "Phoebe / S18", "每册抽 2 个单元，重点看知识顺序、术语、例题梯度、与北师大体系的一致性。"),
        ("2", "跑教材 DeepSeek smoke/full QA", "S18 / S19", "需配置 approved DeepSeek V4 Pro endpoint；当前本地 endpoint 被识别为 DashScope，runner 已拒绝。"),
        ("3", "完成 DOCX 渲染 QA", "S18 / S10", "用 LibreOffice 渲染六册教材 DOCX，检查分页、表格、公式、中文字体、是否有裁切。"),
        ("4", "形成教材上线判定", "Phoebe / S05 / S11", "人工通过后，再做课页接入、浏览器 smoke、学生端路由验证。"),
        ("5", "对外口径保持清晰", "产品负责人", "题库可说已通过 QA 并接入；教材只能说“初中六册草案已生成，正在专家审核”。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [0.46, 1.65, 1.2, 4.29])
    for i, h in enumerate(["序号", "动作", "建议 owner", "完成标准"]):
        set_cell_shading(table.cell(0, i), HEADER_FILL)
        set_cell_text(table.cell(0, i), h, size=10.2, color=DARK_BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(table.rows[0])
    for no, action, owner, done in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], no, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], action, bold=True)
        set_cell_text(cells[2], owner, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[3], done, size=10.1)


def add_evidence_table(doc):
    rows = [
        ("RAG 回归", "北师大定向 RAG 测试", "75/75 passed", "2026-06-02 本地复核"),
        ("题库回归", "`npm run test:question-bank`", "67/67 passed", "2026-06-02 本地复核"),
        ("初中教材结构", "`validate-lessons.mjs`", "105 课，0 issues", "2026-06-01 S18 记录"),
        ("初中教材 preflight", "`deepseek-v4-pro-lesson-qa.mjs --preflight`", "315 sections，0 issues", "只做 preflight，未调用 provider"),
        ("教材 DOCX", "六册 DOCX 结构 audit", "结构可打开", "渲染 QA 尚待完成"),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [1.2, 2.3, 1.55, 2.45])
    for i, h in enumerate(["证据项", "检查 / 文件", "结果", "备注"]):
        set_cell_shading(table.cell(0, i), HEADER_FILL)
        set_cell_text(table.cell(0, i), h, size=10.2, color=DARK_BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(table.rows[0])
    for item, check, result, note in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], item, bold=True)
        set_cell_text(cells[1], check, size=9.9)
        set_cell_text(cells[2], result, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[3], note, size=9.9)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    set_para(header, after=0, line=1.0)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "MAIS 内容质量简报 | 北师大版", size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    set_para(footer, after=0, line=1.0)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "仅供内部课程审核使用", size=9, color=MUTED)


def build_doc():
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    set_para(p, before=10, after=4, line=1.0)
    add_text(p, "课程内容进度简报", size=12, color=BLUE, bold=True)

    title = doc.add_paragraph()
    set_para(title, after=4, line=1.05)
    add_text(title, "北师大版 RAG、AI 教材与题库 QA 进度", size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    set_para(subtitle, after=14, line=1.1)
    add_text(subtitle, "给中国数学课程专家 Phoebe 的状态说明（简体中文）", size=13, color=MUTED, bold=True)

    for label, value in [
        ("日期：", "2026 年 6 月 2 日"),
        ("项目：", "MAIS-MVP 中国大陆北师大版数学内容线"),
        ("阅读对象：", "Phoebe，中国数学课程专家"),
        ("报告口径：", "基于本地项目 RAG、生成内容、QA 证据文件与 2026-06-02 回归复核"),
    ]:
        p = doc.add_paragraph()
        set_para(p, after=2, line=1.05)
        add_text(p, label, size=10.8, color=INK, bold=True)
        add_text(p, value, size=10.8, color=INK)

    rule = doc.add_paragraph()
    set_para(rule, before=10, after=12)
    add_rule(rule)

    add_callout(
        doc,
        "一句话结论",
        "北师大版课程标准 / RAG 与 P1-S6 题库已经达到当前产品可用状态；AI 生成教材目前只完成初中六册 review-only 草案，尚未完成专家审核、DeepSeek full QA、DOCX 渲染 QA 与公开课页接入。",
    )

    add_status_table(doc)
    doc.add_page_break()

    add_heading(doc, "1. 总体判断", 1)
    add_body(doc, "当前最重要的边界是：题目已经通过 QA 并接入；教材正文还不能按“已上线教材”对外表述。")
    add_body(doc, "5 月 28 日的北师大 AI-generated content launch approval，主要覆盖已批准题库、现有课页支架与 P1-S6 学生端可见性；它不等同于后续 6 月 1 日生成的初中六册富教材正文已经通过专家级 QA。")

    add_heading(doc, "2. 分项进度", 1)
    add_inventory_table(doc)

    add_heading(doc, "2.1 课程标准 / RAG", 2)
    add_body(doc, "北师大版安全 RAG 已覆盖小学、初中、高中三段，并包含对应测评模式卡。RAG 层只记录安全抽象、知识点、能力目标、题型/测评模式与复用限制，不保存官方教材原文、OCR 页、页码定位或可复刻文本。")
    add_body(doc, "2026-06-02 定向复核结果为 75/75 通过，说明当前 RAG 检索、publisher 隔离、证据包拼装与 source-distance guardrails 均处于绿色状态。")

    add_heading(doc, "2.2 AI 生成题目与题库 QA", 2)
    add_body(doc, "北师大版题库当前总量为 6000 题：小学 P1-P6 3000 题，初中 S1-S3 1500 题，高中 S4-S6 1500 题。学生端 Roadmap、Practice 与认证题目 API 已能按 MAINLAND_BNU publisher 隔离读取。")
    add_body(doc, "初中题库曾有 153 个 DeepSeek V4 Pro 问题行，已全部修复或裁决；高中早期原始候选包曾被 full-RAG QA 判为未通过，但产品实际接入的是后续确定性重建并批准的 `question-pack.approved.json`。")
    add_body(doc, "2026-06-02 题库回归 `npm run test:question-bank` 结果为 67/67 通过，覆盖北师大 P1-S6 可见性、publisher 隔离、题量、题型、答案可解性与全题库 solvability。")

    add_heading(doc, "2.3 AI 生成教材", 2)
    add_body(doc, "初中教材线已生成 review-only 草案：35 个 RAG 单元，每单元 3 课，共 105 课；并导出七年级上/下、八年级上/下、九年级上/下六册 DOCX。")
    add_body(doc, "这套教材是 MAIS 原创安全对齐版本，不是官方北师大教材的复制或改写。当前状态明确为 review-only，不应接入公开课页或直接给学生使用。")
    add_body(doc, "小学目前已有 lesson seeds 与练习 checkpoints，可支持现有课页体验；高中完整教材正文包在当前证据中尚未形成。")

    add_heading(doc, "2.4 QA 教材", 2)
    add_body(doc, "初中教材本地结构校验已通过：105 课、0 issue；DeepSeek QA preflight 已通过：315 sections，RAG coverage 315/315，0 issue。")
    add_body(doc, "但 preflight 不等于 full QA。当前 DeepSeek smoke/full QA 被 provider 配置挡住：runner 要求 DeepSeek V4 Pro endpoint，但本地 endpoint 被识别为 DashScope，因此按规则拒绝执行。")
    add_body(doc, "六册 DOCX 已完成结构 audit，但仍需要渲染到 PNG 检查中文字体、分页、表格、公式、裁切和可读性。")

    add_heading(doc, "3. Phoebe 建议重点审核", 1)
    add_callout(
        doc,
        "建议 Phoebe 的第一轮审核目标",
        "先不把任务定义为“逐字校对全部 105 课”，而是做课程结构与样本质量判定：每册抽 2 个单元，优先看知识顺序、例题梯度、术语一致性、难度是否符合北师大体系，以及是否存在看起来像官方教材措辞的高风险文本。",
        fill="F7FBFF",
    )
    add_action_table(doc)

    add_heading(doc, "4. 当前证据清单", 1)
    add_evidence_table(doc)

    add_heading(doc, "5. 建议对外口径", 1)
    add_body(doc, "可以说：北师大版 P1-S6 题库与课程标准 RAG 已完成当前 QA 并接入产品，题目层可作为学生练习内容使用。")
    add_body(doc, "应避免说：北师大版 AI 教材已经全部上线或已经通过课程专家审核。")
    add_body(doc, "更准确的表述是：初中六册 AI 原创教材草案已生成，正在进入课程专家审核、模型辅助 QA 和版式渲染 QA；审核通过后再进入 S05 课页接入与 S11 学生端回归。")

    add_heading(doc, "附录：关键证据文件", 1)
    for text in [
        "课程标准 / RAG：`data/rag/mainlandBnuPrimary.ts`、`mainlandBnuJunior.ts`、`mainlandBnuHigh.ts` 及 assessment pattern files。",
        "题库：`data/generated-content/mainland-bnu-primary-generated-bank-v1/v2-1500`、`mainland-bnu-junior-generated-bank-v1-1500/approved-question-pack.json`、`mainland-bnu-high-generated-bank-v1-1500/question-pack.approved.json`。",
        "初中教材草案：`coordination/content-qa/mainland-bnu-junior-lessons-v1/lessons.json`。",
        "初中六册 DOCX：`coordination/content-qa/mainland-bnu-junior-textbooks-v1/docx/`。",
        "上线批准记录：`coordination/content-qa/mainland-bnu-ai-generated-content-launch-approval-2026-05-28.md`。",
    ]:
        add_body(doc, text)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_doc()
    print(out)
