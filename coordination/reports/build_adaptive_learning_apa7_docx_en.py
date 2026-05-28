from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm

from build_adaptive_learning_apa7_docx import (
    ROOT,
    add_caption,
    add_formula,
    add_heading,
    add_note,
    add_page_number,
    add_paragraph,
    references,
    set_cell_margins,
    set_cell_shading,
    set_paragraph_spacing,
    set_run_font,
    set_table_borders,
    set_table_width,
    set_update_fields_on_open,
)


OUT = ROOT / "coordination/reports/2026-05-15-MAIS-v1-v2-adaptive-learning-APA7-EN.docx"


def add_toc_en(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "The table of contents will update automatically in Microsoft Word or LibreOffice."
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run._r.append(placeholder)
    run._r.append(fld_char3)
    set_run_font(run, size=11)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    set_paragraph_spacing(p, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.38)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_paragraph_spacing(p, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


def add_comparison_table(doc):
    add_caption(doc, "Table 1\nCore comparison of v1 and v2 adaptive learning mechanisms")
    rows = [
        ("Decision authority", "Deterministic BKT/rule engine using mastery probability, prerequisites, spaced review, and item mapping.", "BKT generates candidates first; the LLM can only rerank candidates and generate explanations."),
        ("Theoretical base", "Knowledge tracing, knowledge components, ITS student modeling, and mastery learning.", "Transformers/LLMs, instruction tuning, educational dialogue, constrained optimization, and human-AI complementarity."),
        ("Explainability", "High; every action can be traced to mastery, attempts, streaks, due reviews, and prerequisites.", "Moderate to high; explanation depends on teacher audit notes, signals used, validator logs, and the deterministic baseline."),
        ("Safety boundary", "Embedded in deterministic rules; it cannot generate items or bypass prerequisites.", "Enforced through the BKT candidate set and validator; any out-of-bound LLM output is rejected."),
        ("Personalization depth", "Stable but relatively mechanical; personalization is mainly probabilistic and rule based.", "Can combine recent performance, mistakes, learning events, and candidate evidence into a richer bilingual rationale."),
        ("Cost and latency", "Low cost, low latency, and usable without external services.", "Requires provider cost, network latency, JSON validity handling, caching, and rate limits."),
        ("Educational deployment", "Best for MVP, school pilots, and low-risk core learning loops.", "Best for premium AI-assisted recommendations, teacher review, and parent/school-facing explanations."),
        ("Main risks", "Weak KC mapping, static parameters, and limited semantic diagnosis of misconceptions.", "Hallucination, bias, privacy, over-trust, and service unavailability; fallback is essential."),
        ("MAIS-MVP role", "Stable core: review, repair, practice, lesson, and challenge.", "Enhanced layer: BKT guardrails plus LLM rerank, cache/rate-limit, and fallback."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Dimension", "v1: Deterministic adaptive learning without LLM", "v2: Hybrid adaptive learning with LLM"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9.4, bold=True, color="0B2545")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=8.9, bold=(i == 0))
    set_table_borders(table)
    set_table_width(table, [3.0, 6.65, 6.85])
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=8, line=1.15)
    r = p.add_run("Note. v2 refers to guarded LLM reranking, not unrestricted LLM generation of lessons, questions, or answers.")
    set_run_font(r, size=9.5, italic=True, color="555555")


def add_v3_table(doc):
    add_caption(doc, "Table 2\nv3 development path: From hybrid recommendation to auditable instructional decision-making")
    rows = [
        ("v3.0", "Knowledge graph + RAG audit layer", "Organize curriculum goals, item bank, lessons, mistake tags, and prerequisites as retrievable evidence; LLM responses must cite retrieved evidence.", "source coverage, unsupported claim rate, retrieval precision"),
        ("v3.1", "Teacher-in-the-loop console", "Allow teachers to inspect candidates, approve or override recommendations, batch assign repair/review work, and feed decisions back into evaluation.", "teacher override rate, teacher agreement, time-to-assignment"),
        ("v3.2", "Multi-model learner-state fusion", "Keep BKT as the safety floor while adding DKT/DKVMN or item-difficulty signals as offline ranking features.", "AUC, calibration, learning gain, harmful recommendation rate"),
        ("v3.3", "Evaluation and deployment loop", "Use A/B tests, quasi-experiments, and classroom pilots to validate learning gain, retention, mistake repair speed, and fairness.", "nLG, retention, latency, cost per active learner, bias gap"),
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Stage", "Module", "Method", "Primary metrics"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9.1, bold=True, color="0B2545")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=8.7, bold=(i == 0))
    set_table_borders(table)
    set_table_width(table, [1.7, 3.8, 7.15, 3.85])


def build_doc():
    doc = Document()
    set_update_fields_on_open(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    normal.font.size = Pt(11)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.color.rgb = RGBColor.from_string("1F4D78")
    ref_style = styles.add_style("APA Reference", 1)
    ref_style.base_style = styles["Normal"]
    ref_style.font.size = Pt(9.5)
    ref_style.paragraph_format.left_indent = Inches(0.5)
    ref_style.paragraph_format.first_line_indent = Inches(-0.5)
    ref_style.paragraph_format.space_after = Pt(5)
    ref_style.paragraph_format.line_spacing = 1.12

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = header.add_run("MAIS-MVP Adaptive Learning Version Comparison Report")
    set_run_font(hrun, size=9, color="606A78")
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=36, after=10, line=1.2)
    r = p.add_run("MAIS-MVP v1/v2 Adaptive Learning Comparative Report")
    set_run_font(r, size=21, bold=True, color="0B2545")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=26, line=1.25)
    r = p.add_run("From Bayesian Knowledge Tracing to Guarded LLM Reranking")
    set_run_font(r, size=14, italic=True, color="1F4D78")

    meta = [
        ("Project", "MAIS-MVP: A bilingual mathematics learning application for Hong Kong P1-S6 students"),
        ("Report date", "2026-05-15"),
        ("Report type", "Academic technical and educational analysis report"),
        ("Citation style", "APA 7th author-date; references limited to academic journals or conferences"),
        ("Scope", "Theory review plus MAIS-MVP implementation mapping; no feature-code modification"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta:
        row = table.add_row()
        for i, value in enumerate((k, v)):
            cell = row.cells[i]
            set_cell_margins(cell, top=110, bottom=110, start=140, end=140)
            if i == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=10.5, bold=(i == 0), color="0B2545" if i == 0 else "000000")
    set_table_borders(table)
    set_table_width(table, [4.2, 12.2])

    add_note(
        doc,
        "Core judgment",
        "v1 is an interpretable, low-cost, offline-capable deterministic adaptive core; v2 adds a guarded LLM reranking layer on top of v1. The two versions are not substitutes. They form a layered architecture that combines rule-grounded trust with generative explanation and ranking capability.",
    )

    doc.add_page_break()
    add_heading(doc, "Abstract", level=1)
    add_paragraph(doc, "This report compares two adaptive learning versions in MAIS-MVP: v1, adaptive learning without an LLM, and v2, adaptive learning with an LLM. v1 is defined as an adaptive engine based on Bayesian Knowledge Tracing (BKT), knowledge components, prerequisite relations, spaced review, and deterministic policy rules. v2 is defined as a hybrid adaptive engine that introduces a Large Language Model (LLM) on top of the v1 candidate set for constrained reranking, explanation generation, and teacher audit notes.")
    add_paragraph(doc, "From an educational measurement perspective, v1 is strongest in explainability, stability, low cost, and auditable mastery probabilities. Its limitations are relatively static parameters and weaker semantic diagnosis of misconceptions. v2 can integrate learning events, mistakes, candidate evidence, and natural-language feedback into a more teacher-like bilingual rationale. Its risks include hallucination, bias, privacy exposure, cost, latency, and provider dependency. Therefore, v2 should not allow the LLM to directly generate instructional pathways; it should rely on BKT guardrails, candidate validation, fallback, and teacher-in-the-loop governance.")
    add_paragraph(doc, "The report concludes with a v3 roadmap: BKT as the interpretable safety floor, combined with retrieval-augmented generation (RAG), an educational knowledge graph, teacher-in-the-loop controls, offline evaluation, and classroom A/B testing. The goal is to move from recommending a next step to building a verifiable, auditable, and reversible instructional decision-making system.")

    add_heading(doc, "Table of Contents", level=1)
    add_toc_en(doc.add_paragraph())
    add_paragraph(doc, "Note. If page numbers do not appear automatically, update fields in Microsoft Word or LibreOffice. The document uses true Heading styles.", after=10)

    doc.add_page_break()
    add_heading(doc, "Section 1. v1: Adaptive Learning Without LLM", level=1)
    add_heading(doc, "1.1 Conceptual introduction: Deterministic adaptive learning is not non-adaptive", level=2)
    add_paragraph(doc, "In the MAIS-MVP context, v1 adaptive learning without an LLM is not a fixed syllabus or a simple rule-based jump table. It is deterministic adaptive learning centered on a student model. Its core unit is not an entire mathematics course but a knowledge component (KC). A KC may represent the foundation, fluency, or transfer stage of a topic. This approach is consistent with the knowledge-learning-instruction framework, which argues that instructional design should connect knowledge, learning processes, and instructional action rather than merely follow textbook chapters (Koedinger et al., 2012).")
    add_paragraph(doc, "BKT treats whether a student has mastered a KC as a latent state and treats correct answers, incorrect answers, hint use, and mistake records as observable evidence. Every new item response updates the mastery probability and maps that probability to instructional actions such as review, repair, practice, lesson, or challenge. The original knowledge tracing model proposed by Corbett and Anderson (1995) was designed precisely for modeling the acquisition of procedural knowledge; its value is to transform item-level correctness into a time-varying probability of skill mastery.")
    add_paragraph(doc, "In the current MAIS-MVP implementation, v1 can be summarized as follows: construct KCs for P1-S6 mathematics topics; update pMastery, correctStreak, wrongStreak, nextReviewAt, and related states from responses; then choose the next step through deterministic candidate ranking. This version does not call an LLM, does not incur provider cost, and does not stop working when an external model is unavailable.")

    add_heading(doc, "1.2 Literature review: BKT, KT extensions, and ITS effectiveness", level=2)
    add_paragraph(doc, "BKT's classic contribution is to formulate student modeling as an updateable probabilistic inference problem. Early BKT assumes each KC has an initial mastery probability, a learning transition probability, a slip probability, and a guess probability. Later work relaxes the assumption that all students share the same parameters. Pardos and Heffernan (2010) modeled individualization in a Bayesian-network implementation of knowledge tracing. Yudelson et al. (2013) further developed individualized BKT so student and skill differences can enter parameter estimation. Pardos and Heffernan's KT-IDEM model introduced item difficulty into knowledge tracing, addressing the fact that items within the same KC can vary substantially in difficulty (Pardos & Heffernan, 2011).")
    add_paragraph(doc, "Beyond BKT, Performance Factors Analysis represents practice history through cumulative successes and failures rather than latent states (Pavlik et al., 2009). The deep-learning path was advanced by Deep Knowledge Tracing, which uses recurrent neural networks to learn representations from long student-response sequences (Piech et al., 2015). Dynamic Key-Value Memory Networks then stored knowledge concepts in addressable memory structures (Zhang et al., 2017). However, deep KT is not universally superior. Khajah et al. (2016) questioned how deep knowledge tracing needs to be, and Gervet et al. (2020) showed that whether deep learning is best depends on data scale, feature quality, and evaluation design.")
    add_paragraph(doc, "Evidence from intelligent tutoring systems (ITS) also supports the viability of non-LLM adaptive learning. VanLehn (2011) compared human tutoring, ITS, and other tutoring systems, showing that high-quality ITS can approach some human tutoring conditions. Meta-analyses by Kulik and Fletcher (2016) and Ma et al. (2014) also report generally positive effects of ITS on learning outcomes. For MAIS-MVP, these findings support the claim that a system can deliver meaningful adaptive learning without an LLM if the student model, item mapping, and feedback loop are well designed.")
    add_paragraph(doc, "The literature also cautions against treating BKT as a universal solution. Desmarais and Baker (2012) emphasize that learner and skill modeling depend on KC definition, data quality, and evaluation methods. Pelanek (2017) further argues that BKT, logistic models, and other learner models have different use cases. In real educational products, interpretability, parameter stability, and instructional usability may be as important as predictive accuracy.")

    add_heading(doc, "1.3 Mathematical formulation: Core BKT probability updates", level=2)
    add_paragraph(doc, "Let \\(L_t\\) denote whether the student has mastered a KC before the \\(t\\)-th practice opportunity, and let \\(X_t\\) denote whether the \\(t\\)-th response is correct. Classical BKT uses four core parameters: \\(P(L_0)\\), the initial mastery probability; \\(T\\), the probability of learning from not mastered to mastered; \\(S\\), the slip probability, meaning the student knows the skill but answers incorrectly; and \\(G\\), the guess probability, meaning the student does not know the skill but answers correctly.")
    add_formula(doc, r"P(X_t = 1 \mid L_t = 1) = 1 - S, \qquad P(X_t = 1 \mid L_t = 0) = G", "1")
    add_paragraph(doc, "A correct response does not prove mastery because the student may have guessed. An incorrect response does not prove non-mastery because the student may have slipped. BKT therefore uses Bayesian posterior updating.")
    add_formula(doc, r"P(L_t = 1 \mid X_t = 1) = \frac{P(L_t)(1-S)}{P(L_t)(1-S) + (1-P(L_t))G}", "2")
    add_formula(doc, r"P(L_t = 1 \mid X_t = 0) = \frac{P(L_t)S}{P(L_t)S + (1-P(L_t))(1-G)}", "3")
    add_paragraph(doc, "After evidence updating, the model incorporates the probability that learning occurred during the practice opportunity.")
    add_formula(doc, r"P(L_{t+1} = 1) = P(L_t = 1 \mid X_t) + \left[1 - P(L_t = 1 \mid X_t)\right]T", "4")
    add_paragraph(doc, "In MAIS-MVP v1, these probabilities are mapped to instructional actions. If mastery is high and review is due, the policy prioritizes review. If a prerequisite is weak or wrongStreak is high, it selects repair. If evidence is thin for a new skill, it selects lesson. If mastery is below the mastery threshold, it selects practice. If mastery is high and the correct streak is stable, it selects challenge.")
    add_formula(doc, r"a_t = \operatorname{policy}(pMastery_t, prereq_t, dueReview_t, attempt_t, streak_t)", "5")
    add_paragraph(doc, "This policy matters because it transforms prediction into instructional action rather than merely optimizing next-item correctness. It is aligned with mastery learning at scale: the system should continuously identify unmastered skills and route learners back to repairable pathways (Ritter et al., 2016).")

    add_heading(doc, "1.4 Educational application: The value of v1 in mathematics learning", level=2)
    add_paragraph(doc, "For Hong Kong P1-S6 mathematics, v1's first value is stability. Mathematics has clear prerequisite structures: fractions affect ratio, ratio affects functions, and functions support later calculus readiness; algebraic manipulation affects equations, coordinates, and proof. Deterministic BKT can encode these structures through KCs and prerequisites, preventing the system from prematurely assigning challenge work when foundations are weak.")
    add_paragraph(doc, "Its second value is teacher explainability. Teachers and parents can understand why a student is recommended to review a foundation skill: mastery probability is low, wrongStreak is high, a prerequisite threshold is unmet, or spaced review is due. This explanation is easier for schools to trust than the vague claim that an AI model decided the next lesson.")
    add_paragraph(doc, "Its third value is low-cost deployment. v1 does not require an API key, external model, token quota, or network latency. It is therefore suitable for MVP launch, classroom pilots, offline demonstrations, and cost-sensitive schools. Its primary risks are KC mapping quality, item coverage, and parameter calibration. If a KC has too few items, item difficulty is uneven, or evidence is logged incorrectly, the recommendation can still diverge from the true learning state.")

    doc.add_page_break()
    add_heading(doc, "Section 2. v2: Adaptive Learning With LLM", level=1)
    add_heading(doc, "2.1 Conceptual introduction: The LLM is a guarded reranker, not the authority", level=2)
    add_paragraph(doc, "The key idea in v2 is not to let the LLM decide the curriculum directly. Instead, the LLM sits after BKT as a constrained reranker and explanation generator. The flow is not user data to LLM to next step. It is user data to BKT candidate set, then to constrained LLM reranking, then to validation, and finally to fallback if validation fails. This architecture combines BKT's explainability with the LLM's natural-language and synthesis capabilities.")
    add_paragraph(doc, "The LLM layer builds on the Transformer architecture (Vaswani et al., 2017), large-scale language modeling (Brown et al., 2020), instruction tuning and reinforcement learning from human feedback (Ouyang et al., 2022), and chain-of-thought prompting for complex reasoning (Wei et al., 2022). These techniques allow a model to summarize learning evidence, generate bilingual rationales, simulate teacher feedback, and rank candidates using multiple weak signals.")
    add_paragraph(doc, "Educational systems must not confuse linguistic fluency with instructional reliability. LLMs can generate fluent but false explanations, may introduce bias, and can be overconfident when evidence is thin. Reviews by Kasneci et al. (2023) and Yan et al. (2023) both identify strong opportunities for personalized learning support while emphasizing ethics, privacy, accuracy, and the role of teachers. MAIS-MVP v2 should therefore remain LLM-assisted, not LLM-authoritative.")

    add_heading(doc, "2.2 Literature review: LLMs in education, AI tutor evaluation, and risk governance", level=2)
    add_paragraph(doc, "Research on LLMs in education is expanding rapidly. Kasneci et al. (2023) discuss both opportunities and challenges for LLM-based learning support, feedback generation, and teacher assistance. Yan et al. (2023) provide a systematic scoping review of practical and ethical challenges. Tlili et al. (2023) examine ChatGPT as a case study of educational chatbots and their double-edged effects. Broader AI-in-education reviews by Crompton and Burke (2023), Zawacki-Richter et al. (2019), and Chen et al. (2020) also show that technical promise must be grounded in instructional design, educator participation, and empirical evaluation.")
    add_paragraph(doc, "For AI tutor capability, Tack and Piech (2022) proposed the AI Teacher Test, evaluating pedagogical ability in educational dialogue rather than language fluency alone. Macina et al. (2023) introduced MathDial, a dialogue tutoring dataset grounded in mathematical reasoning problems with rich pedagogical properties. These studies imply that the LLM layer in v2 should be evaluated by whether it produces appropriate instructional behavior, not whether it merely sounds like a teacher.")
    add_paragraph(doc, "Risk governance is central. Ji et al. (2023) survey hallucination in natural language generation, and Maynez et al. (2020) analyze faithfulness and factuality in generated text. Weidinger et al. (2022) classify language-model risks including bias, misuse, privacy, and social harms. Baker and Hawn (2021) show that algorithmic bias in education can amplify inequities. For MAIS-MVP, these risks do not imply that LLMs should be avoided; they imply that every LLM output must be constrained by evidence, candidates, and audit mechanisms.")
    add_paragraph(doc, "Retrieval-augmented generation provides a methodological bridge from v2 to v3. Lewis et al. (2020) introduced retrieval-augmented generation for knowledge-intensive tasks, and Asai et al. (2024) extended the idea through self-reflection over retrieval, generation, and critique. Although current v2 focuses on reranking rather than open-domain tutoring, RAG suggests an important design principle: educational LLM reliability should come from retrievable, citeable, and verifiable curriculum evidence.")

    add_heading(doc, "2.3 Mathematical formulation: Constrained reranking, validation, and fallback", level=2)
    add_paragraph(doc, "Let \\(C_t = \\{c_1, c_2, \\ldots, c_k\\}\\) denote the set of instructional candidates generated by BKT at time \\(t\\). Each candidate includes an action, skill, topic, questionIds, guard flags, and evidence. The deterministic v1 recommendation can be written as:")
    add_formula(doc, r"d_t = \operatorname*{arg\,max}_{c_i \in C_t} \operatorname{baseScore}(c_i) \quad \text{subject to } \operatorname{hardGuard}(c_i)", "6")
    add_paragraph(doc, "In v2, the LLM does not search the entire instructional space. It can only score or rerank candidates inside \\(C_t\\). If \\(x_t\\) denotes recent performance, mistakes, learning-event summaries, and candidate features, LLM reranking can be abstracted as:")
    add_formula(doc, r"r_t = \operatorname*{arg\,max}_{c_i \in C_t} s_{\phi}(c_i \mid x_t, C_t, \operatorname{policy})", "7")
    add_paragraph(doc, "The validator then applies hard constraints to the LLM output. Validation fails if the candidateId is not generated by BKT, if questionIds are outside the selected candidate, if due-review or repair-required guards are bypassed, or if required bilingual learnerReason and teacherAuditNote fields are missing.")
    add_formula(doc, r"V(r_t) = \mathbb{1}[\operatorname{candidateId}(r_t) \in C_t] \cdot \mathbb{1}[\operatorname{questionIds}(r_t) \subseteq Q(c_i)] \cdot \mathbb{1}[\operatorname{guardrails}(r_t)=\operatorname{satisfied}]", "8")
    add_paragraph(doc, "The final decision follows a fallback rule. The system adopts the LLM-assisted recommendation only when the provider succeeds, the JSON is valid, and V(r_t)=1. Otherwise it returns the deterministic decision.")
    add_formula(doc, r"D_t = \begin{cases} r_t, & \text{if providerReady} \land \text{validJSON} \land V(r_t)=1 \\ d_t, & \text{otherwise} \end{cases}", "9")
    add_paragraph(doc, "The current MAIS-MVP design can also compute a local AI confidence score for LLM-assisted outputs. This score is not self-reported by the LLM. It is derived from evidence depth, agreement with BKT guardrails, distance from mastery thresholds, and question-ID integrity.")
    add_formula(doc, r"\operatorname{conf}_{AI} = \operatorname{cap}_{evidence}(0.36E + 0.30G_a + 0.22M + 0.12Q)", "10")
    add_paragraph(doc, "This formula expresses the central principle of v2: the LLM may help with ranking and explanation, but confidence must be constrained by locally auditable evidence. This reduces overconfidence and gives teachers access to signalsUsed, teacherAuditNote, and confidenceExplanation.")

    add_heading(doc, "2.4 Educational application: New value and limitations of v2", level=2)
    add_paragraph(doc, "The largest educational gain of v2 is explanation quality and contextual integration. v1 can state that a skill has 42% mastery and requires repair. v2 can explain, using recent mistakes, wrong streaks, item-type distribution, and learning activities, why repairing foundations now is better than moving into challenge work. This matters for student motivation, teacher communication, and parent understanding.")
    add_paragraph(doc, "A second gain is fine-grained ranking among candidates. BKT may generate review, repair, and practice candidates. Within the guardrails, the LLM can consider recent mistake topics, item-type diversity, and learning pace to choose a more suitable question order. But v2 must not allow the LLM to generate new items, reveal answers, or modify mastery probabilities directly; otherwise the system loses verifiability.")
    add_paragraph(doc, "A third gain is teacher auditability. v2 can output teacherAuditNote and signalsUsed, helping teachers decide whether a recommendation is appropriate. Holstein et al. (2019) show the importance of teacher-AI complementarity: educational AI should not displace teachers but help them see learner states, coordinate classroom action, and retain human judgment.")
    add_paragraph(doc, "The limitations are equally clear. v2 introduces provider cost, latency, rate limits, privacy exposure, and service-availability risk. If prompts or validators are weak, the model may produce invalid JSON, hallucinated explanations, or non-compliant recommendations. v2 should therefore be positioned as a premium AI-assisted layer, not a replacement for the v1 core engine.")

    doc.add_page_break()
    add_heading(doc, "Section 3. Integrated Comparison of v1 and v2", level=1)
    add_comparison_table(doc)
    add_heading(doc, "3.1 Core differences", level=2)
    add_paragraph(doc, "The relationship between v1 and v2 is best understood as a foundation layer and an enhanced explanation/ranking layer. v1 reliably transforms learning evidence into instructional actions. v2 makes recommendations more nuanced, more communicable, and more auditable without crossing the instructional boundaries defined by v1. If v2 is incorrectly designed as direct LLM decision-making, the product gives up the safety, explainability, and consistency that matter most in education.")
    add_paragraph(doc, "From a product strategy perspective, v1 should be the default baseline for MVP launch and school pilots. It proves that MAIS-MVP has genuine adaptive learning rather than an AI label. v2 can serve as an advanced version for teachers, parents, and older students by providing more natural explanations, more personalized item sequencing, and richer audit trails.")
    add_paragraph(doc, "From a technical-governance perspective, v2's value is only realized when validator, cache, rate limit, fallback, prompt versioning, and evaluation dashboards are present together. Without these mechanisms, the uncertainty introduced by the LLM can exceed its benefits. Research on LLM risks and AI in education supports this conclusion: generative capability must be embedded in explicit instructional constraints and human oversight (Baker & Hawn, 2021; Ji et al., 2023; Weidinger et al., 2022; Yan et al., 2023).")

    add_heading(doc, "3.2 Recommended positioning for MAIS-MVP", level=2)
    add_numbered(doc, "Name v1 the Deterministic Adaptive Core or BKT Adaptive Core, and make it available to all students as the stable default engine.")
    add_numbered(doc, "Name v2 AI-assisted Adaptive Rerank, and clearly state that the LLM only ranks and explains BKT candidates; it does not directly generate instructional pathways.")
    add_numbered(doc, "Demonstrate fallback in every important demo: when the LLM is disabled, failed, or rejected, the system still returns a safe BKT recommendation.")
    add_numbered(doc, "Show teacherAuditNote, signalsUsed, selectedCandidateId, deterministicCandidateId, and errorKind in teacher-facing views, while keeping student-facing explanations friendly and simple.")
    add_numbered(doc, "For product packaging, treat v1 as the core capability in the base version and v2 as a premium or pilot enhancement that requires cost budgeting and privacy disclosure.")

    doc.add_page_break()
    add_heading(doc, "Section 4. v3 Outlook: A Verifiable, Auditable, and Reversible Instructional Decision System", level=1)
    add_heading(doc, "4.1 Development principles", level=2)
    add_paragraph(doc, "v3 should not simply mean using a stronger LLM. The direction for an educational product should be that every recommendation can answer four questions: What evidence supports it? Why did the model choose it? Can a teacher override it? How does the system fail safely? In other words, v3 should be governance-first adaptive learning, not prompt-first tutoring.")
    add_paragraph(doc, "The first principle is a BKT safety floor. Whether MAIS-MVP later introduces DKT, DKVMN, or additional LLM signals, BKT mastery probability, prerequisite thresholds, spaced review, and repair guards should remain the minimum safety boundary. Deep models may become ranking features, but they should not directly cancel due review or repair-required constraints.")
    add_paragraph(doc, "The second principle is source-grounded generation. v3 should organize Hong Kong curriculum goals, item bank content, lesson blocks, mistake explanations, knowledge graph relations, and teacher notes as retrievable evidence. LLM explanations must cite this evidence rather than inventing curriculum content, new items, or hidden answers. This direction can draw on RAG and Self-RAG methods (Asai et al., 2024; Lewis et al., 2020).")
    add_paragraph(doc, "The third principle is teacher-in-the-loop governance. Teachers should be able to inspect candidates, model evidence, risk labels, and recommendations. They should also be able to approve or override recommendations. The system becomes not only a student-facing practice recommender but also a teacher-facing classroom orchestration and remediation assistant.")

    add_heading(doc, "4.2 Method roadmap", level=2)
    add_v3_table(doc)
    add_paragraph(doc, "Technically, v3 can use a layered pipeline: learning events enter a learning-state store; BKT and optional KT models update skill state; the knowledge graph returns prerequisites and curriculum objectives; the candidate generator produces instructional candidates; the retriever supplies evidence; the LLM generates explanations and ranking only inside the evidence pack; the validator enforces hard constraints; and the teacher console exposes audit and override controls.")
    add_formula(doc, r"state_t = f_{BKT}(attempts_t, mistakes_t, lessonProgress_t, events_t)", "11")
    add_formula(doc, r"evidence_t = \operatorname{retrieve}(KG, curriculum, lessons, questions, state_t)", "12")
    add_formula(doc, r"decision_t = \operatorname{validate}(\operatorname{LLM\_rerank}(C_t, evidence_t), hardGuards) \lor \operatorname{fallback}(d_t)", "13")
    add_paragraph(doc, "Evaluation should use three layers of metrics. The first layer is offline model quality: AUC, calibration, candidate agreement, and guardrail rejection rate. The second layer is product quality: latency, cost per active learner, fallback rate, and teacher override rate. The third layer is educational impact: normalized learning gain, mistake repair speed, retention, and fairness gaps across student groups.")
    add_paragraph(doc, "Deployment should begin with shadow mode, where the LLM produces recommendations but does not affect the student pathway; the team compares it with v1 and teacher judgments. Next comes teacher-approved mode, where teachers confirm recommendations before release. Only after that should MAIS-MVP consider limited automatic mode, and even then it should preserve rollback switches.")

    add_heading(doc, "4.3 Research and product hypotheses for v3", level=2)
    add_bullet(doc, "Hypothesis 1: Deterministic BKT boundaries will substantially reduce the instructional risk of LLM recommendations by converting free generation into constrained explanation and ranking.")
    add_bullet(doc, "Hypothesis 2: RAG will reduce hallucination and curriculum drift only if the retrieval corpus is structured, versioned, and grounded in curriculum and item sources.")
    add_bullet(doc, "Hypothesis 3: Teacher overrides and audit data are not only safety mechanisms but also model-improvement data; they reveal blind spots in the candidate generator.")
    add_bullet(doc, "Hypothesis 4: The commercial value of v3 comes from trustworthy AI personalization, not merely adding a chatbot to the interface.")

    add_heading(doc, "Conclusion", level=1)
    add_paragraph(doc, "MAIS-MVP's v1 and v2 should be designed as a progression. v1 uses BKT, knowledge components, and deterministic policy to establish a trustworthy adaptive-learning foundation. v2 adds LLM reranking on top of that foundation, making recommendation rationales, teacher audit notes, and item sequencing more nuanced. The real product advantage is not whether the system has an LLM; it is whether the LLM is embedded in a verifiable, auditable, and reversible instructional decision system.")
    add_paragraph(doc, "In the short term, v1 should ensure that the MVP is usable, stable, low cost, and explainable. In the medium term, v2 should demonstrate premium AI-assisted adaptive learning. In the long term, v3 should build a trusted platform that combines BKT, RAG, an educational knowledge graph, and teacher-in-the-loop governance. This path is aligned with the research tradition in educational measurement and intelligent tutoring systems while addressing the accuracy, ethics, and governance challenges of LLMs in education.")

    doc.add_page_break()
    add_heading(doc, "References", level=1)
    for ref in references:
        p = doc.add_paragraph(style="APA Reference")
        r = p.add_run(ref)
        set_run_font(r, size=9.2, east_asia="Times New Roman", latin="Times New Roman")

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "Appendix A. Citation Verification Note", level=1)
    add_paragraph(doc, "This report includes 35 formal references. All are academic journal articles, academic conference papers, or conference-proceedings entries. DOI entries were checked against Crossref or publisher metadata; NeurIPS, EDM, and ICLR entries without DOI were checked against official conference pages. In-text citations use APA 7th author-date style, and the References section uses a hanging indent.")
    add_paragraph(doc, "Sources excluded from formal references include blogs, vendor documentation, white papers, arXiv-only or OSF-only preprints, and non-academic webpages. MAIS-MVP implementation details are based on local read-only code inspection and are treated as project evidence, not academic references.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
