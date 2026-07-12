from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-11-president-report.docx")
REPORT_DATE = "2026-07-11"
WINDOW = "2026-07-10 08:00 to 2026-07-11 08:00 Asia/Hong_Kong"
UTC_WINDOW = "2026-07-10 00:00 to 2026-07-11 00:00 UTC"
FORMAT_VERSION = "President Report Business Brief v1"

FONT_LATIN = "Hiragino Sans GB"
FONT_CJK = "Hiragino Sans GB"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().get_or_add_rFonts()
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        size, color = 16, BLUE
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        size, color = 13, BLUE
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        size, color = 12, DEEP_BLUE
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.2, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header)
        set_run_font(r, size=9.1, bold=True)
        set_cell_shading(cell, header_fill)
    row_fills = row_fills or {}
    for row_index, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        fill = row_fills.get(row_index)
        for idx, value in enumerate(row):
            cell = table_row.cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx in (0, 1) and len(str(value)) < 24:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=body_size)
            if fill:
                set_cell_shading(cell, fill)
    set_table_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=10.1)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=10.1)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    hr = header.add_run("MAIS-MVP | Daily coordination report | A01-A25")
    set_run_font(hr, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    fr = footer.add_run(f"{REPORT_DATE} | {FORMAT_VERSION} | Page ")
    set_run_font(fr, size=8.5, color=MUTED)
    add_page_field(footer)


def add_title_block(doc):
    prepared = datetime.now(ZoneInfo("Asia/Hong_Kong")).strftime("%Y-%m-%d %H:%M HKT")
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", REPORT_DATE),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Prepared during run", prepared),
        ("Reporting window", WINDOW),
        ("UTC filter window", UTC_WINDOW),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Format", FORMAT_VERSION),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL, body_size=8.7)


def agent_rows():
    no_log = "窗口内无 fresh session log。"
    statuses = {
        "A01": ("mtime-only", f"{no_log} layout gate 文件有 mtime 变化。", "需 A01 确认 GuestLoginPromptGate 来源与可见行为；无 targeted check。"),
        "A02": ("无 fresh 日志", no_log, "未发现 dashboard/progress owner log；学生 smoke tests 有 mtime 证据但未归属 A02。"),
        "A03": ("无 fresh 日志", no_log, "未发现 roadmap/topic owner log；A25 dirty-map 仍有 A03 bucket。"),
        "A04": ("content QA evidence", "Practice Arena 全量 QA 与 post-fix verdicts 记录 23,753 题三项核心标准绿。", "仍需 A04 owner log 绑定 practice UI/data changes 与 targeted regression。"),
        "A05": ("content review", "A18 lesson visualization realignment review 记录 105 production lessons 结构对齐。", "4 个 analytics source category 仍需 A18/A05 curriculum-fit sanity check。"),
        "A06": ("mtime-only / report evidence", "Visualization components、3D/manim tests 与 premium smoke 文件有窗口 mtime。", "需 A06 isolated visualization validation；A22 top candidate build refresh 仍 blocked。"),
        "A07": ("无 fresh 日志", no_log, "无 AI Tutor provider 行为、speech/voice 或 env readiness 日志。"),
        "A08": ("green fresh gate", "Root type-check 与 analytics fresh gates 通过；types/index.ts 有 mtime 证据。", "Package/worktree gate 仍报 9,907 TS error lines；不要混淆 evidence layers。"),
        "A09": ("无 fresh 日志", no_log, "无 copy/i18n/accessibility owner log；A25 bucket 有少量待确认项。"),
        "A10": ("Active", "生成 2026-07-10 report；本次生成 2026-07-11 DOCX。", "报告-only；无 feature code edit。"),
        "A11": ("mtime-only / QA consumer", "9 个 tests/e2e 文件有 mtime；fresh root gates 绿。", "需 A11 对 affected regression clusters 做正式 split/rerun；Playwright 未在本次报告运行。"),
        "A12": ("mtime-only / backend", "answerMatching、questionStore/userStore、teacher assignment 与 route 文件有 mtime。", "需 A12 session log、backend tests 与 clean-source evidence。"),
        "A13": ("mtime-only / teacher", "TeacherManagement/Reports/ResourceAssessment views 与 saved-reports route 有 mtime。", "需 A13 targeted teacher-console checks；当前仅为 mtime evidence。"),
        "A14": ("无 fresh 日志", no_log, "无 parent console owner evidence；A25 work-order bucket 存在少量项。"),
        "A15": ("green analytics", "Adaptive/analytics deterministic suite 27/27 pass。", "无 fresh adaptive engine owner log；LLM/live provider 未测。"),
        "A16": ("无 fresh 日志", no_log, "无 research/learning-science report in window。"),
        "A17": ("game evidence", "Fishing/Adventure report 覆盖 reward wiring；components/gamification 有 mtime。", "奖励经济 server path 单测绿；E2E 有环境型未绿记录，需 quiet-machine rerun。"),
        "A18": ("Active evidence", "15 个 content-QA files；S18 legacy reports 映射 A18。", "HJB English display、45 warn-level grading notes 与 189 pass-sample manual queue remain。"),
        "A19": ("无 fresh 日志", no_log, "无 env/Vercel/provider readiness 证据；未读取 secret。"),
        "A20": ("game evidence", "Fishing Master + Adventure Island 三项标准审计与修复报告。", "两条 game E2E 在高负载/ENOSPC 下未形成最终绿记录，需安静机器复跑。"),
        "A21": ("mtime-only / content", "7 个 generated-content packs 有 mtime。", "候选内容需 A18 final QA、A23 promotion、A11/A22 release gates。"),
        "A22": ("Evidence active", "Clean-source selection review、top build snapshot、fallback sweep 与 clean runway 在 release-intake 中更新。", "Release source eligible=false；top candidate build refresh required；fallback green reviewed-not-selected。"),
        "A23": ("无 fresh 日志", no_log, "无 candidate-to-live promotion record；内容链不得绕过 A23。"),
        "A24": ("content support", "Question/lesson figure spec 与 illustration-related data evidence present。", "Exact-layer/provenance handoff 未形成 fresh A24 log。"),
        "A25": ("Active", "A25 log 与 1,834 pre-refresh release-intake files；post-window dirty-map 7,635 entries。", "40 canonical authorization rows pending；validation hold 等 owner input。"),
    }
    return [(f"A{i:02d}", *statuses[f"A{i:02d}"]) for i in range(1, 26)]


def build_doc():
    doc = Document()
    setup_document(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "本窗口不是空窗口：有 2 个 fresh session logs（A10、A25）、15 个 content-QA artifacts、3 个 report artifacts，以及 pre-refresh 1,834 个 A25/A22 release-intake files；blockers/decisions/integration 目录无新增。"),
            ("2", "内容质量推进明显：Practice Arena 23,753 题三项核心标准 post-fix 全绿；full bank deterministic audit 23,756/23,756 pass；Mainland PEP 7,200/7,200 pass；California 1,992/1,992 pass（含 11 个 K-G5 decimal wording fix）。"),
            ("3", "Fresh post-window health gates 转绿：root `npm run type-check` pass，`npm run test:analytics` 27/27 pass，no-staged pass，`git diff --check` pass。"),
            ("4", "发布仍未 ready：fresh A25 dirty-map 为 7,635 expanded entries；A22 runway 仍显示 releaseSourceClean=false、readyForMerge=false、pendingCanonicalAuthorizationRows=40，top candidate 需要 fresh build observation，fallback green 仍未被选为 release source。"),
            ("5", "Dr. Peter Hu 今日重点决策：继续禁止 dirty-root deploy；决定 A22 top candidate 修复还是绿色 fallback 评审；处理 validation-hold worktree；要求 mtime-only 变更补 owner logs/checks；确认 HJB English 与 grading warning 的修复优先级。"),
        ],
        [700, 8660],
        body_size=8.4,
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "The window is active: two fresh session logs (A10 and A25), 15 content-QA artifacts, three report artifacts, and 1,834 pre-refresh A25/A22 release-intake files. No new blocker, decision, or integration records were found."),
            ("2", "Content quality moved forward: Practice Arena has 23,753 post-fix questions green on the three core standards; the full-bank deterministic audit is 23,756/23,756 pass; Mainland PEP is 7,200/7,200 pass; California is 1,992/1,992 pass after 11 wording fixes."),
            ("3", "Fresh post-window gates are green: root type-check passed, analytics passed 27/27, no-staged passed, and git diff --check passed."),
            ("4", "Release is still blocked. The fresh A25 dirty map has 7,635 expanded entries; A22 still has releaseSourceClean=false, readyForMerge=false, and 40 pending canonical authorization rows. The top candidate needs a fresh build observation, while the green fallback is reviewed but not selected."),
            ("5", "Dr. Peter Hu should preserve no-dirty-root-deploy, choose the A22 repair-versus-fallback path, resolve the validation-hold worktree, require owner logs/checks for mtime-only changes, and prioritize remaining HJB English/grading-warning cleanup."),
        ],
        [700, 8660],
        body_size=8.4,
    )

    add_heading(doc, "报告窗口摘要（中文）", 1)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ("窗口", WINDOW),
            ("Fresh session logs", "2 个：coordination/session-logs/2026-07-10-A10.md 与 coordination/session-logs/2026-07-10-A25.md。"),
            ("Blockers / decisions / integration", "窗口内新增 0 个。"),
            ("Content QA", "15 个文件：Practice Arena、California focused QA、Mainland PEP/full bank deterministic QA、lesson visualization realignment review。"),
            ("Reports", "3 个：2026-07-10-president-report.docx、build_2026_07_10_president_report.py、2026-07-10-games-three-standard-audit-and-fixes.md。"),
            ("Release intake", "Pre-refresh 窗口扫描 1,834 个；fresh post-window dirty-map 更新后当前 latest map 为 7,635 expanded entries。"),
            ("非 coordination mtime 证据", "57 个：components 18、data 11、tests 9、lib 8、scripts 2、app 2，另有 AGENTS/package/tsconfig/types/docs/local 文件。"),
            ("空窗口 fallback", "不适用；本窗口存在 A10/A25 日志和 A18/S18 content-QA evidence。"),
        ],
        [2450, 6910],
        body_size=8.1,
    )

    add_heading(doc, "整体项目进展（中文）", 1)
    add_bullet(doc, "A18/S18 内容 QA 形成更强 evidence：Practice Arena post-fix 23,753 题核心三项标准全绿，full-bank deterministic audit 23,756/23,756 pass。")
    add_bullet(doc, "A04/A12/A18 相关修复链覆盖 wrong answer keys、acceptedAnswers alias poisoning、mixed-number parsing、HK-EASE prompt spacing、US-CA/US-FL grade-aligned regeneration。")
    add_bullet(doc, "A20/A17 game report 完成 Fishing Master 与 Adventure Island obvious bugs、question solvability、answer correctness 审计与多项 client-side fixes；最终 E2E 绿记录仍需 quiet-machine rerun。")
    add_bullet(doc, "A22/A25 release control plane 更清楚：root green 与 package/worktree red 是不同 evidence layer；clean release source 尚未选定。")
    add_bullet(doc, "A10 daily reporting 格式保持一致；本次仅写 report builder 与 DOCX，未改 feature code。")

    add_heading(doc, "A01-A25 会话状态表（中文）", 1)
    add_table(
        doc,
        ["ID", "状态", "窗口内活动", "当前重点 / 阻塞"],
        agent_rows(),
        [620, 1300, 3650, 3790],
        body_size=6.9,
        row_fills={3: OK_FILL, 4: PENDING_FILL, 9: OK_FILL, 17: OK_FILL, 21: PENDING_FILL, 24: OK_FILL},
    )

    add_heading(doc, "已完成工作（中文）", 1)
    add_bullet(doc, "A10 交付 2026-07-10 president report，并保留 DOCX-only、A01-A25 全表、双语 executive summary 与中文主体格式。")
    add_bullet(doc, "A18/S18 全库 deterministic QA：23,756/23,756 pass，0 failing rows；Mainland PEP 7,200/7,200 pass，review rows 0，manual pass-sample queue 189。")
    add_bullet(doc, "A18/A04/A12 Practice Arena audit 后修复并复核核心三项：grade appropriateness、solvability、answer correctness 均 0 hard failures；live /api/questions -> /api/attempts 27/27 pass。")
    add_bullet(doc, "A18 California focused QA：1,992/1,992 pass；11 个 K-G5 hundredths/tenths wording bugs 已修为 tenths/十分之一，保留 key/options。")
    add_bullet(doc, "A20/A17 Fishing/Adventure report：修复 Fishing timer/aim/reward-cycle 问题与 Adventure trophy duplicate soft-lock、retry spam、duration clamp、question-cycle、restart race。")
    add_bullet(doc, "Fresh A10 checks：root type-check pass、analytics 27/27 pass、no-staged pass、git diff --check pass、A25 dirty-map pass。")

    add_heading(doc, "进行中工作（中文）", 1)
    add_bullet(doc, "A25 closure loop 仍在 validate：40 canonical authorization rows pending，focusBatchPendingRows=0，cleanupAuthorizedRows=0，executableRows=0。")
    add_bullet(doc, "A22 clean-source selection 仍未完成：codex/A22-us-region-alignment 需要 fresh build observation；codex/s22-release-hygiene-2026-06-15 为 reviewed fallback green-not-selected。")
    add_bullet(doc, "A18/A04/A12 content fixes 需要 owning session logs 与 release-slice ownership，不能仅凭 mtime 进入 release。")
    add_bullet(doc, "A11 需要为 tests/e2e mtime cluster、Practice Arena、games、teacher/backend 等变化安排 targeted regression rerun。")
    add_bullet(doc, "A18/A21/A23/A24 内容链需处理 HJB English display corruption、45 warn-level grading notes、US/CA/FL regenerated packs 的 promotion documentation。")

    add_heading(doc, "阻塞项（中文）", 1)
    add_table(
        doc,
        ["Blocker", "Evidence", "Owner / next action"],
        [
            ("Dirty root 不可发布", "Fresh A25 map：1,506 collapsed；7,635 expanded；401 modified；1 deleted；7,233 untracked files。", "A22/A25/A10：只从 clean worktree/clone/reviewed slice/pruned staging 发布。"),
            ("Clean release source 未选", "A22 runway：releaseSourceClean=false；readyForMerge=false；allowed source options 4/4 blocked。", "Owner/A22：选择 top repair 或 fallback review path。"),
            ("Top candidate build stale", "A22 top candidate：codex/A22-us-region-alignment；build status blocked-build-refresh-required；.next 948 MiB；releaseSourceSelected=false。", "A22：fresh build observation 需 owner-approved clean-source workflow。"),
            ("Fallback green 未授权", "Selection review：codex/s22-release-hygiene-2026-06-15 fallback validation/type-check/build passed，但 reviewed-fallback-green-not-selected。", "Owner：明确是否进入 clean-source promotion review。"),
            ("Owner authorization backlog", "pendingCanonicalAuthorizationRows=40；validation hold 等待 owner compose deletion confirmation。", "A25/Owner：继续 guarded focus batches；不要授权 cleanup/destructive action。"),
            ("Package/worktree gates 仍红", "Current cleanup snapshot：root type-check green，但 packageWorktreeTypeCheckErrorLines=9,907。", "A08/A10/A22 + owning agents：分包修复，不用 root green 替代 package proof。"),
            ("Open content/display issues", "HJB English `term-xxxx` display corruption 2,785 prompts；45 warn-level grading robustness notes；189 PEP pass-sample manual queue。", "A18/A21/A23/A24：决定修复/下线/zh fallback 与 manual review order。"),
        ],
        [1900, 4200, 3260],
        body_size=7.3,
        row_fills={0: RISK_FILL, 1: RISK_FILL, 2: RISK_FILL, 4: PENDING_FILL, 5: PENDING_FILL, 6: PENDING_FILL},
    )

    add_heading(doc, "风险（中文）", 1)
    add_bullet(doc, "Release risk：root gates 绿不等于 release ready；dirty root、未选 clean source、package/worktree red 与 pending owner authorization 仍阻断 merge/deploy。")
    add_bullet(doc, "Attribution risk：57 个非 coordination mtime files 涉及 app/components/data/lib/tests/scripts；没有 owner logs 时只能作为 evidence，不能算完成或签核。")
    add_bullet(doc, "Content risk：Practice Arena 和 California 修复质量证据强，但 external quality claim 仍需 pass-sample manual review、HJB English cleanup 与 A23 promotion chain。")
    add_bullet(doc, "Regression risk：Game E2E 在高负载/ENOSPC 下未拿到最终绿色记录；需要 quiet-machine rerun，避免把环境 caveat 当成 full release proof。")
    add_bullet(doc, "Cleanup risk：generated artifacts/residual targets 仍无 cleanup authorization；当前 cleanupAuthorizedRows=0、executableRows=0。")

    add_heading(doc, "测试 / 构建状态（中文）", 1)
    add_table(
        doc,
        ["Check", "Result", "Interpretation"],
        [
            ("Fresh npm run type-check", "PASS。", "Root dirty install 当前 type-green。"),
            ("Fresh npm run test:analytics", "PASS：27 tests passed，0 failed。", "A08/A15 deterministic analytics/adaptive checks 绿色。"),
            ("Fresh A25 dirty-map", "PASS：7,635 expanded entries。", "当前 inventory baseline；不代表 release ready。"),
            ("Fresh no-staged gate", "PASS：stagedEntryCount=0。", "报告运行未 staging。"),
            ("Fresh git diff --check", "PASS：无输出。", "未发现 whitespace diagnostics。"),
            ("Content QA", "PASS claims：23,756/23,756 full bank；7,200/7,200 Mainland PEP；1,992/1,992 California；Practice Arena post-fix 0 hard failures。", "强 QA 证据，但仍需 owner logs/release path。"),
            ("A20 game checks", "Unit/API path green；Playwright under high load/ENOSPC did not yield final clean record。", "需 quiet-machine rerun。"),
            ("A22 top candidate", "blocked-build-refresh-required；releaseSourceSelected=false。", "不能 promotion。"),
            ("npm run build / broad Playwright by A10", "Not run。", "A10 report-only；dirty-root build/E2E 不满足 A22 clean-source release rule。"),
        ],
        [2500, 3250, 3610],
        body_size=7.2,
        row_fills={0: OK_FILL, 1: OK_FILL, 2: OK_FILL, 5: OK_FILL, 6: PENDING_FILL, 7: RISK_FILL},
    )

    add_heading(doc, "变更文件（中文）", 1)
    add_table(
        doc,
        ["Category", "Count / examples", "Meaning"],
        [
            ("窗口扫描总量", "1,911 files（排除 .git、node_modules、.next、.tmp、coverage、Playwright outputs）。", "Mtime 证据；不等于 Git ownership 或 acceptance。"),
            ("coordination", "1,854 files：1,834 release-intake、15 content-QA、3 reports、2 session logs。", "主要 fresh logged/evidence activity。"),
            ("components", "18：visualizations 10、teacher 3、gamification 2、practice 2、layout 1。", "跨 A01/A04/A06/A12/A13/A17/A20；需 owner checks。"),
            ("data", "11：7 generated-content packs、questions、visualizationLabs、mainlandPepJuniorQuestions、hjbQuestionLocalization。", "内容/题库/visualization changes 需 A18/A21/A23/A24 与 A04/A06/A12 签核。"),
            ("tests/lib", "tests 9、lib 8。", "QA evidence 和 code fixes 需 A11/A12/A08 owners 归档。"),
            ("scripts/app/root", "scripts 2、app 2、package/tsconfig/types/AGENTS/docs/local files。", "配置与 route mtime 需 A10/A12/A13/A25 provenance。"),
            ("Fresh dirty owner buckets", "A25 5,702；A10 489；A06 460；A12 178；A24 168；A05 129；A21 76。", "仍需按 owner/slice review；不得整体 merge。"),
        ],
        [1900, 4460, 3000],
        body_size=7.3,
    )

    add_heading(doc, "明日优先事项（中文）", 1)
    add_numbered(doc, "A22/Owner 决定 clean-source 路线：修 top candidate 并取得 fresh build observation，或把绿色 fallback 进入正式 selection review。")
    add_numbered(doc, "A25 继续缩小 40-row canonical authorization backlog，并保持 cleanupAuthorizedRows=0、executableRows=0，直到 owner 明确授权。")
    add_numbered(doc, "A18/A04/A12/A21/A23/A24 把 Practice Arena/California/full-bank fixes 归档为 owner-approved content promotion packet，列明 still-open display/warning items。")
    add_numbered(doc, "A11 在安静机器复跑 game E2E、Practice Arena targeted regressions、teacher/backend affected flows。")
    add_numbered(doc, "A01/A06/A08/A12/A13 对 mtime-only feature/code files 补 session logs、diff scope、checks 与 keep/quarantine/release-slice decision。")

    add_heading(doc, "需要 Owner 决策（中文）", 1)
    add_bullet(doc, "请 Dr. Peter Hu 确认：dirty root 继续只作 inventory/reporting source，不作为 production deploy source。")
    add_bullet(doc, "请决定 A22 路线：继续修 codex/A22-us-region-alignment，或优先评审已 green-for-review-not-selected 的 codex/s22-release-hygiene-2026-06-15。")
    add_bullet(doc, "请确认 validation-hold compose worktree 是否已完成 owner-side deletion；当前 evidence 仍显示 worktree registered、path exists、896 dirty status entries。")
    add_bullet(doc, "请确定 HJB English display corruption、45 warning rows、189 PEP pass-sample review 的优先级：修复、zh fallback、下线或延后。")
    add_bullet(doc, "请要求 mtime-only owner sessions 在下一轮补齐 logs 与 checks，尤其是 A04/A06/A08/A11/A12/A13/A18/A21/A24。")
    add_bullet(doc, "请继续禁止 cleanup/destructive Git/deploy：当前 cleanupAuthorizedRows=0、executableRows=0、deployAuthorized=false。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
