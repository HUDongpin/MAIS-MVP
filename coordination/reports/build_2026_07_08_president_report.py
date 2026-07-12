from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-08-president-report.docx")

FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.4, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, bold=True)
        set_cell_shading(cell, header_fill)
    row_fills = row_fills or {}
    for row_index, row in enumerate(rows):
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        fill = row_fills.get(row_index)
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=body_size)
            if fill:
                set_cell_shading(cells[idx], fill)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DEEP_BLUE),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    prepared = datetime.now(ZoneInfo("Asia/Hong_Kong")).strftime("%Y-%m-%d %H:%M HKT")
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-08"),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Prepared during run", prepared),
        ("Reporting window", "2026-07-07 08:00 to 2026-07-08 08:00 Asia/Hong_Kong"),
        ("UTC filter window", "2026-07-07 00:00 to 2026-07-08 00:00 UTC"),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL, body_size=8.8)


def agent_rows():
    no_fresh = "No fresh in-window session log; represented only by A25 release-intake evidence where applicable."
    statuses = {
        "A01": ("No fresh log", no_fresh, "App shell package remains part of dirty-root closure; no direct A01 changes found in window.", "No A01-owned check."),
        "A02": ("No fresh log", no_fresh, "Dashboard/progress package remains represented through dirty-map owner buckets.", "No A02-owned check."),
        "A03": ("Routed blocker", "No direct A03 log. A25 type-check frontier routes `data/topics.ts` to A03/A18.", "Package/worktree type-check frontier reports 98 error lines on `data/topics.ts`.", "No A03-owned rerun."),
        "A04": ("Routed blocker", "No direct A04 log. A25 type-check frontier routes `data/questions.ts` to A04/A18.", "Package/worktree type-check frontier reports 182 error lines on `data/questions.ts`.", "No A04-owned rerun."),
        "A05": ("No fresh log", no_fresh, "Lesson package remains in dirty-root closure; A22 build blockers include California lesson illustration data parity.", "No A05-owned check."),
        "A06": ("Routed blocker", "No direct A06 log. A25/A22 evidence routes visualization clean-candidate and package type-check blockers to A06.", "Largest package frontier: `TemplatePrimitiveScene.tsx`, `ThreeDGraphCanvas.tsx`, and `VisualizationLabPage.tsx`.", "A22 candidate type-check red; A06 package rerun needed."),
        "A07": ("No fresh log", no_fresh, "AI tutor package present in dirty map; no provider behavior or live key action in window.", "No live provider check."),
        "A08": ("No fresh log", no_fresh, "Shared state/types remain part of package/worktree closure, not a fresh A08 session.", "No A08-owned check."),
        "A09": ("Owner input pending", "No direct A09 log. A25 owner-input packet still includes copy/i18n/accessibility package authorization rows.", "Needs explicit owner final-state decision before closure.", "No A09-owned check."),
        "A10": ("Report owner", "Generated the 2026-07-07 president report inside the window; this run prepares the 2026-07-08 report.", "Reporting only; no feature code edits.", "DOCX build/render QA for current report."),
        "A11": ("Release QA pending", "No direct A11 log. A25 approval packet keeps A11 QA/release-quality as a P2 owner package.", "Regression work waits behind release-source and package closure gates.", "No A11 Playwright run."),
        "A12": ("P2 package pending", "No direct A12 log. A25 approval packet routes 175 backend/API entries to A12.", "Backend/API package requires reviewed final state or blocker.", "No A12 backend rerun."),
        "A13": ("Routed blocker", "No direct A13 log. A25 type-check frontier routes teacher-console files to A13.", "Teacher files account for major package/worktree type-check errors.", "No A13-owned rerun."),
        "A14": ("No fresh log", no_fresh, "Parent console package remains in dirty map with no fresh log.", "No A14-owned check."),
        "A15": ("No fresh log", no_fresh, "Adaptive engine package remains in dirty map; no fresh adaptive behavior work.", "No A15-owned check."),
        "A16": ("No fresh log", no_fresh, "No research implementation or literature work logged in this window.", "No A16 check."),
        "A17": ("Owner input pending", "No direct A17 log. A25 owner-input packet includes gamification/motivation final-state rows.", "Needs explicit owner decision if package is to advance.", "No A17-owned check."),
        "A18": ("Routed content QA", "No direct A18 log. A25 evidence ties A18 to `data/questions.ts` and `data/topics.ts` content correctness risks.", "Content QA signoff remains separate from A21 generation and A04/A03 live integration.", "No A18 QA rerun."),
        "A19": ("No fresh log", no_fresh, "No API environment placement or provider credential work in window.", "No live env check."),
        "A20": ("Routed blocker", "No direct A20 log. A25 type-check frontier routes `MathVirusBlasterGame.tsx` to A20.", "Game package remains dirty and type-check-red in owner package frontier.", "No A20 game rerun."),
        "A21": ("Content package pending", "No direct A21 log. A25 approval packet routes 74 content/RAG entries to A21.", "Candidate/RAG package still needs final-state decision and QA/release gates.", "No package-local A21 check."),
        "A22": ("Evidence consumer", "A25/A22 release-intake files cover clean-source runway, top candidate type-check/build snapshot, root parity, and cleanup gates.", "Top clean candidate is not promotable: type-check failed, build failed, releaseSourceEligibleNow=false.", "Focused smoke 2/2 passed; type-check/build red."),
        "A23": ("Owner input pending", "No direct A23 log. A25 owner-input packet includes integration/promotion rows.", "Needs explicit owner final-state decision.", "No A23-owned check."),
        "A24": ("No fresh log", no_fresh, "Exact-layer package appears only in dirty-map owner bucket; no fresh exact-layer QA or asset integration.", "No A24-owned check."),
        "A25": ("Active", "Primary fresh activity: release-intake gates, A22 root-parity runway, validation-hold scaffold/recorder, owner-response packet, post-response execution plan, and dirty maps.", "Closure remains validate/blocked; no cleanup, deploy, staging, merge, or executable rows.", "Fresh dirty-map pass; no-staged pass; git diff --check pass."),
    }
    return [(f"A{i:02d}", *statuses[f"A{i:02d}"]) for i in range(1, 26)]


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "报告窗口内有 fresh AI session activity，但集中在 A25-owned git hygiene / release intake；A01-A24 大多没有直接 fresh session log。"),
            ("2", "A25/A22 release-intake 证据明显推进：root-parity owner-input runway、candidate-mutation dry-run、validation-hold scaffold/recorder、owner-response packet 与 post-response execution plan 都已 fail-closed 接入。"),
            ("3", "A10 在窗口内生成了 2026-07-07 president report 及 render artifacts；本次运行继续只写 A10-owned report artifacts，不改 feature code。"),
            ("4", "发布仍 blocked：fresh A25 dirty-map 显示 1,464 collapsed / 6,007 expanded dirty entries；A22 top clean candidate type-check failed with 652 TypeScript error lines，build failed with 5 confirmed module blockers。"),
            ("5", "下一步 owner focus：维持 no-dirty-root-deploy，决定 validation-hold / A22 root-parity owner response，并让 A06/A13/A04/A20/A03 等 owning sessions 处理 package/worktree type-check frontier。"),
        ],
        [700, 8660],
        body_size=8.7,
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "Fresh activity exists, but it is concentrated in A25-owned git hygiene and release-intake evidence; most A01-A24 owners have no direct fresh session log."),
            ("2", "A25/A22 release intake advanced materially: root-parity runway, candidate-mutation dry-run, validation-hold scaffolding/recording, owner-response packet, and post-response execution plan are now fail-closed evidence gates."),
            ("3", "A10 generated the July 7 president report and render artifacts within the window. This July 8 run remains report-only and does not edit feature code."),
            ("4", "Release remains blocked. The fresh A25 dirty map reports 1,464 collapsed and 6,007 expanded dirty entries. The A22 top clean candidate failed type-check with 652 TypeScript error lines and failed build on five confirmed module blockers."),
            ("5", "Owner attention should stay on no-dirty-root-deploy discipline, validation-hold / A22 root-parity responses, and routed owner sessions for the package/worktree type-check frontier."),
        ],
        [700, 8660],
        body_size=8.7,
    )

    add_heading(doc, "报告窗口摘要（中文）", 1)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ("窗口", "2026-07-07 08:00 至 2026-07-08 08:00 Asia/Hong_Kong。"),
            ("Fresh session logs", "`coordination/session-logs/2026-07-06-A25.md` 与 `coordination/session-logs/2026-07-07-A25.md` 在窗口内更新；A25 为唯一 fresh active agent log。"),
            ("Blockers / decisions / integration / content QA", "窗口内 `coordination/blockers/`、`coordination/decisions/`、`coordination/integration/`、`coordination/content-qa/` 无新增文件。"),
            ("Reports", "窗口内生成/更新 12 个 report artifacts，主要是 `2026-07-07-president-report.docx`、render PDF/PNG 与 builder cache。"),
            ("Release intake volume", "窗口内约 1,701 个 release-intake 文件更新，反映 A25/A22 evidence control plane 的持续推进。"),
            ("Fresh post-window checks", "本次报告运行后执行 A25 dirty-map refresh、`git diff --check`、`assert-no-staged-changes`；三者均通过，dirty-map 当前为 6,007 expanded entries。"),
        ],
        [2450, 6910],
        body_size=8.4,
    )

    add_heading(doc, "整体项目进展（中文）", 1)
    add_bullet(doc, "A25-owned release intake 已从单纯 dirty inventory 推进到可审查的 owner-response / post-response plan 结构，所有 recorder 与 executor 默认保持 dry-run / fail-closed。")
    add_bullet(doc, "A22-owned release engineering 路线更清楚：`codex/A22-us-region-alignment` 是 top candidate，但 promotion 仍被 type-check 与 build blockers 阻断。")
    add_bullet(doc, "A10-owned reporting 连续性保持：7 月 7 日 report 与 render artifacts 已在窗口内生成，本次 7 月 8 日 report 继续只使用 coordination/report scope。")
    add_bullet(doc, "A01-A24 product/content owners 的真实工作仍需要 isolated owner worktrees；root `main` 应继续作为 inventory/reporting source，而不是 deploy source。")

    add_heading(doc, "A01-A25 会话状态表（中文）", 1)
    add_table(
        doc,
        ["ID", "状态", "窗口内活动", "当前重点 / 阻塞", "Checks"],
        agent_rows(),
        [720, 1100, 3180, 2860, 1500],
        body_size=6.7,
        row_fills={9: OK_FILL, 21: RISK_FILL, 24: OK_FILL},
    )

    add_heading(doc, "已完成工作（中文）", 1)
    add_bullet(doc, "A25 integrated A22 root-parity owner-input landing runway, selectedAction canonical preview, owner-input recorder, candidate-mutation dry-run executor, and validate-frontier capsule/forecast/runway evidence.")
    add_bullet(doc, "A25 integrated validation-hold release confirmation scaffold and recorder; both remain blocked until exact owner confirmation exists.")
    add_bullet(doc, "A25 produced owner-response and post-response execution-plan packets that combine A25 focus rows, A22 root-parity rows, and validation-hold confirmation into one decision path.")
    add_bullet(doc, "A25 refreshes repeatedly passed in-window aggregate/currentness gates in the 475-493 step range, with no staged entries and no whitespace diagnostics reported in the session logs.")
    add_bullet(doc, "A10 generated the prior daily president report (`2026-07-07-president-report.docx`) and render artifacts inside this window.")

    add_heading(doc, "进行中工作（中文）", 1)
    add_bullet(doc, "A25 closure remains in `validate`: completedSteps=2, blockedSteps=2, validateExitReady=false, readyForMerge=false, releaseSourceEligibleNow=false。")
    add_bullet(doc, "A22 clean-source queue remains `waiting-top-candidate-root-parity-owner-input`; top candidate needs owner response plus type/build recovery before promotion.")
    add_bullet(doc, "A25 owner-input action packet remains `waiting-for-owner-input` with 39 pending canonical authorization rows, 1 held Wave 01 row, 123 owner-closure pending items, and 0 executable rows.")
    add_bullet(doc, "A06/A13/A04/A20/A03 are the main routed owner sessions for the package/worktree type-check frontier; A11/A22 final regression/release checks should wait for clean source.")

    add_heading(doc, "阻塞项（中文）", 1)
    add_table(
        doc,
        ["Blocker", "Evidence", "Owner / next action"],
        [
            ("Dirty root not releasable", "Fresh dirty map: 1,464 collapsed status entries, 6,007 expanded entries, 391 modified, 1 deleted, 1,072 untracked status entries.", "A22/A25/A10: keep root inventory-only; release only from clean worktree, clean clone, reviewed clean slice, or owner-approved pruned staging."),
            ("Top clean candidate red", "`codex/A22-us-region-alignment` type-check failed with 652 TS error lines; build failed with 5 webpack module blockers.", "A22 plus routed owners: resolve root-parity/type/build issues before promotion."),
            ("Package/worktree frontier red", "A25 frontier reports 9,907 package/worktree TypeScript error lines; top owners include A06, A13, A18/A04, A20, A03.", "Owning agents fix in isolated worktrees and return targeted checks."),
            ("Owner response pending", "Owner-input action packet reports owner inputs ready=no, 39 pending canonical authorization rows, 123 owner-closure pending items.", "Owner: decide next rows or explicitly defer with blocker text."),
            ("Validation hold", "Validation hold remains waiting for compose-worktree deletion confirmation; cleanupAuthorizedRows=0 and executableRows=0.", "Owner/A25/A22: confirm hold status before any validation-hold release assessment."),
        ],
        [2100, 3700, 3560],
        body_size=7.35,
        row_fills={0: RISK_FILL, 1: RISK_FILL, 2: RISK_FILL, 3: PENDING_FILL, 4: PENDING_FILL},
    )

    add_heading(doc, "风险（中文）", 1)
    add_bullet(doc, "Release risk: dirty-root deployment would mix 6,007 expanded entries and violate A22 release-source policy.")
    add_bullet(doc, "Evidence drift risk: the fresh dirty-map run after the 08:00 window changes the current baseline; downstream packets must be regenerated before release decisions.")
    add_bullet(doc, "Candidate risk: A22 top candidate has focused smoke 2/2 but fails type-check/build; smoke pass alone is not promotion evidence.")
    add_bullet(doc, "Coordination risk: A25 generated a large volume of evidence artifacts; owner decisions should focus on the next executable decision point rather than broad cleanup.")
    add_bullet(doc, "Scope risk: A10 report automation must not repair feature code; red owner-package files need A06/A13/A04/A20/A03 sessions.")

    add_heading(doc, "测试 / 构建状态（中文）", 1)
    add_table(
        doc,
        ["Check", "Result", "Interpretation"],
        [
            ("Fresh `npm run release:dirty-map -- --reason ...`", "Pass; wrote `2026-07-08-A25-dirty-tree-map-20260708T020728Z.md`; 6,007 expanded entries.", "Current post-window A25 inventory for this report."),
            ("Fresh `git diff --check`", "Pass; no diagnostics.", "No whitespace-conflict diagnostics in current dirty tree."),
            ("Fresh `node coordination/release-intake/assert-no-staged-changes.mjs`", "Pass; staged entries 0.", "No staged Git mutation."),
            ("A25 in-window aggregate refreshes", "Passed 475-493 step refresh/currentness sequences in session evidence.", "Release-intake evidence gates were internally consistent during A25 work."),
            ("A25 root type-check status", "Recorded green: `npm run type-check -- --pretty false` status 0, error lines 0.", "Dirty-root TypeScript green is useful but not release-source eligibility."),
            ("A25 package/worktree type-check frontier", "Red: 9,907 error lines; 12 frontier files; 6 critical owner rows.", "Owner worktrees remain blocked."),
            ("A22 top clean candidate type-check", "Failed with 652 TypeScript error lines.", "Candidate cannot promote."),
            ("A22 top clean candidate build", "Failed with 5 confirmed module blockers.", "Candidate cannot release."),
            ("Build / Playwright by report owner", "Not run.", "Dirty-root build/E2E would not satisfy clean-source release policy; A10 report-only scope."),
        ],
        [3350, 2250, 3760],
        body_size=7.25,
        row_fills={0: OK_FILL, 1: OK_FILL, 2: OK_FILL, 3: OK_FILL, 4: OK_FILL, 5: RISK_FILL, 6: RISK_FILL, 7: RISK_FILL, 8: PENDING_FILL},
    )

    add_heading(doc, "变更文件（中文）", 1)
    add_table(
        doc,
        ["Category", "Count / examples", "Meaning"],
        [
            ("Fresh window coordination files", "2 session logs, 12 report artifacts, 1,701 release-intake files; no fresh blockers/decisions/integration/content-QA files.", "Window activity is coordination/release-intake heavy."),
            ("Fresh post-window dirty map", "Owner buckets: A25 4,283; A10 476; A06 460; A12 175; A05 128; A21 74; A22 68; A11 62; A04 59.", "Largest unresolved dirty ownership remains release-intake plus visualization/backend/lesson/content/release QA."),
            ("Slice buckets", "Docs/coordination 4,832; runtime app/API/data/public 688; tests/regression 388; generated/content/RAG 65; release hygiene config 20; manual 13; secret/env quarantine 1.", "Dirty tree should be reviewed in slices, not as one merge."),
            ("Current `git status` top dirs", "components 560; coordination 455; lib 132; app 110; data 81; tests 61; scripts 34.", "Runtime and test surfaces still need owner-package closure."),
            ("Report deliverable", "`coordination/reports/2026-07-08-president-report.docx` plus builder and render artifacts from this run.", "A10-owned final artifact for Dr. Peter Hu."),
        ],
        [1880, 4460, 3020],
        body_size=7.25,
    )

    add_heading(doc, "明日优先事项（中文）", 1)
    add_numbered(doc, "A25 refresh downstream release-intake current gates against the 6,007-entry dirty map before any merge or release planning.")
    add_numbered(doc, "Owner decide validation-hold / A22 root-parity owner response, or explicitly keep the hold with dated rationale.")
    add_numbered(doc, "A22 either fix `codex/A22-us-region-alignment` type/build blockers or propose a different clean-source candidate.")
    add_numbered(doc, "A06 and A13 start the largest type-check frontier fixes in isolated worktrees; A04/A03/A20 follow with smaller routed files.")
    add_numbered(doc, "A11 prepare targeted regression packages only after clean-source/type/build gates become plausible.")

    add_heading(doc, "需要 Owner 决策（中文）", 1)
    add_bullet(doc, "请 Dr. Peter Hu 确认：dirty root 继续只作 inventory/reporting source，不作为 production deploy source。")
    add_bullet(doc, "请决定 validation-hold 所需的 compose-worktree deletion confirmation 是否可以记录；若不能，请说明缺口。")
    add_bullet(doc, "请对 A22 root-parity selectedAction / owner-response packet 给出明确 approve、defer 或 reject。")
    add_bullet(doc, "请决定 39 个 pending canonical authorization rows 的下一批重点；避免一次性批准过大范围。")
    add_bullet(doc, "请指定 A06/A13/A04/A20/A03 的 next owner worktree order，以降低 9,907 package/worktree TS error frontier。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
