from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/dongpinhu/Desktop/MAIS-MVP")
OUT = ROOT / "coordination/reports/2026-05-15-MAIS-v1-v2-adaptive-learning-APA7.docx"


def set_run_font(run, size=None, bold=None, italic=None, color=None, east_asia="Songti SC", latin="Times New Roman"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B8C1CC", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color="606A78")


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Microsoft Word 或 LibreOffice 中自动更新。"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run._r.append(placeholder)
    run._r.append(fld_char3)
    set_run_font(run, size=11)


def set_update_fields_on_open(doc):
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_paragraph(doc, text="", style=None, before=0, after=6, line=1.35, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=11)
    set_paragraph_spacing(p, before=before, after=after, line=line)
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    color = "1F4D78" if level <= 2 else "2C5F78"
    size = {1: 16, 2: 13, 3: 12}.get(level, 11)
    set_run_font(run, size=size, bold=True, color=color, east_asia="Heiti SC", latin="Times New Roman")
    set_paragraph_spacing(p, before=14 if level == 1 else 9, after=6, line=1.18)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.35 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    set_paragraph_spacing(p, after=4, line=1.25)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.38)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_paragraph_spacing(p, after=4, line=1.25)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_formula(doc, formula, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.18)
    set_paragraph_spacing(p, before=5, after=5, line=1.15)
    r = p.add_run(f"\\[\n{formula} \\tag{{{number}}}\n\\]")
    set_run_font(r, size=9.4, east_asia="Courier New", latin="Courier New")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=4, line=1.2)
    r = p.add_run(text)
    set_run_font(r, size=10, bold=True, color="1F4D78", east_asia="Heiti SC")
    return p


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D9E1EA", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F7FA")
    set_cell_margins(cell, top=120, start=180, bottom=120, end=180)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=3, line=1.25)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color="1F4D78", east_asia="Heiti SC")
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.3)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5)
    add_paragraph(doc, "", after=4)


def add_comparison_table(doc):
    add_caption(doc, "表 1\nv1 与 v2 自适应学习机制核心比较")
    rows = [
        ("决策主体", "确定性 BKT/规则引擎，依据掌握概率、先备知识、间隔复习和题目映射。", "BKT 先生成候选；LLM 只能在候选内重排并生成解释。"),
        ("理论基础", "知识追踪、知识组件、ITS 学生模型、掌握学习。", "Transformer/LLM、instruction tuning、教育对话、受限优化与人机互补。"),
        ("可解释性", "强；每个动作可由 mastery、attempt、streak、review due 等变量解释。", "中高；需要 teacher audit note、signals used、validator log 共同解释。"),
        ("安全边界", "边界内生于规则，不能生成题目或越过先备知识。", "边界由 BKT candidate set 与 validator 保证；LLM 越界即拒绝。"),
        ("个性化深度", "稳定但表达较机械，主要是概率和规则层面的个性化。", "可吸收近期表现、错题和学习事件，生成更像教师的 bilingual rationale。"),
        ("成本与延迟", "低成本、低延迟、离线可用。", "有 provider 成本、网络延迟、JSON 格式失败和限流管理。"),
        ("教育部署", "适合 MVP、学校试点、低风险核心闭环。", "适合高级 AI-assisted 版本、教师审计和家长/学校展示。"),
        ("主要风险", "知识点映射不准、参数固定、难以解释深层错误原因。", "幻觉、偏差、隐私、过度信任、服务不可用；需 fallback。"),
        ("MAIS 当前定位", "稳定基座：review/repair/practice/lesson/challenge。", "增强层：BKT guardrails + LLM rerank + cache/rate-limit/fallback。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["维度", "v1：无 LLM 的确定性自适应", "v2：有 LLM 的混合自适应"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9.5, bold=True, color="0B2545", east_asia="Heiti SC")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=9.2, bold=(i == 0))
    set_table_borders(table)
    set_table_width(table, [3.0, 6.65, 6.85])
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=8, line=1.15)
    r = p.add_run("注：本表的 v2 指 guarded LLM reranking，不指 LLM 直接生成课程、题目或答案。")
    set_run_font(r, size=9.5, italic=True, color="555555")


def add_v3_table(doc):
    add_caption(doc, "表 2\nv3 开发路线：从混合推荐到可审计教学决策系统")
    rows = [
        ("v3.0", "知识图谱 + RAG 审计层", "把香港课程、题库、lesson、错题标签和先备关系组织为可检索证据；LLM 回答必须引用检索片段。", "source coverage、unsupported claim rate、retrieval precision。"),
        ("v3.1", "教师在环控制台", "教师可查看候选、覆盖/确认推荐、批量布置 repair/review，并把反馈写回策略评估。", "teacher override rate、teacher agreement、time-to-assignment。"),
        ("v3.2", "多模型学习状态融合", "保留 BKT 为安全阈值，引入 DKT/DKVMN 或 item difficulty 特征作离线排序信号。", "AUC、calibration、learning gain、harmful recommendation rate。"),
        ("v3.3", "评估与部署闭环", "用 A/B 测试、准实验或课堂试点验证 learning gain、留存、错题修复速度和公平性。", "nLG、retention、latency、cost per active learner、bias gap。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["阶段", "模块", "方法", "主要指标"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9.2, bold=True, color="0B2545", east_asia="Heiti SC")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=8.9, bold=(i == 0))
    set_table_borders(table)
    set_table_width(table, [1.9, 3.7, 7.0, 3.9])


references = [
    ("Asai, A., Wu, Z., Wang, Y., Sil, A., & Hajishirzi, H. (2024). Self-RAG: Learning to retrieve, generate, and critique through self-reflection. In International Conference on Learning Representations. https://proceedings.iclr.cc/paper_files/paper/2024/hash/25f7be9694d7b32d5cc670927b8091e1-Abstract-Conference.html"),
    ("Baker, R. S., & Hawn, A. (2021). Algorithmic bias in education. International Journal of Artificial Intelligence in Education, 32(4), 1052-1092. https://doi.org/10.1007/s40593-021-00285-9"),
    ("Brown, T. B., Mann, B., Ryder, N., Subbiah, M., Kaplan, J., Dhariwal, P., Neelakantan, A., Shyam, P., Sastry, G., Askell, A., Agarwal, S., Herbert-Voss, A., Krueger, G., Henighan, T., Child, R., Ramesh, A., Ziegler, D. M., Wu, J., Winter, C., ... Amodei, D. (2020). Language models are few-shot learners. In Advances in Neural Information Processing Systems, 33, 1877-1901. https://proceedings.neurips.cc/paper/2020/hash/1457c0d6bfcb4967418bfb8ac142f64a-Abstract.html"),
    ("Chen, L., Chen, P., & Lin, Z. (2020). Artificial intelligence in education: A review. IEEE Access, 8, 75264-75278. https://doi.org/10.1109/ACCESS.2020.2988510"),
    ("Corbett, A. T., & Anderson, J. R. (1995). Knowledge tracing: Modeling the acquisition of procedural knowledge. User Modelling and User-Adapted Interaction, 4(4), 253-278. https://doi.org/10.1007/BF01099821"),
    ("Crompton, H., & Burke, D. (2023). Artificial intelligence in higher education: The state of the field. International Journal of Educational Technology in Higher Education, 20, Article 22. https://doi.org/10.1186/s41239-023-00392-8"),
    ("Desmarais, M. C., & Baker, R. S. J. d. (2012). A review of recent advances in learner and skill modeling in intelligent learning environments. User Modeling and User-Adapted Interaction, 22(1-2), 9-38. https://doi.org/10.1007/s11257-011-9106-8"),
    ("Gervet, T., Koedinger, K., Schneider, J., & Mitchell, T. (2020). When is deep learning the best approach to knowledge tracing? Journal of Educational Data Mining, 12(3), 31-54. https://doi.org/10.5281/zenodo.4143614"),
    ("Holstein, K., McLaren, B. M., & Aleven, V. (2019). Co-designing a real-time classroom orchestration tool to support teacher-AI complementarity. Journal of Learning Analytics, 6(2), 27-52. https://doi.org/10.18608/jla.2019.62.3"),
    ("Ji, Z., Lee, N., Frieske, R., Yu, T., Su, D., Xu, Y., Ishii, E., Bang, Y., Madotto, A., & Fung, P. (2023). Survey of hallucination in natural language generation. ACM Computing Surveys, 55(12), Article 248. https://doi.org/10.1145/3571730"),
    ("Kasneci, E., Sessler, K., Kuchemann, S., Bannert, M., Dementieva, D., Fischer, F., Gasser, U., Groh, G., Gunnemann, S., Hullermeier, E., Krusche, S., Kutyniok, G., Michaeli, T., Nerdel, C., Pfeffer, J., Poquet, O., Sailer, M., Schmidt, A., Seidel, T., ... Kasneci, G. (2023). ChatGPT for good? On opportunities and challenges of large language models for education. Learning and Individual Differences, 103, Article 102274. https://doi.org/10.1016/j.lindif.2023.102274"),
    ("Khajah, M. M., Lindsey, R. V., & Mozer, M. C. (2016). How deep is knowledge tracing? In Proceedings of the 9th International Conference on Educational Data Mining (pp. 94-101). https://educationaldatamining.org/EDM2016/proceedings/paper_122.pdf"),
    ("Koedinger, K. R., Corbett, A. T., & Perfetti, C. (2012). The knowledge-learning-instruction framework: Bridging the science-practice chasm to enhance robust student learning. Cognitive Science, 36(5), 757-798. https://doi.org/10.1111/j.1551-6709.2012.01245.x"),
    ("Kulik, J. A., & Fletcher, J. D. (2016). Effectiveness of intelligent tutoring systems: A meta-analytic review. Review of Educational Research, 86(1), 42-78. https://doi.org/10.3102/0034654315581420"),
    ("Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Kuttler, H., Lewis, M., Yih, W.-t., Rocktaschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. In Advances in Neural Information Processing Systems, 33, 9459-9474. https://proceedings.neurips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html"),
    ("Ma, W., Adesope, O. O., Nesbit, J. C., & Liu, Q. (2014). Intelligent tutoring systems and learning outcomes: A meta-analysis. Journal of Educational Psychology, 106(4), 901-918. https://doi.org/10.1037/a0037123"),
    ("Macina, J., Daheim, N., Chowdhury, S., Sinha, T., Kapur, M., Gurevych, I., Sachan, M., & Nouri, E. (2023). MathDial: A dialogue tutoring dataset with rich pedagogical properties grounded in math reasoning problems. In Findings of the Association for Computational Linguistics: EMNLP 2023 (pp. 5602-5621). https://doi.org/10.18653/v1/2023.findings-emnlp.372"),
    ("Maynez, J., Narayan, S., Bohnet, B., & McDonald, R. (2020). On faithfulness and factuality in abstractive summarization. In Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics (pp. 1906-1919). https://doi.org/10.18653/v1/2020.acl-main.173"),
    ("Ouyang, L., Wu, J., Jiang, X., Almeida, D., Wainwright, C. L., Mishkin, P., Zhang, C., Agarwal, S., Slama, K., Ray, A., Schulman, J., Hilton, J., Kelton, F., Miller, L., Simens, M., Askell, A., Welinder, P., Christiano, P., Leike, J., & Lowe, R. (2022). Training language models to follow instructions with human feedback. In Advances in Neural Information Processing Systems, 35, 27730-27744. https://proceedings.neurips.cc/paper_files/paper/2022/hash/b1efde53be364a73914f58805a001731-Abstract-Conference.html"),
    ("Pardos, Z. A., & Heffernan, N. T. (2010). Modeling individualization in a Bayesian networks implementation of knowledge tracing. In P. De Bra, A. Kobsa, & D. Chin (Eds.), User Modeling, Adaptation, and Personalization (pp. 255-266). Springer. https://doi.org/10.1007/978-3-642-13470-8_24"),
    ("Pardos, Z. A., & Heffernan, N. T. (2011). KT-IDEM: Introducing item difficulty to the knowledge tracing model. In G. Biswas, S. Bull, J. Kay, & A. Mitrovic (Eds.), Artificial Intelligence in Education (pp. 243-254). Springer. https://doi.org/10.1007/978-3-642-22362-4_21"),
    ("Pavlik, P. I., Jr., Cen, H., & Koedinger, K. R. (2009). Performance factors analysis: A new alternative to knowledge tracing. In V. Dimitrova, R. Mizoguchi, B. du Boulay, & A. Graesser (Eds.), Artificial Intelligence in Education (pp. 531-538). IOS Press. https://doi.org/10.3233/978-1-60750-028-5-531"),
    ("Pelanek, R. (2017). Bayesian knowledge tracing, logistic models, and beyond: An overview of learner modeling techniques. User Modeling and User-Adapted Interaction, 27(3-5), 313-350. https://doi.org/10.1007/s11257-017-9193-2"),
    ("Piech, C., Spencer, J., Huang, J., Ganguli, S., Sahami, M., Guibas, L. J., & Sohl-Dickstein, J. (2015). Deep knowledge tracing. In Advances in Neural Information Processing Systems, 28. https://papers.nips.cc/paper/5654-deep-knowledge-tracing"),
    ("Ritter, S., Yudelson, M., Fancsali, S. E., & Berman, S. R. (2016). How mastery learning works at scale. In Proceedings of the Third (2016) ACM Conference on Learning @ Scale (pp. 71-79). https://doi.org/10.1145/2876034.2876039"),
    ("Tack, A., & Piech, C. (2022). The AI teacher test: Measuring the pedagogical ability of Blender and GPT-3 in educational dialogues. In Proceedings of the 15th International Conference on Educational Data Mining. https://educationaldatamining.org/edm2022/proceedings/2022.EDM-short-papers.54/index.html"),
    ("Tlili, A., Shehata, B., Adarkwah, M. A., Bozkurt, A., Hickey, D. T., Huang, R., & Agyemang, B. (2023). What if the devil is my guardian angel: ChatGPT as a case study of using chatbots in education. Smart Learning Environments, 10, Article 15. https://doi.org/10.1186/s40561-023-00237-x"),
    ("VanLehn, K. (2011). The relative effectiveness of human tutoring, intelligent tutoring systems, and other tutoring systems. Educational Psychologist, 46(4), 197-221. https://doi.org/10.1080/00461520.2011.611369"),
    ("Vaswani, A., Shazeer, N., Parmar, N., Uszkoreit, J., Jones, L., Gomez, A. N., Kaiser, L., & Polosukhin, I. (2017). Attention is all you need. In Advances in Neural Information Processing Systems, 30. https://papers.nips.cc/paper/7181-attention-is-all-you-need"),
    ("Weidinger, L., Uesato, J., Rauh, M., Griffin, C., Huang, P.-S., Mellor, J., Glaese, A., Cheng, M., Balle, B., Kasirzadeh, A., Biles, C., Brown, S., Kenton, Z., Hawkins, W., Stepleton, T., Birhane, A., Haas, J., Rimell, L., Hendricks, L. A., ... Gabriel, I. (2022). Taxonomy of risks posed by language models. In 2022 ACM Conference on Fairness, Accountability, and Transparency (pp. 214-229). https://doi.org/10.1145/3531146.3533088"),
    ("Wei, J., Wang, X., Schuurmans, D., Bosma, M., Xia, F., Chi, E. H., Le, Q. V., Zhou, D., & Narang, S. (2022). Chain-of-thought prompting elicits reasoning in large language models. In Advances in Neural Information Processing Systems, 35, 24824-24837. https://proceedings.neurips.cc/paper_files/paper/2022/hash/9d5609613524ecf4f15af0f7b31abca4-Abstract-Conference.html"),
    ("Yan, L., Sha, L., Zhao, L., Li, Y., Martinez-Maldonado, R., Chen, G., Li, X., Jin, Y., & Gašević, D. (2023). Practical and ethical challenges of large language models in education: A systematic scoping review. British Journal of Educational Technology, 55(1), 90-112. https://doi.org/10.1111/bjet.13370"),
    ("Yudelson, M. V., Koedinger, K. R., & Gordon, G. J. (2013). Individualized Bayesian knowledge tracing models. In H. C. Lane, K. Yacef, J. Mostow, & P. Pavlik (Eds.), Artificial Intelligence in Education (pp. 171-180). Springer. https://doi.org/10.1007/978-3-642-39112-5_18"),
    ("Zawacki-Richter, O., Marin, V. I., Bond, M., & Gouverneur, F. (2019). Systematic review of research on artificial intelligence applications in higher education: Where are the educators? International Journal of Educational Technology in Higher Education, 16, Article 39. https://doi.org/10.1186/s41239-019-0171-0"),
    ("Zhang, J., Shi, X., King, I., & Yeung, D.-Y. (2017). Dynamic key-value memory networks for knowledge tracing. In Proceedings of the 26th International Conference on World Wide Web (pp. 765-774). https://doi.org/10.1145/3038912.3052580"),
]


def build_doc():
    doc = Document()
    set_update_fields_on_open(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    normal.font.size = Pt(11)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.color.rgb = RGBColor.from_string("1F4D78")

    ref_style = styles.add_style("APA Reference", 1)
    ref_style.base_style = styles["Normal"]
    ref_style.font.size = Pt(9.5)
    ref_style.paragraph_format.left_indent = Inches(0.5)
    ref_style.paragraph_format.first_line_indent = Inches(-0.5)
    ref_style.paragraph_format.space_after = Pt(5)
    ref_style.paragraph_format.line_spacing = 1.12

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = header.add_run("MAIS-MVP 自适应学习版本比较报告")
    set_run_font(hrun, size=9, color="606A78", east_asia="Heiti SC")
    add_page_number(section.footer.paragraphs[0])

    # Title page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=36, after=10, line=1.2)
    r = p.add_run("MAIS-MVP v1/v2 自适应学习对比报告")
    set_run_font(r, size=22, bold=True, color="0B2545", east_asia="Heiti SC")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=26, line=1.25)
    r = p.add_run("从 Bayesian Knowledge Tracing 到受防护的 LLM Reranking")
    set_run_font(r, size=14, italic=True, color="1F4D78", east_asia="Heiti SC")

    meta = [
        ("项目", "MAIS-MVP：面向香港 P1-S6 学生的双语数学学习应用"),
        ("报告日期", "2026-05-15"),
        ("报告类型", "学术式技术与教育分析报告"),
        ("引用格式", "APA 7th author-date；参考文献限学术期刊或学术会议论文"),
        ("范围", "理论综述 + MAIS-MVP 当前实现映射；不修改功能代码"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta:
        row = table.add_row()
        for i, value in enumerate((k, v)):
            cell = row.cells[i]
            set_cell_margins(cell, top=110, bottom=110, start=140, end=140)
            if i == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=10.5, bold=(i == 0), color="0B2545" if i == 0 else "000000", east_asia="Heiti SC" if i == 0 else "Songti SC")
    set_table_borders(table)
    set_table_width(table, [4.2, 12.2])

    add_note(
        doc,
        "核心判断",
        "v1 是可解释、低成本、可离线运行的确定性自适应基座；v2 是在 v1 之上加入受防护 LLM reranking 的增强层。二者不是替代关系，而是“规则可信度 + 生成式解释与排序能力”的分层架构。"
    )

    doc.add_page_break()
    add_heading(doc, "摘要", level=1)
    add_paragraph(doc, "本报告比较 MAIS-MVP 的两个自适应学习版本：v1「无 LLM 自适应」与 v2「有 LLM 自适应」。本文将 v1 定义为以 Bayesian Knowledge Tracing（BKT）、知识组件、先备关系、间隔复习和确定性规则为核心的 adaptive engine；将 v2 定义为在 v1 候选集之上引入 Large Language Model（LLM）进行受限重排、解释生成和教师审计备注的 hybrid adaptive engine。")
    add_paragraph(doc, "从教育测量角度看，v1 的优势在于可解释性、稳定性、低成本和可审计的掌握概率；其弱点是参数较静态、对错误原因的语义解释较弱。v2 的优势在于能够整合学习事件、错题、候选依据和自然语言反馈，形成更贴近教师语言的双语解释；其风险在于幻觉、偏差、隐私、成本和服务依赖。因此，v2 不应让 LLM 直接生成教学路径，而应采用 BKT guardrails、candidate validation、fallback 和 teacher-in-the-loop。")
    add_paragraph(doc, "本文最后提出 v3 路线：以 BKT 为可解释安全底座，引入检索增强生成（RAG）、教育知识图谱、教师在环控制台、离线评估与课堂 A/B 测试，把系统从“推荐一个下一步”推进到“可验证、可审计、可回滚的教学决策系统”。")

    add_heading(doc, "目录", level=1)
    add_toc(doc.add_paragraph())
    add_paragraph(doc, "注：如目录页码未自动显示，请在 Word 或 LibreOffice 中更新域。正文结构已使用真实 Heading 样式。", style=None, after=10)

    doc.add_page_break()
    add_heading(doc, "Section 1. v1：无 LLM 自适应", level=1)
    add_heading(doc, "1.1 概念介绍：确定性自适应不是“没有智能”", level=2)
    add_paragraph(doc, "在 MAIS-MVP 语境中，v1「无 LLM 自适应」并不等于固定课程表或简单规则跳转，而是以学生模型为中心的确定性自适应。其核心对象不是整门数学课，而是知识组件（knowledge components, KCs）：一个 KC 可以对应某个课题的基础、熟练或迁移应用阶段。知识组件思想与 knowledge-learning-instruction framework 一致，即教学设计应把知识、学习过程和教学行动连接起来，而不是只按教材章节线性推进（Koedinger et al., 2012）。")
    add_paragraph(doc, "BKT 将“学生是否掌握某个知识点”视为不可直接观察的 latent state，将答对、答错、hint 使用、错题记录等视为 observable evidence。系统每得到一次新证据，就更新掌握概率，并把结果映射为 review、repair、practice、lesson 或 challenge 等教学动作。Corbett 和 Anderson（1995）提出的知识追踪模型正是面向程序性知识习得的学生建模方法，其价值在于把学生状态从“做对几题”转化为“某知识点掌握概率随时间变化”。")
    add_paragraph(doc, "在 MAIS-MVP 当前实现中，v1 的教育逻辑可概括为：先建立 P1-S6 数学课题的知识组件；再依据作答结果更新 pMastery、correctStreak、wrongStreak、nextReviewAt 等状态；最后由确定性候选排序选择下一步。该版本不调用 LLM，不产生外部 API 成本，也不会因 provider 不可用而中断学习路径。")

    add_heading(doc, "1.2 文献综述：BKT、KT 扩展与 ITS 效果证据", level=2)
    add_paragraph(doc, "BKT 的经典贡献在于把学生建模问题转化为可更新的概率推断。早期 BKT 假设每个 KC 有初始掌握概率、学习转移概率、slip 与 guess 参数。后续研究不断放宽“全体学生共享同一参数”的限制。Pardos 和 Heffernan（2010）提出在贝叶斯网络实现中建模个体化差异；Yudelson 等（2013）进一步发展 individualized BKT，使学生和技能的差异能够进入参数估计。Pardos 和 Heffernan（2011）的 KT-IDEM 则把题目难度引入知识追踪，使“同一知识点下不同题目难度”不再被模型忽略。")
    add_paragraph(doc, "BKT 之外，Performance Factors Analysis 用累计成功与失败次数替代隐状态建模，强调可解释的练习历史特征（Pavlik et al., 2009）。深度学习路线则由 Deep Knowledge Tracing 推动，它用循环神经网络从长序列中学习学生状态表示（Piech et al., 2015），后来 Dynamic Key-Value Memory Networks 将知识概念存入可寻址记忆结构（Zhang et al., 2017）。但是，深度 KT 并不总是优于经典模型。Khajah 等（2016）质疑“更深”是否真正带来稳定提升；Gervet 等（2020）也指出深度学习是否最佳取决于数据规模、特征质量和评估设置。")
    add_paragraph(doc, "从教育系统效果看，智能导师系统（ITS）已有相当多实证证据。VanLehn（2011）比较了人类导师、ITS 与其他辅导系统的效果，指出高质量 ITS 能接近某些人类辅导情境。Kulik 和 Fletcher（2016）以及 Ma 等（2014）的 meta-analysis 也表明 ITS 对学习结果通常有正向影响。对 MAIS-MVP 来说，这些证据支持 v1 的基本方向：即使没有 LLM，只要学生模型、题目映射和反馈闭环设计得当，系统仍可形成有效的个性化学习支持。")
    add_paragraph(doc, "不过，文献也提醒我们不要把 BKT 视为万能。Desmarais 和 Baker（2012）综述 learner and skill modeling 时强调，学生模型的有效性依赖知识点划分、数据质量和评估方式。Pelanek（2017）进一步指出，BKT、logistic models 与其他学生模型各有适用边界；在真实教育产品中，模型可解释性、参数估计稳定性、教学可用性往往与预测精度同样重要。")

    add_heading(doc, "1.3 数学公式：BKT 的核心概率更新", level=2)
    add_paragraph(doc, "令 \\(L_t\\) 表示学生在第 \\(t\\) 次练习前是否已经掌握某知识组件，\\(X_t\\) 表示第 \\(t\\) 次作答是否正确。经典 BKT 设定四个核心参数：\\(P(L_0)\\) 为初始掌握概率，\\(T\\) 为从未掌握到掌握的学习转移概率，\\(S\\) 为 slip，即已掌握但答错的概率，\\(G\\) 为 guess，即未掌握但猜对的概率。")
    add_formula(doc, r"P(X_t = 1 \mid L_t = 1) = 1 - S, \qquad P(X_t = 1 \mid L_t = 0) = G", "1")
    add_paragraph(doc, "当学生答对时，系统不能直接断定其已经掌握，因为可能是 guess；当学生答错时，也不能直接断定其没有掌握，因为可能是 slip。BKT 因此使用贝叶斯后验更新。")
    add_formula(doc, r"P(L_t = 1 \mid X_t = 1) = \frac{P(L_t)(1-S)}{P(L_t)(1-S) + (1-P(L_t))G}", "2")
    add_formula(doc, r"P(L_t = 1 \mid X_t = 0) = \frac{P(L_t)S}{P(L_t)S + (1-P(L_t))(1-G)}", "3")
    add_paragraph(doc, "完成证据更新后，模型再考虑学生在此次练习后发生学习的可能性。")
    add_formula(doc, r"P(L_{t+1} = 1) = P(L_t = 1 \mid X_t) + \left[1 - P(L_t = 1 \mid X_t)\right]T", "4")
    add_paragraph(doc, "在 MAIS-MVP v1 中，可将这些概率量映射为教学动作。若知识点已经掌握且 nextReviewAt 到期，则优先 review；若先备知识低于阈值或 wrongStreak 较高，则 repair；若新知识证据不足，则 lesson；若 pMastery 低于掌握阈值，则 practice；若 pMastery 高且正确连击稳定，则 challenge。")
    add_formula(doc, r"a_t = \operatorname{policy}(pMastery_t, prereq_t, dueReview_t, attempt_t, streak_t)", "5")
    add_paragraph(doc, "这一 policy 的重要性在于，它不是单纯追求下一题预测正确率，而是把预测结果转化为教学上可解释的行动。该设计与 mastery learning at scale 的思路相容：系统应不断识别尚未掌握的技能，并把学习者引回可修复的路径（Ritter et al., 2016）。")

    add_heading(doc, "1.4 教育中的应用：v1 在数学学习中的价值", level=2)
    add_paragraph(doc, "对香港 P1-S6 数学学习而言，v1 的第一项价值是稳定。小学至高中数学有清晰的先备结构：分数影响比例，比例影响函数，函数影响微积分准备；代数变形影响方程、坐标与证明。确定性 BKT 可把这种结构写入知识组件和 prerequisites 中，让系统避免在基础薄弱时过早推送挑战题。")
    add_paragraph(doc, "第二项价值是教师可解释性。教师和家长可以理解“为什么学生被推荐回顾分数基础”：因为 mastery probability 低、wrongStreak 高、先备知识未达阈值或间隔复习到期。这种解释比“AI 认为你该学这个”更容易被学校接受，也更适合低年级学生。")
    add_paragraph(doc, "第三项价值是低成本部署。v1 不依赖 API key、外部模型、token quota 或网络延迟，因此适合 MVP、课堂试点、离线演示和成本敏感学校。它的主要风险在于知识点映射、题库覆盖和参数校准：如果某知识组件题目太少、题目难度不均或 wrong/correct evidence 被错误记录，推荐结果仍会偏离真实学习状态。")

    doc.add_page_break()
    add_heading(doc, "Section 2. v2：有 LLM 自适应", level=1)
    add_heading(doc, "2.1 概念介绍：LLM 不取代 BKT，而是受防护的重排器", level=2)
    add_paragraph(doc, "v2 的关键不是让 LLM 直接决定课程，而是把 LLM 放在 BKT 之后，作为受约束的 reranker 和解释生成器。换言之，v2 的决策流程不是“用户数据 -> LLM -> 下一步”，而是“用户数据 -> BKT 候选集 -> LLM 受限重排 -> validator -> fallback”。这个架构把 BKT 的可解释性和 LLM 的自然语言能力结合起来。")
    add_paragraph(doc, "LLM 的能力基础来自 Transformer 架构（Vaswani et al., 2017）、大规模语言建模（Brown et al., 2020）、instruction tuning 与人类反馈强化学习（Ouyang et al., 2022），以及 chain-of-thought prompting 对复杂推理的促进（Wei et al., 2022）。这些技术使模型能够总结学习证据、生成双语理由、模拟教师反馈，并根据多个弱信号做排序。")
    add_paragraph(doc, "但是，教育场景不能把这种语言能力误认为教学可靠性。LLM 可以生成流畅但错误的解释，可能引入偏差，也可能在没有证据时过度自信。Yan 等（2023）和 Kasneci 等（2023）的综述都指出，LLM 在教育中有个性化反馈和学习支持潜力，同时必须处理伦理、隐私、准确性和教师角色问题。MAIS-MVP v2 因此应保持“LLM-assisted”而非“LLM-authoritative”。")

    add_heading(doc, "2.2 文献综述：LLM 教育应用、AI tutor 评估与风险治理", level=2)
    add_paragraph(doc, "LLM 在教育中的研究正快速扩展。Kasneci 等（2023）从机会与挑战两方面讨论了大型语言模型用于学习支持、反馈生成和教师辅助的可能性；Yan 等（2023）系统梳理了教育中使用 LLM 的实践与伦理挑战；Tlili 等（2023）以 ChatGPT 为案例讨论聊天机器人在教育中的双刃剑属性。Crompton 和 Burke（2023）、Zawacki-Richter 等（2019）以及 Chen 等（2020）也从更广义的人工智能教育应用中指出，技术价值必须落到教学设计、教师参与和评估证据上。")
    add_paragraph(doc, "对 AI tutor 的具体能力评估，Tack 和 Piech（2022）提出 AI Teacher Test，用教育对话评估模型的教学能力，而不是只看语言流畅度。Macina 等（2023）的 MathDial 数据集进一步把数学推理问题与富教学属性对话结合，为数学 tutor 的评估提供了更贴近课堂的基准。这些研究说明，v2 的 LLM 层应被评估为“是否产生合适教学行为”，而不是“是否说得像老师”。")
    add_paragraph(doc, "风险治理方面，Ji 等（2023）综述了自然语言生成中的 hallucination，Maynez 等（2020）讨论摘要生成中的 factuality 与 faithfulness；Weidinger 等（2022）提出语言模型风险分类，覆盖偏差、误用、隐私和社会影响。Baker 和 Hawn（2021）则从教育算法偏差角度提醒，学生模型和推荐系统可能放大群体差异。对 MAIS-MVP 而言，这些风险并不意味着不能使用 LLM，而是要求 LLM 输出必须被证据、候选集和审计机制约束。")
    add_paragraph(doc, "检索增强生成（RAG）为 v2 到 v3 的演进提供了方法基础。Lewis 等（2020）提出 retrieval-augmented generation，使生成过程能够结合外部知识；Asai 等（2024）的 Self-RAG 进一步强调模型在检索、生成和自我批判之间的循环。虽然当前 v2 重点是 rerank 而不是开放问答，但 RAG 思想提示我们：教育 LLM 的可靠性应来自可检索、可引用、可验证的课程证据。")

    add_heading(doc, "2.3 数学公式：受限重排、校验与回退", level=2)
    add_paragraph(doc, "令 \\(C_t = \\{c_1, c_2, \\ldots, c_k\\}\\) 表示 BKT 在时刻 \\(t\\) 生成的候选教学动作集合，每个候选包含 action、skill、topic、questionIds、guard flags 和 evidence。v1 的确定性推荐可表示为：")
    add_formula(doc, r"d_t = \operatorname*{arg\,max}_{c_i \in C_t} \operatorname{baseScore}(c_i) \quad \text{subject to } \operatorname{hardGuard}(c_i)", "6")
    add_paragraph(doc, "v2 中，LLM 不拥有全集搜索空间，只能对 \\(C_t\\) 内的候选打分或重排。若 \\(x_t\\) 表示学习者近期表现、错题、学习事件摘要和候选特征，LLM reranking 可抽象为：")
    add_formula(doc, r"r_t = \operatorname*{arg\,max}_{c_i \in C_t} s_{\phi}(c_i \mid x_t, C_t, \operatorname{policy})", "7")
    add_paragraph(doc, "随后 validator 对 LLM 输出进行硬约束校验。若输出 candidateId 不存在、questionIds 不属于该 candidate、due-review 被绕过、repair-required 被跳过、缺少中英文 learnerReason 或 teacherAuditNote，则校验失败。")
    add_formula(doc, r"V(r_t) = \mathbb{1}[\operatorname{candidateId}(r_t) \in C_t] \cdot \mathbb{1}[\operatorname{questionIds}(r_t) \subseteq Q(c_i)] \cdot \mathbb{1}[\operatorname{guardrails}(r_t)=\operatorname{satisfied}]", "8")
    add_paragraph(doc, "最终决策遵循回退规则：只有 provider 成功、JSON 格式有效且 V(r_t)=1 时采用 LLM-assisted recommendation；否则返回 deterministic decision。")
    add_formula(doc, r"D_t = \begin{cases} r_t, & \text{if providerReady} \land \text{validJSON} \land V(r_t)=1 \\ d_t, & \text{otherwise} \end{cases}", "9")
    add_paragraph(doc, "MAIS-MVP 当前实现还可为 LLM-assisted 输出计算本地信心分。该分数不是 LLM 自报信心，而是由证据深度、与 BKT guardrail 的一致性、mastery threshold 距离和 question-id 完整性组成。")
    add_formula(doc, r"\operatorname{conf}_{AI} = \operatorname{cap}_{evidence}(0.36E + 0.30G_a + 0.22M + 0.12Q)", "10")
    add_paragraph(doc, "这个公式体现了 v2 的原则：LLM 可以帮助排序和解释，但信心必须由本地可审计证据来约束。它降低了 LLM 过度自信的风险，也让教师能看到 signalsUsed、teacherAuditNote 和 confidenceExplanation。")

    add_heading(doc, "2.4 教育中的应用：v2 的新增价值与限制", level=2)
    add_paragraph(doc, "v2 对数学学习的最大增益是解释质量和情境整合。v1 可以说“该知识点掌握概率为 42%，需要 repair”；v2 可以进一步用学生近期错题、连续答错、题型分布和学习活动解释“为什么现在先修补基础比进入挑战题更好”。这对学生动机、教师沟通和家长理解都有价值。")
    add_paragraph(doc, "第二个增益是候选间细排序。例如 BKT 可能生成 review、repair、practice 多个候选；LLM 可在 guardrails 内结合近期错题主题、题型多样性和学习节奏，选择更符合学生当前状态的一组题目顺序。但 v2 必须禁止 LLM 生成新题、泄露答案或直接修改掌握概率，否则系统会失去可验证性。")
    add_paragraph(doc, "第三个增益是教师审计。v2 输出 teacherAuditNote 和 signalsUsed，可帮助教师判断系统推荐是否合理。Holstein 等（2019）的 teacher-AI complementarity 研究提示，教育 AI 不应把教师排除在外，而应帮助教师看见学生状态、协调课堂行动并保留人类判断。")
    add_paragraph(doc, "v2 的限制同样清楚：它带来 provider 成本、延迟、rate limit、隐私和服务可用性问题；若 prompt 或校验器设计薄弱，模型可能产生格式错误、幻觉或不合规建议。因此 v2 的产品定位应是高级 AI-assisted 层，而不是替代 v1 的核心引擎。")

    doc.add_page_break()
    add_heading(doc, "Section 3. v1 与 v2 的综合比较", level=1)
    add_comparison_table(doc)
    add_heading(doc, "3.1 核心差异", level=2)
    add_paragraph(doc, "v1 和 v2 的关系更像“基础设施层”和“增强解释层”。v1 负责把学习证据稳定地转换成教学动作；v2 负责在不越过教学边界的前提下，让推荐更细腻、更可沟通、更适合教师审计。若把 v2 错误设计为 LLM 直接决策，就会把教育产品中最重要的安全性、可解释性和一致性让渡给不可完全预测的生成模型。")
    add_paragraph(doc, "从产品路线看，v1 应作为 MVP 和学校试点的默认基线。它证明 MAIS-MVP 已具备真实 adaptive learning，而不是营销上的 AI label。v2 则可作为更高阶版本：面向教师、家长和高年级学生提供更自然的解释、更个性化的题组排序和更完整的 audit trail。")
    add_paragraph(doc, "从技术治理看，v2 的收益只有在 validator、cache、rate limit、fallback、prompt versioning 和 evaluation dashboard 同时存在时才成立。没有这些机制，LLM 带来的不确定性会超过收益。已有 LLM 风险研究与教育 AI 综述都支持这一判断：教育场景需要把生成能力嵌入明确的教学约束和人类监督中（Baker & Hawn, 2021; Ji et al., 2023; Weidinger et al., 2022; Yan et al., 2023）。")

    add_heading(doc, "3.2 对 MAIS-MVP 的建议定位", level=2)
    add_numbered(doc, "v1 命名为 Deterministic Adaptive Core 或 BKT Adaptive Core，作为所有学生默认可用的稳定引擎。")
    add_numbered(doc, "v2 命名为 AI-assisted Adaptive Rerank，明确说明 LLM 只在 BKT 候选内排序和解释，不直接生成课程路径。")
    add_numbered(doc, "所有演示中应展示 fallback：当 LLM disabled、failed 或 rejected 时，系统仍能返回安全的 BKT 推荐。")
    add_numbered(doc, "教师端应显示 teacherAuditNote、signalsUsed、selectedCandidateId、deterministicCandidateId 和 errorKind，而学生端只显示友好的学习理由。")
    add_numbered(doc, "商业分层上，v1 可作为基础版核心能力；v2 可作为高级版或学校试点增强能力，但必须配套成本预算和隐私说明。")

    doc.add_page_break()
    add_heading(doc, "Section 4. v3 版本展望：可验证、可审计、可回滚的教学决策系统", level=1)
    add_heading(doc, "4.1 开发原则", level=2)
    add_paragraph(doc, "v3 不应简单升级为“更强 LLM”。教育产品的方向应是让每个推荐都能够回答四个问题：证据是什么，模型为何这样判断，教师是否能覆盖，系统失败时如何安全回退。换言之，v3 是 governance-first adaptive learning，而不是 prompt-first tutoring。")
    add_paragraph(doc, "第一个原则是 BKT safety floor。无论是否引入 DKT、DKVMN 或 LLM，BKT 的掌握概率、先备阈值、间隔复习和 repair guard 都应保留为最低安全边界。深度模型可以作为排序特征，但不应直接取消 due review 或 repair-required。")
    add_paragraph(doc, "第二个原则是 source-grounded generation。v3 应将香港课程目标、题库、lesson blocks、错题解释、知识图谱和教师备注构建为检索语料。LLM 生成解释时必须引用检索证据，不能凭空扩展教材、创造题目或泄露答案。这一路线可借鉴 RAG 与 Self-RAG 的检索-生成-批判框架（Asai et al., 2024; Lewis et al., 2020）。")
    add_paragraph(doc, "第三个原则是 teacher-in-the-loop。教师应能看到候选、模型证据、风险标签和推荐理由，也能覆盖或确认推荐。这样，系统不再只是给学生推题，而是成为教师的课堂编排与补救教学助手。")

    add_heading(doc, "4.2 方法路线", level=2)
    add_v3_table(doc)
    add_paragraph(doc, "技术上，v3 可采用分层 pipeline：学习事件进入 learning state store；BKT 和可选 KT 模型更新 skill state；知识图谱返回先备关系和课程目标；candidate generator 生成教学候选；retriever 提供证据；LLM 仅在 evidence pack 内生成解释与排序；validator 执行 hard constraints；teacher console 暴露审计与覆盖入口。")
    add_formula(doc, r"state_t = f_{BKT}(attempts_t, mistakes_t, lessonProgress_t, events_t)", "11")
    add_formula(doc, r"evidence_t = \operatorname{retrieve}(KG, curriculum, lessons, questions, state_t)", "12")
    add_formula(doc, r"decision_t = \operatorname{validate}(\operatorname{LLM\_rerank}(C_t, evidence_t), hardGuards) \lor \operatorname{fallback}(d_t)", "13")
    add_paragraph(doc, "评估上，v3 需要三层指标。第一层是离线模型指标：AUC、calibration、candidate agreement、guardrail rejection rate。第二层是产品指标：latency、cost per active learner、fallback rate、teacher override rate。第三层是教育指标：normalized learning gain、错题修复速度、知识点留存、不同学生群体之间的公平性差距。")
    add_paragraph(doc, "部署上，建议先做 shadow mode：LLM 产生建议但不影响学生路径，只比较它与 v1 的一致性和教师评价；随后做 teacher-approved mode，让教师确认后推送；最后才进入 limited automatic mode，并保留可回滚开关。")

    add_heading(doc, "4.3 v3 的研究与产品假设", level=2)
    add_bullet(doc, "假设一：BKT 的确定性边界能显著降低 LLM 推荐的教育风险，使 LLM 输出从“自由生成”变成“受限解释与排序”。")
    add_bullet(doc, "假设二：RAG 能降低幻觉和教材偏离，但前提是检索语料本身结构化、版本化并包含课程与题目来源。")
    add_bullet(doc, "假设三：教师覆盖和审计数据不仅是安全机制，也是模型改进数据；teacher override 可用于发现 candidate generator 的盲点。")
    add_bullet(doc, "假设四：v3 的商业价值来自“可信 AI 个性化”，而不是单纯在界面上加入聊天机器人。")

    add_heading(doc, "结论", level=1)
    add_paragraph(doc, "MAIS-MVP 的 v1 和 v2 应被设计成递进关系。v1 用 BKT、知识组件和确定性策略建立可信自适应基础；v2 在这个基础上加入 LLM reranking，使推荐理由、教师审计和题组排序更加细腻。真正的产品优势不是“有无 LLM”，而是是否能把 LLM 放入可验证、可审计、可回滚的教育决策系统中。")
    add_paragraph(doc, "因此，短期内应以 v1 确保 MVP 可用、稳定、低成本和可解释；中期用 v2 展示 AI-assisted adaptive learning 的高级能力；长期以 v3 构建 BKT + RAG + 教育知识图谱 + teacher-in-the-loop 的可信学习平台。这个路线既符合教育测量与 ITS 研究传统，也回应了 LLM 教育应用中的准确性、伦理和治理挑战。")

    doc.add_page_break()
    add_heading(doc, "References", level=1)
    for ref in references:
        p = doc.add_paragraph(style="APA Reference")
        r = p.add_run(ref)
        set_run_font(r, size=9.2, east_asia="Times New Roman", latin="Times New Roman")

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "Appendix A. 引用核验说明", level=1)
    add_paragraph(doc, "本报告正式参考文献共 35 条，均为学术期刊、学术会议论文或会议论文集条目。DOI 条目通过 Crossref/出版方元数据核对；无 DOI 的 NeurIPS、EDM、ICLR 条目使用会议官方永久页面。正文采用 APA 7th author-date 引用，References 列表采用悬挂缩进。")
    add_paragraph(doc, "未纳入正式参考文献的来源包括博客、厂商文档、白皮书、纯 arXiv/OSF 预印本和非学术网页。MAIS-MVP 当前实现信息来自本地代码只读检查，用于项目映射，不作为学术文献引用。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_doc()
    print(path)
