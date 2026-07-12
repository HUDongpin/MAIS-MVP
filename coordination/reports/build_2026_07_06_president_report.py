from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-06-president-report.docx")

FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.4, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, bold=True)
        set_cell_shading(cell, header_fill)
    row_fills = row_fills or {}
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        fill = row_fills.get(row_index)
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=body_size)
            if fill:
                set_cell_shading(cells[idx], fill)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DEEP_BLUE),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-06"),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Prepared during run", "2026-07-06 09:59 HKT"),
        ("Reporting window", "2026-07-05 08:00 to 2026-07-06 08:00 Asia/Hong_Kong"),
        ("UTC filter window", "2026-07-05 00:00 to 2026-07-06 00:00 UTC"),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL, body_size=8.8)


def agent_rows():
    inactive = "No fresh in-window session log; covered by A25 owner package/work-order evidence where applicable."
    rows = []
    statuses = {
        "A01": ("No fresh log", inactive, "Pending owner-package closure item; no direct work recorded.", "None in window."),
        "A02": ("No fresh log", inactive, "Pending owner-package closure item; no direct work recorded.", "None in window."),
        "A03": ("No fresh log", inactive, "A25 Wave 03/04 readiness records roadmap/package blockers; no direct log.", "No fresh A03 checks."),
        "A04": ("No fresh log", inactive, "A25 Wave 04 records practice package not ready; no direct log.", "No fresh A04 checks."),
        "A05": ("No fresh log", inactive, "A25 Wave 04 records lesson package not ready; no direct log.", "No fresh A05 checks."),
        "A06": ("No fresh log", inactive, "A25 focus batch includes one A06 owner authorization row; visualization package remains blocked.", "Wave 05 readiness: not ready."),
        "A07": ("No fresh log", inactive, "AI tutor package remains in A25 work-order queue; no direct log.", "Wave 05 readiness: not ready."),
        "A08": ("No fresh log", inactive, "A08/A12 shared-contract package remains blocked in A25 Wave 02.", "Wave 02 readiness: not ready."),
        "A09": ("No fresh log", inactive, "Copy/i18n package in A25 queue; no direct log.", "Wave 05 readiness: not ready."),
        "A10": ("Report owner", "No fresh in-window session log; current run creates this DOCX only.", "A10 supports A25/A22 release-intake synthesis and no feature code changes.", "Report-only; no app checks required."),
        "A11": ("No fresh log", inactive, "A25 focus batch includes one A11 owner authorization row; regression package remains blocked.", "Wave 05 readiness: not ready."),
        "A12": ("No fresh log", inactive, "A25 focus batch includes one A12 owner authorization row; shared/backend package remains blocked.", "Wave 02 readiness: not ready."),
        "A13": ("No fresh log", inactive, "Teacher console package queued in A25 work orders; no direct log.", "Wave 05 readiness: not ready."),
        "A14": ("No fresh log", inactive, "Parent console package queued in A25 work orders; no direct log.", "No fresh A14 checks."),
        "A15": ("No fresh log", inactive, "Adaptive engine package queued in A25 work orders; no direct log.", "Wave 03 readiness includes A02/A15 not ready."),
        "A16": ("Completed", "Created and repaired the TRUST/COMPASS cognitive-affective algorithm DOCX outside the repo.", "Only traceability log changed in repo; no feature code.", "DOCX render QA completed; 15 page PNGs inspected."),
        "A17": ("No fresh log", inactive, "Gamification package queued in A25 work orders; no direct log.", "Wave 05 readiness: not ready."),
        "A18": ("No fresh log", inactive, "Content QA package remains blocked through A25 Wave 04 content evidence.", "Wave 04 readiness: not ready."),
        "A19": ("No fresh log", inactive, "No fresh API/env work in the reporting window.", "No live provider checks."),
        "A20": ("No fresh log", inactive, "Game package remains blocked through A25 Wave 05 game/motivation readiness.", "Wave 05 readiness: not ready."),
        "A21": ("No fresh log", inactive, "Content/RAG package remains blocked through A25 Wave 04 content evidence.", "Wave 04 readiness: not ready."),
        "A22": ("Evidence only", "No fresh A22 session log; A22 blocker evidence was generated under A25 release-intake.", "Release-source clean remains blocked; generated residual cleanup not executable.", "Fresh release-source clean gate failed."),
        "A23": ("No fresh log", inactive, "Integration package queued; no direct promotion work in window.", "No fresh A23 checks."),
        "A24": ("No fresh log", inactive, "Exact-layer package queued through A25 content evidence; no direct log.", "No fresh A24 checks."),
        "A25": ("Active", "Built release-intake guardrails, owner-input packets, focus-batch gates, closure loop snapshots, and work-order bundles.", "Closure loop remains validate/blocked; no cleanup, deploy, staging, or executable rows authorized.", "In-window aggregate reached 107/107; fresh currentness later failed because root moved."),
    }
    for i in range(1, 26):
        agent = f"A{i:02d}"
        status, activity, next_step, checks = statuses[agent]
        rows.append((agent, status, activity, next_step, checks))
    return rows


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "报告窗口内有 fresh AI session activity，主要集中在 A25-owned git hygiene/release intake；A16 完成了 repo 外 TRUST/COMPASS cognitive-affective 算法 DOCX，并在 repo 内留下 traceability log。"),
            ("2", "A25 将 dirty-root cleanup/release-intake 证据推进到 owner-input action packet、focus batch、work-order bundle、Wave 01-06 readiness 与 A22 blocker evidence，但所有结果仍是 evidence-only。"),
            ("3", "当前发布状态仍 blocked：root `main` 不能作为发布源；A22 release-source clean gate 失败；A25 strict worktree lifecycle 未清；cleanupAuthorizedRows=0，executableRows=0，deploy authorized=false。"),
            ("4", "Fresh reporting checks显示 root 仍在移动：`git status --short` 约 1456 collapsed entries；`release:dirty-map --assert-current` 失败，最新 map saved 5056 expanded entries、current 5057；A25 aggregate currentness 仅 40/109 通过。"),
            ("5", "Dr. Peter Hu 今天最需要处理：确认 compose worktree exact deletion 是否完成；决定五个 focus-batch owner-package final-state rows；决定六个 A25 artifact-clean rows 是否给 separate execution instruction；保持 dirty-root 不发布。"),
        ],
        [700, 8660],
        body_size=8.8,
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "Fresh work exists in the reporting window, led by A25-owned git hygiene and release-intake evidence. A16 also completed a research document outside the repo and logged traceability."),
            ("2", "A25 produced current owner-input, focus-batch, work-order, and Wave readiness evidence, but it remains non-executable and does not authorize cleanup, deploy, staging, reset, or deletion."),
            ("3", "Release readiness remains blocked: root `main` is not a release source, A22 release-source clean fails, strict worktree lifecycle is not clean, and cleanup/executable/deploy rows remain zero."),
            ("4", "Fresh report checks found the dirty tree moving after the latest A25 map: dirty-map currentness failed by one expanded entry, while no staged changes were present."),
            ("5", "Owner decisions should focus on the compose-worktree validation hold, five focus-batch final-state rows, six A25 artifact-clean execution-instruction rows, and clean-source release discipline."),
        ],
        [700, 8660],
        body_size=8.8,
    )

    add_heading(doc, "报告窗口摘要（中文）", 1)
    add_table(
        doc,
        ["项目", "结果"],
        [
            ("窗口", "2026-07-05 08:00 至 2026-07-06 08:00 Asia/Hong_Kong。"),
            ("Fresh session logs", "2 个 in-window log：`coordination/session-logs/2026-07-05-A16.md` 与 `coordination/session-logs/2026-07-05-A25.md`。`2026-07-06-A25.md` 的 mtime 为 09:41 HKT，作为窗口后最新状态参考，不作为窗口内成果。"),
            ("Blockers / decisions", "`coordination/blockers/` 与 `coordination/decisions/` 在窗口内没有新增文件。主要 blocker 来自 A25/A22 release-intake evidence。"),
            ("Changed files", "窗口内 380 个 filesystem-changed files：约 379 个在 `coordination/`，另有 `.DS_Store`。其中 `coordination/release-intake/` 363 个，`coordination/reports/` 13 个，`coordination/session-logs/` 2 个。"),
            ("Fresh checks by report owner", "No staged changes gate passed；release-source clean failed；dirty-map currentness failed because root changed after latest map；aggregate currentness failed because multiple packets became stale against the moving root."),
        ],
        [2400, 6960],
        body_size=8.6,
    )

    add_heading(doc, "整体项目进展（中文）", 1)
    add_bullet(doc, "A25-owned release intake made progress in governance depth, not release readiness: owner-input surfaces, work orders, and focus-batch gates are more explicit and auditable.")
    add_bullet(doc, "A16-owned research work delivered a cognitive-affective adaptive-learning algorithm specification outside the app; this is research evidence, not a product integration.")
    add_bullet(doc, "A22-owned release engineering remains blocked because root `main` has thousands of dirty expanded entries and must stay inventory-only until packages are reviewed and closed.")
    add_bullet(doc, "A01-A24 feature/content owners mostly had no fresh direct logs in the window; their current state is represented through A25 work-order and readiness matrices.")

    add_heading(doc, "A01-A25 会话状态表（中文）", 1)
    add_table(
        doc,
        ["ID", "状态", "窗口内活动", "当前重点 / 阻塞", "Checks"],
        agent_rows(),
        [720, 1040, 3040, 3050, 1510],
        body_size=6.8,
        row_fills={15: OK_FILL, 21: RISK_FILL, 24: PENDING_FILL},
    )

    add_heading(doc, "已完成工作（中文）", 1)
    add_bullet(doc, "A16 completed `/Users/dongpinhu/Desktop/TRUST Algorithm/MAIS_algorithm_both cog & aff.docx`, then repaired 41 displayed formulas into Word Office Math and verified a 15-page render.")
    add_bullet(doc, "A25 generated/updated owner-input action packet, owner-closure work-order bundle, focus-batch acceptance/recording gates, closure loop state, and A22 release-source/generated-artifact blocker evidence.")
    add_bullet(doc, "A25 recorded 62 final-state ledger rows as approved/finalized evidence, but each remains non-executable without later scoped action and gates.")
    add_bullet(doc, "A25 kept all destructive operations out of scope: no staging, commit, branch, merge, push, reset, clean, file deletion, worktree removal, cleanup apply, or deploy recorded in the window evidence.")

    add_heading(doc, "进行中工作（中文）", 1)
    add_bullet(doc, "A25 closure loop remains at validate: 2/5 steps completed, 2 steps blocked, pending canonical authorization rows remain 64, and five focus-batch rows are waiting for owner authorization.")
    add_bullet(doc, "A22 release-source cleanliness remains in progress and blocked until root dirty entries are drained through reviewed packages or clean slices.")
    add_bullet(doc, "Wave 02-05 owner packages remain in progress/not ready, with shared-contract, shell/dashboard/roadmap, practice/lesson/content, visualization/AI/runtime packages carrying check failures or stale evidence.")

    add_heading(doc, "阻塞项（中文）", 1)
    add_table(
        doc,
        ["Blocker", "Evidence", "Owner / next action"],
        [
            ("Release source dirty", "Fresh `assert-release-source-clean.mjs` failed; root is dirty and cannot be release source.", "A22/A25/A10: use clean worktree, clean clone, reviewed slice, or owner-approved pruned staging only."),
            ("Dirty-map/currentness stale", "Fresh `release:dirty-map --assert-current` failed: saved 5056 expanded entries, current 5057.", "A25: refresh dirty-map and downstream evidence before any release/preflight decision."),
            ("Validation hold", "A25 snapshot says waiting-for-owner-compose-deletion-confirmation for compose worktree.", "Owner: confirm exact deletion in compose worktree is complete before deferred aggregate validation."),
            ("No executable cleanup rows", "cleanupAuthorizedRows=0 and executableRows=0 in owner packets and work-order bundle.", "Owner/A25/A22: record exact approvals and separate execution instructions before any cleanup."),
            ("Generated artifact residuals", "A22 evidence: `.tmp` and `.next` residual targets, ~48.7 GB, not executable.", "A22/Owner: decide exact cleanup authorization path; preserve evidence first."),
        ],
        [2150, 3600, 3610],
        body_size=7.4,
        row_fills={0: RISK_FILL, 1: RISK_FILL, 2: PENDING_FILL, 3: PENDING_FILL, 4: PENDING_FILL},
    )

    add_heading(doc, "风险（中文）", 1)
    add_bullet(doc, "Release risk: using dirty root for deploy would violate A22 policy and mix unreviewed owner packages.")
    add_bullet(doc, "Coordination risk: many A25 packets are current only relative to their source map; once root changes, downstream packet currentness must be refreshed before acting.")
    add_bullet(doc, "Scope risk: Wave 02-05 readiness failures span shared contracts, visualizations, teacher review, content/question data, games, and tests; broad fixes should remain owner-routed, not patched from the report run.")
    add_bullet(doc, "Generated artifact risk: `.tmp` and `.next` are large residual targets, but cleanup is not authorized and may contain evidence or active output if not checked immediately before apply.")

    add_heading(doc, "测试 / 构建状态（中文）", 1)
    add_table(
        doc,
        ["Check", "Result", "Interpretation"],
        [
            ("`node coordination/release-intake/assert-no-staged-changes.mjs`", "Pass: staged entries 0", "No staged Git mutation during report scan."),
            ("`npm run release:dirty-map -- --assert-current --max-age-minutes 120`", "Fail: saved 5056 vs current 5057 expanded entries", "A25 map is stale because root changed; refresh before release planning."),
            ("`node coordination/release-intake/assert-release-source-clean.mjs`", "Fail: release source dirty", "A22 release source is blocked; dirty root remains inventory-only."),
            ("`node coordination/release-intake/assert-dirty-worktree-remediation-current.mjs --json`", "Fail: 40/109 current checks passed", "Many A25/A22 packets stale against newest dirty map; this is a coordination-refresh blocker."),
            ("`npm run type-check`, `npm run build`, Playwright", "Not run by report owner", "Not practical as fresh release evidence from dirty root; A25 wave evidence already records package-level red checks."),
        ],
        [3300, 2200, 3860],
        body_size=7.8,
        row_fills={0: OK_FILL, 1: RISK_FILL, 2: RISK_FILL, 3: RISK_FILL, 4: PENDING_FILL},
    )

    add_heading(doc, "变更文件（中文）", 1)
    add_table(
        doc,
        ["Category", "Count / examples", "Meaning"],
        [
            ("Session logs", "2 in-window files: `2026-07-05-A16.md`, `2026-07-05-A25.md`", "Fresh direct work evidence."),
            ("Release intake", "363 files under `coordination/release-intake/`; 296 A25-named, 8 A22-named, 59 `.mjs`, 46 `.pathspec`, 105 `.json`, 153 `.md`", "A25/A22 governance evidence, current gates, owner packets, and work orders."),
            ("Reports", "13 files under `coordination/reports/`, including prior `2026-07-05-president-report.docx` and render artifacts", "Previous report deliverable/render outputs inside the window."),
            ("Local metadata", "`.DS_Store`, `coordination/.DS_Store`", "Local/generated metadata; do not use as product evidence."),
            ("Feature code", "No feature-code filesystem changes were detected within the window filter; current root remains dirty from older accumulated work.", "Report run did not edit feature code."),
        ],
        [1900, 4200, 3260],
        body_size=7.5,
    )

    add_heading(doc, "明日优先事项（中文）", 1)
    add_numbered(doc, "A25 refresh dirty-map and downstream currentness only after the owner confirms the compose-worktree deletion hold is resolved, then rerun no-staged and release-source gates.")
    add_numbered(doc, "A22 keep release path clean-source only; do not publish from dirty root.")
    add_numbered(doc, "A06/A12/A11/A22/A25 process the five focus-batch owner-package final-state rows with exact owner decisions.")
    add_numbered(doc, "A25/A22 decide whether six A25 artifact-clean rows should receive separate execution instructions; keep `wave01-resync-01-tsconfig-json` held unless explicitly changed.")
    add_numbered(doc, "A08/A12/A10/A22 prioritize shared-contract and build/type-check blockers before wider runtime package closure.")

    add_heading(doc, "需要 Owner 决策（中文）", 1)
    add_bullet(doc, "请 Dr. Peter Hu 确认 `/Users/dongpinhu/.config/superpowers/worktrees/MAIS-MVP/A10-A22-A08-A12-A06-compose-20260628` 的 exact deletion 是否已完成，以解除 A25 validation hold。")
    add_bullet(doc, "请决定五个 focus-batch owner-package final-state rows：A25、A22、A06、A12、A11，是 reviewed commit、exact-path discard、evidence archive 还是 blocker。")
    add_bullet(doc, "请决定六个 A25 artifact-clean rows 是否给 separate execution instruction；目前 cleanupAuthorizedRows=0，executableRows=0。")
    add_bullet(doc, "请决定 A22 generated-artifact residual cleanup（`.tmp` 与 `.next`，约 48.7 GB）是否进入 exact approval；不要使用 broad cleanup。")
    add_bullet(doc, "请确认 dirty-root 继续只作为 inventory/reporting source，不作为 production deploy source。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
