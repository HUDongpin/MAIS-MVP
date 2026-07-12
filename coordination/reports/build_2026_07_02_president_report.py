from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-02-president-report.docx")


BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_para(doc, text="", style=None, size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    hdr = table.rows[0].cells
    for idx, h in enumerate(headers):
        hdr[idx].text = ""
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=9.5, bold=True, color=BLACK)
        set_cell_shading(hdr[idx], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=9.2, color=BLACK)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DEEP_BLUE)]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-02"),
        ("Reporting window", "2026-07-01 08:00 to 2026-07-02 08:00 Asia/Hong_Kong"),
        ("Prepared by", "A10 reporting automation; source evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL)


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "A22 已将获批的 P0 lesson checklist 候选包部署到 `mais.hk`、`www.mais.hk`、`mais.ac`、`www.mais.ac`；现场 API 与浏览器 smoke 均通过，原始 P0 路由已修复。"),
            ("2", "A05/A12 修复了另一个 US-CA lesson attempt feedback 问题：Postgres 行写入失败不再阻断已计算的正确反馈；聚焦测试通过，但尚未由 A22 清洁切片发布。"),
            ("3", "A20/A04/A12/A22 修复了 Adventure Island 从 `/practice` 触发重复资格请求导致 Vercel 300 秒超时的风险；源码回归、Playwright 聚焦用例通过，仍需清洁发布路径。"),
            ("4", "A06 在 Manim v2/Visualization closure 上推进明显：目录解耦、final-audit evidence ID、review-slice manifest、A11 root-attribute guard 等通过聚焦与安全全量测试；但 A11/A18/A22/final 4/4 gate 仍未完成。"),
            ("5", "A25 最新非破坏性 dirty-map 通过， expanded status entries 为 3387；owner package matrix 仅 A16 ready，15 个 package blocked，root 仍禁止作为发布源。"),
        ],
        [700, 8660],
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "A22 deployed the approved P0 lesson-checklist candidate to all custom production domains; live authenticated API and browser smoke passed."),
            ("2", "A05/A12 fixed a separate lesson attempt feedback failure so computed feedback survives best-effort persistence outages; the fix is locally verified but not yet clean-slice released."),
            ("3", "A20/A04/A12/A22 addressed the Adventure Island Vercel timeout mechanism with a focused source regression and one targeted Playwright pass."),
            ("4", "A06 advanced Manim v2 closure architecture and evidence propagation, but final A11 browser, A18 teaching, A22 clean-release, and final 4/4 audit gates remain open."),
            ("5", "A25 release intake shows 3387 expanded dirty entries. The root is still inventory-only and cannot be used as a production source."),
        ],
        [700, 8660],
    )

    add_heading(doc, "汇报窗口摘要", 1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("窗口", "2026-07-01 08:00 至 2026-07-02 08:00 Asia/Hong_Kong。报告生成时间晚于窗口，用于补做 DOCX 与 fresh dirty-map。"),
            ("活跃证据", "新鲜/窗口内更新的 session logs：A05/A12、A06、A10、A20、A22、A25；另有 A25 release-intake 大量 currentness/approval artifacts。"),
            ("阻塞/决策文件", "窗口内无新的 `coordination/blockers/` 或 `coordination/decisions/` 文件；活跃阻塞来自 session logs 与 A25 release-intake reports。"),
            ("新鲜检查", "`npm run release:dirty-map -- --reason \"2026-07-02 8 AM president report fresh release-intake baseline\"` 通过，生成 `2026-07-02-A25-dirty-tree-map-20260702T014952Z.md`。"),
        ],
        [1800, 7560],
    )

    add_heading(doc, "整体项目进展", 1)
    add_para(doc, "项目继续从功能修复、可视化闭环和 release-intake 三条线推进，但整体发布状态仍由 A25/A22 clean-source gate 控制。生产上，A22 已完成 P0 lesson checklist 的获批发布；开发上，A05/A12 与 A20/A04/A12/A22 分别完成两个生产相关问题的局部修复；A06 完成多个 Manim v2 closure architecture 切片。")
    add_para(doc, "治理上，A25 将 dirty-root remediation 扩展为 owner package blocker matrix、owner closure queue、next approval packet 和 no-dirty-root-deploy evidence。当前状态不是“可发布”，而是“证据更清楚、下一步授权更明确”。")

    add_heading(doc, "A01-A25 会话状态表", 1)
    status_rows = [
        ("A01", "App shell", "无新日志；A25 matrix blocked", "root package/Playwright/type-check closure需owner处理。"),
        ("A02", "Dashboard", "无新日志；A02/A15 package blocked", "dashboard/adaptive package仍需 owner package approval 与修复。"),
        ("A03", "Roadmap", "无新日志；package blocked", "A25列为 P3 owner approval，roadmap check仍红。"),
        ("A04", "Practice", "与 A20 联动完成局部修复", "Adventure Island timeout focused regression通过；仍需清洁发布切片。"),
        ("A05", "Lesson", "完成 A05/A12 lesson feedback bugfix", "practice attempt fast path best-effort persistence；聚焦 tests green。"),
        ("A06", "Visualization", "高活跃；in progress/blocked", "Manim evidence architecture推进；A11/A18/A22/final 4/4 gate未关。"),
        ("A07", "AI tutor", "无新日志；A25 matrix blocked", "AI tutor type/playwright package still blocked in Wave 05。"),
        ("A08", "State/analytics", "无新日志；shared-contract package blocked", "A08/A12 shared-contract package仍有 test/type/build failures。"),
        ("A09", "Copy/i18n/a11y", "无新日志；package blocked", "Wave 05 copy/i18n/accessibility type-check blocked。"),
        ("A10", "Tooling/report", "活跃", "生成本报告；消费 A25/A22/A06/A05/A12/A20 evidence。"),
        ("A11", "QA", "无新日志；被 A06/A22 消费", "P0 checklist旧证据已被A22消费；A06 broad browser evidence仍待A11。"),
        ("A12", "Backend/API", "与 A05/A20 联动活跃", "lesson feedback persistence path 与 Adventure Island API/persistence风险被局部处理。"),
        ("A13", "Teacher console", "无新日志；package blocked", "A25列出 console type/playwright blocked。"),
        ("A14", "Parent console", "无新日志；package blocked with A13", "需与 A13 console package一起处理。"),
        ("A15", "Adaptive", "无新日志；与 A02 package blocked", "dashboard/adaptive package仍在 owner closure queue。"),
        ("A16", "Research", "无新日志；A25 row ready", "Wave 05 中唯一 ready package。"),
        ("A17", "Motivation", "无新日志；A17/A20 package blocked", "game/motivation type/playwright package仍红。"),
        ("A18", "Curriculum QA", "无新日志；A06 gate consumer", "A06 final teaching acceptance仍需 A18。"),
        ("A19", "API env", "无新日志", "owner closure queue仍有 1 pending item；无新credential工作。"),
        ("A20", "Games", "完成 Adventure Island timeout局部修复", "node source regression 5/5；Playwright focused 1/1。"),
        ("A21", "Content/RAG", "无新日志；package pending", "A18/A21 content evidence仍 blocked，需要签核/清洁切片。"),
        ("A22", "Release engineering", "完成 P0 production deploy；release gates继续阻塞root", "clean candidate发布成功；root dirty release仍禁止。"),
        ("A23", "Integration", "无新日志；package pending", "content promotion chain仍需A23计划，不可绕过。"),
        ("A24", "Exact layer", "无新日志；package pending", "A18/A21/A23/A24 content evidence仍 blocked。"),
        ("A25", "Git hygiene/release intake", "高活跃；blocked by owner decisions", "dirty-map 3387；owner closure queue 134 pending，0 executable。"),
    ]
    add_table(doc, ["Agent", "职责", "窗口状态", "证据/下一步"], status_rows, [700, 1900, 2300, 4460])

    add_heading(doc, "已完成工作", 1)
    add_table(
        doc,
        ["Owner", "完成项", "验证"],
        [
            ("A22", "P0 lesson checklist candidate deployed to production custom domains.", "preflight, npm ci, type-check, build, live API/browser smokes passed."),
            ("A05/A12", "Lesson attempt feedback now returns computed feedback even if row persistence fails.", "3/3 store tests, 21/21 focused suite, exact invalid-Postgres mode returned correct feedback."),
            ("A20/A04/A12/A22", "Adventure Island repeated eligibility refetch risk reduced in `/practice`.", "source regression 5/5, targeted Playwright CTA test 1/1, diff check clean."),
            ("A06", "Manim closure tests decoupled from live catalog; final-audit evidence and review-slice propagation improved.", "latest safe Manim/teaching suites reported 359/359 and focused/adjacent suites green."),
            ("A25", "Release-intake artifacts regenerated: owner package matrix, closure queue, no-dirty-root-deploy evidence, next approval packet.", "A25 currentness gates reported pass where applicable; no staging/cleanup/deploy executed."),
        ],
        [1600, 4300, 3460],
    )

    add_heading(doc, "进行中工作", 1)
    add_table(
        doc,
        ["Owner", "状态"],
        [
            ("A06/A11/A18/A22", "Manim v2 overall owner goal still open until A18 teaching acceptance, A11 broad browser evidence, A22 clean release evidence, and final 4/4 objective audit pass."),
            ("A25/A10/A22", "Dirty-root remediation remains in approval/routing state; no package rows are executable without owner-selected final states."),
            ("A05/A12", "Lesson feedback fix is locally verified but needs clean release packaging if it should reach production."),
            ("A20/A04/A12/A22", "Adventure Island timeout fix is locally verified but needs clean slice/release decision before production."),
        ],
        [1800, 7560],
    )

    add_heading(doc, "阻塞项", 1)
    add_table(
        doc,
        ["阻塞", "影响", "下一步"],
        [
            ("Dirty root", "3387 expanded status entries; root cannot be production source.", "A22 only releases clean worktree/clean clone/reviewed slice/pruned staging。"),
            ("Owner package matrix", "16 rows, only 1 ready; 15 blocked; 35 failed checks and 9907 type-check error lines in A25 evidence.", "Owner sessions must handle package-specific failures or write blocker reports。"),
            ("Physical lifecycle", "Wave 06 final closure not ready; 33 physical lifecycle approvals represented globally.", "Owner must pick final states/commands before cleanup or closure。"),
            ("A06 final acceptance", "Visualization/Manim proof is strong locally but not release-accepted.", "A18/A11/A22/final 4/4 evidence required。"),
            ("A25 owner-routing blocker", "A25 cannot fix A01/A02/A04/A15/A18/A21/A23/A24 failures inside A25 write scope.", "Route package rows to owning sessions。"),
        ],
        [2300, 3300, 3760],
    )

    add_heading(doc, "风险", 1)
    add_table(
        doc,
        ["风险", "说明", "控制"],
        [
            ("Dirty-root release risk", "Many feature and coordination changes coexist; full type/build claims from root are not reliable.", "继续执行 clean-source release rule。"),
            ("Type-check signal过宽", "A25 matrix显示大量 type errors，且 A06说明 scoped tsc 会穿透 forbidden/out-of-scope topic graph。", "按 owner package 切片，用 focused tests 支撑局部结论。"),
            ("未发布修复风险", "A05/A12 与 A20 fix 已验证但未说明已 production deploy。", "若要上线，交由 A22 clean slice发布。"),
            ("Deployment protection blind spot", "A22 raw Vercel URL受SSO保护，未使用 bypass secret。", "以 custom-domain live smoke为生产证据；如需raw URL smoke需owner-approved bypass。"),
            ("Generated artifacts/disk pressure", "A25窗口内记录过 ENOSPC；当前 df 显示约 26GiB available。", "cleanup apply仍需A22/owner exact authorization。"),
        ],
        [2300, 3800, 3260],
    )

    add_heading(doc, "测试 / 构建状态", 1)
    add_table(
        doc,
        ["Scope", "Result"],
        [
            ("Fresh A25 dirty-map", "Passed; `2026-07-02-A25-dirty-tree-map-20260702T014952Z.md`; expanded entries 3387."),
            ("A22 P0 production deploy", "Passed: release preflight, npm ci, type-check, build, built-server smoke, Vercel Ready, custom-domain live API/browser smoke."),
            ("A05/A12 lesson feedback", "Passed focused suites; broad `npm run type-check` not clean in dirty root due unrelated A06/tmp validator drift."),
            ("A20 Adventure Island", "Passed source regression 5/5 and targeted Playwright 1/1; broad type-check still red on unrelated dirty root files."),
            ("A06 Manim/Visualization", "Focused and safe Manim suites green; repo-wide type-check intentionally not rerun because dirty-root graph is unsafe/out-of-scope."),
            ("Report automation", "Broad build/type-check not run from dirty root; risk recorded. DOCX validation performed after generation."),
        ],
        [2700, 6660],
    )

    add_heading(doc, "变更文件", 1)
    add_table(
        doc,
        ["类别", "代表文件/范围"],
        [
            ("生产发布证据", "`coordination/reports/2026-07-01-A22-p0-lesson-checklist-production-deploy.md`; A22 session log。"),
            ("A05/A12 bugfix", "`lib/server/practiceAttemptStore.ts`, `lib/server/practiceAttemptStore.test.ts`。"),
            ("A20 timeout fix", "`app/practice/page.tsx`, `app/practice/practiceArenaPageRegressions.test.ts`。"),
            ("A06 Manim/Visualization", "`components/visualizations/three/manim/*`, `coordination/reports/2026-06-28-manim-v2-completion-status.md`。"),
            ("A25 release intake", "`coordination/release-intake/2026-07-02-A25-*`, `latest-A25-*`, owner pathspec/work-order artifacts。"),
            ("当前 dirty-map 摘要", "390 tracked modified, 1 tracked deleted, 1032 untracked status entries, 2996 untracked files, 3387 expanded status entries。"),
        ],
        [2200, 7160],
    )

    doc.add_page_break()
    add_heading(doc, "明日优先级", 1)
    add_table(
        doc,
        ["优先级", "事项"],
        [
            ("P0", "Owner decide Wave 01 package-resync approvals or keep them blocked; no physical cleanup without exact approval IDs and commands."),
            ("P0", "A22 continue rejecting dirty-root deploy; any new production work must use clean/pruned source."),
            ("P1", "Package A05/A12 lesson feedback fix into a clean release slice if production promotion is desired."),
            ("P1", "Package A20 Adventure Island timeout fix into a clean release slice if production promotion is desired."),
            ("P1", "A06 close remaining final acceptance evidence with A18, A11, A22, and final 4/4 audit."),
            ("P2", "A25 route pending owner blocker reports to actual owner sessions and refresh the matrix after each closure."),
        ],
        [900, 8460],
    )

    add_heading(doc, "需要 Owner 决策", 1)
    add_table(
        doc,
        ["Decision", "Owner action needed"],
        [
            ("Wave 01 package resync", "Approve or reject the seven exact resync approvals listed in `2026-07-02-A25-next-owner-approval-packet.md`."),
            ("Owner package final states", "Choose selected final state for the first owner package approvals, starting with A25/A22, A06, A12, A11, A10, A18/A21, A05."),
            ("Physical lifecycle", "Choose final state for dirty/open worktrees and clean-diverged branches; no branch/worktree cleanup is authorized yet."),
            ("A05/A12 production path", "Decide whether the lesson feedback bugfix should be promoted via A22 clean slice."),
            ("A20 production path", "Decide whether the Adventure Island timeout fix should be promoted via A22 clean slice."),
            ("A06 acceptance", "Confirm required A18 teaching review, A11 browser scope, A22 clean release evidence, and final objective audit path."),
        ],
        [2600, 6760],
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("MAIS-MVP daily report - 2026-07-02 - A10/A25 evidence based")
    set_run_font(r, size=9, color=MUTED)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
