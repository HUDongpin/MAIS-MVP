from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("coordination/reports/2026-07-05-president-report.docx")

FONT = "Arial Unicode MS"
BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"
RISK_FILL = "FCE4D6"
OK_FILL = "E2F0D9"
PENDING_FILL = "FFF2CC"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_para(doc, text="", size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.10
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DEEP_BLUE)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, body_size=8.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    hdr = table.rows[0].cells
    repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        hdr[idx].text = ""
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, bold=True, color=BLACK)
        set_cell_shading(hdr[idx], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=body_size, color=BLACK)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.0)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, size=10.0)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DEEP_BLUE),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(doc):
    add_para(doc, "MAIS-MVP PRESIDENT REPORT", size=23, bold=True, after=3)
    add_para(doc, "Daily bilingual coordination report for Dr. Peter Hu", size=13, color=MUTED, after=10)
    rows = [
        ("Report date", "2026-07-05"),
        ("Report time", "8:00 AM Asia/Hong_Kong"),
        ("Reporting window", "2026-07-04 08:00 to 2026-07-05 08:00 Asia/Hong_Kong"),
        ("Prepared by", "A10 reporting automation; evidence from A01-A25 coordination files"),
        ("Automation ID", "mais-mvp-9-am-president-report"),
        ("Repository", "/Users/dongpinhu/Desktop/MAIS-MVP on main; root remains inventory-only"),
    ]
    add_table(doc, ["Field", "Value"], rows, [2100, 7260], header_fill=LIGHT_BLUE_FILL)


def build_doc():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_heading(doc, "Chinese Executive Summary", 1)
    add_table(
        doc,
        ["序号", "重点"],
        [
            ("1", "本窗口有 fresh AI session activity，主要集中在 A25-owned git hygiene/release intake。A25 将 dirty-root closure loop 从普通状态盘点推进到 owner-authorization、blocker-report、A16 guarded extraction、release-source blocker 的可验证链路。"),
            ("2", "A16 research evidence package 在 owner approval 与 separate execution instruction 后被 A25 守护式提交：`ce2ae5258 Add A16 research evidence package`，只包含 6 个 `coordination/research/` 文件；提交后这些 A16 路径 `git status --short -- <paths>` 无输出。"),
            ("3", "报告时 fresh checks 通过：`npm run type-check -- --pretty false` 通过；`npm run test:analytics` 通过 27/27；`git diff --check` 无输出；`node coordination/release-intake/assert-no-staged-changes.mjs` 显示 staged entries 0。"),
            ("4", "发布仍 blocked：Fresh A25 dirty map 显示 4,625 expanded status entries；A22 release-source clean gate failed；A25 aggregate currentness 在最新 dirty map 后 44/92 currentness checks，48 failures，说明多个 release-intake packet 需要重新同步。Root `main` 不能作为生产发布源。"),
            ("5", "今天最需要 Dr. Peter Hu 关注：A22/A25 clean-source 发布路径、A25 authorization rows 与 Wave01 rows 的下一轮决策、A22 generated-artifact residual cleanup 是否进入 exact approval、以及 A08/A12/A10/A22 shared-contract/package-baseline 修复优先级。"),
        ],
        [700, 8660],
    )

    add_heading(doc, "English Executive Summary", 1)
    add_table(
        doc,
        ["No.", "Summary"],
        [
            ("1", "Fresh activity exists in this window, led by A25-owned git hygiene and release intake. The work did not change feature code; it hardened the dirty-root closure process and owner-decision evidence."),
            ("2", "The concrete completion is the owner-approved A16 research-evidence extraction: commit `ce2ae5258 Add A16 research evidence package`, containing exactly six `coordination/research/` evidence files."),
            ("3", "Fresh report checks are green for type-check, analytics tests, whitespace diagnostics, and no staged entries. The analytics suite passed 27/27."),
            ("4", "Release readiness is still blocked. The fresh dirty map has 4,625 expanded entries, A22 release-source-clean fails, and the A25 aggregate currentness gate is stale against the newest map."),
            ("5", "Owner decisions should focus on clean-source release policy, remaining authorization rows, Wave01/package resync, generated-artifact cleanup approval, and shared-contract/package-baseline remediation sequencing."),
        ],
        [700, 8660],
    )

    add_heading(doc, "汇报窗口摘要", 1)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("窗口", "2026-07-04 08:00 至 2026-07-05 08:00 Asia/Hong_Kong。"),
            ("活跃 session logs", "窗口内 fresh formal session log：`coordination/session-logs/2026-07-04-A25-dirty-worktree-governance-status.md`。A01-A24 无 fresh standalone session log。"),
            ("No assigned work flag", "No assigned work in this reporting window: No。窗口内存在明确 A25 execution/verification activity，并有 owner-approved A16 evidence package commit。"),
            ("Blockers/decisions folders", "`coordination/blockers/` 与 `coordination/decisions/` 在窗口内无新增文件；blocking evidence 主要来自 `coordination/release-intake/` 与 A25 session log。"),
            ("文件时间片", "窗口内 mtime 变动 342 files，全部在 `coordination/` 下；主要为 release-intake packets、reports/render artifacts、A25 status log。"),
            ("Fresh report run after window", "报告生成时又运行 fresh checks；这些是 report-owner checks，不计入窗口内 session activity。"),
        ],
        [1900, 7460],
    )

    add_heading(doc, "整体项目进展", 1)
    add_para(doc, "项目在窗口内主要推进的是 release governance，而不是新功能。A25 把 dirty-root closure loop 拆成更明确的 owner-decision surfaces：ready candidate、Wave01 package resync、remaining owner-package final states、physical-lifecycle states、A22 generated-artifact cleanup。")
    add_para(doc, "A16 research evidence package 已完成最小可审查提交：commit `ce2ae5258` 只含 6 个 research/provenance/CSV 文件，没有 feature code、cleanup、deploy、push、branch deletion 或 worktree removal。")
    add_para(doc, "质量状态分裂明显：root type-check 与 analytics 在当前报告环境中通过，但 release-source 仍不干净，A25 aggregate evidence 需要按最新 dirty map 重新同步后才能作为当前 release-intake gate 使用。")

    add_heading(doc, "A01-A25 会话状态表", 1)
    status_rows = [
        ("A01", "App shell", "无 fresh log", "无窗口内 A01 session activity。", "受 dirty-root release/source gate 影响。", "未单独运行。"),
        ("A02", "Dashboard", "无 fresh log", "无窗口内 A02 session activity。", "A02/A15 dashboard-adaptive closure 仍依赖 shared contracts。", "未单独运行。"),
        ("A03", "Roadmap", "无 fresh log；A25 blocker referenced", "A25 blocker reports 仍记录 roadmap package 受 shared schema/i18n/package baseline 影响。", "A03 local remediation 不应绕过 A08/A09/A10/A22。", "未在本报告窗口新跑。"),
        ("A04", "Practice", "无 fresh log", "上一窗口 bug-fix 仍是 dirty-root evidence；本窗口无 A04 新动作。", "clean release slice 未形成。", "未单独运行。"),
        ("A05", "Lesson", "无 fresh log", "无窗口内 A05 session activity。", "lesson/content packages 仍需 A18/A21/A23/A24 gates。", "未单独运行。"),
        ("A06", "Visualization", "无 fresh log；A25 blocker referenced", "A25 blocker reports 仍显示 A06 visualization package 依赖 A10/A22 package baseline resync。", "Manim/visualization closure 不应在缺 shared/package gate 时单独推进。", "未在本报告窗口新跑 A06 suite。"),
        ("A07", "AI tutor", "无 fresh log；A25 blocker referenced", "A25 blocker reports 仍显示 A07 AI tutor package 依赖 A12 userStore/aiGovernance 与 A08 shared types。", "A07/A12 split 仍是治理风险。", "未单独运行。"),
        ("A08", "State/analytics", "无 fresh log", "Fresh root type-check 与 analytics passed，但 isolated owner worktrees 仍暴露 shared schema/storage drift。", "A08 shared schema remains critical path。", "Root `npm run type-check` PASS；`npm run test:analytics` 27/27 PASS。"),
        ("A09", "Copy/i18n/a11y", "无 fresh log", "无窗口内 A09 activity。", "US grade/copy helper gaps remain routed through A03/A15 blockers。", "未单独运行。"),
        ("A10", "Reporting", "活跃：本报告", "读取 AGENTS、automation memory、A25 log、release-intake packets、current git state；生成本 DOCX。", "报告-only；未编辑 feature code；build 未跑。", "DOCX build/validation/render QA。"),
        ("A11", "QA", "无 fresh log；A25 blocker referenced", "A25 blocker reports 仍显示 A11 regression evidence package 受 generated content/shared curriculum/A06/A13/A20 blockers影响。", "A11 clean regression cannot proceed until upstream gates current。", "未在窗口新跑 A11 package。"),
        ("A12", "Backend/API", "无 fresh log；A25 routing referenced", "A25 已在此前修正 A12 recommended worktree；本窗口仍作为 blocker dependency。", "A12/A07 ownership split 与 userStore/API exports remain critical。", "未单独运行。"),
        ("A13", "Teacher console", "无 fresh log；A25 blocker referenced", "A25 blocker reports 仍显示 A13/A14 console package 受 shared contracts/package baseline 影响。", "A13 cannot resolve alone without A08/A12/A10/A22。", "未单独运行。"),
        ("A14", "Parent console", "无 fresh log", "无窗口内 A14 session activity。", "Physical lifecycle decisions still in A25 backlog。", "未单独运行。"),
        ("A15", "Adaptive engine", "无 fresh log；A25 blocker referenced", "A25 blocker reports 仍显示 A02/A15 dashboard-adaptive package受 A08/A12/A07/A10/A22 影响。", "Adaptive local fixes需 shared/provider/package baseline先行。", "未单独运行。"),
        ("A16", "Research", "完成 package extraction", "A16 research evidence package committed as `ce2ae5258` with six `coordination/research/` files。", "No cleanup or worktree lifecycle action authorized。", "A25 post-extraction verification 7/7；closeout 8/8；path status clean。"),
        ("A17", "Gamification", "无 fresh log", "无窗口内 A17 session activity。", "A17/A20 game-motivation package still routed as blocker evidence。", "未单独运行。"),
        ("A18", "Content QA", "无 fresh log；A25 blocker referenced", "A25 blocker reports still hold A18/A21 content evidence until shared gates are current。", "Candidate-to-live chain remains required。", "未单独运行。"),
        ("A19", "API env", "无 fresh log", "无窗口内 A19 activity；无 secret/env work performed。", "Provider/env changes仍需 A19 redacted process。", "未单独运行。"),
        ("A20", "Games", "无 fresh log；A25 blocker referenced", "A25 blocker reports still show game-motivation package routing/shared gate blockers。", "A20 should not fix beyond game scope until routing correct。", "未单独运行。"),
        ("A21", "Content/RAG", "无 fresh log", "无窗口内 A21 activity；content/RAG remains candidate-only unless promoted through A18/A23。", "Live data promotion blocked without QA/integration/release evidence。", "未单独运行。"),
        ("A22", "Release", "Gate checked; blocked", "Fresh A22 release-source clean gate failed closed at 4,625 expanded entries。", "Dirty root cannot deploy；clean worktree/clone/reviewed slice/pruned staging required。", "`assert-release-source-clean` FAIL；no staged entries PASS。"),
        ("A23", "Integration", "无 fresh log", "无窗口内 A23 activity。", "Candidate-to-live integration not ready while owner packages dirty。", "未单独运行。"),
        ("A24", "Illustration exact-layer", "无 fresh log", "无窗口内 A24 activity。", "Exact-layer/live promotion depends on A18/A21/A23 gates。", "未单独运行。"),
        ("A25", "Git hygiene/release intake", "高活跃", "Built owner-decision packets; recorded owner approvals/instructions; applied guarded A16 extraction; refreshed dirty maps; no staged entries。", "Latest A25 aggregate currentness stale after new dirty map: 44/92, 48 failures；release source still dirty。", "Dirty map PASS；A16 verification PASS；aggregate currentness FAIL after latest map。"),
    ]
    add_table(doc, ["ID", "工作流", "状态", "完成/进展", "阻塞/风险", "检查"], status_rows, [650, 1050, 1000, 2920, 2660, 1080], body_size=6.9)

    add_heading(doc, "已完成工作", 1)
    for item in [
        "A25/A16：在 owner approval 与 separate execution instruction 后，guarded extraction 提交 `ce2ae5258 Add A16 research evidence package`；提交只包含 6 个 `coordination/research/` 文件。",
        "A25：完成 A16 post-extraction verification、extraction closeout、execution readiness、guarded plan/executor 证据同步；A16 package dirty rows 变为 0。",
        "A25/A22：生成并刷新 release-source blocker、generated-artifact residual acceptance、authorization gap/round、owner-decision focus 等 evidence packets。",
        "A10：完成本次 daily president report evidence collection、fresh checks、DOCX artifact generation；未修改 feature code。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "进行中工作", 1)
    for item in [
        "A25 closure loop：仍处于 validate。A16 extraction closed one small owner package，但 root dirty tree 和 release gates remain open。",
        "A22 release source：需要 clean worktree、clean clone、reviewed clean release slice 或 owner-approved pruned staging package；dirty root deploy remains blocked。",
        "A08/A12/A10/A22 shared-contract/package-baseline work：仍是 A03/A04/A07/A11/A13/A15/A18/A20 owner-package closure 的前置条件。",
        "A22 generated-artifact cleanup：2 个 residual targets 共 45.35 GiB 已有 evidence docket，但 0 cleanup-authorized rows、0 executable rows。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "阻塞项", 1)
    blocker_rows = [
        ("A22", "Release-source clean blocked", "Fresh gate failed：4,625 expanded status entries；Branch main；HEAD `ce2ae5258`；allowed sources only clean/pruned/reviewed release paths。"),
        ("A25", "Aggregate currentness stale", "`assert-dirty-worktree-remediation-current` after fresh map: 44/92 currentness checks, 48 failures；many packets need refresh to match 4,625-entry map。"),
        ("A25/A22", "Cleanup not authorized", "A22 residual `.tmp`/`.next` total 45.35 GiB；no cleanup apply can run until exact owner/A22 authorization and validation。"),
        ("A25", "Remaining authorizations", "A16 package consumed, but owner-package/final-state/physical-lifecycle/Wave01 rows remain the next decision surface。"),
        ("A08/A12/A10/A22", "Shared/package drift", "Multiple owner packages remain blocked by shared type, storage/API export, and package dependency baseline drift from prior A25 blocker evidence。"),
    ]
    add_table(doc, ["Owner", "阻塞", "证据/影响"], blocker_rows, [1300, 2300, 5760], header_fill=RISK_FILL, body_size=8.4)

    add_heading(doc, "风险", 1)
    for item in [
        "Root dirty surface remains large：current git status has 1,451 status entries; fresh A25 map expands this to 4,625 entries。Any release from root risks mixing unrelated owner/agent changes。",
        "A25 aggregate packets can become stale quickly because report/check artifacts themselves move the dirty-map signature and expanded-entry count。",
        "A16 extraction is a narrow success, not a broad cleanup: it does not authorize cleanup, branch deletion, worktree removal, push, deploy, or unrelated staging。",
        "A22 generated-artifact residual cleanup has a large disk footprint but must remain blocked until exact approval; using broad `git clean` or manual deletion would violate current governance。",
        "`npm run build` was not run from dirty root in this report run; build readiness must come from A22 clean release path, not dirty-root evidence。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "测试与构建状态", 1)
    checks = [
        ("Fresh root type-check", "`npm run type-check -- --pretty false`", "PASS。"),
        ("Fresh analytics", "`npm run test:analytics`", "PASS；27/27。"),
        ("Fresh whitespace check", "`git diff --check`", "PASS；无输出。"),
        ("Fresh A25 dirty map", "`npm run release:dirty-map -- --reason \"2026-07-05 8 AM president report fresh release-intake baseline\"`", "PASS；report `2026-07-05-A25-dirty-tree-map-20260705T012052Z.md`；4,625 expanded entries。"),
        ("No staged changes", "`node coordination/release-intake/assert-no-staged-changes.mjs`", "PASS；staged entries 0。"),
        ("A22 release-source clean", "`node coordination/release-intake/assert-release-source-clean.mjs`", "FAIL；release source dirty with 4,625 expanded entries。"),
        ("A25 aggregate currentness", "`node coordination/release-intake/assert-dirty-worktree-remediation-current.mjs`", "FAIL；44/92 currentness checks, 48 failures after fresh dirty map。"),
        ("Build", "`npm run build`", "Not run；report-only run in dirty root。A22 build/release evidence should come from clean worktree/clone/reviewed slice/pruned staging。"),
    ]
    add_table(doc, ["检查", "命令/范围", "结果"], checks, [1700, 4200, 3460], header_fill=OK_FILL, body_size=8.2)

    add_heading(doc, "变更文件", 1)
    add_para(doc, "窗口内 mtime 变动：342 files，全部位于 `coordination/`。Git commit 变动另计：A16 package extraction commit touched 6 `coordination/research/` files.")
    changed_rows = [
        ("Committed during window", "coordination/research/2026-06-20-3blue1brown-skill-distillation-handbook.md; .provenance.md", "`ce2ae5258` A16 research evidence package。"),
        ("Committed during window", "coordination/research/2026-06-20-3blue1brown-style-visualization-effects-report.md", "`ce2ae5258` A16 research evidence package。"),
        ("Committed during window", "coordination/research/2026-06-21-ke-to-mais-behavior-event-dictionary.csv", "`ce2ae5258` A16 research evidence package。"),
        ("Committed during window", "coordination/research/2026-06-21-ke-to-mais-stealth-assessment-integration-spec.md; .provenance.md", "`ce2ae5258` A16 research evidence package。"),
        ("A25 session log", "coordination/session-logs/2026-07-04-A25-dirty-worktree-governance-status.md", "2,411-line status log; latest section records A16 guarded extraction applied and verified。"),
        ("A25 release intake", "coordination/release-intake/2026-07-04-A25-* and 2026-07-05-A25-*", "Owner decision packets, authorization rounds, A16 execution/readiness/closeout, release-source blocker, residual cleanup dockets。"),
        ("Report artifacts", "coordination/reports/2026-07-04-president-report.docx and render folder", "Prior day report artifacts touched after 2026-07-04 08:00 and counted in the mtime slice。"),
        ("Current dirty status", "Repository-wide", "Fresh map: 391 tracked modified, 1 tracked deleted, 1,059 untracked status entries, 4,233 untracked files, 4,625 expanded entries。"),
    ]
    add_table(doc, ["类别", "文件/区域", "说明"], changed_rows, [1500, 4450, 3410], body_size=7.9)

    add_heading(doc, "明日优先级", 1)
    for item in [
        "A25/A10：先刷新/重跑 A25 aggregate currentness after the latest dirty map；不要基于 stale 4,616/4,623-entry packets 做 release decision。",
        "A22：从 clean worktree、clean clone、reviewed clean release slice 或 owner-approved pruned staging 跑 build/release preflight；继续阻止 dirty-root deploy。",
        "Dr. Peter Hu/A25：审核下一轮 owner authorization rows，尤其 Wave01 package resync、remaining owner-package final states、physical lifecycle final states、A22 residual cleanup approvals。",
        "A08/A12/A10/A22：优先解决 shared schema/storage/API/package baseline drift，使 A03/A04/A07/A11/A13/A15/A18/A20 owner worktrees可以各自收敛。",
        "A22/A25：若要清理 `.tmp`/`.next` residuals，先运行 dry-run/current evidence，再记录 exact approval；不要使用 broad destructive cleanup。",
    ]:
        add_numbered(doc, item)

    add_heading(doc, "需要 Dr. Peter Hu 决策", 1)
    decision_rows = [
        ("1", "是否允许 A25/A22 将当前 release-intake packets 重新同步到 4,625-entry dirty map，并将 stale packet refresh 作为下一步优先任务。"),
        ("2", "对 Wave01 package-resync rows、remaining owner-package final states、physical-lifecycle final states 的下一轮 approve/reject/defer 顺序做决策。"),
        ("3", "是否授权 A22 generated-artifact residual cleanup 进入 exact approval 流程；当前 `.tmp` + `.next` residual evidence 是 45.35 GiB，但 0 rows executable。"),
        ("4", "是否要求 A22 建立 clean release slice 来验证/发布之前 A11/A04/A06 等已验证但仍在 dirty-root inventory 中的工作。"),
        ("5", "确认 A08/A12 shared-contract 和 A10/A22 package-baseline resync 是否作为明日最高优先级，以解除多个 owner package blocker。"),
    ]
    add_table(doc, ["序号", "决策事项"], decision_rows, [700, 8660], header_fill=PENDING_FILL, body_size=8.6)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
